"""Los cinco en duda de GC5 por HTTP (fase C6, pista A): los llamadores de la API del nucleo de C6.

- INV-093 (UC-06 4a): `POST /partidas/{codigo}/rendimientos` registra la justificacion de la
  advertencia de RF-27, solo si hubo advertencia, ligada al rendimiento de esa misma peticion.
- INV-090 (UC-02 2a y 2b): `POST /listas-precios` da de alta los desconocidos elegidos y reporta
  aparte los ignorados (2a) y los rechazados por unidad (2b).
- INV-086 (UC-01 1): `POST /proyectos` crea el proyecto sobre el que luego se elabora.
- INV-076 (RF-24, UC-05 1 y 2): `POST /presupuestos/importar` importa el XLSX de un tercero, lo
  audita y lo persiste; el informe y el libro salen por los `GET` que ya existian.
- INV-099 (UC-07 5): la actualizacion de precios pasa el detector de `ml.anomaly` al nucleo y los
  atipicos salen en el informe de la version revalorada.

Oraculos independientes de la ruta: la linea base (`tests/fixtures/apu_linea_base.py`), las siete
inconsistencias documentadas (`INCONSISTENCIAS`), el presupuesto del tercero de la fixture tabular,
las variaciones de la muestra de UC-02 calculadas a mano (`tests/fixtures/actualizacion_uc02.py`)
y la base leida directamente con SQLAlchemy. La base vive en memoria (fixture `cliente`) y los
archivos en `tmp_path`: nada escribe en `data/`.
"""

from __future__ import annotations

import io
from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import func, select

from core import models
from tests.fixtures import apu_linea_base as linea_base
from tests.fixtures.actualizacion_uc02 import CEMENTO
from tests.fixtures.presupuesto_tabular import (
    PREFIJO_LINEA_BASE,
    codigo_tercero,
    escribir_presupuesto_tabular,
    presupuesto_tercero,
)

TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CENTIMO = Decimal("0.01")


def _dos_decimales(valor: Decimal) -> str:
    return str(valor.quantize(CENTIMO, ROUND_HALF_UP))


def _contar(sesion, modelo) -> int:
    return sesion.scalar(select(func.count()).select_from(modelo))


def _filas_de_hallazgos(informe: str) -> list[tuple[str, str, str | None]]:
    """(regla, severidad, partida) de cada fila de la tabla «## Hallazgos» del markdown.

    Lee el texto que devuelve `GET /presupuestos/{codigo}/informe`, no el objeto del nucleo: es lo
    que ve el cliente de la API. Una partida «-» es un hallazgo sin partida.
    """
    _, _, tabla = informe.partition("## Hallazgos")
    filas = []
    for linea in tabla.splitlines():
        celdas = [celda.strip() for celda in linea.strip().strip("|").split("|")]
        if len(celdas) >= 4 and celdas[0].isdigit():
            filas.append((celdas[1], celdas[2], None if celdas[3] == "-" else celdas[3]))
    return filas


# ---------------------------------------------------------------------------------------------
# INV-093 (UC-06 4a): la justificacion de la advertencia de rendimiento
# ---------------------------------------------------------------------------------------------

EXCAVACION = linea_base.APU_EXCAVACION.codigo_partida
RUTA_RENDIMIENTOS = f"/partidas/{EXCAVACION}/rendimientos"
JUSTIFICACION = "terreno rocoso en la zanja norte: se excava con martillo"


def _registrar(cliente, valor: str, fecha: str = "2026-07-01", **extra):
    cuerpo = {"valor": valor, "tipo": "estimado", "fecha": fecha, **extra}
    return cliente.post(RUTA_RENDIMIENTOS, json=cuerpo)


def _justificaciones(sesion) -> list[tuple[Decimal, str, str, date]]:
    """(valor del rendimiento, texto, advertencia, fecha) de cada justificacion en la base."""
    consulta = select(
        models.Rendimiento.valor,
        models.JustificacionRendimiento.texto,
        models.JustificacionRendimiento.advertencia,
        models.JustificacionRendimiento.fecha,
    ).join(
        models.JustificacionRendimiento,
        models.JustificacionRendimiento.rendimiento_id == models.Rendimiento.id,
    )
    return [tuple(fila) for fila in sesion.execute(consulta)]


def test_la_justificacion_con_advertencia_queda_ligada_al_rendimiento_registrado(cliente, sesion):
    """UC-06 4a: con advertencia, la justificacion se guarda junto al rendimiento de la peticion."""
    assert _registrar(cliente, "75").status_code == 201  # rango observado: [75, 80]

    r = _registrar(cliente, "300", fecha="2026-07-05", justificacion=JUSTIFICACION)

    assert r.status_code == 201  # advertir nunca es impedir
    cuerpo = r.json()
    assert cuerpo["advertencia"] is not None and "rango" in cuerpo["advertencia"]
    assert cuerpo["justificacion"] == JUSTIFICACION
    assert _justificaciones(sesion) == [
        (Decimal("300"), JUSTIFICACION, cuerpo["advertencia"], date(2026, 7, 5))
    ]


def test_cada_peticion_justifica_su_propio_rendimiento(cliente, sesion):
    """Dos advertencias seguidas: cada justificacion va al rendimiento de su peticion, no a otro."""
    _registrar(cliente, "75")
    primera = _registrar(cliente, "300", fecha="2026-07-05", justificacion="primera")
    segunda = _registrar(cliente, "10", fecha="2026-07-06", justificacion="segunda")

    assert primera.status_code == segunda.status_code == 201
    assert segunda.json()["advertencia"] is not None
    assert [(valor, texto) for valor, texto, _, _ in _justificaciones(sesion)] == [
        (Decimal("300"), "primera"),
        (Decimal("10"), "segunda"),
    ]


def test_la_justificacion_sin_advertencia_se_ignora_y_el_registro_se_hace(cliente, sesion):
    """Sin advertencia no hay nada que justificar: se registra el rendimiento, no la justificacion
    (decision de la pista A: ignorarla, no 422, porque el cliente no sabe de antemano si habra
    advertencia y un 422 impediria un registro que RF-27 no impide)."""
    r = _registrar(cliente, "75", justificacion=JUSTIFICACION)

    assert r.status_code == 201
    assert r.json()["advertencia"] is None
    assert r.json()["justificacion"] is None
    assert _justificaciones(sesion) == []
    historico = cliente.get(RUTA_RENDIMIENTOS).json()["historico"]
    assert [h["valor"] for h in historico][-1] == "75"


@pytest.mark.parametrize("texto", ["", "   "])
def test_una_justificacion_en_blanco_es_422_sin_registrar_nada(cliente, sesion, texto):
    _registrar(cliente, "75")
    antes = _contar(sesion, models.Rendimiento)

    r = _registrar(cliente, "300", fecha="2026-07-05", justificacion=texto)

    assert r.status_code == 422
    assert "justificacion" in r.json()["detalle"]
    assert _contar(sesion, models.Rendimiento) == antes
    assert _justificaciones(sesion) == []


def test_sin_justificacion_la_respuesta_la_trae_nula(cliente, sesion):
    _registrar(cliente, "75")
    r = _registrar(cliente, "300", fecha="2026-07-05")

    assert r.status_code == 201
    assert r.json()["advertencia"] is not None
    assert r.json()["justificacion"] is None
    assert _justificaciones(sesion) == []


# ---------------------------------------------------------------------------------------------
# INV-090 (UC-02 2a y 2b): dar de alta los desconocidos elegidos
# ---------------------------------------------------------------------------------------------

ADITIVO = "Aditivo plastificante"
MALLA = "Malla electrosoldada"
#: «Arena lavada» esta en el catalogo en m3: en sacos es el flujo 2b (unidad no equivalente).
ARENA_EN_SACOS = "Arena lavada"


def _cargar_lista(cliente, tmp_path, filas: str, dar_de_alta=None, nombre="Lista C6"):
    archivo = tmp_path / "lista_c6.csv"
    archivo.write_text("tipo,insumo,unidad,precio\n" + filas, encoding="utf-8")
    datos: dict[str, object] = {
        "nombre": nombre,
        "moneda": linea_base.MONEDA,
        "fecha_vigencia": "2026-06-01",
    }
    if dar_de_alta is not None:
        datos["dar_de_alta"] = dar_de_alta
    with archivo.open("rb") as f:
        return cliente.post(
            "/listas-precios", files={"archivo": (archivo.name, f, "text/csv")}, data=datos
        )


FILAS_CON_DESCONOCIDOS = (
    f"material,{CEMENTO},saco,18\n"
    f"material,{ADITIVO},lt,12\n"
    f"material,{MALLA},m2,9\n"
    f"material,{ARENA_EN_SACOS},saco,4\n"
)


def _insumo(sesion, descripcion: str) -> models.Insumo | None:
    return sesion.scalars(
        select(models.Insumo).where(models.Insumo.descripcion == descripcion)
    ).one_or_none()


def test_dar_de_alta_crea_solo_el_elegido_y_separa_ignorados_de_rechazados(
    cliente, sesion, tmp_path
):
    insumos_antes = _contar(sesion, models.Insumo)

    r = _cargar_lista(cliente, tmp_path, FILAS_CON_DESCONOCIDOS, dar_de_alta=[ADITIVO])

    assert r.status_code == 201
    cuerpo = r.json()
    assert cuerpo["dados_de_alta"] == [ADITIVO]
    assert cuerpo["ignorados"] == [MALLA]  # 2a: desconocido que no se eligio
    assert cuerpo["rechazados"] == [ARENA_EN_SACOS]  # 2b: unidad no equivalente
    assert sorted(cuerpo["desconocidos"]) == sorted([MALLA, ARENA_EN_SACOS])

    assert _contar(sesion, models.Insumo) == insumos_antes + 1
    aditivo = _insumo(sesion, ADITIVO)
    assert aditivo is not None and aditivo.codigo.startswith("MAT-")
    assert _insumo(sesion, MALLA) is None
    precios = {
        (precio.lista.nombre, precio.insumo.descripcion): precio.precio
        for precio in sesion.scalars(select(models.PrecioInsumo))
    }
    assert precios[("Lista C6", ADITIVO)] == Decimal("12")
    assert not any(insumo == ADITIVO for lista, insumo in precios if lista != "Lista C6")


def test_sin_dar_de_alta_nada_se_crea_y_los_desconocidos_se_reportan_aparte(
    cliente, sesion, tmp_path
):
    insumos_antes = _contar(sesion, models.Insumo)

    r = _cargar_lista(cliente, tmp_path, FILAS_CON_DESCONOCIDOS)

    assert r.status_code == 201
    cuerpo = r.json()
    assert cuerpo["dados_de_alta"] == []
    assert cuerpo["ignorados"] == [ADITIVO, MALLA]
    assert cuerpo["rechazados"] == [ARENA_EN_SACOS]
    assert _contar(sesion, models.Insumo) == insumos_antes


@pytest.mark.parametrize(
    ("filas", "elegidos", "fragmento"),
    [
        # Un insumo del catalogo no es desconocido.
        (FILAS_CON_DESCONOCIDOS, [CEMENTO], CEMENTO),
        # Un rechazado por unidad (2b) no se ofrece para el alta.
        (FILAS_CON_DESCONOCIDOS, [ARENA_EN_SACOS], ARENA_EN_SACOS),
        # Una descripcion que ni esta en el archivo.
        (FILAS_CON_DESCONOCIDOS, ["Insumo fantasma"], "Insumo fantasma"),
        # El mismo desconocido en dos filas con precios distintos (nucleo, ronda 1 de N).
        (f"material,{ADITIVO},lt,12\nmaterial,{ADITIVO},lt,13\n", [ADITIVO], "filas"),
    ],
    ids=["catalogado", "rechazado-2b", "ausente", "precios-distintos"],
)
def test_un_alta_invalida_es_422_sin_escribir_nada(
    cliente, sesion, tmp_path, filas, elegidos, fragmento
):
    antes = (_contar(sesion, models.Insumo), _contar(sesion, models.ListaPrecios))

    r = _cargar_lista(cliente, tmp_path, filas, dar_de_alta=elegidos)

    assert r.status_code == 422
    assert fragmento in r.json()["detalle"]
    assert (_contar(sesion, models.Insumo), _contar(sesion, models.ListaPrecios)) == antes


# ---------------------------------------------------------------------------------------------
# INV-086 (UC-01 1): crear el proyecto
# ---------------------------------------------------------------------------------------------

PROYECTO_NUEVO = "Ampliacion del drenaje"


def test_crear_proyecto_y_elaborar_sobre_el(cliente, sesion, items_payload):
    cuerpo_presupuesto = {
        "codigo": "C6-001",
        "fecha": "2026-04-28",
        "moneda": linea_base.MONEDA,
        "proyecto": PROYECTO_NUEVO,
        "items": items_payload,
    }
    assert cliente.post("/presupuestos", json=cuerpo_presupuesto).status_code == 404

    r = cliente.post(
        "/proyectos",
        json={"nombre": PROYECTO_NUEVO, "descripcion": "segunda etapa", "dominio": "civil"},
    )

    assert r.status_code == 201
    assert r.json() == {
        "nombre": PROYECTO_NUEVO,
        "descripcion": "segunda etapa",
        "dominio": "civil",
    }
    creado = cliente.post("/presupuestos", json=cuerpo_presupuesto)
    assert creado.status_code == 201
    assert creado.json()["total"] == str(linea_base.TOTAL_PRESUPUESTO_AUDITADO)
    modelo = sesion.scalars(
        select(models.Presupuesto).where(models.Presupuesto.codigo == "C6-001")
    ).one()
    assert modelo.proyecto.nombre == PROYECTO_NUEVO


def test_crear_proyecto_sin_dominio_ni_descripcion(cliente):
    r = cliente.post("/proyectos", json={"nombre": "Solo el nombre"})
    assert r.status_code == 201
    assert r.json() == {"nombre": "Solo el nombre", "descripcion": "", "dominio": ""}


@pytest.mark.parametrize(
    ("cuerpo", "fragmento"),
    [
        ({"nombre": PROYECTO_NUEVO}, "ya existe"),
        ({"nombre": "   "}, "vacio"),
        ({"nombre": "Otro", "dominio": "naval"}, "naval"),
    ],
    ids=["repetido", "vacio", "dominio-desconocido"],
)
def test_un_proyecto_invalido_es_422_sin_crear_nada(cliente, sesion, cuerpo, fragmento):
    assert cliente.post("/proyectos", json={"nombre": PROYECTO_NUEVO}).status_code == 201
    antes = _contar(sesion, models.Proyecto)

    r = cliente.post("/proyectos", json=cuerpo)

    assert r.status_code == 422
    assert fragmento in r.json()["detalle"]
    assert _contar(sesion, models.Proyecto) == antes


# ---------------------------------------------------------------------------------------------
# INV-076 (RF-24, UC-05 1 y 2): importar el presupuesto tabular de un tercero
# ---------------------------------------------------------------------------------------------

PROYECTO_TERCERO = "Auditoria del drenaje de un tercero"
#: Las siete inconsistencias documentadas de la linea base, con los codigos del tercero.
SIETE = {
    (i.regla, codigo_tercero(i.codigo_partida) if i.codigo_partida else None)
    for i in linea_base.INCONSISTENCIAS
}


def _importar(cliente, ruta: Path, proyecto: str = PROYECTO_TERCERO, tipo: str = TIPO_XLSX):
    with ruta.open("rb") as f:
        return cliente.post(
            "/presupuestos/importar",
            files={"archivo": (ruta.name, f, tipo)},
            data={"proyecto": proyecto},
        )


@pytest.fixture
def proyecto_tercero(cliente):
    r = cliente.post("/proyectos", json={"nombre": PROYECTO_TERCERO, "dominio": "civil"})
    assert r.status_code == 201
    return PROYECTO_TERCERO


def test_importar_audita_siete_de_siete_y_persiste_informe_y_libro(
    cliente, sesion, tmp_path, proyecto_tercero
):
    esperado = presupuesto_tercero()

    r = _importar(cliente, escribir_presupuesto_tabular(tmp_path))

    assert r.status_code == 201
    cuerpo = r.json()
    assert cuerpo["codigo"] == esperado.codigo
    assert cuerpo["total"] == _dos_decimales(esperado.total)
    assert cuerpo["hallazgos"] >= len(SIETE)
    assert cuerpo["pendientes"] == []

    consulta = {"proyecto": proyecto_tercero}
    informe = cliente.get(f"/presupuestos/{esperado.codigo}/informe", params=consulta)
    assert informe.status_code == 200
    detectadas = {
        (regla, partida)
        for regla, severidad, partida in _filas_de_hallazgos(informe.text)
        if severidad in ("ERROR", "CRITICO")
    }
    assert len(SIETE) == 7 and SIETE <= detectadas
    assert "mts" in informe.text  # RF-24: R3 cita la unidad del autor

    libro = cliente.get(f"/presupuestos/{esperado.codigo}/excel", params=consulta)
    assert libro.status_code == 200
    hojas = openpyxl.load_workbook(io.BytesIO(libro.content)).sheetnames
    assert "Auditoria" in hojas

    lista = sesion.scalars(
        select(models.ListaPrecios).where(models.ListaPrecios.origen == "presupuesto_tercero.xlsx")
    ).one()
    assert lista.nombre == "Importada: presupuesto_tercero.xlsx"


def test_importar_se_declara_antes_de_las_rutas_con_codigo():
    """`importar` no debe poder leerse como un `{codigo}`. Hoy ninguna ruta POST lleva solo
    `/presupuestos/{codigo}`, asi que por HTTP el orden no se nota (Starlette sigue buscando tras
    una coincidencia parcial de otro metodo); la prueba fija la convencion del plan de C6 (§2.4)
    para que una ruta futura con ese patron no la tape."""
    from api.rutas.presupuestos import router

    rutas = [ruta.path for ruta in router.routes]
    importar = rutas.index("/presupuestos/importar")
    con_codigo = [i for i, ruta in enumerate(rutas) if ruta.startswith("/presupuestos/{codigo}")]
    assert con_codigo and importar < min(con_codigo)


def test_importar_sobre_un_proyecto_inexistente_es_404(cliente, tmp_path):
    r = _importar(cliente, escribir_presupuesto_tabular(tmp_path), proyecto="No existe")
    assert r.status_code == 404


def _con_moneda(ruta: Path, moneda: str) -> Path:
    libro = openpyxl.load_workbook(ruta)
    for fila in libro["Datos"].iter_rows(min_row=2):
        if fila[0].value == "moneda":
            fila[1].value = moneda
    libro.save(ruta)
    return ruta


def _sin_hoja(ruta: Path, hoja: str) -> Path:
    libro = openpyxl.load_workbook(ruta)
    libro.remove(libro[hoja])
    libro.save(ruta)
    return ruta


@pytest.mark.parametrize(
    ("preparar", "fragmento"),
    [
        # Codigos de partida que ya estan en el catalogo sembrado.
        (lambda d: escribir_presupuesto_tabular(d, prefijo=PREFIJO_LINEA_BASE), "catalogo"),
        # O-2: la moneda se compara exacta; «usd» no es la «USD» de la lista vigente.
        (lambda d: _con_moneda(escribir_presupuesto_tabular(d), "usd"), "usd"),
        # Error de forma: falta la hoja de los APU.
        (lambda d: _sin_hoja(escribir_presupuesto_tabular(d), "APU"), "APU"),
    ],
    ids=["codigo-existente", "moneda-distinta", "hoja-ausente"],
)
def test_un_archivo_que_no_se_puede_importar_es_422_sin_escribir_nada(
    cliente, sesion, tmp_path, proyecto_tercero, preparar, fragmento
):
    antes = (_contar(sesion, models.Presupuesto), _contar(sesion, models.ListaPrecios))

    r = _importar(cliente, preparar(tmp_path))

    assert r.status_code == 422
    assert fragmento in r.json()["detalle"]
    assert (_contar(sesion, models.Presupuesto), _contar(sesion, models.ListaPrecios)) == antes


def test_si_la_importacion_falla_despues_de_escribir_la_ruta_no_deja_nada(
    cliente, sesion, tmp_path, proyecto_tercero, monkeypatch
):
    """La ruta es una sola operacion: si el nucleo falla con la lista importada y los APU ya
    escritos (aqui, al guardar el presupuesto), deshace todo y responde 422."""
    import core.budget.importacion as importacion

    def _falla_al_guardar(*_args, **_kwargs):
        raise ValueError("fallo simulado al guardar el presupuesto")

    monkeypatch.setattr(importacion, "guardar_presupuesto", _falla_al_guardar)
    modelos = (models.Partida, models.ListaPrecios, models.Presupuesto)
    antes = [_contar(sesion, modelo) for modelo in modelos]

    r = _importar(cliente, escribir_presupuesto_tabular(tmp_path))

    assert r.status_code == 422
    assert "fallo simulado" in r.json()["detalle"]
    assert [_contar(sesion, modelo) for modelo in modelos] == antes


def test_importar_dos_veces_el_mismo_codigo_es_422(cliente, sesion, tmp_path, proyecto_tercero):
    assert _importar(cliente, escribir_presupuesto_tabular(tmp_path)).status_code == 201
    otra = tmp_path / "otra"
    otra.mkdir()
    # Otros codigos de partida (no chocan con los TER-* ya importados), mismo codigo de presupuesto.
    r = _importar(cliente, escribir_presupuesto_tabular(otra, prefijo="OTR-"))
    assert r.status_code == 422
    assert presupuesto_tercero().codigo in r.json()["detalle"]


def test_importar_un_archivo_que_no_es_xlsx_es_422(cliente, tmp_path, proyecto_tercero):
    ruta = tmp_path / "presupuesto.csv"
    ruta.write_text("codigo_partida,cantidad\nTER-01,1\n", encoding="utf-8")
    r = _importar(cliente, ruta, tipo="text/csv")
    assert r.status_code == 422
