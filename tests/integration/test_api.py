"""API de F.1: el nucleo expuesto por HTTP. Montos como texto (Decimal), nunca float."""

from __future__ import annotations

from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from types import SimpleNamespace

import pytest
from sqlalchemy import select

from tests.fixtures import apu_linea_base as linea_base
from tests.fixtures.actualizacion_uc02 import INCREMENTO_TOTAL
from tests.fixtures.computo_civil import CODIGOS as CODIGOS_CIVIL

RAIZ = Path(__file__).resolve().parents[2]
MUESTRA_PRECIOS = RAIZ / "data" / "samples" / "precios" / "lista_2026-06-01.csv"


def test_salud(cliente):
    r = cliente.get("/salud")
    assert r.status_code == 200 and r.json()["estado"] == "ok"


def test_partidas_lista_las_cinco_de_la_linea_base(cliente):
    r = cliente.get("/partidas")
    assert r.status_code == 200
    codigos = {p["codigo"] for p in r.json()}
    assert set(linea_base.CODIGOS_LINEA_BASE) <= codigos


def test_insumos_traen_precio_vigente_como_texto(cliente):
    r = cliente.get("/insumos")
    cemento = next(i for i in r.json() if i["descripcion"] == "Cemento Portland")
    assert cemento["precio_vigente"] == "15"  # str exacto, jamas 15.0


def test_cargar_lista_de_precios_reporta_desconocidos_vacios(cliente):
    with MUESTRA_PRECIOS.open("rb") as f:
        r = cliente.post(
            "/listas-precios",
            files={"archivo": (MUESTRA_PRECIOS.name, f, "text/csv")},
            data={"nombre": "Lista API", "moneda": "USD", "fecha_vigencia": "2026-06-01"},
        )
    assert r.status_code == 201
    assert r.json()["desconocidos"] == []
    assert cliente.get("/listas-precios").status_code == 200


def test_computo_civil_desde_la_muestra(cliente):
    ruta = RAIZ / "data" / "samples" / "civil" / "tanquillas_y_zanja.csv"
    import json

    with ruta.open("rb") as f:
        r = cliente.post(
            "/computos/civil",
            files={"archivo": (ruta.name, f, "text/csv")},
            data={"codigos": json.dumps(CODIGOS_CIVIL)},
        )
    assert r.status_code == 200 and len(r.json()) == 6
    assert all(i["origen_id"] for i in r.json())


def test_elaborar_presupuesto_por_http_reproduce_el_total_auditado(cliente, items_payload):
    r = cliente.post(
        "/presupuestos",
        json={
            "codigo": "API-001",
            "fecha": "2026-04-28",
            "moneda": "USD",
            "proyecto": "Drenaje de la clínica",
            "items": items_payload,
        },
    )
    assert r.status_code == 201
    assert r.json()["total"] == str(linea_base.TOTAL_PRESUPUESTO_AUDITADO)  # texto exacto
    assert r.json()["hallazgos"] >= 7  # audita siempre
    detalle = cliente.get("/presupuestos/API-001").json()
    assert detalle["total"] == str(linea_base.TOTAL_PRESUPUESTO_AUDITADO)
    assert len(detalle["partidas"]) == len(linea_base.PRESUPUESTO_AUDITADO)
    assert "Informe" in cliente.get("/presupuestos/API-001/informe").text
    assert cliente.get("/presupuestos/API-001/excel").content[:2] == b"PK"


def test_actualizacion_por_http_solo_revalora_el_concreto(cliente, items_payload):
    cliente.post(
        "/presupuestos",
        json={
            "codigo": "API-001",
            "fecha": "2026-04-28",
            "moneda": "USD",
            "proyecto": "Drenaje de la clínica",
            "items": items_payload,
        },
    )
    with MUESTRA_PRECIOS.open("rb") as f:
        cliente.post(
            "/listas-precios",
            files={"archivo": (MUESTRA_PRECIOS.name, f, "text/csv")},
            data={"nombre": "L2", "moneda": "USD", "fecha_vigencia": "2026-06-01"},
        )
    r = cliente.post(
        "/presupuestos/API-001/actualizacion", json={"lista": "L2", "codigo_nuevo": "API-002"}
    )
    assert r.status_code == 200
    cuerpo = r.json()
    total_nuevo = linea_base.TOTAL_EXACTO_LINEA_BASE + INCREMENTO_TOTAL
    assert cuerpo["total_nuevo"] == str(total_nuevo.quantize(Decimal("0.01"), ROUND_HALF_UP))
    filas = {f["codigo_partida"]: f for f in cuerpo["filas"]}
    assert filas[linea_base.APU_CONCRETO.codigo_partida]["variacion_pct"] != "0"
    assert sum(1 for f in cuerpo["filas"] if f["variacion_pct"] not in ("0", "0.00")) == 1


def test_cambios_precio_por_http_con_filtro_de_insumo(cliente, items_payload):
    """GET /cambios-precio (sin prueba propia hasta esta sesion): tras una actualizacion, el
    historial responde con montos como texto exacto y respeta el filtro `insumo`. La ruta
    delega en `core.catalog.cambios_precio`, la unica consulta del historial.
    """
    cliente.post(
        "/presupuestos",
        json={
            "codigo": "API-001",
            "fecha": "2026-04-28",
            "moneda": "USD",
            "proyecto": "Drenaje de la clínica",
            "items": items_payload,
        },
    )
    with MUESTRA_PRECIOS.open("rb") as f:
        cliente.post(
            "/listas-precios",
            files={"archivo": (MUESTRA_PRECIOS.name, f, "text/csv")},
            data={"nombre": "L2", "moneda": "USD", "fecha_vigencia": "2026-06-01"},
        )
    cliente.post(
        "/presupuestos/API-001/actualizacion", json={"lista": "L2", "codigo_nuevo": "API-002"}
    )

    r = cliente.get("/cambios-precio", params={"insumo": "Cemento Portland"})
    assert r.status_code == 200
    cambios = r.json()
    assert len(cambios) == 1
    assert cambios[0]["precio_anterior"] == "15" and cambios[0]["precio_nuevo"] == "18"

    todos = cliente.get("/cambios-precio").json()
    assert {c["insumo"] for c in todos} == {"Cemento Portland", "Arena lavada"}


def test_cantidad_decimal_malformada_responde_422_no_500(cliente, items_payload):
    """`Decimal("12,50")` lanza `InvalidOperation` (que NO hereda de `ValueError`): sin manejo en
    la frontera seria un 500 crudo (hallazgo menor diferido de la bitacora F.1). La API debe
    responder 422 con el campo y el texto rechazado en el detalle.
    """
    items = items_payload
    items[0]["cantidad"] = "12,50"  # coma decimal: malformado para Decimal
    r = cliente.post(
        "/presupuestos",
        json={
            "codigo": "API-MAL",
            "fecha": "2026-04-28",
            "moneda": "USD",
            "proyecto": "Drenaje de la clínica",
            "items": items,
        },
    )
    assert r.status_code == 422
    assert "12,50" in r.json()["detalle"]


def test_parametro_decimal_no_finito_responde_422(cliente, items_payload):
    """`Decimal("NaN")` SI parsea, pero no es un monto: debe rechazarse en la frontera, no
    propagarse hasta que una comparacion del nucleo lance `InvalidOperation` (500 crudo).
    """
    r = cliente.post(
        "/presupuestos",
        json={
            "codigo": "API-NAN",
            "fecha": "2026-04-28",
            "moneda": "USD",
            "proyecto": "Drenaje de la clínica",
            "fcas": "NaN",
            "items": items_payload,
        },
    )
    assert r.status_code == 422
    assert "fcas" in r.json()["detalle"]


def test_codigo_de_presupuesto_repetido_en_dos_proyectos_no_revienta_en_500(
    cliente, sesion, items_payload
):
    """`codigo` solo es unico por proyecto (`UniqueConstraint(proyecto_id, codigo)`): si dos
    proyectos comparten codigo, las rutas que resuelven `{codigo}` deben responder 409 con
    detalle claro (nunca `MultipleResultsFound` sin manejar) y desambiguar con `?proyecto=`
    (hallazgo Important de la revision de la Tarea 2).
    """
    from core import models

    payload_base = {
        "codigo": "COL-001",
        "fecha": "2026-04-28",
        "moneda": "USD",
        "items": items_payload,
    }
    r1 = cliente.post("/presupuestos", json={**payload_base, "proyecto": "Drenaje de la clínica"})
    assert r1.status_code == 201

    sesion.add(models.Proyecto(nombre="Segundo proyecto", descripcion="", dominio="civil"))
    sesion.commit()
    r2 = cliente.post("/presupuestos", json={**payload_base, "proyecto": "Segundo proyecto"})
    assert r2.status_code == 201

    # Sin 'proyecto': el codigo es ambiguo entre los dos proyectos. 409 claro, nunca un 500 crudo.
    for r in (
        cliente.get("/presupuestos/COL-001"),
        cliente.get("/presupuestos/COL-001/informe"),
        cliente.get("/presupuestos/COL-001/excel"),
        cliente.post(
            "/presupuestos/COL-001/actualizacion",
            json={"lista": "no importa", "codigo_nuevo": "COL-002"},
        ),
    ):
        assert r.status_code == 409
        assert "COL-001" in r.json()["detalle"]

    # Con 'proyecto': cada ruta resuelve sin ambiguedad, sobre el proyecto indicado.
    detalle = cliente.get("/presupuestos/COL-001", params={"proyecto": "Segundo proyecto"})
    assert detalle.status_code == 200
    assert detalle.json()["total"] == str(linea_base.TOTAL_PRESUPUESTO_AUDITADO)
    assert (
        cliente.get(
            "/presupuestos/COL-001/informe", params={"proyecto": "Drenaje de la clínica"}
        ).status_code
        == 200
    )


def test_actualizacion_sin_cambios_de_precio_responde_409_sin_persistir(
    cliente, sesion, items_payload
):
    """UC-02, flujo 3a (rama sin prueba dedicada hasta esta sesion, bitacora F.1): si ningun
    insumo cambio de precio entre la lista del presupuesto y la nueva, no hay nada que
    confirmar: 409 y el `codigo_nuevo` NO queda guardado (la ruta no hace `commit`).
    """
    cliente.post(
        "/presupuestos",
        json={
            "codigo": "API-001",
            "fecha": "2026-04-28",
            "moneda": "USD",
            "proyecto": "Drenaje de la clínica",
            "items": items_payload,
        },
    )
    with MUESTRA_PRECIOS.open("rb") as f:
        cliente.post(
            "/listas-precios",
            files={"archivo": (MUESTRA_PRECIOS.name, f, "text/csv")},
            data={"nombre": "L2", "moneda": "USD", "fecha_vigencia": "2026-06-01"},
        )
    r1 = cliente.post(
        "/presupuestos/API-001/actualizacion", json={"lista": "L2", "codigo_nuevo": "API-002"}
    )
    assert r1.status_code == 200  # L2 si cambia precios respecto de la lista sembrada

    # API-002 quedo valorado con L2: revalorarlo con la misma L2 no afecta ningun insumo.
    r2 = cliente.post(
        "/presupuestos/API-002/actualizacion", json={"lista": "L2", "codigo_nuevo": "API-003"}
    )
    assert r2.status_code == 409
    assert "nada que confirmar" in r2.json()["detalle"]

    # En produccion la sesion por peticion se cierra sin commit (rollback implicito); aqui la
    # sesion compartida del fixture obliga a reproducir ese cierre a mano antes de comprobar.
    sesion.rollback()
    assert cliente.get("/presupuestos/API-003").status_code == 404
    codigos = {p["codigo"] for p in cliente.get("/presupuestos").json()}
    assert "API-003" not in codigos and {"API-001", "API-002"} <= codigos


def test_computo_ifc_sin_extra_civil_responde_400(cliente, monkeypatch):
    """La rama `ImportError -> 400` de `/computos/civil` no tenia evidencia RED real (bitacora
    F.1: no se podia desinstalar el extra del entorno compartido). Bloquear el modulo en
    `sys.modules` reproduce el entorno sin `ifcopenshell` sin tocar el venv.
    """
    import sys

    monkeypatch.setitem(sys.modules, "adapters.civil.ifc", None)  # import -> ImportError
    ruta = RAIZ / "data" / "samples" / "tanquilla.ifc"
    with ruta.open("rb") as f:
        r = cliente.post(
            "/computos/civil",
            files={"archivo": (ruta.name, f, "application/octet-stream")},
        )
    assert r.status_code == 400
    assert "extra civil" in r.json()["detalle"]


def test_computo_civil_desde_un_modelo_ifc(cliente):
    """Compuerta G1 por HTTP: `.ifc` como entrada de `/computos/civil`, decidido por la extension
    del archivo (docs/superpowers/specs/2026-08-31-f1-api-simulador-design.md, seccion 7). No hace
    falta el mapeo `codigos` que exige la entrada tabular: `Pset_APU` ya trae el codigo de partida.
    """
    ruta = RAIZ / "data" / "samples" / "tanquilla.ifc"
    with ruta.open("rb") as f:
        r = cliente.post(
            "/computos/civil",
            files={"archivo": (ruta.name, f, "application/octet-stream")},
        )
    assert r.status_code == 200
    items = r.json()
    assert len(items) == 1
    concreto = linea_base.APU_CONCRETO.codigo_partida
    assert items[0]["cantidad"] == str(linea_base.CANTIDAD_POR_TANQUILLA_MEMORIA[concreto])
    assert items[0]["codigo_partida"] == concreto
    assert items[0]["origen_tipo"] == "ifc"


CODIGO_SIN_PARTIDA = "XX-99-NUE"


def _item_sin_partida(descripcion):
    """Un item cuya partida no existe en el catalogo (UC-01, flujo 3a)."""
    return {
        "codigo_partida": CODIGO_SIN_PARTIDA,
        "descripcion": descripcion,
        "unidad": "m3",
        "cantidad": "2.5",
        "origen_id": "PEND-01",
        "origen_tipo": "manual",
        "dominio": "civil",
    }


#: El submodulo del criterio unico de UC-03 (INV-081): los dobles se parchan ahi, porque
#: `proponedor` resuelve `NormalizadorPartidas` y `_cargar_modelo` en su propio modulo.
NORMALIZADOR = "ml.normalization.normalizador"


def _sin_extra(nombre):
    """Como falla de verdad el extra `ml` ausente: al importar sentence-transformers."""
    raise ImportError("simulado: sin sentence-transformers")


def _cuerpo(codigo, items):
    return {
        "codigo": codigo,
        "fecha": "2026-04-28",
        "moneda": "USD",
        "proyecto": "Drenaje de la clínica",
        "items": items,
    }


def test_un_item_sin_partida_queda_pendiente_por_http(cliente, sesion, items_payload, monkeypatch):
    """UC-01, flujo 3a, por HTTP (INV-072): el item sin partida no aborta la elaboracion. Queda
    pendiente, no suma al total, el informe lo declara (y se persiste) y la respuesta lo lista con
    las partidas similares que propone UC-03. El normalizador es un doble: esta prueba verifica el
    cableado, no la similitud (esa la cubre la prueba con el modelo real, mas abajo).
    """
    from core import models

    recibidas = {}
    propuesta = SimpleNamespace(codigo="LB-01-EXC", descripcion="Excavacion", puntaje=0.91)

    class _NormalizadorDoble:
        def __init__(self, partidas):
            recibidas.update(partidas)

        def similares(self, descripcion):
            return [propuesta]

    monkeypatch.setattr(f"{NORMALIZADOR}.NormalizadorPartidas", _NormalizadorDoble)
    from core.catalog import Catalogo

    consultas = []
    partidas = Catalogo.partidas

    def _partidas_contadas(self, *args, **kwargs):
        consultas.append(args)
        return partidas(self, *args, **kwargs)

    base = cliente.post("/presupuestos", json=_cuerpo("API-BASE", items_payload))
    monkeypatch.setattr(Catalogo, "partidas", _partidas_contadas)
    r = cliente.post(
        "/presupuestos",
        json=_cuerpo("API-3A", [*items_payload, _item_sin_partida("Excavacion de un foso nuevo")]),
    )
    monkeypatch.undo()

    assert r.status_code == 201
    cuerpo = r.json()
    assert cuerpo["total"] == str(linea_base.TOTAL_PRESUPUESTO_AUDITADO)  # no suma al total
    assert cuerpo["hallazgos"] == base.json()["hallazgos"] + 1  # el informe lo declara
    assert base.json()["pendientes"] == []
    (pendiente,) = cuerpo["pendientes"]
    assert pendiente["codigo_partida"] == CODIGO_SIN_PARTIDA
    assert pendiente["origen_id"] == "PEND-01" and pendiente["cantidad"] == "2.5"
    assert pendiente["propuestas"] == [
        {"codigo": "LB-01-EXC", "descripcion": "Excavacion", "puntaje": 0.91}
    ]
    assert set(linea_base.CODIGOS_LINEA_BASE) <= set(recibidas)  # propone sobre el catalogo
    assert len(consultas) == 1  # la propuesta reusa las partidas ya consultadas (GC2, A3)
    # Postcondicion de UC-01: el informe persistido lleva la declaracion del pendiente.
    guardado = sesion.scalars(
        select(models.Presupuesto).where(models.Presupuesto.codigo == "API-3A")
    ).one()
    assert [h.regla for h in guardado.hallazgos if h.regla == "PENDIENTE"] == ["PENDIENTE"]


def test_sin_el_extra_ml_el_pendiente_queda_sin_propuestas(cliente, items_payload, monkeypatch):
    """La propuesta es una ayuda, no una condicion: sin `ml` el item sigue pendiente."""
    monkeypatch.setattr(f"{NORMALIZADOR}._cargar_modelo", _sin_extra)

    r = cliente.post(
        "/presupuestos",
        json=_cuerpo("API-3A-SIN-ML", [*items_payload, _item_sin_partida("Foso nuevo")]),
    )

    assert r.status_code == 201
    assert r.json()["pendientes"][0]["propuestas"] == []


@pytest.mark.parametrize(
    "error",
    [
        OSError("sin modelo en la cache local"),
        ValueError("identificador de modelo invalido"),
        RuntimeError("archivo de pesos corrupto"),
    ],
    ids=["sin-modelo", "identificador-invalido", "pesos-corruptos"],
)
def test_sin_el_modelo_el_pendiente_queda_sin_propuestas(
    cliente, items_payload, monkeypatch, caplog, error
):
    """INV-072: el extra `ml` esta, pero el modelo no se puede cargar. Sin cache y sin red,
    sentence-transformers lanza `OSError`; un identificador que Hugging Face rechaza lanza
    `HFValidationError`, subclase de `ValueError`, que `api/main.py` convertiria en 422; un
    archivo de pesos corrupto hace que `torch.load` lance `RuntimeError`. La ruta no responde 500
    ni 422: el item sigue pendiente y sin propuestas."""

    class _NormalizadorSinModelo:
        def __init__(self, partidas):
            raise error

    monkeypatch.setattr(f"{NORMALIZADOR}.NormalizadorPartidas", _NormalizadorSinModelo)

    with caplog.at_level("WARNING", logger=NORMALIZADOR):
        r = cliente.post(
            "/presupuestos",
            json=_cuerpo(
                f"API-3A-{type(error).__name__}", [*items_payload, _item_sin_partida("Foso nuevo")]
            ),
        )

    assert r.status_code == 201
    (pendiente,) = r.json()["pendientes"]
    assert pendiente["codigo_partida"] == CODIGO_SIN_PARTIDA
    assert pendiente["propuestas"] == []
    # Observabilidad (GC2): la falla no se traga en silencio; el aviso nombra su tipo.
    (aviso,) = [a for a in caplog.records if a.name == NORMALIZADOR]
    assert aviso.levelname == "WARNING" and type(error).__name__ in aviso.getMessage()


def test_la_propuesta_real_reconoce_la_partida_de_la_misma_descripcion(
    cliente, items_payload, monkeypatch
):
    """Con el modelo de similitud real (RF-15): una descripcion identica a la de una partida del
    catalogo la pone primera. Oraculo independiente del modelo: textos iguales, coseno 1."""
    pytest.importorskip("sentence_transformers")
    monkeypatch.setenv("HF_HUB_DISABLE_SYMLINKS_WARNING", "1")
    from ml.normalization import NormalizadorPartidas

    try:
        NormalizadorPartidas({})
    except OSError as error:  # mismo criterio que tests/unit/test_normalizacion.py
        pytest.skip(f"el modelo de similitud no esta disponible localmente: {error}")

    codigo = linea_base.APU_EXCAVACION.codigo_partida
    excavacion = next(p for p in cliente.get("/partidas").json() if p["codigo"] == codigo)

    r = cliente.post(
        "/presupuestos",
        json=_cuerpo("API-3A-REAL", [*items_payload, _item_sin_partida(excavacion["descripcion"])]),
    )

    assert r.status_code == 201  # antes de leer el cuerpo: un 404 no tiene `pendientes`
    propuestas = r.json()["pendientes"][0]["propuestas"]
    assert 1 <= len(propuestas) <= 3
    assert propuestas[0]["codigo"] == codigo
    puntajes = [p["puntaje"] for p in propuestas]
    assert puntajes == sorted(puntajes, reverse=True)


@pytest.mark.parametrize(
    "error",
    [
        ImportError("dependencia perezosa ausente"),
        OSError("tokenizador ausente"),
        ValueError("entrada rechazada"),
        RuntimeError("fallo del modelo al codificar"),
    ],
    ids=["import", "os", "valor", "ejecucion"],
)
def test_si_similares_falla_el_pendiente_queda_sin_propuestas(
    cliente, items_payload, monkeypatch, caplog, error
):
    """El normalizador se construye, pero `similares` falla al codificar la descripcion: la
    ruta no responde 500 despues de guardar. Ese item queda sin propuestas y sigue pendiente."""

    class _NormalizadorQueFallaAlProponer:
        def __init__(self, partidas):
            pass

        def similares(self, descripcion):
            raise error

    monkeypatch.setattr(f"{NORMALIZADOR}.NormalizadorPartidas", _NormalizadorQueFallaAlProponer)

    with caplog.at_level("WARNING", logger=NORMALIZADOR):
        r = cliente.post(
            "/presupuestos",
            json=_cuerpo(
                f"API-3A-SIM-{type(error).__name__}",
                [*items_payload, _item_sin_partida("Foso nuevo")],
            ),
        )

    assert r.status_code == 201
    (pendiente,) = r.json()["pendientes"]
    assert pendiente["codigo_partida"] == CODIGO_SIN_PARTIDA
    assert pendiente["propuestas"] == []
    (aviso,) = [a for a in caplog.records if a.name == NORMALIZADOR]
    assert aviso.levelname == "WARNING" and type(error).__name__ in aviso.getMessage()


def test_si_todos_los_items_quedan_pendientes_el_presupuesto_sale_vacio_y_no_cumple(
    cliente, sesion, monkeypatch
):
    """UC-01 3a en el extremo: ningun item tiene partida. La elaboracion no aborta (201), el
    total es 0, todos los items se listan como pendientes y el informe emitido no cumple.
    Oraculo del veredicto: el informe guardado lleva un `PENDIENTE` de severidad ERROR (3) por
    item, y un solo hallazgo ERROR basta para NO CUMPLE (CLAUDE.md, principio 7)."""
    from core import models

    monkeypatch.setattr(f"{NORMALIZADOR}._cargar_modelo", _sin_extra)  # sin propuestas aqui
    items = [
        {**_item_sin_partida("Foso nuevo"), "origen_id": "PEND-01"},
        {**_item_sin_partida("Muro nuevo"), "codigo_partida": "XX-98-NUE", "origen_id": "PEND-02"},
    ]

    r = cliente.post("/presupuestos", json=_cuerpo("API-3A-TODOS", items))

    assert r.status_code == 201
    cuerpo = r.json()
    assert cuerpo["total"] == "0.00"
    assert [(p["codigo_partida"], p["origen_id"]) for p in cuerpo["pendientes"]] == [
        (CODIGO_SIN_PARTIDA, "PEND-01"),
        ("XX-98-NUE", "PEND-02"),
    ]
    guardado = sesion.scalars(
        select(models.Presupuesto).where(models.Presupuesto.codigo == "API-3A-TODOS")
    ).one()
    assert guardado.partidas == []
    pendientes = [h for h in guardado.hallazgos if h.regla == "PENDIENTE"]
    assert len(pendientes) == 2
    assert all(h.severidad >= 3 for h in pendientes)  # ERROR: el informe no cumple


def test_el_informe_recargado_declara_el_pendiente(cliente, items_payload, monkeypatch):
    """INV-080 por HTTP: el informe que se vuelve a pedir (`GET .../informe`) es el del documento
    emitido. El presupuesto se emitio con un item sin partida y su informe lo declaraba; recargado,
    tiene que seguir declarandolo (RF-02, «al producir o cargar»)."""
    monkeypatch.setattr(f"{NORMALIZADOR}._cargar_modelo", _sin_extra)  # sin propuestas aqui
    creado = cliente.post(
        "/presupuestos",
        json=_cuerpo("API-REC-INF", [*items_payload, _item_sin_partida("Foso nuevo")]),
    )
    assert creado.status_code == 201

    r = cliente.get("/presupuestos/API-REC-INF/informe")

    assert r.status_code == 200
    assert "PENDIENTE" in r.text
    assert CODIGO_SIN_PARTIDA in r.text and "PEND-01" in r.text


def test_el_excel_recargado_declara_el_pendiente_en_la_auditoria(
    cliente, items_payload, monkeypatch
):
    """INV-080 por HTTP: la hoja de auditoria del libro descargado (`GET .../excel`) declara el
    pendiente del documento emitido. El libro se abre desde los bytes de la respuesta (QA-16)."""
    from io import BytesIO

    from openpyxl import load_workbook

    from core.budget.excel import HOJA_AUDITORIA

    monkeypatch.setattr(f"{NORMALIZADOR}._cargar_modelo", _sin_extra)
    creado = cliente.post(
        "/presupuestos",
        json=_cuerpo("API-REC-XLS", [*items_payload, _item_sin_partida("Foso nuevo")]),
    )
    assert creado.status_code == 201

    r = cliente.get("/presupuestos/API-REC-XLS/excel")

    assert r.status_code == 200
    hoja = load_workbook(BytesIO(r.content))[HOJA_AUDITORIA]
    celdas = [str(celda) for fila in hoja.iter_rows(values_only=True) for celda in fila]
    assert "PENDIENTE" in celdas
    assert any(CODIGO_SIN_PARTIDA in celda and "PEND-01" in celda for celda in celdas)
