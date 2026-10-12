# tests/integration/test_importar_valuaciones.py
"""SL-E4: importación de valuaciones y costos desde XLSX (spec S2 §6.2, RF-73).

El XLSX versionado de OB-S2 da las mismas valuaciones que la entrada a mano; cada malformación de
§7.2 (fila RF-73) da su mensaje, sin *traceback* y sin escribir nada.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import func, select

from core import models
from core.budget.valuaciones import (
    LIMITE_BYTES,
    anular_valuacion,
    leer_valuaciones_excel,
    registrar_lectura,
    valuaciones_registradas,
)
from tests.fixtures import control_ob_s2 as ob

D = Decimal


def _cuenta(sesion, modelo) -> int:
    return sesion.scalar(select(func.count()).select_from(modelo))


def _libro(tmp_path, cambiar):
    """El XLSX de OB-S2 con una malformación aplicada por `cambiar(libro)`."""
    ruta = ob.escribir_xlsx(tmp_path / "malo.xlsx")
    libro = openpyxl.load_workbook(ruta)
    cambiar(libro)
    libro.save(ruta)
    return ruta


def test_el_xlsx_versionado_de_la_memoria_existe_y_se_regenera_igual(tmp_path):
    regenerado = ob.escribir_xlsx(tmp_path / "otro.xlsx")

    def celdas(ruta):
        libro = openpyxl.load_workbook(ruta)
        return {h.title: [[c.value for c in fila] for fila in h.iter_rows()] for h in libro}

    assert celdas(ob.RUTA_XLSX) == celdas(regenerado)


def test_el_xlsx_de_ob_s2_da_las_mismas_valuaciones_que_la_entrada_a_mano(obra):
    lectura = leer_valuaciones_excel(ob.RUTA_XLSX)
    assert lectura.informe == (
        "Valuaciones: 2 valuacion(es)",
        "Lineas: 3 linea(s)",
        "Costos: 1 costo(s)",
    )
    registrar_lectura(obra, ob.NOMBRE, lectura, fuente="XLSX ficticio", responsable="Ing.")
    del_excel = valuaciones_registradas(obra, ob.NOMBRE)
    a_mano = ob.v01_v02()
    assert [(v.numero, v.desde, v.hasta, v.fecha, v.inspector, v.lineas) for v in del_excel] == [
        (v.numero, v.desde, v.hasta, v.fecha, v.inspector, v.lineas) for v in a_mano
    ]
    assert all(v.origen == "excel:ob_s2.xlsx" for v in del_excel)
    costo = obra.scalars(select(models.CostoReal)).one()
    assert (costo.monto, costo.moneda, costo.fecha) == (D("1280.00"), "USD", date(2026, 11, 13))


def test_un_numero_ya_registrado_se_rechaza_entero_y_no_escribe(obra):
    ob.registrar(obra, ob.V01)
    antes = (_cuenta(obra, models.Valuacion), _cuenta(obra, models.CostoReal))
    with pytest.raises(ValueError, match="V-01 ya estan registradas"):
        registrar_lectura(
            obra, ob.NOMBRE, leer_valuaciones_excel(ob.RUTA_XLSX), fuente="x", responsable="I"
        )
    assert (_cuenta(obra, models.Valuacion), _cuenta(obra, models.CostoReal)) == antes


def test_el_archivo_no_corrige_se_anula_y_se_vuelve_a_registrar(obra):
    primera = ob.registrar(obra, ob.V01)
    anular_valuacion(obra, primera.id, motivo="error de medicion", responsable="Ing.")
    registrar_lectura(
        obra, ob.NOMBRE, leer_valuaciones_excel(ob.RUTA_XLSX), fuente="x", responsable="I"
    )
    assert [v.numero for v in valuaciones_registradas(obra, ob.NOMBRE)] == ["V-01", "V-02"]


def _quitar_hoja(nombre):
    return lambda libro: libro.remove(libro[nombre])


def _quitar_columna(hoja, columna):
    def cambiar(libro):
        fila = libro[hoja][1]
        celda = next(c for c in fila if c.value == columna)
        libro[hoja].delete_cols(celda.column)

    return cambiar


def _escribir(hoja, fila, columna, valor):
    def cambiar(libro):
        libro[hoja].cell(row=fila, column=columna, value=valor)

    return cambiar


def _repetir_numero(libro):
    libro["Valuaciones"].append(["V-01", "2026-11-16", "2026-11-20", "", "Ing. Ficticio"])


@pytest.mark.parametrize(
    ("cambiar", "mensaje"),
    [
        (_quitar_hoja("Valuaciones"), "malo.xlsx: falta la hoja 'Valuaciones'"),
        (_quitar_hoja("Lineas"), "malo.xlsx: falta la hoja 'Lineas'"),
        (_quitar_columna("Lineas", "cantidad"), "hoja Lineas: faltan las columnas cantidad"),
        (_quitar_columna("Valuaciones", "hasta"), "hoja Valuaciones: faltan las columnas hasta"),
        (_escribir("Valuaciones", 2, 2, "02/11/2026"), "hoja Valuaciones, fila 2: desde"),
        (_escribir("Lineas", 3, 3, "-5"), "hoja Lineas, fila 3: la cantidad de 'OB-A' no puede"),
        (_escribir("Lineas", 2, 3, "cuarenta"), "hoja Lineas, fila 2: cantidad 'cuarenta'"),
        (_escribir("Lineas", 2, 1, "V-77"), "hoja Lineas, fila 2: la valuacion 'V-77' no esta"),
        (_repetir_numero, "hoja Valuaciones, fila 4: el numero V-01 esta repetido"),
        (_escribir("Valuaciones", 2, 1, ""), "hoja Valuaciones, fila 2: falta el numero"),
        (_escribir("Costos", 2, 4, "EUR"), "hoja Costos, fila 2: moneda 'EUR'"),
        (_escribir("Costos", 2, 6, ""), "hoja Costos, fila 2: falta la fuente"),
        (_escribir("Costos", 2, 3, "-1"), "hoja Costos, fila 2: el monto no puede ser negativo"),
    ],
)
def test_cada_malformacion_da_su_mensaje_sin_traceback_y_sin_escribir(
    obra, tmp_path, cambiar, mensaje
):
    ruta = _libro(tmp_path, cambiar)
    with pytest.raises(ValueError) as error:
        registrar_lectura(
            obra, ob.NOMBRE, leer_valuaciones_excel(ruta), fuente="x", responsable="I"
        )
    assert mensaje in str(error.value)
    assert "Traceback" not in str(error.value)
    assert _cuenta(obra, models.Valuacion) == _cuenta(obra, models.CostoReal) == 0


def test_una_valuacion_sin_lineas_se_rechaza(tmp_path):
    ruta = _libro(tmp_path, lambda libro: libro["Lineas"].delete_rows(3, 2))
    with pytest.raises(ValueError, match="V-02 no tienen lineas"):
        leer_valuaciones_excel(ruta)


def test_un_archivo_que_no_es_xlsx_o_pesa_mas_de_20_mb_se_rechaza(tmp_path):
    csv = tmp_path / "valuaciones.csv"
    csv.write_text("numero\n", encoding="utf-8")
    with pytest.raises(ValueError, match="libro .xlsx"):
        leer_valuaciones_excel(csv)
    grande = tmp_path / "grande.xlsx"
    with grande.open("wb") as archivo:
        archivo.truncate(LIMITE_BYTES + 1)
    with pytest.raises(ValueError, match="pesa mas de 20 MB"):
        leer_valuaciones_excel(grande)


def test_las_cantidades_decimales_se_leen_exactas_como_texto(tmp_path):
    ruta = _libro(tmp_path, _escribir("Lineas", 2, 3, "40.125"))
    primera = leer_valuaciones_excel(ruta).valuaciones[0]
    assert primera.lineas[0].cantidad == D("40.125")
