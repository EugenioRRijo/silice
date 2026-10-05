"""Prueba de aceptacion del sprint final (SF): del DXF de OF-001 al Excel, sin Streamlit.

Recorre, sobre una base temporal y con `ui.hoja` (la misma logica que pintan las pantallas), la
historia completa de UC-13 a UC-16 con las tres ramas integradas:

1. la hoja MANUAL de OF-001 da el total de la memoria, 2 703,86 USD;
2. un proyecto NUEVO (creado sin versiones, como lo deja la barra) importa el DXF de OF-001
   (`tests/fixtures/planos/of001.dxf`, dibujado desde las tablas de la memoria y no desde la
   hoja), pide las alturas, declara los circuitos, fusiona
   con el borrador y elabora las tres especialidades: mismas partidas y cantidades que la hoja
   manual a +-0,001 m, y el mismo total;
3. guarda su PRIMERA version, la `001` (`crear_primera_version`, RI-7), con fuente y responsable, la
   reabre y la encuentra igual: el registro de su creacion es exacto y el espacio se copio;
4. cambia la cantidad de AF-LLAVE (3 a 5), intenta guardar sin fuente (rechazo, nada se escribe) y
   guarda con fuente la `002`: el registro de cambios es EXACTO (una celda, valor anterior y nuevo,
   fuente, responsable y hora);
5. restaura la `001` como version nueva, la `003` (nada se borra);
6. el punto de equilibrio sale de la aritmetica de la version guardada (CV, CF y V del contrato),
   no de la funcion que se prueba;
7. el Excel se exporta, se evalua con `formulas` y coincide con el sistema a +-0,005 por celda.

Oraculos independientes del codigo bajo prueba: el literal 2 703,86 USD de la memoria; el 3 de
AF-LLAVE (tres llaves de paso en la memoria); el punto de equilibrio como 0,15 / 0,265 de las ventas
(administracion 15 % y utilidad 10 % en cascada: V = CV x 1,265, de modo que PE = CF x V / (V - CV)
= 0,15 x V / 0,265, 56,60 %, igual que en TI-001); y las cantidades por la via manual.
"""

from __future__ import annotations

import shutil
from dataclasses import dataclass, replace
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from adapters.espacio import cargar_espacio
from core import models
from core.budget import crear_proyecto
from core.catalog import Catalogo, abrir_sesion, crear_motor
from tests.fixtures import caso_of001, caso_ti001
from tests.fixtures.espacio_of001 import espacio_of001
from tests.fixtures.evaluador_excel import _celda, evaluar
from tests.fixtures.hoja_en_base import NOMBRE_OF001, plantilla_of001
from ui import hoja

D = Decimal
DXF = Path(__file__).resolve().parents[1] / "fixtures" / "planos" / "of001.dxf"
NOMBRE_PLANO = "OF-001 desde el plano (aceptacion SF)"
FUENTE_PLANO = "plano of001.dxf y memoria de calculo rev. C"
FUENTE_EDICION = "criterio del ingeniero residente: cinco llaves de paso, acta 12"
RESPONSABLE = "Ing. Prueba"
AHORA = datetime(2026, 10, 5, 16, 30, 0)
TOLERANCIA_M = D("0.001")
TOLERANCIA_EXCEL = D("0.005")
TOTAL_MEMORIA = D("2703.86")


@dataclass
class Historia:
    """Lo que dejo recorrer la historia completa, para que cada prueba compruebe una parte."""

    ruta: Path
    manual_total: Decimal
    importacion: object
    importacion_cruda: object
    borrador_dxf: hoja.Borrador
    total_dxf: Decimal
    guardada_001: hoja.VersionGuardada
    cambios_001: int
    clave_llave: str
    rechazo: str
    guardada_002: hoja.VersionGuardada
    restaurada: str
    excel: Path
    resultado_003: hoja.ResultadoHoja
    borrador_003: hoja.Borrador


def _lista_de(sesion) -> models.ListaPrecios:
    return sesion.scalars(select(models.ListaPrecios)).first()


def _proyecto_nuevo(sesion) -> None:
    """El proyecto nuevo, como lo deja la herramienta Proyecto de la barra: SIN versiones."""
    crear_proyecto(sesion, NOMBRE_PLANO)
    sesion.commit()


def _altura_de_la_memoria() -> str:
    """La altura de los seis ambientes, como la teclearia el usuario (la dice la memoria)."""
    return "2,50"


@pytest.fixture(scope="module")
def historia(tmp_path_factory) -> Historia:
    carpeta = tmp_path_factory.mktemp("aceptacion_sf")
    ruta = plantilla_of001(carpeta / "caso.db")
    motor = crear_motor(f"sqlite:///{ruta.as_posix()}")
    try:
        with abrir_sesion(motor) as sesion:
            catalogo = Catalogo(sesion)

            # 1. la hoja manual: el total de la memoria
            manual = hoja.abrir_version(sesion, catalogo, NOMBRE_OF001, "001")
            manual_total = manual.guardada.presupuesto.total

            # 2. proyecto nuevo (sin versiones) + importar el plano sobre un borrador en blanco
            _proyecto_nuevo(sesion)
            en_blanco = hoja.Borrador(
                fecha=caso_of001.FECHA,
                moneda=caso_of001.MONEDA,
                parametros=caso_ti001.PARAMETROS,
            )
            cruda = hoja.leer_plano(DXF)
            ambientes = {i: _altura_de_la_memoria() for i, _ in hoja.alturas_pendientes(cruda)}
            con_alturas = hoja.cargar_alturas_desde_texto(cruda, ambientes)
            circuitos = {
                c.id: tuple(str(n) for n in (*c.conductores, c.puntos_maximos))
                for c in hoja.circuitos_supuestos(con_alturas)
            }
            lista_para_declarar = hoja.declarar_circuitos_desde_texto(con_alturas, circuitos)
            borrador, _informe = hoja.fusionar_plano(en_blanco, lista_para_declarar)
            borrador = hoja.con_plan_secuencial(
                hoja.completar_composiciones(catalogo, borrador, lista=_lista_de(sesion))
            )
            total_dxf = hoja.elaborar_borrador(borrador, codigo="001").elaboracion.presupuesto.total

            # 3. guardar la PRIMERA version (RI-7): la 001
            guardada_001 = hoja.crear_primera_version(
                sesion,
                catalogo,
                NOMBRE_PLANO,
                borrador,
                fuente=FUENTE_PLANO,
                responsable=RESPONSABLE,
                ahora=AHORA,
            )
            cambios_001 = len(hoja.registros_de_version(sesion, NOMBRE_PLANO, "001"))

            # 4. editar la cantidad de AF-LLAVE, guardar sin fuente (rechazo) y con fuente
            reabierta = hoja.abrir_version(sesion, catalogo, NOMBRE_PLANO, "001")
            items = hoja.items_sin_escribir(reabierta.borrador)
            llave = next(
                c
                for c, i in zip(hoja.claves_de_items(items), items, strict=True)
                if i.codigo_partida == "AF-LLAVE"
            )
            editado = replace(reabierta.borrador, cantidades={llave: D("5")})
            try:
                hoja.guardar_version(
                    sesion, catalogo, reabierta, editado, fuente="", responsable=RESPONSABLE
                )
                rechazo = ""
            except hoja.HojaInvalida as error:
                rechazo = str(error)
            guardada_002 = hoja.guardar_version(
                sesion,
                catalogo,
                reabierta,
                editado,
                fuente=FUENTE_EDICION,
                responsable=RESPONSABLE,
                ahora=AHORA,
            )

            # 5. restaurar la 001
            restaurada = hoja.restaurar(
                sesion,
                catalogo,
                NOMBRE_PLANO,
                "001",
                fuente="la edicion de las llaves no se aprobo",
                responsable=RESPONSABLE,
                ahora=AHORA,
            )

            # 7. exportar el Excel de la version restaurada
            final = hoja.abrir_version(sesion, catalogo, NOMBRE_PLANO, restaurada)
            resultado = hoja.elaborar_borrador(final.borrador, codigo=restaurada)
            excel = hoja.exportar_excel(
                carpeta / "aceptacion.xlsx",
                resultado,
                final.borrador,
                trazabilidad=hoja.tabla_de_trazabilidad(sesion, NOMBRE_PLANO, restaurada),
                precios=hoja.precios_de_version(sesion, NOMBRE_PLANO, restaurada),
            )
            return Historia(
                ruta=ruta,
                manual_total=manual_total,
                importacion=lista_para_declarar,
                importacion_cruda=cruda,
                borrador_dxf=borrador,
                total_dxf=total_dxf,
                guardada_001=guardada_001,
                cambios_001=cambios_001,
                clave_llave=llave,
                rechazo=rechazo,
                guardada_002=guardada_002,
                restaurada=restaurada,
                excel=excel,
                resultado_003=resultado,
                borrador_003=final.borrador,
            )
    finally:
        motor.dispose()


@pytest.fixture
def sesion_de(historia, tmp_path):
    """Una sesion sobre una copia de la base que dejo la historia (las lecturas no la tocan)."""
    copia = tmp_path / "copia.db"
    shutil.copy(historia.ruta, copia)
    motor = crear_motor(f"sqlite:///{copia.as_posix()}")
    with abrir_sesion(motor) as sesion:
        yield sesion, Catalogo(sesion)
    motor.dispose()


def _modelo(sesion, proyecto: str, codigo: str) -> models.Presupuesto:
    return sesion.scalars(
        select(models.Presupuesto)
        .join(models.Proyecto)
        .where(models.Proyecto.nombre == proyecto, models.Presupuesto.codigo == codigo)
    ).one()


def _firma(items):
    return sorted((i.codigo_partida, i.unidad, i.cantidad) for i in items)


# --- 1 y 2: la via manual y la via del plano dan lo mismo ----------------------------------------


def test_la_hoja_manual_de_of001_da_el_total_de_la_memoria(historia):
    assert abs(historia.manual_total - TOTAL_MEMORIA) <= D("0.01")


def test_el_plano_pidio_las_seis_alturas_y_las_dos_declaraciones_de_circuito(historia):
    cruda = historia.importacion_cruda
    assert len(hoja.alturas_pendientes(cruda)) == 6
    assert [c.id for c in hoja.circuitos_supuestos(cruda)] == ["C1", "C2"]
    # ya cargadas y declaradas, no queda nada pendiente ni supuesto
    assert hoja.alturas_pendientes(historia.importacion) == []
    assert hoja.circuitos_supuestos(historia.importacion) == []


@pytest.mark.parametrize("especialidad", hoja.TODAS)
def test_las_cantidades_por_el_plano_son_las_de_la_hoja_manual_a_una_milesima(
    historia, especialidad
):
    del_plano = _firma(hoja.items_del_borrador(historia.borrador_dxf.espacio, especialidad))
    manual = _firma(hoja.items_del_borrador(espacio_of001(), especialidad))
    assert del_plano
    assert [(c, u) for c, u, _ in del_plano] == [(c, u) for c, u, _ in manual]
    for (codigo, _, a), (_, _, b) in zip(del_plano, manual, strict=True):
        assert abs(a - b) <= TOLERANCIA_M, (especialidad, codigo, a, b)


def test_el_total_por_el_plano_es_el_de_la_memoria(historia):
    assert abs(historia.total_dxf - TOTAL_MEMORIA) <= D("0.01")


# --- 3: la primera version, con fuente y responsable, y reabrir -----------------------------------


def test_la_primera_version_se_guardo_con_su_registro_y_se_reabre_igual(historia, sesion_de):
    sesion, catalogo = sesion_de
    assert historia.guardada_001.codigo == "001"
    assert historia.guardada_001.cambios == historia.cambios_001 == 1
    assert abs(historia.guardada_001.total - TOTAL_MEMORIA) <= D("0.01")

    reabierta = hoja.abrir_version(sesion, catalogo, NOMBRE_PLANO, "001")
    assert reabierta.borrador.espacio == historia.borrador_dxf.espacio
    assert reabierta.guardada.presupuesto.total == historia.guardada_001.total
    # el espacio quedo guardado con la clave de la version nueva
    nueva = _modelo(sesion, NOMBRE_PLANO, "001")
    assert len(cargar_espacio(sesion, nueva.id).ambientes) == 6


def test_el_registro_de_creacion_es_exacto_con_fuente_responsable_y_hora(historia, sesion_de):
    sesion, _ = sesion_de
    v1 = _modelo(sesion, NOMBRE_PLANO, "001")
    registros = list(
        sesion.scalars(
            select(models.RegistroCambio).where(
                models.RegistroCambio.version_resultante_id == v1.id
            )
        )
    )
    assert [
        (r.objeto, r.clave, r.campo, r.valor_anterior, r.valor_nuevo, r.version_origen_id)
        for r in registros
    ] == [("version", "001", "creacion", "", "001", None)]
    assert {(r.fuente, r.responsable, r.fecha_hora) for r in registros} == {
        (FUENTE_PLANO, RESPONSABLE, AHORA)
    }


# --- 4: editar una cantidad, guardar sin fuente y con fuente -------------------------------------


def test_guardar_sin_fuente_se_rechazo_con_el_motivo_y_no_creo_version(historia, sesion_de):
    sesion, _ = sesion_de
    assert "fuente" in historia.rechazo
    codigos = sorted(
        p.codigo
        for p in sesion.scalars(
            select(models.Presupuesto)
            .join(models.Proyecto)
            .where(models.Proyecto.nombre == NOMBRE_PLANO)
        )
    )
    assert codigos == ["001", "002", "003"]  # la 002 es la guardada con fuente


def test_el_registro_de_la_edicion_es_exactamente_la_celda_que_se_cambio(historia, sesion_de):
    sesion, _ = sesion_de
    assert historia.guardada_002.codigo == "002" and historia.guardada_002.cambios == 1
    registros = hoja.registros_de_version(sesion, NOMBRE_PLANO, "002")
    assert [
        (r.objeto, r.clave, r.campo, r.valor_anterior, r.valor_nuevo, r.fuente, r.responsable)
        for r in registros
    ] == [
        (
            "cantidad",
            historia.clave_llave,
            "cantidad",
            "3",
            "5",
            FUENTE_EDICION,
            RESPONSABLE,
        )
    ]
    assert registros[0].fecha_hora == AHORA


def test_la_cantidad_escrita_a_mano_sube_el_total_en_la_aritmetica_de_dos_llaves(
    historia, sesion_de
):
    """Dos llaves mas (de 3 a 5) suman 2 x CD x 1,15 x 1,10 al total de la 001."""
    sesion, catalogo = sesion_de
    v1 = hoja.abrir_version(sesion, catalogo, NOMBRE_PLANO, "001").guardada.presupuesto
    v2 = hoja.abrir_version(sesion, catalogo, NOMBRE_PLANO, "002").guardada.presupuesto
    llave = next(p for p in v1.partidas if p.item.codigo_partida == "AF-LLAVE")
    assert llave.item.cantidad == D("3")
    esperado = 2 * llave.resultado.costo_directo * D("1.15") * D("1.10")
    assert abs((v2.total - v1.total) - esperado) <= D("0.0001")


# --- 5: restaurar --------------------------------------------------------------------------------


def test_restaurar_la_001_crea_la_003_igual_a_ella_y_no_borra_nada(historia, sesion_de):
    sesion, catalogo = sesion_de
    assert historia.restaurada == "003"
    assert hoja.codigo_vigente(sesion, NOMBRE_PLANO) == "003"
    v1 = hoja.abrir_version(sesion, catalogo, NOMBRE_PLANO, "001")
    v3 = hoja.abrir_version(sesion, catalogo, NOMBRE_PLANO, "003")
    assert v3.guardada.presupuesto.total == v1.guardada.presupuesto.total
    assert v3.borrador.espacio == v1.borrador.espacio
    assert dict(v3.borrador.cantidades) == {}  # las cinco llaves de la 002 no volvieron
    filas = hoja.historial_de_versiones(sesion, NOMBRE_PLANO)
    assert [(f.codigo, f.motivo, f.origen, f.vigente) for f in filas] == [
        ("001", "", "", False),
        ("002", "edicion", "001", False),
        ("003", "restauracion", "001", True),
    ]


# --- 6: el punto de equilibrio -------------------------------------------------------------------


def test_el_punto_de_equilibrio_es_el_de_la_aritmetica_de_la_version(historia):
    pe = historia.resultado_003.equilibrio
    total = historia.resultado_003.elaboracion.presupuesto.total
    # oraculo cerrado: V = CV x 1,265 (administracion 15 % y utilidad 10 % en cascada), asi que
    # PE = CF x V / (V - CV) = 0,15 x V / 0,265 y es el 56,60 % de las ventas, como en TI-001
    assert abs(pe.porcentaje - D("56.60")) <= D("0.01")
    assert abs(pe.punto - total * D("0.15") / D("0.265")) <= D("0.01")
    # y por el camino largo, con las cifras de cada renglon que trae la version guardada
    cv = sum(
        (
            p.item.cantidad * p.resultado.costo_directo
            for p in historia.resultado_003.elaboracion.presupuesto.partidas
        ),
        D(0),
    )
    assert abs(pe.costo_variable - cv) <= D("0.0001")
    assert abs(pe.punto - cv * D("0.15") * total / (total - cv)) <= D("0.01")
    assert pe.periodo is not None  # la version restaurada trae su plan: el PE se ubica en la curva


# --- 7: el Excel ---------------------------------------------------------------------------------


def test_el_excel_evaluado_coincide_con_el_sistema_a_cinco_milesimas(historia):
    libro = load_workbook(historia.excel)
    assert libro.sheetnames[:9] == [
        "Espacio y metrajes",
        "APU",
        "Presupuesto",
        "Plan y curva",
        "Punto de equilibrio",
        "Precios",
        "Parametros",
        "Auditoria",
        "Trazabilidad",
    ]
    celdas = evaluar(historia.excel)
    pe = historia.resultado_003.equilibrio
    total = historia.resultado_003.elaboracion.presupuesto.total

    total_xl = celdas[("PRESUPUESTO", _celda(libro, "Presupuesto", "TOTAL", "G"))]
    assert abs(D(str(total_xl)) - total) <= TOLERANCIA_EXCEL
    for etiqueta, esperado in (
        ("Ventas (V)", pe.venta),
        ("Costo variable (CV)", pe.costo_variable),
        ("Costo fijo (CF)", pe.costo_fijo),
        ("Punto de equilibrio (PE)", pe.punto),
    ):
        valor = celdas[("PUNTO DE EQUILIBRIO", _celda(libro, "Punto de equilibrio", etiqueta))]
        assert abs(D(str(valor)) - esperado) <= TOLERANCIA_EXCEL, etiqueta
    porcentaje = celdas[
        ("PUNTO DE EQUILIBRIO", _celda(libro, "Punto de equilibrio", "PE como % de las ventas"))
    ]
    assert abs(D(str(porcentaje)) - pe.porcentaje) <= TOLERANCIA_EXCEL


def test_el_excel_lleva_el_espacio_importado_y_la_trazabilidad_de_la_restauracion(historia):
    libro = load_workbook(historia.excel)
    metrajes = libro["Espacio y metrajes"]
    assert metrajes.max_row == 1 + len(historia.resultado_003.items)
    assert len(libro["Esp ambientes"]["A"]) == 1 + 6
    traza = libro["Trazabilidad"]
    assert traza.max_row >= 2  # la restauracion deja su registro
