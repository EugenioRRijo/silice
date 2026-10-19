"""Ola 2 de la migración: pruebas que cierran la revisión de `silice-rev-pruebas`.

Cada prueba apunta a un hallazgo de ese dictamen (se cita entre corchetes). Oráculos a mano o de
fixtures independientes: el informe del presupuesto auditado sale de sus defectos de diseño
(`tests/fixtures/apu_linea_base.py`), la curva de los totales por partida de ese mismo archivo y el
responsable de `seguridad.cuentas.texto_responsable` (`Nombre <login>`, RF-82).
"""

from __future__ import annotations

import io
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import select

from api.dependencias import USUARIO_SIN_CLAVE
from core import models
from core.budget import plan_de_version
from core.verification.informe import DECIMALES_PRESENTACION
from core.verification.texto import formatear_decimal
from scripts.seed import NOMBRE_PROYECTO
from tests.fixtures import apu_linea_base as linea_base
from tests.fixtures.sin_rendimiento_vigente import (
    FECHA_PRESUPUESTO,
    dejar_sin_estimado,
    dejar_sin_historico,
)

RAIZ = Path(__file__).resolve().parents[2]
MUESTRA = RAIZ / "data" / "samples" / "precios" / "lista_2026-06-01.csv"
CABECERA = "tipo,insumo,unidad,precio\n"
CODIGOS = list(linea_base.CODIGOS_LINEA_BASE)
DATOS = {
    "proyecto": NOMBRE_PROYECTO,
    "codigo": "API-001",
    "codigo_nuevo": "API-002",
    "nombre": "Lista de prueba",
    "moneda": "USD",
    "fecha_vigencia": "2026-06-01",
}


def _csv(texto: str, nombre: str = "lista.csv"):
    return {"archivo": (nombre, texto.encode("utf-8"), "text/csv")}


def _revalorar(cliente, texto: str | None = None, dar_de_alta=(), **cambios):
    contenido = MUESTRA.read_text(encoding="utf-8") if texto is None else texto
    datos = {**DATOS, **cambios}
    if dar_de_alta:
        datos["dar_de_alta"] = list(dar_de_alta)
    return cliente.post("/presupuestos/revalorar", files=_csv(contenido), data=datos)


def _guardar_base(cliente, items, **extra):
    respuesta = cliente.post(
        "/presupuestos",
        json={
            "codigo": "API-001",
            "fecha": "2026-04-28",
            "moneda": "USD",
            "proyecto": NOMBRE_PROYECTO,
            "items": items,
            **extra,
        },
    )
    assert respuesta.status_code == 201, respuesta.text
    return respuesta.json()


def _calculo(items, **extra):
    return {"codigo": "P-001", "fecha": "2026-04-28", "moneda": "USD", "items": items, **extra}


def _modelo(sesion, codigo: str) -> models.Presupuesto:
    return sesion.scalars(
        select(models.Presupuesto).where(models.Presupuesto.codigo == codigo)
    ).one()


def _item_pendiente(codigo="NO-EXISTE-99"):
    return {
        "codigo_partida": codigo,
        "descripcion": "Excavación manual de zanja",
        "unidad": "m3",
        "cantidad": "2.5",
        "origen_id": "T-99",
        "origen_tipo": "manual",
        "dominio": "civil",
    }


@pytest.fixture
def base(cliente, items_payload):
    return _guardar_base(cliente, items_payload)


@pytest.fixture
def navegador(sesion, tmp_path_factory):
    """Un navegador con la cookie de `ana` (rol arma) sobre la base en memoria de `cliente`."""
    from api.main import app
    from tests.fixtures.sesion_web import (
        con_cuenta,
        copiar_instalacion,
        plantilla_con_cuentas,
    )

    copiar_instalacion(plantilla_con_cuentas(tmp_path_factory.mktemp("cuentas")))
    return con_cuenta(app, "ana")


# ---------------------------------------------------------------------------------------------
# [1] RF-82: quien confirma la revaloración queda sellado en la cadena
# ---------------------------------------------------------------------------------------------


def _usuario_sellado(sesion, codigo: str) -> str:
    modelo = _modelo(sesion, codigo)
    return sesion.get(models.EslabonVersion, modelo.id).usuario


def test_revalorar_con_clave_de_api_sella_al_usuario_sin_clave(cliente, sesion, base):
    assert _revalorar(cliente).status_code == 200

    assert _usuario_sellado(sesion, "API-002") == USUARIO_SIN_CLAVE


def test_revalorar_con_la_cuenta_sella_nombre_y_login_de_la_persona(
    cliente, sesion, navegador, base
):
    from tests.fixtures.sesion_web import protegida

    persona, ficha = navegador
    respuesta = persona.post(
        "/presupuestos/revalorar",
        files=_csv(MUESTRA.read_text(encoding="utf-8")),
        data=DATOS,
        headers=protegida(ficha),
    )

    assert respuesta.status_code == 200, respuesta.text
    assert _usuario_sellado(sesion, "API-002") == "Ana <ana>"


# ---------------------------------------------------------------------------------------------
# [2] SF3-R4: revalorar conserva el plan guardado de la versión de origen
# ---------------------------------------------------------------------------------------------


def test_revalorar_conserva_el_plan_de_la_version_de_origen(cliente, sesion):
    """El origen se guarda por el núcleo con un plan de tres días que parte la excavación en dos
    (`POST /presupuestos` aún no guarda el plan en `PlanVersion`: hallazgo anotado en el plan)."""
    from core.budget import elaborar, guardar_presupuesto
    from core.catalog import Catalogo
    from tests.fixtures.computo_auditado import items_auditados

    c = CODIGOS
    plan = [
        ("Dia 1", [(c[0], Decimal("0.5"))]),
        ("Dia 2", [(c[0], Decimal("0.5")), (c[1], Decimal("1"))]),
        ("Dia 3", [(c[2], Decimal("1")), (c[3], Decimal("1")), (c[4], Decimal("1"))]),
    ]
    catalogo = Catalogo(sesion)
    resultado = elaborar(
        items_auditados(),
        catalogo.composiciones(c, fecha=linea_base.FECHA_LINEA_BASE),
        linea_base.PARAMETROS_LINEA_BASE,
        codigo="API-001",
        fecha=linea_base.FECHA_LINEA_BASE,
        moneda=linea_base.MONEDA,
        plan=plan,
    )
    guardar_presupuesto(
        sesion,
        resultado.presupuesto,
        resultado.informe,
        proyecto=sesion.scalars(select(models.Proyecto)).one(),
        lista=catalogo.lista_vigente(linea_base.FECHA_LINEA_BASE),
        parametros=linea_base.PARAMETROS_LINEA_BASE,
        plan=plan,
    )
    sesion.commit()
    assert plan_de_version(sesion, _modelo(sesion, "API-001")) == [(e, tuple(r)) for e, r in plan]

    assert _revalorar(cliente).status_code == 200

    # Con `plan=None` la versión nueva tendría 5 periodos (una partida por día), no estos 3.
    assert [p.periodo.split(" [")[0] for p in _modelo(sesion, "API-002").curva] == [
        "Dia 1",
        "Dia 2",
        "Dia 3",
    ]


# ---------------------------------------------------------------------------------------------
# [3] el plan enviado a `calcular` se usa: se ve en la curva del libro de Excel
# ---------------------------------------------------------------------------------------------


def _curva(respuesta) -> list[tuple[str, float, float]]:
    libro = openpyxl.load_workbook(io.BytesIO(respuesta.content))
    filas = list(libro["Curva"].iter_rows(min_row=2, values_only=True))
    return [(fila[0], fila[1], fila[2]) for fila in filas]


def test_el_plan_por_defecto_da_una_partida_por_dia_en_la_curva(cliente, items_payload):
    respuesta = cliente.post("/presupuestos/calcular/excel", json=_calculo(items_payload))

    curva = _curva(respuesta)
    assert [codigo for _, codigo in ((f[0], f[0].split("[")[1][:-1]) for f in curva)] == CODIGOS
    assert [f[1] for f in curva] == [
        float(linea.total) for linea in linea_base.PRESUPUESTO_AUDITADO
    ]
    assert curva[-1][2] == float(linea_base.TOTAL_PRESUPUESTO_AUDITADO)


def test_un_plan_de_un_solo_periodo_cambia_la_curva_y_cierra_en_el_total(cliente, items_payload):
    filas = [{"periodo": "S1", "codigo_partida": c, "fraccion": "1"} for c in CODIGOS]

    respuesta = cliente.post(
        "/presupuestos/calcular/excel", json=_calculo(items_payload, plan=filas)
    )

    curva = _curva(respuesta)
    assert len(curva) == 1
    assert curva[0][0].startswith("S1 [")
    assert curva[0][1] == curva[0][2] == float(linea_base.TOTAL_PRESUPUESTO_AUDITADO)


def test_un_plan_de_dos_semanas_parte_cada_partida_a_la_mitad(cliente, items_payload):
    filas = [
        {"periodo": s, "codigo_partida": c, "fraccion": "0.5"}
        for s in ("S1", "S2")
        for c in CODIGOS
    ]

    curva = _curva(
        cliente.post("/presupuestos/calcular/excel", json=_calculo(items_payload, plan=filas))
    )

    assert [f[0].split(" ")[0] for f in curva] == ["S1", "S2"]
    assert curva[0][1] == pytest.approx(float(linea_base.TOTAL_PRESUPUESTO_AUDITADO) / 2, abs=0.02)
    assert curva[1][2] == float(linea_base.TOTAL_PRESUPUESTO_AUDITADO)


# ---------------------------------------------------------------------------------------------
# [4] el informe del presupuesto auditado: conjunto exacto de hallazgos
# ---------------------------------------------------------------------------------------------

#: Los defectos de diseño del presupuesto auditado (`apu_linea_base.py`): la tubería se computa en
#: «mts» y su APU está en «pieza» (R3), el encofrado en m3 con un APU en m2 (R3), el vehículo de
#: transporte lleva dos factores de depreciación (R6) y las cinco cantidades son manuales (R1).
HALLAZGOS_AUDITADO = {
    ("R3", "ERROR", "LB-02-TUB"),
    ("R3", "ERROR", "LB-03-ENC"),
    ("R6", "ERROR", "LB-02-TUB"),
    *{("R1", "ADVERTENCIA", codigo) for codigo in CODIGOS},
}


def test_el_informe_del_presupuesto_auditado_declara_exactamente_sus_defectos(
    cliente, items_payload
):
    cuerpo = cliente.post("/presupuestos/calcular", json=_calculo(items_payload)).json()

    informe = cuerpo["informe"]
    assert {
        (h["regla"], h["severidad"], h["partida"]) for h in informe["hallazgos"]
    } == HALLAZGOS_AUDITADO
    assert informe["total_hallazgos"] == len(informe["hallazgos"]) == 8
    assert informe["veredicto"] == "NO CUMPLE" and informe["cumple"] is False
    assert informe["sin_hallazgos"] is None
    assert cuerpo["por_severidad"] == [
        {"severidad": "ERROR", "total": 3},
        {"severidad": "ADVERTENCIA", "total": 5},
    ]


def test_el_informe_de_la_version_revalorada_conserva_los_defectos_del_base(cliente, base):
    informe = _revalorar(cliente).json()["informe"]

    assert informe["codigo_presupuesto"] == "API-002"
    assert {
        (h["regla"], h["severidad"], h["partida"]) for h in informe["hallazgos"]
    } == HALLAZGOS_AUDITADO


# ---------------------------------------------------------------------------------------------
# [5] y [6] pendientes con propuestas, y aceptar una propuesta con total exacto
# ---------------------------------------------------------------------------------------------


def test_un_pendiente_declara_su_hallazgo_pendiente_y_la_forma_de_sus_propuestas(
    cliente, items_payload
):
    cuerpo = cliente.post(
        "/presupuestos/calcular", json=_calculo([*items_payload, _item_pendiente()])
    ).json()

    (pendiente,) = cuerpo["pendientes"]
    assert pendiente["codigo_partida"] == "NO-EXISTE-99"
    for propuesta in pendiente["propuestas"]:
        assert set(propuesta) == {"codigo", "descripcion", "puntaje"}
        assert propuesta["codigo"] in CODIGOS
    pendientes = [h for h in cuerpo["informe"]["hallazgos"] if h["regla"] == "PENDIENTE"]
    assert len(pendientes) == 1
    assert pendientes[0]["severidad"] == "ERROR"
    assert "T-99" in pendientes[0]["origen"]


def test_un_normalizador_que_falla_deja_al_pendiente_sin_propuestas_y_con_su_falla(
    cliente, items_payload, monkeypatch
):
    from ml.normalization import normalizador

    class _Propuesta:
        similares: tuple = ()
        falla = "el modelo de similitud no está disponible"

    monkeypatch.setattr(
        "ml.normalization.propuestas_para",
        lambda partidas, pendientes: [(item, _Propuesta()) for item in pendientes],
    )
    assert normalizador  # el módulo existe: el parche es sobre el nombre que importa la lógica

    cuerpo = cliente.post(
        "/presupuestos/calcular", json=_calculo([*items_payload, _item_pendiente()])
    ).json()

    (pendiente,) = cuerpo["pendientes"]
    assert pendiente["propuestas"] == []
    assert pendiente["falla"] == "el modelo de similitud no está disponible"


def test_aceptar_una_propuesta_suma_exactamente_cantidad_por_precio_unitario(
    cliente, items_payload
):
    elegido = "LB-01-EXC"

    despues = cliente.post(
        "/presupuestos/calcular", json=_calculo([*items_payload, _item_pendiente(elegido)])
    ).json()

    esperado = (
        linea_base.TOTAL_EXACTO_LINEA_BASE
        + Decimal("2.5") * linea_base.PRECIO_UNITARIO_EXACTO[elegido]
    )
    assert despues["pendientes"] == []
    assert despues["total"] == formatear_decimal(esperado, DECIMALES_PRESENTACION)


# ---------------------------------------------------------------------------------------------
# [7] el base de escenarios devuelve los parámetros congelados, no los de por defecto
# ---------------------------------------------------------------------------------------------


def test_el_base_trae_los_parametros_congelados_del_presupuesto(cliente, items_payload):
    _guardar_base(cliente, items_payload, fcas="5.50", utilidad="0.12", bono_alimentacion="2.5")

    cuerpo = cliente.get("/presupuestos/API-001/escenarios/base").json()

    assert cuerpo["parametros"] == {
        "fcas": "5.50",
        "bono_alimentacion": "2.5",
        "administracion": "0.15",
        "utilidad": "0.12",
    }


def test_el_base_de_un_proyecto_inexistente_es_404(cliente, base):
    respuesta = cliente.get(
        "/presupuestos/API-001/escenarios/base", params={"proyecto": "No existe"}
    )

    assert respuesta.status_code == 404


def test_el_base_con_el_proyecto_correcto_responde(cliente, base):
    respuesta = cliente.get(
        "/presupuestos/API-001/escenarios/base", params={"proyecto": NOMBRE_PROYECTO}
    )

    assert respuesta.status_code == 200 and respuesta.json()["proyecto"] == NOMBRE_PROYECTO


# ---------------------------------------------------------------------------------------------
# [8] el CSV de escenarios
# ---------------------------------------------------------------------------------------------


def test_el_csv_calcula_el_escenario_de_utilidad_con_valores_exactos(cliente, base):
    import csv

    respuesta = cliente.post(
        "/presupuestos/API-001/escenarios/csv",
        json={"escenarios": [{"nombre": "utilidad 20", "utilidad": "0.20"}]},
    )

    filas = list(csv.DictReader(io.StringIO(respuesta.text)))
    total_base = Decimal(filas[0]["total"])
    assert total_base == linea_base.TOTAL_EXACTO_LINEA_BASE
    # La utilidad encadena en cascada sobre todo el costo: total × 1,20 / 1,10.
    assert Decimal(filas[1]["total"]) == pytest.approx(
        total_base * Decimal("1.2") / Decimal("1.1"), abs=Decimal("1e-9")
    )
    esperado_pct = Decimal(100) * (Decimal("1.2") / Decimal("1.1") - 1)
    assert Decimal(filas[1]["variacion_pct"]) == pytest.approx(esperado_pct, abs=Decimal("1e-9"))


def test_el_csv_de_un_presupuesto_inexistente_es_404(cliente, base):
    respuesta = cliente.post(
        "/presupuestos/NO-EXISTE/escenarios/csv", json={"escenarios": [{"nombre": "x"}]}
    )

    assert respuesta.status_code == 404


def test_el_csv_con_un_precio_que_no_es_decimal_es_422(cliente, base):
    respuesta = cliente.post(
        "/presupuestos/API-001/escenarios/csv",
        json={"escenarios": [{"nombre": "x", "precios": {"Cemento Portland": "doce"}}]},
    )

    assert respuesta.status_code == 422
    assert "Cemento Portland" in respuesta.json()["detalle"]


def test_el_csv_de_un_base_con_pendientes_no_los_cuenta(cliente_con_pendiente):
    import csv

    respuesta = cliente_con_pendiente.post(
        f"/presupuestos/{linea_base.CODIGO_PRESUPUESTO}/escenarios/csv",
        json={"escenarios": [{"nombre": "utilidad 20", "utilidad": "0.20"}]},
    )

    filas = list(csv.DictReader(io.StringIO(respuesta.text)))
    assert Decimal(filas[0]["total"]) == linea_base.TOTAL_EXACTO_LINEA_BASE  # el pendiente no suma


# ---------------------------------------------------------------------------------------------
# [9] descartar: errores y estado persistido
# ---------------------------------------------------------------------------------------------


def test_descartar_deja_la_marca_con_la_fecha_de_hoy(cliente, sesion, base):
    from datetime import date

    _revalorar(cliente)

    respuesta = cliente.post(
        "/presupuestos/API-002/descartar", params={"proyecto": NOMBRE_PROYECTO}
    )

    marca = sesion.get(models.DescarteVersion, _modelo(sesion, "API-002").id)
    assert marca is not None and marca.fecha == date.today()
    assert respuesta.json()["fecha"] == date.today().isoformat()


def test_descartar_un_codigo_inexistente_es_404(cliente, base):
    assert cliente.post("/presupuestos/NO-EXISTE/descartar").status_code == 404


def test_un_descarte_rechazado_no_escribe_nada(cliente, sesion, base):
    _revalorar(cliente)

    rechazado = cliente.post(
        "/presupuestos/API-001/descartar", params={"proyecto": NOMBRE_PROYECTO}
    )

    assert rechazado.status_code == 422
    assert sesion.scalars(select(models.DescarteVersion)).all() == []


def test_descartar_sin_proyecto_con_un_codigo_en_dos_proyectos_es_409(
    cliente, sesion, items_payload
):
    from core.budget import crear_proyecto

    _guardar_base(cliente, items_payload)
    crear_proyecto(sesion, "Otro proyecto ficticio", "", "civil")
    sesion.commit()
    otro = cliente.post(
        "/presupuestos",
        json={
            "codigo": "API-001",
            "fecha": "2026-04-28",
            "moneda": "USD",
            "proyecto": "Otro proyecto ficticio",
            "items": items_payload,
        },
    )
    assert otro.status_code == 201, otro.text

    respuesta = cliente.post("/presupuestos/API-001/descartar")

    assert respuesta.status_code == 409
    assert "Otro proyecto ficticio" in respuesta.json()["detalle"]


# ---------------------------------------------------------------------------------------------
# [10] y [11] revalorar: conflicto con rollback y detector de atípicos
# ---------------------------------------------------------------------------------------------


def test_el_conflicto_de_un_codigo_repetido_no_deja_otra_lista_ni_altas(cliente, sesion, base):
    assert _revalorar(cliente).status_code == 200
    listas = cliente.get("/listas-precios").json()
    insumos = cliente.get("/insumos").json()
    lista = CABECERA + "material,Cemento Portland,saco,19\nmaterial,Cemento blanco,saco,20\n"

    segunda = _revalorar(cliente, lista, dar_de_alta=["Cemento blanco"], nombre="Otra lista")

    assert segunda.status_code in (409, 422)
    assert cliente.get("/listas-precios").json() == listas
    assert cliente.get("/insumos").json() == insumos


def test_un_error_de_base_de_datos_en_revalorar_es_409_con_rollback(cliente, base, monkeypatch):
    from sqlalchemy.exc import IntegrityError

    def _falla(*args, **kwargs):
        raise IntegrityError(
            "INSERT", {}, Exception("UNIQUE constraint failed: presupuesto.codigo")
        )

    monkeypatch.setattr("aplicacion.actualizacion.revalorar", _falla)
    listas = cliente.get("/listas-precios").json()

    respuesta = _revalorar(cliente)

    assert respuesta.status_code == 409
    assert "ya existe una version con ese codigo" in respuesta.json()["detalle"]
    assert cliente.get("/listas-precios").json() == listas


def test_revalorar_pasa_a_actualizar_precios_el_detector_y_el_usuario(cliente, base, monkeypatch):
    import aplicacion.actualizacion as aplicacion

    visto = {}
    original = aplicacion.actualizar_precios

    def _espia(*args, **kwargs):
        visto.update(kwargs)
        return original(*args, **{**kwargs, "detectar_atipicos": None})

    detector = object()
    monkeypatch.setattr(aplicacion, "detector_de_atipicos", lambda: detector)
    monkeypatch.setattr(aplicacion, "actualizar_precios", _espia)

    assert _revalorar(cliente).status_code == 200
    assert visto["detectar_atipicos"] is detector
    assert visto["usuario"] == USUARIO_SIN_CLAVE


def test_sin_el_extra_ml_revalorar_no_falla(cliente, base, monkeypatch):
    monkeypatch.setattr("importlib.util.find_spec", lambda nombre: None)

    assert _revalorar(cliente).status_code == 200


# ---------------------------------------------------------------------------------------------
# [12] altas y vista previa
# ---------------------------------------------------------------------------------------------


def test_un_alta_ajena_no_deja_lista_ni_alta(cliente, base):
    listas = cliente.get("/listas-precios").json()
    insumos = cliente.get("/insumos").json()

    respuesta = _revalorar(
        cliente,
        CABECERA + "material,Cemento Portland,saco,18\n",
        dar_de_alta=["Cemento Portland"],
    )

    assert respuesta.status_code == 422
    assert cliente.get("/listas-precios").json() == listas
    assert cliente.get("/insumos").json() == insumos


def test_dos_filas_del_mismo_desconocido_con_precios_distintos_son_422_sin_escribir(cliente, base):
    lista = (
        CABECERA
        + "material,Cemento Portland,saco,18\n"
        + "material,Cemento blanco,saco,20\n"
        + "material,Cemento blanco,saco,21\n"
    )
    insumos = cliente.get("/insumos").json()

    respuesta = _revalorar(cliente, lista, dar_de_alta=["Cemento blanco"])

    assert respuesta.status_code == 422
    assert cliente.get("/insumos").json() == insumos


def test_pedir_dos_veces_el_mismo_alta_da_un_solo_insumo(cliente, base):
    lista = CABECERA + "material,Cemento Portland,saco,18\nmaterial,Cemento blanco,saco,20\n"

    respuesta = _revalorar(cliente, lista, dar_de_alta=["Cemento blanco", "Cemento blanco"])

    assert respuesta.status_code == 200, respuesta.text
    assert len(respuesta.json()["dados_de_alta"]) == 1


def test_la_vista_previa_de_un_archivo_vacio_no_ofrece_nada(cliente, base):
    respuesta = cliente.post("/listas-precios/vista-previa", files=_csv(""))

    assert respuesta.status_code == 200
    assert respuesta.json()["desconocidos"] == []


def test_la_vista_previa_con_nombre_inutilizable_es_422(cliente, base):
    respuesta = cliente.post(
        "/listas-precios/vista-previa", files={"archivo": ("..", b"tipo,insumo\n", "text/csv")}
    )

    assert respuesta.status_code == 422
    assert "nombre utilizable" in respuesta.json()["detalle"]


def test_el_aviso_de_cambios_que_no_afectan_al_presupuesto(cliente, base):
    """Un insumo dado de alta y que ninguna partida usa cambia de precio: se guarda la versión y
    el total no varía (distinto del flujo 3a, donde ningún precio cambia)."""
    sin_uso = "Insumo ficticio sin uso"
    primera = _revalorar(
        cliente,
        CABECERA + f"material,Cemento Portland,saco,18\nmaterial,{sin_uso},saco,5\n",
        dar_de_alta=[sin_uso],
        codigo="API-001",
        codigo_nuevo="API-002",
        nombre="Con alta",
    )
    assert primera.status_code == 200, primera.text

    segunda = _revalorar(
        cliente,
        CABECERA + f"material,{sin_uso},saco,6\n",
        codigo="API-002",
        codigo_nuevo="API-003",
        nombre="Solo el insumo sin uso",
    )

    assert segunda.status_code == 200, segunda.text
    cuerpo = segunda.json()
    assert cuerpo["version_guardada"] == "API-003"
    assert cuerpo["insumos_afectados"] == 1
    assert cuerpo["variacion"] == "0.00"
    assert cuerpo["avisos"] == [
        {
            "nivel": "info",
            "texto": "Cambiaron precios de insumos que no afectan a este presupuesto: se guardo la "
            "version nueva y su historial de cambios, pero el total no vario.",
        }
    ]


# ---------------------------------------------------------------------------------------------
# [13] declarar rendimiento: casos que faltaban
# ---------------------------------------------------------------------------------------------

DECLARAR = "/presupuestos/calcular/declarar-rendimiento"


def test_declarar_una_partida_inexistente_es_404(cliente):
    respuesta = cliente.post(
        DECLARAR, json={"codigo_partida": "NO-EXISTE", "tipo": "estimado", "fecha": "2026-06-01"}
    )

    assert respuesta.status_code == 404


def test_declarar_con_un_tipo_invalido_es_422(cliente):
    respuesta = cliente.post(
        DECLARAR, json={"codigo_partida": "LB-02-TUB", "tipo": "inventado", "fecha": "2026-06-01"}
    )

    assert respuesta.status_code == 422


def test_declarar_una_partida_que_ya_tiene_estimado_vigente_es_422(cliente):
    respuesta = cliente.post(
        DECLARAR,
        json={
            "codigo_partida": "LB-02-TUB",
            "tipo": "estimado",
            "fecha": FECHA_PRESUPUESTO.isoformat(),
        },
    )

    assert respuesta.status_code == 422
    assert "ya tiene un rendimiento estimado vigente" in respuesta.json()["detalle"]


def test_declarar_una_partida_sin_historico_es_422(cliente, sesion):
    dejar_sin_historico(sesion, "LB-02-TUB")
    sesion.commit()

    respuesta = cliente.post(
        DECLARAR,
        json={
            "codigo_partida": "LB-02-TUB",
            "tipo": "estimado",
            "fecha": FECHA_PRESUPUESTO.isoformat(),
        },
    )

    assert respuesta.status_code == 422
    assert "no tiene historico de rendimientos que proponer" in respuesta.json()["detalle"]


def test_un_medido_rechazado_no_registra_nada_y_calcular_sigue_en_409(
    cliente, sesion, items_payload
):
    dejar_sin_estimado(sesion)
    sesion.commit()
    cuerpo = _calculo(items_payload, fecha=FECHA_PRESUPUESTO.isoformat())
    antes = len(sesion.scalars(select(models.Rendimiento)).all())

    rechazo = cliente.post(
        DECLARAR,
        json={
            "codigo_partida": "LB-02-TUB",
            "tipo": "medido",
            "fecha": FECHA_PRESUPUESTO.isoformat(),
        },
    )

    assert rechazo.status_code == 422
    assert len(sesion.scalars(select(models.Rendimiento)).all()) == antes
    assert cliente.post("/presupuestos/calcular", json=cuerpo).status_code == 409


def test_el_409_propone_el_rendimiento_de_cada_partida_sin_estimado_en_su_orden(
    cliente, sesion, items_payload
):
    dejar_sin_estimado(sesion)
    dejar_sin_estimado(sesion, "LB-05-REL")
    sesion.commit()

    respuesta = cliente.post(
        "/presupuestos/calcular", json=_calculo(items_payload, fecha=FECHA_PRESUPUESTO.isoformat())
    )

    assert respuesta.status_code == 409
    assert [e["codigo_partida"] for e in respuesta.json()["rendimientos_por_declarar"]] == [
        "LB-02-TUB",
        "LB-05-REL",
    ]


# ---------------------------------------------------------------------------------------------
# [14] el Excel de calcular: contenido, 409 y que nada se escribe
# ---------------------------------------------------------------------------------------------


def test_el_excel_trae_la_auditoria_y_el_total_del_presupuesto(cliente, items_payload):
    respuesta = cliente.post("/presupuestos/calcular/excel", json=_calculo(items_payload))

    libro = openpyxl.load_workbook(io.BytesIO(respuesta.content))
    assert libro.sheetnames == ["Presupuesto", "APU", "Curva", "Auditoria"]
    valores = [celda for fila in libro["Presupuesto"].iter_rows(values_only=True) for celda in fila]
    assert float(linea_base.TOTAL_PRESUPUESTO_AUDITADO) in valores


def test_el_excel_sin_lista_vigente_es_409(cliente, items_payload):
    respuesta = cliente.post(
        "/presupuestos/calcular/excel", json=_calculo(items_payload, fecha="2026-01-01")
    )

    assert respuesta.status_code == 409
    assert "lista de precios" in respuesta.json()["detalle"]


def test_el_excel_sin_rendimiento_estimado_es_409_con_la_propuesta(cliente, sesion, items_payload):
    dejar_sin_estimado(sesion)
    sesion.commit()

    respuesta = cliente.post(
        "/presupuestos/calcular/excel",
        json=_calculo(items_payload, fecha=FECHA_PRESUPUESTO.isoformat()),
    )

    assert respuesta.status_code == 409
    assert respuesta.json()["rendimientos_por_declarar"][0]["codigo_partida"] == "LB-02-TUB"


def test_calcular_no_escribe_listas_ni_insumos_ni_rendimientos(cliente, sesion, items_payload):
    listas = cliente.get("/listas-precios").json()
    insumos = cliente.get("/insumos").json()
    rendimientos = len(sesion.scalars(select(models.Rendimiento)).all())

    cliente.post("/presupuestos/calcular", json=_calculo(items_payload))
    cliente.post("/presupuestos/calcular/excel", json=_calculo(items_payload))

    assert cliente.get("/listas-precios").json() == listas
    assert cliente.get("/insumos").json() == insumos
    assert len(sesion.scalars(select(models.Rendimiento)).all()) == rendimientos


# ---------------------------------------------------------------------------------------------
# [15] validaciones del cuerpo de calcular
# ---------------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    "cambio",
    [
        {"fecha": "28/04/2026"},
        {"items": [{"codigo_partida": "LB-01-EXC"}]},
    ],
)
def test_un_cuerpo_mal_formado_es_422(cliente, items_payload, cambio):
    cuerpo = _calculo(items_payload) | cambio

    assert cliente.post("/presupuestos/calcular", json=cuerpo).status_code == 422


def test_una_cantidad_que_no_es_decimal_es_422_con_su_partida(cliente, items_payload):
    items = [dict(i) for i in items_payload]
    items[0]["cantidad"] = "nueve,74"

    respuesta = cliente.post("/presupuestos/calcular", json=_calculo(items))

    assert respuesta.status_code == 422
    assert items[0]["codigo_partida"] in respuesta.json()["detalle"]


def test_un_dominio_desconocido_es_422(cliente, items_payload):
    items = [dict(i) for i in items_payload]
    items[0]["dominio"] = "marciano"

    assert cliente.post("/presupuestos/calcular", json=_calculo(items)).status_code == 422
