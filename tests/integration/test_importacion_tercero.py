"""RF-24 y UC-05, pasos 1 y 2 (INV-076, expediente C6.4): importar el presupuesto de un tercero.

El archivo es el XLSX de `tests/fixtures/presupuesto_tabular.py`, que deriva del presupuesto de las
siete inconsistencias con los codigos de partida del tercero (`TER-*`). Oraculos independientes de
la importacion: el presupuesto de la fixture (`presupuesto_tercero`), las siete inconsistencias
documentadas de la linea base (`INCONSISTENCIAS`) y su total exacto. Los archivos viven en
`tmp_path` y la base en memoria: nada escribe en `data/`.
"""

from __future__ import annotations

from dataclasses import replace
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import func, select

from core import models
from core.budget import (
    actualizar_precios,
    cargar_elaboracion,
    crear_proyecto,
    importar_presupuesto,
    leer_presupuesto_tabular,
)
from core.catalog import Catalogo, abrir_sesion, crear_lista_desde_archivo
from core.contracts import Dominio, ModalidadManoObra, OrigenTipo, Severidad
from core.costing import calcular_apu
from core.verification import auditar
from core.verification.directivas import CLAVE_UNIDAD_ORIGINAL
from core.verification.reglas import NO_APLICABLE
from tests.fixtures import apu_linea_base as linea_base
from tests.fixtures.auditoria import detectadas
from tests.fixtures.presupuesto_tabular import (
    NOMBRE_ARCHIVO,
    PREFIJO_LINEA_BASE,
    codigo_tercero,
    escribir_presupuesto_tabular,
    presupuesto_tercero,
)

PROYECTO = "Auditoria del drenaje de un tercero"
TUBERIA = codigo_tercero(linea_base.APU_TUBERIA.codigo_partida)
EXCAVACION = codigo_tercero(linea_base.APU_EXCAVACION.codigo_partida)
#: Las siete inconsistencias documentadas, con los codigos del tercero.
SIETE = {
    (i.regla, codigo_tercero(i.codigo_partida) if i.codigo_partida else None)
    for i in linea_base.INCONSISTENCIAS
}


def _contar(sesion, modelo) -> int:
    return sesion.scalar(select(func.count()).select_from(modelo))


def _conteos(sesion) -> tuple[int, ...]:
    return tuple(
        _contar(sesion, modelo)
        for modelo in (models.ListaPrecios, models.Partida, models.Presupuesto, models.Insumo)
    )


def _editar(ruta: Path, cambio) -> Path:
    """Aplica `cambio(libro)` al XLSX y lo guarda en el mismo sitio."""
    libro = openpyxl.load_workbook(ruta)
    cambio(libro)
    libro.save(ruta)
    return ruta


@pytest.fixture
def archivo(tmp_path):
    return escribir_presupuesto_tabular(tmp_path)


@pytest.fixture
def proyecto(sesion):
    return crear_proyecto(sesion, PROYECTO, dominio=Dominio.CIVIL)


def test_la_lectura_reproduce_partidas_unidades_apu_y_curva(archivo):
    importado = leer_presupuesto_tabular(archivo)

    esperado = presupuesto_tercero()
    assert importado.presupuesto == esperado
    assert importado.presupuesto.total == linea_base.TOTAL_EXACTO_LINEA_BASE
    assert importado.parametros == linea_base.PARAMETROS_LINEA_BASE
    assert importado.dominio is Dominio.CIVIL
    assert set(importado.composiciones) == {p.item.codigo_partida for p in esperado.partidas}
    tuberia = next(
        p.item for p in importado.presupuesto.partidas if p.item.codigo_partida == TUBERIA
    )
    assert (tuberia.unidad, tuberia.especificaciones[CLAVE_UNIDAD_ORIGINAL]) == ("m", "mts")


def test_la_importacion_audita_siete_de_siete_con_la_unidad_del_autor(sesion, archivo, proyecto):
    resultado = importar_presupuesto(sesion, archivo, proyecto)

    assert len(SIETE) == 7
    assert SIETE <= detectadas(resultado.informe)
    (r3,) = [
        h
        for h in resultado.informe.hallazgos
        if h.regla == "R3" and resultado.informe.partida_de(h) == TUBERIA
    ]
    assert "'mts'" in r3.descripcion  # criterio de RF-24: la unidad del autor llega a R3
    assert resultado.pendientes == ()


def test_guardar_y_recargar_da_el_mismo_informe(sesion, archivo, proyecto):
    # Un precio de la lista sembrada que ningun APU del tercero usa: la lista importada debe
    # copiarlo, o dejaria de valorar el catalogo completo.
    catalogo = Catalogo(sesion)
    sembrada = catalogo.lista_vigente(linea_base.FECHA_LINEA_BASE)
    ajeno = catalogo.dar_de_alta_insumo(models.TipoInsumo.MATERIAL, "Cal hidratada", "saco")
    sesion.add(models.PrecioInsumo(lista_id=sembrada.id, insumo_id=ajeno.id, precio=Decimal("7.5")))
    sesion.flush()
    precios_sembrados = {p.insumo_id: p.precio for p in sembrada.precios}

    resultado = importar_presupuesto(sesion, archivo, proyecto)
    sesion.commit()
    sesion.expunge_all()

    recargado = cargar_elaboracion(
        sesion, PROYECTO, linea_base.CODIGO_PRESUPUESTO, Catalogo(sesion)
    )

    assert recargado.presupuesto.total == linea_base.TOTAL_EXACTO_LINEA_BASE
    assert recargado.informe.hallazgos == resultado.informe.hallazgos
    assert recargado.informe.a_markdown() == resultado.informe.a_markdown()
    lista = sesion.scalars(
        select(models.ListaPrecios).where(models.ListaPrecios.origen == NOMBRE_ARCHIVO)
    ).one()
    assert lista.nombre == f"Importada: {NOMBRE_ARCHIVO}"
    assert lista.fecha_vigencia == linea_base.FECHA_LINEA_BASE
    assert precios_sembrados.items() <= {p.insumo_id: p.precio for p in lista.precios}.items()
    # La lista importada es la vigente a su fecha, y copia la sembrada: la linea base del catalogo
    # se sigue reconstruyendo igual con ella.
    catalogo = Catalogo(sesion)
    assert catalogo.lista_vigente(linea_base.FECHA_LINEA_BASE).id == lista.id
    for apu in linea_base.APUS_LINEA_BASE:
        assert catalogo.composicion(apu.codigo_partida, fecha=linea_base.FECHA_LINEA_BASE) == apu


def test_un_codigo_que_ya_esta_en_el_catalogo_se_rechaza_sin_escribir(sesion, tmp_path, proyecto):
    archivo = escribir_presupuesto_tabular(tmp_path, prefijo=PREFIJO_LINEA_BASE)
    antes = _conteos(sesion)

    with pytest.raises(ValueError, match="ya estan en el catalogo"):
        importar_presupuesto(sesion, archivo, proyecto)

    assert _conteos(sesion) == antes


@pytest.mark.parametrize(
    ("cambio", "patron"),
    [
        (lambda libro: libro.remove(libro["APU"]), rf"{NOMBRE_ARCHIVO}: falta la hoja 'APU'"),
        (
            lambda libro: libro["Presupuesto"].delete_cols(4),  # cantidad
            rf"{NOMBRE_ARCHIVO}, hoja Presupuesto: faltan las columnas cantidad",
        ),
        (
            lambda libro: libro["Presupuesto"].append(
                ["TER-99-XXX", "Partida sin APU", "m3", "1", "", "", "", ""]
            ),
            rf"{NOMBRE_ARCHIVO}, hoja Presupuesto, fila 7: .*TER-99-XXX",
        ),
        (
            lambda libro: libro["Datos"].delete_rows(3),  # fecha
            rf"{NOMBRE_ARCHIVO}, hoja Datos: falta la clave fecha",
        ),
        (
            lambda libro: libro["Insumos"].cell(row=2, column=5, value="uno"),
            rf"{NOMBRE_ARCHIVO}, hoja Insumos, fila 2: .*cantidad",
        ),
        (
            lambda libro: libro["Presupuesto"].delete_rows(2, libro["Presupuesto"].max_row),
            rf"{NOMBRE_ARCHIVO}, hoja Presupuesto: no trae renglones",
        ),
        (
            lambda libro: libro["APU"].append([EXCAVACION, "Repetida", "m3", "80"]),
            rf"{NOMBRE_ARCHIVO}, hoja APU, fila 7: la partida {EXCAVACION} ya tiene APU en la "
            "fila 2",
        ),
        (
            lambda libro: libro["Insumos"].append(
                ["TER-99-XXX", "material", "Arena", "m3", "1", "1", "", ""]
            ),
            rf"{NOMBRE_ARCHIVO}, hoja Insumos, fila \d+: la partida 'TER-99-XXX' no tiene APU",
        ),
        (
            lambda libro: libro["Presupuesto"].cell(row=2, column=7, value="{no es json"),
            rf"{NOMBRE_ARCHIVO}, hoja Presupuesto, fila 2: la columna parametros debe ser un "
            "objeto JSON",
        ),
        (
            lambda libro: libro["Datos"].cell(row=3, column=2, value="ayer"),
            rf"{NOMBRE_ARCHIVO}, hoja Datos: la fecha 'ayer' no es una fecha ISO",
        ),
        (
            lambda libro: libro["Datos"].cell(row=5, column=2, value="naval"),
            rf"{NOMBRE_ARCHIVO}, hoja Datos: dominio 'naval' desconocido",
        ),
        (
            lambda libro: libro["Insumos"].cell(row=2, column=2, value="herramienta"),
            rf"{NOMBRE_ARCHIVO}, hoja Insumos, fila 2: tipo de insumo 'herramienta' desconocido",
        ),
        (
            lambda libro: libro["Insumos"].cell(row=2, column=6, value="Infinity"),
            rf"{NOMBRE_ARCHIVO}, hoja Insumos, fila 2: precio 'Infinity' no es un numero finito",
        ),
        (
            lambda libro: libro["APU"].cell(row=2, column=4, value="0"),
            rf"{NOMBRE_ARCHIVO}, hoja APU, fila 2: .*rendimiento",
        ),
    ],
    ids=[
        "hoja",
        "columna",
        "item-sin-apu",
        "clave",
        "numero",
        "vacio",
        "apu-repetida",
        "insumo-sin-apu",
        "json",
        "fecha",
        "dominio",
        "tipo",
        "no-finito",
        "invariante",
    ],
)
def test_un_archivo_mal_formado_se_rechaza_con_archivo_hoja_y_fila(archivo, cambio, patron):
    _editar(archivo, cambio)

    with pytest.raises(ValueError, match=patron):
        leer_presupuesto_tabular(archivo)


def test_sin_curva_r2_y_r7_no_aplican_y_las_demas_reglas_si(sesion, tmp_path, proyecto):
    archivo = escribir_presupuesto_tabular(tmp_path, con_curva=False)

    resultado = importar_presupuesto(sesion, archivo, proyecto)

    assert resultado.presupuesto.curva == ()
    for regla in ("R2", "R7"):
        (declarado,) = [h for h in resultado.informe.hallazgos if h.regla == regla]
        assert declarado.severidad is Severidad.INFO
        assert NO_APLICABLE in declarado.descripcion
    sin_cierre = {(regla, partida) for regla, partida in SIETE if regla != "R2"}
    assert sin_cierre <= detectadas(resultado.informe)


def test_un_item_sin_regla_es_manual_y_r1_lo_advierte(archivo):
    def sin_regla(libro):
        hoja = libro["Presupuesto"]
        (fila,) = [f for f in hoja.iter_rows(min_row=2) if f[0].value == EXCAVACION]
        fila[5].value = None  # regla

    def advertidos(presupuesto) -> dict[tuple[str, ...], str]:
        return {
            h.origen_ids: h.descripcion
            for h in auditar(presupuesto).hallazgos
            if h.regla == "R1" and h.severidad is Severidad.ADVERTENCIA
        }

    antes = advertidos(leer_presupuesto_tabular(archivo).presupuesto)
    importado = leer_presupuesto_tabular(_editar(archivo, sin_regla))

    item = next(
        p.item for p in importado.presupuesto.partidas if p.item.codigo_partida == EXCAVACION
    )
    assert (item.origen_tipo, item.regla) == (OrigenTipo.MANUAL, None)
    despues = advertidos(importado.presupuesto)
    assert (item.origen_id,) not in antes
    assert set(despues) - set(antes) == {(item.origen_id,)}
    assert "no declara la regla" in despues[(item.origen_id,)]


def test_el_formato_es_el_del_plan():
    """QA N-QA-6 (M15): el formato es la interfaz externa de RF-24 y su oraculo es la tabla del
    plan de C6 (§2.4), escrita aqui a mano, no las constantes del modulo que se prueba."""
    from core.budget import importacion as formato

    assert (
        formato.HOJA_DATOS,
        formato.HOJA_PRESUPUESTO,
        formato.HOJA_APU,
        formato.HOJA_INSUMOS,
        formato.HOJA_CURVA,
    ) == ("Datos", "Presupuesto", "APU", "Insumos", "Curva")
    assert formato.COLUMNAS_DATOS == ("clave", "valor")
    assert formato.CLAVES_DATOS == (
        "codigo",
        "fecha",
        "moneda",
        "dominio",
        "fcas",
        "bono_alimentacion",
        "administracion",
        "utilidad",
    )
    assert formato.COLUMNAS_PRESUPUESTO == (
        "codigo_partida",
        "descripcion",
        "unidad",
        "cantidad",
        "origen_id",
        "regla",
        "parametros",
        "especificaciones",
    )
    assert formato.COLUMNAS_APU == ("codigo_partida", "descripcion", "unidad", "rendimiento")
    assert formato.COLUMNAS_INSUMOS == (
        "codigo_partida",
        "tipo",
        "insumo",
        "unidad",
        "cantidad",
        "precio",
        "depreciacion",
        "modalidad",
    )
    assert formato.COLUMNAS_CURVA == ("periodo", "monto", "acumulado")


def test_un_archivo_que_no_es_un_libro_se_rechaza(tmp_path):
    archivo = tmp_path / "presupuesto.csv"
    archivo.write_text("codigo_partida\n", encoding="utf-8")

    with pytest.raises(ValueError, match=r"presupuesto\.csv: .*se importa de un libro"):
        leer_presupuesto_tabular(archivo)


def test_sin_origen_id_el_origen_es_archivo_hoja_y_fila(archivo):
    """Req 6: un renglon sin `origen_id` ni regla (UC-05 2b) recibe `<archivo>:Presupuesto:<n>`."""
    filas = []

    def sin_origen(libro):
        hoja = libro["Presupuesto"]
        (fila,) = [f for f in hoja.iter_rows(min_row=2) if f[0].value == EXCAVACION]
        fila[4].value = None  # origen_id
        fila[5].value = None  # regla
        filas.append(fila[0].row)

    importado = leer_presupuesto_tabular(_editar(archivo, sin_origen))

    (item,) = [
        p.item for p in importado.presupuesto.partidas if p.item.codigo_partida == EXCAVACION
    ]
    assert item.origen_id == f"{NOMBRE_ARCHIVO}:Presupuesto:{filas[0]}"
    assert item.origen_tipo is OrigenTipo.MANUAL


def test_una_linea_a_destajo_se_importa_y_se_valora_a_destajo(archivo):
    """QA N-QA-4 (M9): la columna `modalidad` (D9) se lee. Oraculo: el APU de la fixture con esa
    linea a destajo, valorado por el motor con los parametros de la linea base."""
    esperado = {p.apu.codigo_partida: p.apu for p in presupuesto_tercero().partidas}
    marcada = {}

    def a_destajo(libro):
        hoja = libro["Insumos"]
        fila = next(f for f in hoja.iter_rows(min_row=2) if f[1].value == "mano_obra")
        fila[7].value = "destajo"
        marcada["codigo"], marcada["obrero"] = fila[0].value, fila[2].value

    importado = leer_presupuesto_tabular(_editar(archivo, a_destajo))

    codigo = marcada["codigo"]
    apu = esperado[codigo]
    apu_destajo = replace(
        apu,
        mano_obra=tuple(
            replace(o, modalidad=ModalidadManoObra.DESTAJO)
            if o.descripcion == marcada["obrero"]
            else o
            for o in apu.mano_obra
        ),
    )
    assert importado.composiciones[codigo] == apu_destajo
    (partida,) = [p for p in importado.presupuesto.partidas if p.item.codigo_partida == codigo]
    assert partida.resultado == calcular_apu(apu_destajo, linea_base.PARAMETROS_LINEA_BASE)
    assert partida.resultado != calcular_apu(apu, linea_base.PARAMETROS_LINEA_BASE)


def test_un_presupuesto_repetido_en_el_proyecto_se_rechaza_sin_escribir(sesion, tmp_path, proyecto):
    """QA N-QA-2 (M7): el mismo codigo de presupuesto con otras partidas (prefijo OTR-)."""
    importar_presupuesto(sesion, escribir_presupuesto_tabular(tmp_path), proyecto)
    otro = escribir_presupuesto_tabular(tmp_path, prefijo="OTR-", nombre="otro.xlsx")
    antes = _conteos(sesion)

    with pytest.raises(ValueError, match="ya tiene un presupuesto '001'"):
        importar_presupuesto(sesion, otro, proyecto)

    assert _conteos(sesion) == antes


def test_una_moneda_distinta_de_la_lista_vigente_se_rechaza_sin_escribir(sesion, archivo, proyecto):
    """Req 2 / arq N-2: la siembra esta en USD; un tercero en VES no mezcla monedas en la lista
    importada que copiaria esos precios."""
    _editar(archivo, lambda libro: libro["Datos"].cell(row=4, column=2, value="VES"))
    antes = _conteos(sesion)

    with pytest.raises(ValueError, match="no mezcla monedas"):
        importar_presupuesto(sesion, archivo, proyecto)

    assert _conteos(sesion) == antes


def test_sin_lista_vigente_la_importada_lleva_solo_los_precios_del_tercero(motor, archivo):
    """QA N-QA-4 (M10): una base sin sembrar (sin lista vigente a la fecha) importa igual; la
    lista importada lleva solo los insumos del tercero, y la recarga da el mismo total."""
    with abrir_sesion(motor) as sesion:
        proyecto = crear_proyecto(sesion, PROYECTO)

        resultado = importar_presupuesto(sesion, archivo, proyecto)
        sesion.commit()

        (lista,) = sesion.scalars(select(models.ListaPrecios)).all()
        assert lista.nombre == f"Importada: {NOMBRE_ARCHIVO}"
        assert len(lista.precios) == _contar(sesion, models.Insumo) > 0
        recargado = cargar_elaboracion(
            sesion, PROYECTO, linea_base.CODIGO_PRESUPUESTO, Catalogo(sesion)
        )
        assert recargado.presupuesto.total == resultado.presupuesto.total
        assert resultado.presupuesto.total == linea_base.TOTAL_EXACTO_LINEA_BASE


def test_importar_y_revalorar_parte_del_precio_del_tercero(sesion, tmp_path, proyecto):
    """QA H1 / M22: el presupuesto importado queda ligado a la lista «Importada: <archivo>», y
    una revaloracion (UC-02) calcula sus cambios contra los precios del tercero, no contra el
    catalogo sembrado.

    El tercero paga el cemento a 16 USD/saco (el catalogo, a 15). Oraculo a mano: el vaciado de
    concreto usa 7,5 sacos/m3 y el computo auditado 1,66 m3; cada USD de saco pesa
    7,5 x 1,15 x 1,10 x 1,66 = 15,74925 USD en el presupuesto. Importado: +1 USD sobre la linea
    base. Revalorado a 18: de 16 a 18, +2 USD sobre lo importado (no +3, que seria partir de 15).
    """
    cemento = "Cemento Portland"
    por_saco = Decimal("7.5") * Decimal("1.15") * Decimal("1.10") * Decimal("1.66")
    archivo = escribir_presupuesto_tabular(tmp_path, precios_materiales={cemento: Decimal("16")})
    resultado = importar_presupuesto(sesion, archivo, proyecto)
    sesion.commit()

    assert resultado.presupuesto.total == linea_base.TOTAL_EXACTO_LINEA_BASE + por_saco
    modelo = sesion.scalars(
        select(models.Presupuesto).where(models.Presupuesto.codigo == resultado.presupuesto.codigo)
    ).one()
    assert modelo.lista_precios.nombre == f"Importada: {NOMBRE_ARCHIVO}"
    assert modelo.lista_precios.origen == NOMBRE_ARCHIVO

    nueva = tmp_path / "lista_nueva.csv"
    nueva.write_text(f"tipo,insumo,unidad,precio\nmaterial,{cemento},saco,18\n", encoding="utf-8")
    lista = crear_lista_desde_archivo(
        sesion, nueva, "Lista de junio", "USD", date(2026, 6, 1), origen=nueva.name
    )
    revalorado, comparativo = actualizar_precios(
        sesion, PROYECTO, resultado.presupuesto.codigo, lista.lista, "002"
    )
    sesion.commit()

    assert comparativo.total_anterior == resultado.presupuesto.total
    assert comparativo.total_nuevo == linea_base.TOTAL_EXACTO_LINEA_BASE + 3 * por_saco
    assert revalorado.presupuesto.total == comparativo.total_nuevo
    # La variante de 16 USD que creo la importacion (docs/modelo_datos.md §6) es la que usa el
    # tercero; el cemento de 15 sigue en el catalogo y la lista lo sube igual, pero ninguna
    # partida suya lo usa.
    cambios = {
        (c.precio_anterior, c.precio_nuevo): c for c in sesion.scalars(select(models.CambioPrecio))
    }
    assert set(cambios) == {(Decimal("15"), Decimal("18")), (Decimal("16"), Decimal("18"))}
    del_tercero = cambios[(Decimal("16"), Decimal("18"))]
    assert del_tercero.variacion == Decimal("0.125")  # (18 - 16) / 16; contra el catalogo, 0,2
    (incidencia,) = del_tercero.incidencias
    assert incidencia.precio_unitario_nuevo - incidencia.precio_unitario_anterior == (
        Decimal("7.5") * 2 * Decimal("1.15") * Decimal("1.10")
    )
    assert not cambios[(Decimal("15"), Decimal("18"))].incidencias

    sesion.expunge_all()
    recargado = cargar_elaboracion(sesion, PROYECTO, "002", Catalogo(sesion))
    assert recargado.presupuesto.total == comparativo.total_nuevo
    assert recargado.informe.hallazgos == revalorado.informe.hallazgos
    assert recargado.informe.a_markdown() == revalorado.informe.a_markdown()
