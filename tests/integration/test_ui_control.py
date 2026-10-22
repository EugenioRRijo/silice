"""UC-19 y UC-20 desde la capa sin Streamlit (`ui.control`, spec S2 §5.5).

Registrar una valuación a mano o por XLSX, ver el control al corte, pedir las opciones y elegir una,
y las hojas del Excel; todo sobre OB-S2 y con los mismos números de la memoria.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import func, select

from aplicacion import control
from core import models
from core.budget import exportar_hoja
from core.budget.excel import (
    HOJA_APU,
    HOJA_AUDITORIA,
    HOJA_PRESUPUESTO,
    PESTANA_EQUILIBRIO,
    PESTANA_PLAN,
)
from core.budget.valor_ganado import ParametrosControl
from core.catalog import Catalogo
from core.catalog.indices import Observacion, definicion_tasa, registrar_importacion
from tests.fixtures import control_ob_s2 as ob
from tests.fixtures import reajuste_ob_s2 as rz

D = Decimal
FILAS_V01 = [
    {"Partida": "OB-A", "Cantidad": "40", "Fecha de medicion": "", "Observacion": ""},
    {"Partida": "", "Cantidad": "", "Fecha de medicion": "", "Observacion": ""},
]


def _registrar_v01(sesion, **extra):
    datos = dict(
        numero="V-01", desde="2026-11-02", hasta="2026-11-06", fecha="2026-11-09",
        inspector=ob.INSPECTOR, filas=FILAS_V01, fuente=ob.FUENTE, responsable=ob.RESPONSABLE,
    )  # fmt: skip
    datos.update(extra)
    return control.registrar_y_controlar(sesion, ob.NOMBRE, Catalogo(sesion), **datos)


def test_declarar_la_linea_base_desde_los_textos_de_la_pantalla(obra_sin_linea_base):
    sesion = obra_sin_linea_base
    control.declarar_linea_base(
        sesion, ob.NOMBRE, ob.CODIGO, desde="2026-11-02", perspectiva="costo", inicio="2026-11-02",
        dias_por_periodo="5", feriados="2026-11-04, 2026-11-05\n2026-11-06", fuente=ob.FUENTE,
        responsable=ob.RESPONSABLE,
    )  # fmt: skip
    fila = sesion.scalars(select(models.CalendarioObra)).one()
    assert (fila.dias_por_periodo, fila.feriados) == (
        5,
        '["2026-11-04", "2026-11-05", "2026-11-06"]',
    )


@pytest.mark.parametrize(
    ("campo", "valor", "mensaje"),
    [
        ("desde", "02/11/2026", "desde: '02/11/2026' no es una fecha"),
        ("dias_por_periodo", "cinco", "dias por periodo: 'cinco' no es un numero"),
        ("dias_por_periodo", "2.5", "debe ser un entero"),
        ("dias_por_periodo", "0", "1 o mas"),
        ("perspectiva", "otra", "perspectiva"),
        ("feriados", "2026-13-40", "feriado: '2026-13-40' no es una fecha"),
        ("fuente", "", "obligatoria"),
    ],
)
def test_una_linea_base_mal_declarada_se_rechaza_con_su_mensaje_y_sin_escribir(
    obra_sin_linea_base, campo, valor, mensaje
):
    datos = dict(
        desde="2026-11-02", perspectiva="costo", inicio="2026-11-02", dias_por_periodo="5",
        feriados="", fuente=ob.FUENTE, responsable=ob.RESPONSABLE,
    )  # fmt: skip
    datos[campo] = valor
    with pytest.raises(control.ControlInvalido, match=mensaje):
        control.declarar_linea_base(obra_sin_linea_base, ob.NOMBRE, ob.CODIGO, **datos)
    assert (
        obra_sin_linea_base.scalar(select(func.count()).select_from(models.LineaBaseControl)) == 0
    )


def test_registrar_una_valuacion_devuelve_el_control_al_corte_de_su_hasta(obra):
    resultado = _registrar_v01(obra)
    assert resultado.instantanea.corte == date(2026, 11, 6)
    assert resultado.indicadores.ev == D(400)
    assert resultado.reproyeccion.termino is not None
    assert resultado.informe.hallazgos == ()
    assert obra.scalar(select(func.count()).select_from(models.Valuacion)) == 1  # confirmada


def test_una_cantidad_ilegible_o_negativa_se_rechaza_con_su_fila_y_no_se_escribe(obra):
    for cantidad, mensaje in (("cuarenta", "fila 1, cantidad"), ("-3", "no puede ser negativa")):
        filas = [{**FILAS_V01[0], "Cantidad": cantidad}]
        with pytest.raises(control.ControlInvalido, match=mensaje):
            _registrar_v01(obra, filas=filas)
    with pytest.raises(control.ControlInvalido, match="fila 1: falta la partida"):
        control.lineas_desde_filas([{"Partida": "", "Cantidad": "5"}])
    assert obra.scalar(select(func.count()).select_from(models.Valuacion)) == 0


def test_una_valuacion_sin_fuente_se_rechaza_y_se_deshace(obra):
    with pytest.raises(control.ControlInvalido, match="obligatoria"):
        _registrar_v01(obra, fuente=" ")
    assert obra.scalar(select(func.count()).select_from(models.Valuacion)) == 0


def test_sin_linea_base_el_control_dice_que_hay_que_declararla(obra_sin_linea_base):
    with pytest.raises(control.ControlInvalido, match="linea base de control"):
        control.calcular_el_control(
            obra_sin_linea_base, ob.NOMBRE, Catalogo(obra_sin_linea_base), corte=ob.CORTE
        )


def test_la_tabla_de_valuaciones_muestra_la_fecha_de_registro_y_el_peor_hallazgo(obra):
    ob.registrar_v01_y_v02(obra)
    ob.registrar(obra, ob.VS)
    resultado = control.calcular_el_control(obra, ob.NOMBRE, Catalogo(obra), corte=ob.CORTE_DE_VS)
    filas = control.filas_de_valuaciones(resultado)
    assert [f["Numero"] for f in filas] == ["V-01", "V-02", "V-03"]
    assert filas[0]["Registrada"] == "09/11/2026 09:00" and filas[0]["Estado"] == "sin hallazgos"
    # V-S trae seis hallazgos, todos ERROR (memoria §10); el informe lista primero R13 y, ante un
    # empate de severidad, la columna conserva el primero.
    assert filas[2]["Estado"] == "ERROR: R13"
    assert filas[1]["Estado"] == "sin hallazgos"  # el solape se reporta en V-03, no en V-02
    assert tuple(filas[0]) == control.COLUMNAS_VALUACIONES


def test_los_indicadores_salen_en_las_dos_monedas_con_la_tasa_del_corte(obra):
    ob.registrar_v01_y_v02(obra)
    registrar_importacion(
        obra, definicion_tasa("venta"), [Observacion(ob.CORTE, D("850"))], archivo="t.xls",
        huella="a" * 64, fuente="https://www.bcv.org.ve/", fecha_descarga=date(2026, 11, 20),
        responsable="Admin", filas_leidas=1, filas_rechazadas=0,
    )  # fmt: skip
    resultado = control.calcular_el_control(
        obra, ob.NOMBRE, Catalogo(obra), precios=ob.precios_del_corte()
    )
    filas = {f["Indicador"]: f for f in control.filas_de_indicadores(obra, resultado)}
    assert filas["PV (USD)"]["Valor"] == "1500.00"
    assert filas["PV (USD)"]["En la otra moneda"] == "1275000.00 Bs"
    assert filas["SPI"]["Valor"] == "0.8000" and filas["CPI"]["Valor"] == "0.9375"
    assert filas["EAC (sistema) (USD)"]["Valor"] == "4480.00"
    assert filas["Desvío de precio (USD)"]["Valor"] == "-60.00"
    assert filas["Desvío de consumo (USD)"]["Valor"] == "-20.00"
    assert filas["Método y fuente del precio del corte"]["Valor"] == "comprobacion"


def test_sin_tasa_del_corte_la_otra_moneda_dice_no_disponible(obra):
    ob.registrar_v01_y_v02(obra)
    resultado = control.calcular_el_control(obra, ob.NOMBRE, Catalogo(obra))
    filas = {f["Indicador"]: f for f in control.filas_de_indicadores(obra, resultado)}
    assert filas["PV (USD)"]["En la otra moneda"] == "no disponible"
    assert filas["Tasa del corte"]["Valor"] == "no disponible"
    assert filas["Desvío de precio (USD)"]["Valor"] == "no disponible"  # sin precios del corte
    assert "sin precios del corte" in filas["Desvío de precio (USD)"]["Nota"]


def test_las_alertas_y_la_reproyeccion_se_muestran_siempre(obra):
    ob.registrar_v01_y_v02(obra)
    resultado = control.calcular_el_control(
        obra, ob.NOMBRE, Catalogo(obra), precios=ob.precios_del_corte()
    )
    alertas = {f["Alerta"]: f for f in control.alertas_en_filas(resultado)}
    assert {c: alertas[c]["Severidad"] for c in ("AL-1", "AL-2", "AL-3", "AL-4", "AL-5")} == {
        "AL-1": "ERROR", "AL-2": "ADVERTENCIA", "AL-3": "ERROR", "AL-4": "ERROR", "AL-5": "INFO",
    }  # fmt: skip
    assert "SUPUESTO_DECLARADO" in alertas["AL-1"]["Umbral"]
    filas = {f["Partida"]: f for f in control.filas_de_reproyeccion(resultado)}
    assert filas["OB-A"]["Indice"] == "0.9000" and filas["OB-C"]["Nota"].startswith("sin evidencia")
    assert filas["OB-C"]["Fin (dia habil)"] == "22.50"


def test_las_opciones_se_calculan_a_pedido_y_elegir_una_crea_la_version(obra):
    ob.registrar_v01_y_v02(obra)
    catalogo = Catalogo(obra)
    resultado = control.calcular_el_control(
        obra, ob.NOMBRE, catalogo, precios=ob.precios_del_corte()
    )
    todas = control.calcular_las_opciones(resultado)
    filas = {f["Opcion"]: f for f in control.filas_de_opciones(todas)}
    assert filas["Más cuadrillas"]["Termino"] == "día 17"
    assert filas["Más cuadrillas"]["Costo directo adicional"] == "120.00"
    assert filas["Horas extra"]["Termino"] == "día 21"
    assert filas["Partidas en paralelo"]["Nota"].startswith("ninguna partida sin empezar")
    assert filas["Aceptar el retraso"]["Costo indirecto"] == "90.00"
    assert "eta de cuadrillas: 0.80 (SUPUESTO_DECLARADO)" in filas["Más cuadrillas"]["Nota"]
    version = control.elegir_opcion(
        obra, resultado, todas[0], catalogo, fuente="decision ficticia", responsable="Ing."
    )
    assert version == "OB-002"
    (decision,) = control.filas_de_decisiones(obra, ob.NOMBRE)
    assert (decision["Version"], decision["Opcion"], decision["Termino"]) == (
        "OB-002", "cuadrillas", "24/11/2026",
    )  # fmt: skip


def test_elegir_sin_fuente_se_rechaza_y_no_deja_version(obra):
    ob.registrar_v01_y_v02(obra)
    catalogo = Catalogo(obra)
    resultado = control.calcular_el_control(
        obra, ob.NOMBRE, catalogo, precios=ob.precios_del_corte()
    )
    todas = control.calcular_las_opciones(resultado)
    with pytest.raises(control.ControlInvalido, match="obligatoria"):
        control.elegir_opcion(obra, resultado, todas[0], catalogo, fuente="", responsable="Ing.")
    assert obra.scalar(select(func.count()).select_from(models.Presupuesto)) == 1


def test_un_limite_legal_o_un_numero_ilegible_en_las_opciones_se_rechaza(obra):
    ob.registrar_v01_y_v02(obra)
    resultado = control.calcular_el_control(
        obra, ob.NOMBRE, Catalogo(obra), precios=ob.precios_del_corte()
    )
    with pytest.raises(control.ControlInvalido, match="objetivo"):
        control.calcular_las_opciones(resultado, objetivo="veinte")
    with pytest.raises(control.ControlInvalido, match="horas extra previas"):
        control.calcular_las_opciones(resultado, horas_previas="muchas")
    assert control.calcular_las_opciones(resultado, objetivo="25")[0].aplica is False


def test_anular_y_el_costo_real_desde_los_textos(obra):
    primera, _ = ob.registrar_v01_y_v02(obra)
    control.anular(obra, primera.id, motivo="error de medicion", responsable="Ing.")
    assert obra.scalar(select(func.count()).select_from(models.AnulacionValuacion)) == 1
    with pytest.raises(control.ControlInvalido, match="ya esta anulada"):
        control.anular(obra, primera.id, motivo="otra vez", responsable="Ing.")
    control.registrar_costo_real(
        obra, ob.NOMBRE, monto="12.50", moneda="USD", fecha="2026-11-13", fuente="factura",
        responsable="Ing.", partida="OB-A",
    )  # fmt: skip
    with pytest.raises(control.ControlInvalido, match="monto"):
        control.registrar_costo_real(
            obra,
            ob.NOMBRE,
            monto="mucho",
            moneda="USD",
            fecha="2026-11-13",
            fuente="f",
            responsable="I",
        )


def test_importar_el_xlsx_de_ob_s2_con_informe_previo(obra):
    lectura = control.leer_excel(ob.RUTA_XLSX)
    assert lectura.informe[0] == "Valuaciones: 2 valuacion(es)"
    assert control.importar_excel(obra, ob.NOMBRE, lectura, fuente="XLSX", responsable="Ing.") == 2
    with pytest.raises(control.ControlInvalido, match="ya estan registradas"):
        control.importar_excel(obra, ob.NOMBRE, lectura, fuente="XLSX", responsable="Ing.")


def test_un_archivo_malo_se_rechaza_con_control_invalido(tmp_path):
    csv = tmp_path / "x.csv"
    csv.write_text("a\n", encoding="utf-8")
    with pytest.raises(control.ControlInvalido, match="libro .xlsx"):
        control.leer_excel(csv)


def test_los_precios_del_corte_segun_el_metodo_elegido(obra):
    ob.registrar_v01_y_v02(obra)
    rz.sembrar_revaloracion(obra)
    catalogo = Catalogo(obra)
    assert control.versiones_revaloradas(obra, ob.NOMBRE) == ["OB-R01"]
    assert control.versiones_para_la_linea_base(obra, ob.NOMBRE)[-1] == "OB-001"
    ninguno = control.precios_del_corte(obra, ob.NOMBRE, catalogo, control.PrecioElegido(), "costo")
    assert ninguno is None
    comprobacion = control.precios_del_corte(
        obra, ob.NOMBRE, catalogo,
        control.PrecioElegido(control.METODOS_DE_PRECIOS[1], version="OB-R01"), "costo",
    )  # fmt: skip
    assert comprobacion.unitarios["OB-C"] == D("22.10")
    declarada = control.PrecioElegido(
        control.METODOS_DE_PRECIOS[2], km="1.10", r="0.10", fuente="S1"
    )
    escalatoria = control.precios_del_corte(obra, ob.NOMBRE, catalogo, declarada, "costo")
    assert escalatoria.unitarios["OB-A"] == D(10) * (1 + D("0.10") * D("0.90"))
    for elegido, mensaje in (
        (control.PrecioElegido(control.METODOS_DE_PRECIOS[1]), "elija la version"),
        (control.PrecioElegido(control.METODOS_DE_PRECIOS[1], version="NO-HAY"), "NO-HAY"),
        (control.PrecioElegido(control.METODOS_DE_PRECIOS[2], km="1.1"), "fuente"),
        (control.PrecioElegido(control.METODOS_DE_PRECIOS[2], km="x", fuente="S1"), "Km"),
    ):
        with pytest.raises(control.ControlInvalido, match=mensaje):
            control.precios_del_corte(obra, ob.NOMBRE, catalogo, elegido, "costo")


def test_las_tres_hojas_del_excel_de_la_hoja_conservan_las_propias(obra, tmp_path):
    ob.registrar_v01_y_v02(obra)
    resultado = ob.elaborar_ob_s2()
    # Sin decisiones ni valuaciones raras: las tres hojas, con sus encabezados.
    tablas = control.tablas_de_control(
        obra, ob.NOMBRE, Catalogo(obra), precios=ob.precios_del_corte()
    )
    assert [t.titulo for t in tablas] == ["Valuaciones", "Valor ganado", "Decisiones"]
    ruta = exportar_hoja(
        resultado.presupuesto, resultado.informe, tmp_path / "hoja.xlsx", ob.PARAMETROS,
        hojas_extra=tablas,
    )  # fmt: skip
    libro = openpyxl.load_workbook(ruta)
    propias = {HOJA_PRESUPUESTO, HOJA_APU, PESTANA_PLAN, PESTANA_EQUILIBRIO, HOJA_AUDITORIA}
    assert propias <= set(libro.sheetnames)
    assert {"Valuaciones", "Valor ganado", "Decisiones"} <= set(libro.sheetnames)
    assert [c.value for c in libro["Valuaciones"][1]][:3] == ["Numero", "Desde", "Hasta"]
    assert libro["Valuaciones"].max_row == 4  # encabezado + A 40, A 50, B 60


def test_sin_linea_base_no_hay_hojas_extra(obra_sin_linea_base):
    sesion = obra_sin_linea_base
    assert control.tablas_de_control(sesion, ob.NOMBRE, Catalogo(sesion)) == []


def test_con_linea_base_y_sin_valuaciones_no_hay_hojas_extra(obra):
    assert control.tablas_de_control(obra, ob.NOMBRE, Catalogo(obra)) == []


def test_un_inspector_con_formula_viaja_como_texto_al_excel(obra, tmp_path):
    ob.registrar(obra, {**ob.V01, "inspector": '=HYPERLINK("http://127.0.0.1:9/robo")'})
    tablas = control.tablas_de_control(obra, ob.NOMBRE, Catalogo(obra))
    resultado = ob.elaborar_ob_s2()
    ruta = exportar_hoja(
        resultado.presupuesto, resultado.informe, tmp_path / "hoja.xlsx", ob.PARAMETROS,
        hojas_extra=tablas,
    )  # fmt: skip
    celda = openpyxl.load_workbook(ruta)["Valuaciones"]["E2"]
    assert celda.data_type == "s" and celda.value.startswith("=HYPERLINK")


# --- Agregadas en la auditoría final (2026-10-13): parámetros de control y curva PV/EV/AC --------


def _filas_con(cambios: dict[str, str]) -> list[dict[str, str]]:
    filas = control.filas_de_parametros(ParametrosControl())
    return [{**f, "Valor": cambios.get(f["Parametro"], f["Valor"])} for f in filas]


def test_los_parametros_por_defecto_ida_y_vuelta_por_las_filas():
    filas = control.filas_de_parametros(ParametrosControl())
    assert tuple(filas[0]) == control.COLUMNAS_PARAMETROS
    por_campo = {f["Parametro"]: f for f in filas}
    # Valores por defecto declarados en el spec S2 §4.5 y §4.6 (SUPUESTO_DECLARADO).
    assert por_campo["spi_error"]["Valor"] == "0.90"
    assert por_campo["cuadrillas_max"]["Valor"] == "3"
    assert por_campo["fcas_en_horas_extra"]["Valor"] == "sí"
    assert por_campo["penalizacion_diaria"]["Naturaleza"] == control.NATURALEZA_DATO
    assert control.parametros_desde_filas(filas) == ParametrosControl()


@pytest.mark.parametrize(
    ("campo", "valor", "acepta"),
    [
        ("eta_cuadrilla", "1", True),  # frontera superior de (0, 1]
        ("eta_cuadrilla", "1.01", False),
        ("eta_cuadrilla", "0", False),
        ("cuadrillas_max", "2", True),  # mínimo 2
        ("cuadrillas_max", "1", False),
        ("cuadrillas_max", "2.5", False),
        ("fcas_en_horas_extra", "NO", True),
        ("fcas_en_horas_extra", "quizas", False),
        ("spi_error", "noventa", False),
    ],
)
def test_un_parametro_fuera_de_rango_o_ilegible_se_rechaza(campo, valor, acepta):
    filas = _filas_con({campo: valor})
    if acepta:
        leidos = control.parametros_desde_filas(filas)
        assert getattr(leidos, campo) != getattr(ParametrosControl(), campo)
    else:
        with pytest.raises(control.ControlInvalido, match=campo):
            control.parametros_desde_filas(filas)


def test_falta_un_parametro_se_rechaza_con_su_nombre():
    filas = [f for f in _filas_con({}) if f["Parametro"] != "precio_info"]
    with pytest.raises(control.ControlInvalido, match="falta el parametro precio_info"):
        control.parametros_desde_filas(filas)


def test_el_umbral_de_al_1_en_la_frontera_exacta_del_spi(obra):
    """SPI = 0,8000 (memoria §3). AL-1 es ERROR si SPI < spi_error: con 0,80 pasa a ADVERTENCIA
    (0,80 no es menor que 0,80) y con 0,8001 vuelve a ERROR."""
    ob.registrar_v01_y_v02(obra)

    def severidad_al_1(umbral: str) -> str:
        cambiados = control.guardar_los_parametros(
            obra, ob.NOMBRE, _filas_con({"spi_error": umbral}), fuente="acta", responsable="Ing."
        )
        assert cambiados == 1
        resultado = control.calcular_el_control(obra, ob.NOMBRE, Catalogo(obra), corte=ob.CORTE)
        return {f["Alerta"]: f for f in control.alertas_en_filas(resultado)}["AL-1"]["Severidad"]

    assert severidad_al_1("0.80") == "ADVERTENCIA"
    assert severidad_al_1("0.8001") == "ERROR"
    assert control.parametros_actuales(obra, ob.NOMBRE).spi_error == D("0.8001")
    registros = obra.scalars(
        select(models.RegistroCambio).where(models.RegistroCambio.campo == "spi_error")
    ).all()
    assert len(registros) == 2


def test_guardar_parametros_invalidos_no_escribe(obra):
    with pytest.raises(control.ControlInvalido, match="eta_cuadrilla"):
        control.guardar_los_parametros(
            obra, ob.NOMBRE, _filas_con({"eta_cuadrilla": "2"}), fuente="acta", responsable="I"
        )
    assert control.parametros_actuales(obra, ob.NOMBRE) == ParametrosControl()


def test_la_curva_pv_ev_ac_cierra_cada_semana_hasta_el_corte(obra):
    """A mano (memoria §1 a §3, perspectiva costo): al cierre de la semana 1 (2026-11-06) PV =
    100 × 10 × 0,5 = 500, EV = 40 × 10 = 400 y AC = 0 (el costo real es del 11-13); al corte
    (semana 2) PV 1 500, EV 1 200, AC 1 280."""
    ob.registrar_v01_y_v02(obra)
    resultado = control.calcular_el_control(obra, ob.NOMBRE, Catalogo(obra), corte=ob.CORTE)
    curva = control.curva_pv_ev_ac(obra, resultado)
    esperada = [("Semana 1", D(500), D(400), D(0)), ("Semana 2", D(1500), D(1200), D(1280))]
    assert len(curva) == len(esperada)
    for punto, (periodo, pv, ev, ac) in zip(curva, esperada, strict=True):
        assert punto["Periodo"] == periodo
        for serie, valor in (("PV", pv), ("EV", ev), ("AC", ac)):
            assert abs(punto[serie] - valor) <= D("0.005"), (periodo, serie)
