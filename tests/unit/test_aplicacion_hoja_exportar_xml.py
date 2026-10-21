"""Los caracteres que un .xlsx no puede llevar, saneados a la salida del Excel de la hoja (X1).

Eugenio pidio una recomendacion viable (ola de decisiones, 2026-10-21): con U+FFFE o U+FFFF en el
nombre de un ambiente, E1 lo acepta y X1 entregaba un libro que ni openpyxl podia leer; con un
caracter de control, 422. Se medio con openpyxl (ver la bitacora): el conjunto que rompe un libro es
el que XML 1.0 prohibe, a saber

* los controles C0 salvo tabulador, salto de linea y retorno (U+0000 a U+0008, U+000B, U+000C,
  U+000E a U+001F): openpyxl se niega a escribirlos (`IllegalCharacterError`);
* los sustitutos sueltos (U+D800 a U+DFFF): el libro se escribe pero no se puede leer;
* U+FFFE y U+FFFF: igual.

Todo lo demas se lee bien (U+007F, U+0080 a U+009F, U+FDD0, U+FFFD, U+1FFFE, U+10FFFF, emoji). La
recomendacion es SUSTITUIR cada uno por U+FFFD (el caracter de sustitucion): conserva la posicion y
deja a la vista que ahi habia algo, en vez de borrarlo en silencio o impedir exportar.

Oraculos a mano: los puntos de codigo de arriba; ninguno sale del codigo bajo prueba.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from decimal import Decimal
from enum import StrEnum

import pytest
from openpyxl import load_workbook

from aplicacion import hoja_exportar

SUSTITUTO = chr(0xFFFD)
PROHIBIDOS = [
    *range(0x00, 0x09),
    0x0B,
    0x0C,
    *range(0x0E, 0x20),
    0xD800,
    0xDBFF,
    0xDC00,
    0xDFFF,
    0xFFFE,
    0xFFFF,
]
ADMITIDOS = [
    0x09,
    0x0A,
    0x0D,
    0x20,
    0x7F,
    0x80,
    0x85,
    0x9F,
    0xA0,
    0xD7FF,
    0xE000,
    0xFDD0,
    0xFFFC,
    0xFFFD,
    0x10000,
    0x1F600,
    0x1FFFE,
    0x10FFFF,
]


@pytest.mark.parametrize("punto", PROHIBIDOS, ids=lambda p: f"U+{p:04X}")
def test_cada_caracter_prohibido_se_sustituye_por_u_fffd(punto: int) -> None:
    assert hoja_exportar.sanear_para_xlsx(f"a{chr(punto)}b") == f"a{SUSTITUTO}b"


@pytest.mark.parametrize("punto", ADMITIDOS, ids=lambda p: f"U+{p:04X}")
def test_los_caracteres_que_un_xlsx_admite_no_se_tocan(punto: int) -> None:
    texto = f"a{chr(punto)}b"
    assert hoja_exportar.sanear_para_xlsx(texto) == texto


def test_varios_prohibidos_juntos_se_sustituyen_uno_a_uno() -> None:
    entrada = chr(0) + chr(0xFFFE) + chr(0xFFFF) + "x" + chr(0xD800)
    assert hoja_exportar.sanear_para_xlsx(entrada) == SUSTITUTO * 3 + "x" + SUSTITUTO


@dataclass(frozen=True, slots=True)
class _Fila:
    nombre: str
    cantidad: Decimal
    notas: tuple[str, ...]
    extra: dict[str, str]


class _Tipo(StrEnum):
    UNO = "uno"


def test_el_saneado_llega_a_los_textos_anidados_y_no_toca_lo_que_no_es_texto() -> None:
    sucio = "Ba" + chr(0xFFFE) + "ño"
    limpio = "Ba" + SUSTITUTO + "ño"
    fila = _Fila(
        sucio, Decimal("1.50"), ("ok", "x" + chr(1) + "y"), {"k" + chr(2): "v" + chr(0xFFFF)}
    )
    resultado = hoja_exportar.sanear_para_xlsx([fila, _Tipo.UNO, 7, None, Decimal("2")])
    assert resultado == [
        _Fila(
            limpio,
            Decimal("1.50"),
            ("ok", "x" + SUSTITUTO + "y"),
            {"k" + SUSTITUTO: "v" + SUSTITUTO},
        ),
        _Tipo.UNO,
        7,
        None,
        Decimal("2"),
    ]
    assert fila.nombre == sucio  # el original (el borrador vivo) no cambia


def test_si_no_hay_nada_que_sanear_se_devuelve_el_mismo_objeto() -> None:
    fila = _Fila("Baño", Decimal("1"), ("a",), {"b": "c"})
    contenedor = [fila, (fila, "x")]
    assert hoja_exportar.sanear_para_xlsx(contenedor) is contenedor
    assert hoja_exportar.sanear_para_xlsx(fila) is fila


@pytest.mark.parametrize("punto", PROHIBIDOS[:: len(PROHIBIDOS) // 6], ids=lambda p: f"U+{p:04X}")
def test_un_libro_con_el_texto_saneado_se_escribe_y_se_lee(punto: int) -> None:
    from openpyxl import Workbook

    libro = Workbook()
    libro.active["A1"] = hoja_exportar.sanear_para_xlsx(f"a{chr(punto)}b")
    salida = io.BytesIO()
    libro.save(salida)
    assert load_workbook(io.BytesIO(salida.getvalue())).active["A1"].value == f"a{SUSTITUTO}b"
