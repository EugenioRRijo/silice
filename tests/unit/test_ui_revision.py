"""La vista de quien revisa, sin Streamlit: reauditoría, diferencias, quién cambió qué y Excel
(UC-23, RF-86, RF-87; spec S3 §9 y §13.2).

Oráculos: la memoria del ejemplo mínimo (`tests/fixtures/canonico/ejemplo_v1.py`: PU 41,839875;
total 836,7975), la memoria de diferencias (`diferencias_s3.md` §2 y §4: 33 R8 de OF-001) y
aritmética hecha a mano en cada prueba. Los paquetes salen de `tests/fixtures/paquetes.py`, cuyas
huellas se calculan con `hashlib`, no con el núcleo."""

from __future__ import annotations

import copy
import io
import threading
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from adapters.espacio.serie import espacio_a_objeto
from auditoria import reauditoria
from core.budget.canonico import json_canonico
from core.contracts.verificacion import Severidad
from paquete.formato import PaqueteRechazado
from paquete.verificar import EmpresaConocida, verificar
from seguridad import firma
from tests.fixtures.canonico.ejemplo_v1 import EJEMPLO_V1
from tests.fixtures.espacio_of001 import espacio_of001
from tests.fixtures.paquetes import EMPRESA, paquete
from ui import revision
from ui.revision import ReferenciaDeMercado, Servicios

PAR = firma.generar_par()
RECIBIDO = datetime(2026, 11, 12, 10, 0)
PU = Decimal("41.83987500")  # memoria del ejemplo mínimo
TOTAL = Decimal("836.79750000")  # 20 x 41,839875


def doc_json(**cambios) -> bytes:
    objeto = copy.deepcopy(EJEMPLO_V1)
    for ruta, valor in cambios.items():
        destino = objeto
        *camino, ultimo = ruta.split("__")
        for paso in camino:
            destino = destino[int(paso)] if paso.isdigit() else destino[paso]
        destino[int(ultimo) if ultimo.isdigit() else ultimo] = valor
    return json_canonico(objeto)


def verificado(presupuesto: bytes, **mas):
    datos = paquete(par=PAR, presupuesto=presupuesto, **mas)
    return verificar(
        datos,
        recibido_en=RECIBIDO,
        buscar_empresa=lambda h: EmpresaConocida(EMPRESA, PAR.publica, "activa"),
    )


def test_el_paquete_de_prueba_llega_verificado() -> None:
    assert verificado(doc_json(), anexos={}).estado == "verificado"


def test_la_reauditoria_recalcula_los_precios_en_vez_de_creer_los_declarados() -> None:
    mentira = doc_json(renglones__0__precio_unitario="99.00", renglones__0__total="1980.00")
    r = revision.reauditar(verificado(mentira, anexos={}), Servicios())
    assert r.total_recalculado == TOTAL
    campos = {d.campo: d for d in r.diferencias}
    assert set(campos) == {"precio_unitario", "total"}
    assert campos["precio_unitario"].declarado == Decimal("99.00")
    assert campos["precio_unitario"].recalculado == PU
    # (99,00 - 41,839875) x 20 = 1143,2025 ; 1980,00 - 836,7975 = 1143,2025 (a mano)
    assert campos["precio_unitario"].monto == Decimal("1143.2025")
    assert campos["total"].monto == Decimal("1143.2025")


def test_un_paquete_honesto_no_tiene_diferencias_de_precio() -> None:
    r = revision.reauditar(verificado(doc_json(), anexos={}), Servicios())
    assert r.diferencias == () and r.contextos == ("R1-R7", "R11-R12") and not r.no_termino
    assert r.fecha_evaluacion == date(2026, 11, 10)  # la de emisión del paquete, no el reloj


def test_sin_referencia_ni_tasas_solo_queda_el_aviso_de_no_contrastables() -> None:
    """Conjunto exacto: sin referencia MaPreX, R11 da un único INFO con los tres insumos."""
    r = revision.reauditar(verificado(doc_json(), anexos={}), Servicios())
    assert [(h.regla, h.severidad) for h in r.nuevos] == [("R11", Severidad.INFO)]
    assert "3 insumos no contrastables" in r.nuevos[0].descripcion
    assert r.ya_no == () and r.comunes == ()


def test_los_hallazgos_se_separan_en_no_declarados_ya_no_y_comunes() -> None:
    declarado = {
        "regla": "R11",
        "severidad": 1,
        "descripcion": "algo viejo",
        "impacto": None,
        "origen_ids": [],
        "valor_observado": None,
        "valor_esperado": None,
    }
    r = revision.reauditar(verificado(doc_json(hallazgos=[declarado]), anexos={}), Servicios())
    assert [h.descripcion for h in r.ya_no] == ["algo viejo"]
    assert len(r.nuevos) == 1 and r.nuevos[0].regla == "R11" and r.comunes == ()


def test_un_hallazgo_declarado_igual_al_recalculado_es_comun() -> None:
    propio = revision.reauditar(verificado(doc_json(), anexos={}), Servicios()).nuevos[0]
    declarado = {
        "regla": propio.regla,
        "severidad": int(propio.severidad),
        "descripcion": propio.descripcion,
        "impacto": None,
        "origen_ids": list(propio.origen_ids),
        "valor_observado": format(propio.valor_observado, "f"),
        "valor_esperado": None,
    }
    r = revision.reauditar(verificado(doc_json(hallazgos=[declarado]), anexos={}), Servicios())
    assert len(r.comunes) == 1 and r.nuevos == () and r.ya_no == ()


def test_dos_declarados_identicos_cuentan_dos_veces() -> None:
    """Multiconjunto (nota 3 de la tarea 19): el repetido no se colapsa."""
    declarado = {
        "regla": "R9",
        "severidad": 1,
        "descripcion": "repetido",
        "impacto": None,
        "origen_ids": [],
        "valor_observado": None,
        "valor_esperado": None,
    }
    r = revision.reauditar(
        verificado(doc_json(hallazgos=[declarado, declarado]), anexos={}), Servicios()
    )
    assert [h.descripcion for h in r.ya_no] == ["repetido", "repetido"]


def _con_referencia(precio_arena: str, referencia: str = "10.00") -> list:
    presupuesto = doc_json(
        composiciones__0__lineas__0__referencia="MP-0042",
        composiciones__0__lineas__0__precio=precio_arena,
    )
    servicios = Servicios(
        referencia_de=lambda c: ReferenciaDeMercado("MP-0042", Decimal(referencia), "USD", None)
    )
    r = revision.reauditar(verificado(presupuesto, anexos={}), servicios)
    return [h for h in r.nuevos if h.regla == "R11" and h.severidad is Severidad.ADVERTENCIA]


def test_r11_marca_el_precio_fuera_de_rango_contra_la_referencia_de_quien_revisa() -> None:
    [h] = _con_referencia("14.00")  # 14 / 10 - 1 = +40 %
    assert "+40.00 %" in h.descripcion and "MP-0042" in h.descripcion


def test_r11_fronteras_del_rango_de_menos_20_a_mas_30() -> None:
    """S1-R5: [-20 %, +30 %] pasa; 13,00 (+30 %) y 8,00 (-20 %) no marcan; 13,01 y 7,99 sí."""
    assert _con_referencia("13.00") == []
    assert len(_con_referencia("13.01")) == 1
    assert _con_referencia("8.00") == []
    assert len(_con_referencia("7.99")) == 1


def test_r12_se_evalua_a_la_fecha_de_emision_y_no_cambia_de_un_dia_a_otro() -> None:
    antiguo = doc_json(
        composiciones__0__lineas__0__precio_origen__fecha_precio="2026-09-01",
        composiciones__0__lineas__1__precio_origen__fecha_precio="2026-09-01",
        composiciones__0__lineas__2__precio_origen__fecha_precio="2026-09-01",
    )
    v = verificado(antiguo, anexos={})
    primera = revision.reauditar(v, Servicios())
    segunda = revision.reauditar(v, Servicios())
    r12 = [h for h in primera.nuevos if h.regla == "R12"]
    # del 01/09/2026 al 10/11/2026 (emisión): 29 + 31 + 10 = 70 días
    assert len(r12) == 3 and all("70 dias al 10/11/2026" in h.descripcion for h in r12)
    otras = [h.descripcion for h in segunda.nuevos if h.regla == "R12"]
    assert [h.descripcion for h in r12] == otras


def test_con_el_anexo_de_espacio_corren_r8_a_r10_r15_y_r17() -> None:
    espacio = json_canonico(espacio_a_objeto(espacio_of001()))
    r = revision.reauditar(
        verificado(doc_json(), anexos={"anexos/espacio.json": espacio}), Servicios()
    )
    assert r.contextos == ("R1-R7", "R8-R10", "R15", "R17", "R11-R12")
    assert sum(1 for h in r.nuevos if h.regla == "R8") == 33  # el conjunto exacto de OF-001
    assert not [h for h in r.nuevos if h.regla in {"R9", "R10", "R15", "R17"}]


def test_con_el_anexo_de_control_y_el_servicio_de_s2_corren_r13_y_r14() -> None:
    llamadas = []

    def reglas_de_control(doc, anexo, corte):
        llamadas.append((doc.codigo, anexo["formato"], corte))
        return ()

    control = b'{"formato":"silice-anexo","nombre":"control","entradas":[]}'
    r = revision.reauditar(
        verificado(doc_json(), anexos={"anexos/control.json": control}),
        Servicios(reglas_de_control=reglas_de_control),
    )
    assert "R13-R14" in r.contextos and llamadas == [("001", "silice-anexo", date(2026, 11, 10))]
    assert r.avisos == ()


def test_sin_el_servicio_de_s2_r13_y_r14_no_corren_y_se_avisa() -> None:
    """RP8-21 y nota 4 de la tarea 19: el puente de S2 no existe; la reauditoría lo declara."""
    control = b'{"formato":"silice-anexo","nombre":"control","entradas":[]}'
    sin = revision.reauditar(
        verificado(doc_json(), anexos={"anexos/control.json": control}), Servicios()
    )
    assert "R13-R14" not in sin.contextos
    assert len(sin.avisos) == 1 and "R13 y R14" in sin.avisos[0]


def test_la_reauditoria_se_corta_a_los_60_segundos_y_lo_dice() -> None:
    suelta = threading.Event()
    servicios = Servicios(tasa_en=lambda fecha: (suelta.wait(10), None)[1])
    try:
        r = revision.reauditar(verificado(doc_json(), anexos={}), servicios, tiempo_maximo=0.05)
    finally:
        suelta.set()
    assert r.no_termino and r.resultado is None and "no terminó" in r.avisos[0]
    assert revision.LIMITE_SEGUNDOS == 60


def test_un_documento_con_la_forma_equivocada_es_rechazado_sin_traza() -> None:
    """Nota 2 de la tarea 19: un `KeyError` dentro de la reauditoría (un renglón sin su
    composición) sale como `PaqueteRechazado` con mensaje, no como traza."""
    objeto = copy.deepcopy(EJEMPLO_V1)
    objeto["renglones"][0]["item"]["codigo_partida"] = "NO-EXISTE"
    with pytest.raises(PaqueteRechazado, match="no tiene la forma esperada"):
        revision.reauditar(verificado(json_canonico(objeto), anexos={}), Servicios())


def test_una_cifra_que_falta_es_un_error_de_valor_con_mensaje() -> None:
    """Sin precio en una línea, el lector del documento lo dice (`DocumentoInvalido`, un
    `ValueError` que la pantalla y la API atrapan); tampoco hay traza."""
    objeto = copy.deepcopy(EJEMPLO_V1)
    del objeto["composiciones"][0]["lineas"][0]["precio"]
    with pytest.raises(ValueError, match="linea 0.precio"):
        revision.reauditar(verificado(json_canonico(objeto), anexos={}), Servicios())


def test_un_paquete_con_mas_filas_que_el_tope_se_rechaza_y_con_el_tope_justo_no(
    monkeypatch,
) -> None:
    archivos = verificado(doc_json(), anexos={}).archivos
    monkeypatch.setattr(reauditoria, "LIMITE_FILAS", 4)  # 1 renglón + 3 líneas = 4: pasa
    assert revision.leer_documento(archivos).codigo == "001"
    monkeypatch.setattr(reauditoria, "LIMITE_FILAS", 3)  # 4 > 3
    with pytest.raises(PaqueteRechazado, match="renglones y líneas"):
        revision.leer_documento(archivos)


# --- diferencias ---------------------------------------------------------------------------


def test_diferencia_contra_la_entrega_anterior_y_contra_la_base_del_revisor() -> None:
    actual = revision.leer_documento(
        verificado(
            doc_json(renglones__0__item__cantidad="25", renglones__0__cantidad="25"), anexos={}
        ).archivos
    )
    anterior = revision.leer_documento(verificado(doc_json(), anexos={}).archivos)
    d = revision.contra_la_entrega_anterior(actual, anterior)
    # (25 - 20) x 41,839875 = 209,199375 por cantidad; nada por precio
    assert (d.desvio_cantidad, d.desvio_precio) == (Decimal("209.199375"), Decimal(0))
    base = [revision.PrecioDeLista("MAT-001", "", Decimal("16.00"))]
    diferencia, revalorado = revision.contra_la_base_del_revisor(anterior, base)
    # arena +2,00 -> x 1,15 x 1,10 = 2,53 por m3; x 20 = 50,60 (a mano)
    assert diferencia.desvio_cantidad == 0 and diferencia.desvio_precio == Decimal("50.60")
    assert {n.insumo for n in revalorado.no_contrastables} == {"EQU-001", "MO-001"}


# --- quién cambió qué -----------------------------------------------------------------------


def test_los_cambios_nuevos_son_los_de_las_versiones_que_la_entrega_anterior_no_traia() -> None:
    v = verificado(doc_json(), anexos={})
    cadena = v.cadena
    previas = [cadena["eslabones"][0]["huella_eslabon"]]
    [cambio] = revision.cambios_nuevos(cadena, previas)
    assert (cambio["version"], cambio["responsable"], cambio["despues"], cambio["marca"]) == (
        "005",
        "Ana <ana>",
        "1.50",
        "",
    )
    assert revision.cambios_nuevos(cadena, [e["huella_eslabon"] for e in cadena["eslabones"]]) == []
    assert [c["version"] for c in revision.cambios_nuevos(cadena, [])] == ["004", "005"]


def test_un_cambio_hecho_tras_un_restablecimiento_se_marca() -> None:
    cadena = copy.deepcopy(verificado(doc_json(), anexos={}).cadena)
    cadena["cambios"]["005"][0][6] = "Luis Pérez <luis> (tras restablecimiento)"
    primero, segundo = revision.cambios_nuevos(cadena, [])
    assert primero["marca"] == ""
    assert segundo["marca"] == "hecho tras un restablecimiento del administrador"


def test_la_declaracion_retroactiva_sale_de_los_eslabones_del_paquete() -> None:
    cadena = copy.deepcopy(verificado(doc_json(), anexos={}).cadena)
    assert revision.declaracion_retroactiva(cadena) is None  # todos los eslabones son normales
    for e in cadena["eslabones"]:
        e["tipo"] = "retroactivo"
    assert revision.declaracion_retroactiva(cadena) == (
        "Versiones 004 a 005 selladas el 2026-11-01: prueban que no cambiaron desde esa fecha; "
        "antes no hay garantía."
    )
    cadena["eslabones"] = cadena["eslabones"][:1]
    assert revision.declaracion_retroactiva(cadena).startswith("Versión 004 sellada")


def test_el_login_se_saca_del_responsable() -> None:
    assert revision.login_de("Ana Pérez <ana>") == "ana"
    assert revision.login_de("Luis <luis> (tras restablecimiento)") == "luis"
    assert revision.login_de("script:seed_demo") is None
    assert revision.login_de("") is None


# --- quien revisa es quien armó (P2 ratificada) ---------------------------------------------


HUELLA_PROPIA = firma.huella_completa(PAR.publica)


def test_revisor_igual_al_usuario_emisor_de_un_paquete_de_la_empresa_propia_lo_avisa() -> None:
    v = verificado(doc_json(), anexos={})
    donde = revision.revisor_armo(v, login_revisor="ana", empresa_propia=(EMPRESA, HUELLA_PROPIA))
    assert donde == ["004", "005"]


def test_revisor_igual_al_usuario_de_un_eslabon_anterior_tambien_lo_avisa() -> None:
    v = verificado(doc_json(), anexos={})
    v.cadena["eslabones"][0]["usuario"] = "Luis <luis>"
    donde = revision.revisor_armo(v, login_revisor="luis", empresa_propia=(EMPRESA, HUELLA_PROPIA))
    assert donde == ["004"]


def test_otro_revisor_de_la_misma_empresa_no_ve_el_aviso() -> None:
    v = verificado(doc_json(), anexos={})
    donde = revision.revisor_armo(v, login_revisor="pedro", empresa_propia=(EMPRESA, HUELLA_PROPIA))
    assert donde == []


def test_un_paquete_de_una_empresa_ajena_con_un_login_igual_no_avisa() -> None:
    v = verificado(doc_json(), anexos={})
    assert revision.revisor_armo(v, login_revisor="ana", empresa_propia=("Otra", "0" * 64)) == []
    # mismo nombre de empresa pero otra clave: firmado por otra, no es la propia
    assert revision.revisor_armo(v, login_revisor="ana", empresa_propia=(EMPRESA, "0" * 64)) == []


def test_un_paquete_sin_firma_con_el_nombre_de_la_empresa_propia_cuenta_como_propio() -> None:
    sin_firma = verificar(
        paquete(par=None, presupuesto=doc_json(), anexos={}),
        buscar_empresa=lambda h: None,
        recibido_en=RECIBIDO,
    )
    assert revision.revisor_armo(sin_firma, login_revisor="ana", empresa_propia=(EMPRESA, None))
    assert (
        revision.revisor_armo(sin_firma, login_revisor="ana", empresa_propia=("Otra", None)) == []
    )


# --- Excel ---------------------------------------------------------------------------------


def _libro(datos: bytes):
    return openpyxl.load_workbook(io.BytesIO(datos))


def test_el_excel_se_regenera_desde_el_json_y_lleva_las_hojas_de_la_revision() -> None:
    v = verificado(doc_json(), anexos={})
    r = revision.reauditar(v, Servicios())
    sin_extras = _libro(revision.excel_de_presupuesto(r))
    assert "Presupuesto" in sin_extras.sheetnames and "Verificación" not in sin_extras.sheetnames
    tablas = revision.tablas_de_revision(v, r, None, revision.cambios_nuevos(v.cadena, []))
    con = _libro(revision.excel_de_revision(v, r, None, []))
    assert {"Verificación", "Hallazgos", "Diferencias", "Cambios"} <= set(con.sheetnames)
    assert [t.titulo for t in tablas] == ["Verificación", "Hallazgos", "Diferencias", "Cambios"]
    valores = {fila[0].value: fila[1].value for fila in con["Verificación"].iter_rows(min_row=2)}
    assert valores["Estado"] == "verificado" and valores["SHA-256"] == v.sha256


def test_una_formula_en_el_texto_de_un_paquete_no_se_ejecuta_al_exportar() -> None:
    maligna = '=HYPERLINK("http://127.0.0.1:9/robo","clic")'
    v = verificado(doc_json(renglones__0__item__descripcion=maligna), anexos={})
    r = revision.reauditar(v, Servicios())
    libro = _libro(revision.excel_de_presupuesto(r))
    celdas = [
        c
        for hoja in libro
        for fila in hoja.iter_rows()
        for c in fila
        if isinstance(c.value, str) and c.value.startswith("=HYPERLINK")
    ]
    assert celdas and all(c.data_type == "s" and c.quotePrefix for c in celdas)


def test_sin_reauditoria_no_hay_excel() -> None:
    v = verificado(doc_json(), anexos={})
    no_termino = revision.Reauditoria(
        documento=revision.leer_documento(v.archivos),
        resultado=None,
        contextos=(),
        no_termino=True,
    )
    with pytest.raises(ValueError, match="no terminó"):
        revision.excel_de_presupuesto(no_termino)


# --- presupuesto externo en Excel (23d) ------------------------------------------------------


def test_la_plantilla_descargada_se_audita_y_no_toca_ninguna_base(tmp_path: Path) -> None:
    ruta = tmp_path / "plantilla.xlsx"
    ruta.write_bytes(revision.plantilla_uc05())
    importado, informe = revision.auditar_excel_externo(ruta)
    assert importado.presupuesto.codigo == "EJ-001"
    assert importado.presupuesto.total == TOTAL  # la misma memoria del ejemplo mínimo
    assert informe.codigo_presupuesto == "EJ-001"
    assert sorted(p.name for p in tmp_path.iterdir()) == ["plantilla.xlsx"]


def test_un_excel_con_formula_se_lee_por_su_valor_nunca_por_la_formula(tmp_path: Path) -> None:
    ruta = tmp_path / "externo.xlsx"
    ruta.write_bytes(revision.plantilla_uc05())
    libro = openpyxl.load_workbook(ruta)
    libro["Presupuesto"]["B2"] = '=HYPERLINK("http://127.0.0.1:9/robo","clic")'
    libro.save(ruta)
    importado, _ = revision.auditar_excel_externo(ruta)
    descripcion = importado.presupuesto.partidas[0].item.descripcion
    assert "HYPERLINK" not in descripcion and "127.0.0.1" not in descripcion
