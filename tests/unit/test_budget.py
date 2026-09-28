"""Presupuesto, curva de inversión y exportación a Excel (Sesión I0.5).

El presupuesto se arma con las cantidades del caso auditado y los cinco APU de la línea base, así
que su total debe reproducir los 1 586,61 USD del PDF con la tolerancia de ± 0,01 de CLAUDE.md §4.

La curva es el punto crítico de la sesión: en el presupuesto del caso didáctico cerraba en 99,30 %
(hallazgo 3), y aquí debe cerrar exactamente en el total **por construcción**, incluso repartiendo
con montos redondeados a dos decimales.

Ningún dato se repite: cantidades, APU y totales vienen de `tests.fixtures.apu_linea_base` a través
de `tests.fixtures.computo_auditado` (principio DRY de CLAUDE.md §2).
"""

from __future__ import annotations

from dataclasses import replace
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from adapters.civil.tabular import AdaptadorCivilTabular
from core.budget import (
    CODIGO_PENDIENTE,
    con_curva,
    elaborar,
    exportar_excel,
    generar_curva,
    generar_presupuesto,
    plan_secuencial,
)
from core.budget.curva import PlanInvalido
from core.contracts.apu import ComposicionAPU, LineaManoObra, ModalidadManoObra, ParametrosCosto
from core.contracts.dominio import Dominio
from core.contracts.item_computo import ItemComputo, OrigenTipo
from core.contracts.verificacion import Hallazgo, Severidad
from core.verification import REGLAS
from core.verification.directivas import codigos_en_etiqueta
from core.verification.informe import InformeAuditoria
from tests.fixtures import apu_linea_base as linea_base
from tests.fixtures.computo_auditado import composiciones_linea_base, items_auditados
from tests.fixtures.computo_civil import CODIGOS as CODIGOS_CIVIL

TOLERANCIA = Decimal("0.01")
CENTIMO = Decimal("0.01")
COMPLETO = Decimal("1")
MITAD = Decimal("0.5")

HOJAS_ESPERADAS = ["Presupuesto", "APU", "Curva", "Auditoria"]
MUESTRA_CIVIL = (
    Path(__file__).resolve().parents[2] / "data" / "samples" / "civil" / "tanquillas_y_zanja.csv"
)


def _presupuesto():
    return generar_presupuesto(
        items_auditados(),
        composiciones_linea_base(),
        linea_base.PARAMETROS_LINEA_BASE,
        codigo=linea_base.CODIGO_PRESUPUESTO,
        fecha=linea_base.FECHA_LINEA_BASE,
        moneda=linea_base.MONEDA,
    )


def _plan_auditado():
    """Los seis períodos del plan de trabajo del caso: el encofrado se ejecuta en dos mitades."""
    return [
        ("Dia 1", ((linea_base.APU_EXCAVACION.codigo_partida, COMPLETO),)),
        ("Dia 2", ((linea_base.APU_TUBERIA.codigo_partida, COMPLETO),)),
        ("Dia 3", ((linea_base.APU_ENCOFRADO.codigo_partida, MITAD),)),
        ("Dia 4", ((linea_base.APU_ENCOFRADO.codigo_partida, MITAD),)),
        ("Dia 5", ((linea_base.APU_CONCRETO.codigo_partida, COMPLETO),)),
        ("Dia 6", ((linea_base.APU_RELLENO.codigo_partida, COMPLETO),)),
    ]


def _acumulados_coherentes(curva) -> bool:
    corrido = Decimal(0)
    for punto in curva:
        corrido += punto.monto
        if punto.acumulado != corrido:
            return False
    return True


def test_presupuesto_reproduce_1586_61():
    presupuesto = _presupuesto()

    assert len(presupuesto.partidas) == len(linea_base.PRESUPUESTO_AUDITADO)
    assert abs(presupuesto.total - linea_base.TOTAL_PRESUPUESTO_AUDITADO) <= TOLERANCIA


def test_cada_renglon_reproduce_el_total_auditado():
    presupuesto = _presupuesto()

    for partida, linea in zip(presupuesto.partidas, linea_base.PRESUPUESTO_AUDITADO, strict=True):
        assert partida.item.codigo_partida == linea.codigo_partida
        assert abs(partida.resultado.precio_unitario - linea.precio_unitario) <= TOLERANCIA
        assert abs(partida.total - linea.total) <= TOLERANCIA, linea.codigo_partida


def test_codigo_sin_composicion_lanza_valueerror():
    """`generar_presupuesto` es el paso intermedio: exige composición para todo ítem. El flujo 3a
    de UC‑01 (ítem pendiente) lo implementa `elaborar`: ver la prueba siguiente."""
    composiciones = composiciones_linea_base()
    huerfano = linea_base.APU_RELLENO.codigo_partida
    del composiciones[huerfano]

    with pytest.raises(ValueError, match=huerfano):
        generar_presupuesto(
            items_auditados(),
            composiciones,
            linea_base.PARAMETROS_LINEA_BASE,
            codigo=linea_base.CODIGO_PRESUPUESTO,
            fecha=linea_base.FECHA_LINEA_BASE,
        )


def test_un_item_sin_partida_queda_pendiente_y_el_informe_lo_declara():
    """UC‑01, flujo 3a (INV‑013): no suma al total, no entra al plan, el informe lo declara.

    Los ítems auditados ya dan errores de R1 (hallazgos 1 y 2 de la línea base), así que
    `not informe.cumple` no probaría nada. La prueba compara contra la elaboración de los mismos
    ítems sin el huérfano: el presupuesto es idéntico y el único hallazgo nuevo es el `PENDIENTE`.
    """
    composiciones = composiciones_linea_base()
    huerfano = linea_base.APU_RELLENO.codigo_partida
    del composiciones[huerfano]
    completo = _presupuesto()
    total_huerfano = sum(
        (p.total for p in completo.partidas if p.item.codigo_partida == huerfano), Decimal(0)
    )
    argumentos = dict(
        codigo=linea_base.CODIGO_PRESUPUESTO,
        fecha=linea_base.FECHA_LINEA_BASE,
        plan=plan_secuencial,
    )

    resultado = elaborar(
        items_auditados(), composiciones, linea_base.PARAMETROS_LINEA_BASE, **argumentos
    )
    sin_huerfano = elaborar(
        [item for item in items_auditados() if item.codigo_partida != huerfano],
        composiciones_linea_base(),
        linea_base.PARAMETROS_LINEA_BASE,
        **argumentos,
    )

    assert [item.codigo_partida for item in resultado.pendientes] == [huerfano]
    assert resultado.presupuesto == sin_huerfano.presupuesto
    assert resultado.presupuesto.total == completo.total - total_huerfano
    assert resultado.presupuesto.total_curva == resultado.presupuesto.total
    (declarado,) = resultado.informe.por_regla()[CODIGO_PENDIENTE]
    # `InformeAuditoria` ordena sus hallazgos por gravedad y regla: el `PENDIENTE` no queda al
    # final. Sin él, el informe es exactamente el de la elaboración sin el huérfano.
    otros = tuple(hallazgo for hallazgo in resultado.informe.hallazgos if hallazgo is not declarado)
    assert otros == sin_huerfano.informe.hallazgos
    assert declarado.severidad is Severidad.ERROR
    assert resultado.informe.partida_de(declarado) == huerfano


def test_el_pendiente_se_atribuye_a_su_partida_aunque_comparta_origen_id():
    """El adaptador civil tabular da un mismo `origen_id` (la fila `T-01`) al concreto, al
    encofrado y a la excavación. Sin composición de concreto, el `PENDIENTE` de esa fila debe
    atribuirse al concreto, no a la partida valorada que comparte el origen."""
    items = AdaptadorCivilTabular(codigos=CODIGOS_CIVIL).extraer(MUESTRA_CIVIL)
    concreto = linea_base.APU_CONCRETO.codigo_partida
    composiciones = composiciones_linea_base()
    del composiciones[concreto]

    resultado = elaborar(
        items,
        composiciones,
        linea_base.PARAMETROS_LINEA_BASE,
        codigo=linea_base.CODIGO_PRESUPUESTO,
        fecha=linea_base.FECHA_LINEA_BASE,
        plan=plan_secuencial,
    )

    (declarado,) = resultado.informe.por_regla()[CODIGO_PENDIENTE]
    (pendiente,) = resultado.pendientes
    assert any(
        partida.item.origen_id == pendiente.origen_id for partida in resultado.presupuesto.partidas
    )
    assert resultado.informe.partida_de(declarado) == concreto


def test_plan_secuencial_admite_un_codigo_repetido_y_la_curva_cierra():
    """INV‑071: la muestra civil tabular da `LB-01-EXC` en dos renglones (filas `T-01` y `Z-01`).
    `plan_secuencial` abre un período por código de partida, en el orden de su primera aparición,
    y la curva cierra exactamente en el total del presupuesto: ni R2 ni R7 dan ERROR."""
    items = AdaptadorCivilTabular(codigos=CODIGOS_CIVIL).extraer(MUESTRA_CIVIL)
    codigos = [item.codigo_partida for item in items]
    assert len(codigos) > len(set(codigos)), "la muestra ya no repite ningún código"

    resultado = elaborar(
        items,
        composiciones_linea_base(),
        linea_base.PARAMETROS_LINEA_BASE,
        codigo=linea_base.CODIGO_PRESUPUESTO,
        fecha=linea_base.FECHA_LINEA_BASE,
        plan=plan_secuencial,
    )

    presupuesto = resultado.presupuesto
    assert resultado.pendientes == ()
    assert presupuesto.total_curva == presupuesto.total
    periodos = [codigos_en_etiqueta(punto.periodo) for punto in presupuesto.curva]
    assert periodos == [(codigo,) for codigo in dict.fromkeys(codigos)]
    por_regla = resultado.informe.por_regla()
    for regla in ("R2", "R7"):
        assert all(h.severidad is not Severidad.ERROR for h in por_regla.get(regla, [])), regla


def test_sin_ninguna_partida_valorable_el_informe_se_emite_igual():
    resultado = elaborar(
        items_auditados(),
        {},
        linea_base.PARAMETROS_LINEA_BASE,
        codigo=linea_base.CODIGO_PRESUPUESTO,
        fecha=linea_base.FECHA_LINEA_BASE,
        plan=plan_secuencial,
    )

    assert resultado.presupuesto.partidas == ()
    assert resultado.presupuesto.total == Decimal(0)
    assert len(resultado.pendientes) == len(linea_base.PRESUPUESTO_AUDITADO)
    assert len(resultado.informe.por_regla()[CODIGO_PENDIENTE]) == len(
        linea_base.PRESUPUESTO_AUDITADO
    )


def test_elaborar_acepta_el_plan_como_funcion_del_presupuesto_valorado():
    """INV‑069: el protocolo borrador → `plan_secuencial` → `elaborar` vive en `elaborar`."""
    argumentos = dict(
        codigo=linea_base.CODIGO_PRESUPUESTO,
        fecha=linea_base.FECHA_LINEA_BASE,
        moneda=linea_base.MONEDA,
    )
    con_funcion = elaborar(
        items_auditados(),
        composiciones_linea_base(),
        linea_base.PARAMETROS_LINEA_BASE,
        plan=plan_secuencial,
        **argumentos,
    )
    con_lista = elaborar(
        items_auditados(),
        composiciones_linea_base(),
        linea_base.PARAMETROS_LINEA_BASE,
        plan=plan_secuencial(_presupuesto()),
        **argumentos,
    )

    assert con_funcion.presupuesto == con_lista.presupuesto
    assert con_funcion.informe.hallazgos == con_lista.informe.hallazgos
    assert con_funcion.presupuesto.total_curva == con_funcion.presupuesto.total
    assert con_funcion.pendientes == ()


def test_curva_cierra_exactamente_en_el_total():
    presupuesto = _presupuesto()

    curva = generar_curva(presupuesto, _plan_auditado(), cuantizar=CENTIMO)
    cerrado = con_curva(presupuesto, curva)

    assert len(curva) == len(linea_base.CURVA_AUDITADA)
    assert _acumulados_coherentes(curva)
    assert cerrado.total_curva == cerrado.total
    assert presupuesto.curva == (), "con_curva no debe modificar el presupuesto original"
    assert codigos_en_etiqueta(curva[0].periodo) == (linea_base.APU_EXCAVACION.codigo_partida,)


@pytest.mark.parametrize(
    "fracciones",
    [
        (Decimal("1"),),
        (Decimal("0.5"), Decimal("0.5")),
        (Decimal("0.25"), Decimal("0.25"), Decimal("0.5")),
        (Decimal("0.125"), Decimal("0.375"), Decimal("0.2"), Decimal("0.3")),
    ],
)
def test_curva_con_fracciones_arbitrarias_cierra(fracciones):
    presupuesto = _presupuesto()
    codigos = [partida.item.codigo_partida for partida in presupuesto.partidas]
    plan = [
        (f"Etapa {numero}", tuple((codigo, fraccion) for codigo in codigos))
        for numero, fraccion in enumerate(fracciones, start=1)
    ]

    curva = generar_curva(presupuesto, plan, cuantizar=CENTIMO)

    assert len(curva) == len(fracciones)
    assert _acumulados_coherentes(curva)
    assert curva[-1].acumulado == presupuesto.total


def test_fracciones_que_no_suman_uno_lanzan_valueerror():
    presupuesto = _presupuesto()
    incompleto = linea_base.APU_ENCOFRADO.codigo_partida
    plan = [
        (
            etiqueta,
            tuple(
                (codigo, MITAD if codigo == incompleto else fraccion)
                for codigo, fraccion in reparto
            ),
        )
        for etiqueta, reparto in plan_secuencial(presupuesto)
    ]

    with pytest.raises(PlanInvalido, match=incompleto) as error:
        generar_curva(presupuesto, plan)
    assert isinstance(error.value, ValueError)


def test_codigo_desconocido_en_el_plan_lanza_valueerror():
    presupuesto = _presupuesto()
    plan = [*plan_secuencial(presupuesto), ("Dia 6", (("LB-99-XXX", COMPLETO),))]

    with pytest.raises(PlanInvalido, match="LB-99-XXX"):
        generar_curva(presupuesto, plan)


def test_elaborar_siempre_devuelve_informe():
    sin_plan = elaborar(
        items_auditados(),
        composiciones_linea_base(),
        linea_base.PARAMETROS_LINEA_BASE,
        codigo=linea_base.CODIGO_PRESUPUESTO,
        fecha=linea_base.FECHA_LINEA_BASE,
    )
    con_plan = elaborar(
        items_auditados(),
        composiciones_linea_base(),
        linea_base.PARAMETROS_LINEA_BASE,
        codigo=linea_base.CODIGO_PRESUPUESTO,
        fecha=linea_base.FECHA_LINEA_BASE,
        plan=_plan_auditado(),
    )

    assert isinstance(sin_plan.informe, InformeAuditoria)
    assert isinstance(con_plan.informe, InformeAuditoria)
    assert sin_plan.presupuesto.curva == ()
    assert con_plan.presupuesto.total_curva == con_plan.presupuesto.total
    assert con_plan.informe.codigo_presupuesto == linea_base.CODIGO_PRESUPUESTO


def test_elaborar_con_un_parametro_no_decimal_emite_presupuesto_e_informe():
    """Principio 7 (INV‑003): datos degradados dan hallazgos, nunca un presupuesto sin informe."""
    excavacion = linea_base.APU_EXCAVACION.codigo_partida
    items = tuple(
        replace(item, regla="longitud * 2", parametros={"longitud": 1.5})
        if item.codigo_partida == excavacion
        else item
        for item in items_auditados()
    )

    resultado = elaborar(
        items,
        composiciones_linea_base(),
        linea_base.PARAMETROS_LINEA_BASE,
        codigo=linea_base.CODIGO_PRESUPUESTO,
        fecha=linea_base.FECHA_LINEA_BASE,
    )

    assert abs(resultado.presupuesto.total - linea_base.TOTAL_PRESUPUESTO_AUDITADO) <= TOLERANCIA
    assert any("'longitud'" in h.descripcion for h in resultado.informe.por_regla()["R1"])


def test_exportar_excel_escribe_cuatro_hojas(tmp_path):
    resultado = elaborar(
        items_auditados(),
        composiciones_linea_base(),
        linea_base.PARAMETROS_LINEA_BASE,
        codigo=linea_base.CODIGO_PRESUPUESTO,
        fecha=linea_base.FECHA_LINEA_BASE,
        plan=_plan_auditado(),
    )
    ruta = tmp_path / "presupuesto.xlsx"

    escrita = exportar_excel(resultado.presupuesto, resultado.informe, ruta)

    assert escrita == ruta
    libro = load_workbook(ruta)
    assert libro.sheetnames == HOJAS_ESPERADAS

    hoja = libro["Presupuesto"]
    total_escrito = Decimal(str(hoja.cell(row=hoja.max_row, column=hoja.max_column).value))
    assert abs(total_escrito - linea_base.TOTAL_PRESUPUESTO_AUDITADO) <= TOLERANCIA
    assert libro["Curva"].max_row == len(resultado.presupuesto.curva) + 1
    assert libro["Auditoria"].max_row == len(resultado.informe.hallazgos) + 1


def test_hoja_apu_muestra_la_modalidad_de_mano_de_obra(tmp_path):
    """La columna Modalidad evita el malentendido de la decision D9: bajo destajo, `sueldo` es el
    precio por unidad de partida, no un jornal diario, y quien audite el libro debe poder verlo."""
    item = ItemComputo(
        codigo_partida="X-01",
        descripcion="Partida con mano de obra mixta",
        unidad="m2",
        cantidad=Decimal("1"),
        origen_id="TEST-MODALIDAD-01",
        origen_tipo=OrigenTipo.MANUAL,
        dominio=Dominio.CIVIL,
    )
    apu = ComposicionAPU(
        codigo_partida="X-01",
        descripcion="Partida con mano de obra mixta",
        unidad="m2",
        rendimiento=Decimal("50"),
        mano_obra=(
            LineaManoObra("Ayudante a jornal", Decimal("2"), Decimal("3")),
            LineaManoObra(
                "Instalador a destajo",
                Decimal("1"),
                Decimal("6"),
                ModalidadManoObra.DESTAJO,
            ),
        ),
    )
    resultado = elaborar(
        [item],
        {"X-01": apu},
        ParametrosCosto(),
        codigo="PRES-TEST-MODALIDAD",
        fecha=date(2026, 9, 20),
    )
    ruta = tmp_path / "modalidad.xlsx"

    exportar_excel(resultado.presupuesto, resultado.informe, ruta)

    hoja = load_workbook(ruta)["APU"]
    filas = list(hoja.iter_rows(values_only=True))
    indice_encabezado = next(i for i, fila in enumerate(filas) if fila[0] == "Mano de obra")
    encabezado = filas[indice_encabezado]
    fila_jornal = filas[indice_encabezado + 1]
    fila_destajo = filas[indice_encabezado + 2]

    assert "Modalidad" in encabezado
    assert encabezado[-1] == "Total"
    assert fila_jornal[-2] == "jornal"
    assert fila_destajo[-2] == "destajo"
    # El Total de esta tabla es cantidad x sueldo (igual que en materiales y equipos), no el
    # aporte con FCAS y bono que va en el resumen: aqui coincide en 6.00 para ambas lineas
    # (2 obreros x 3 = 1 destajista x 6), pero por razones distintas segun la modalidad.
    assert fila_jornal[-1] == Decimal("6.00")
    assert fila_destajo[-1] == Decimal("6.00")


# ---------------------------------------------------------------------------------------------
# Reglas evaluadas y atribución con pendientes (INV‑092 e INV‑095, expedientes C5.1 y C5.2)
# ---------------------------------------------------------------------------------------------


def test_la_elaboracion_con_pendientes_conserva_las_reglas_evaluadas_y_la_atribucion():
    """`_declarar_pendientes` no pierde lo que el informe ya sabía (ARQ‑3 de GC4).

    La muestra civil sin composición de concreto: el pendiente comparte la fila `T-01` con el
    encofrado y la excavación, y el R3 de la tubería (fila `Z-01`, que también da la excavación)
    solo se atribuye bien si sobrevive la declaración de la regla.
    """
    items = AdaptadorCivilTabular(codigos=CODIGOS_CIVIL).extraer(MUESTRA_CIVIL)
    composiciones = composiciones_linea_base()
    del composiciones[linea_base.APU_CONCRETO.codigo_partida]

    resultado = elaborar(
        items,
        composiciones,
        linea_base.PARAMETROS_LINEA_BASE,
        codigo=linea_base.CODIGO_PRESUPUESTO,
        fecha=linea_base.FECHA_LINEA_BASE,
        plan=plan_secuencial,
    )

    assert resultado.pendientes
    informe = resultado.informe
    markdown = informe.a_markdown()
    for regla in REGLAS:
        assert f"| {regla.codigo} | {regla.nombre} |" in markdown
    assert f"| {CODIGO_PENDIENTE} |" not in markdown.split("## Hallazgos")[0]
    (r3,) = informe.por_regla()["R3"]
    assert informe.partida_de(r3) == linea_base.APU_TUBERIA.codigo_partida


def test_la_hoja_de_auditoria_enumera_las_reglas_evaluadas(tmp_path):
    resultado = elaborar(
        items_auditados(),
        composiciones_linea_base(),
        linea_base.PARAMETROS_LINEA_BASE,
        codigo=linea_base.CODIGO_PRESUPUESTO,
        fecha=linea_base.FECHA_LINEA_BASE,
        plan=_plan_auditado(),
    )
    ruta = exportar_excel(resultado.presupuesto, resultado.informe, tmp_path / "reglas.xlsx")

    hoja = load_workbook(ruta)["Auditoria"]
    bloque = [
        fila
        for fila in hoja.iter_rows(min_col=7, max_col=9, values_only=True)
        if any(celda is not None for celda in fila)
    ]
    assert bloque, "la hoja de auditoria no enumera las reglas evaluadas"
    assert bloque[0] == ("Regla evaluada", "Nombre", "Hallazgos")
    assert bloque[1:] == [
        (
            regla.codigo,
            regla.nombre,
            sum(1 for hallazgo in resultado.informe.hallazgos if hallazgo.regla == regla.codigo),
        )
        for regla in REGLAS
    ]
    # La tabla de hallazgos (columna A) sigue con una fila por hallazgo, más el encabezado.
    tabla = [fila for (fila,) in hoja.iter_rows(max_col=1, values_only=True) if fila is not None]
    assert len(tabla) == len(resultado.informe.hallazgos) + 1


def test_el_origen_de_un_pendiente_no_atribuye_a_una_regla_que_no_declara_su_partida():
    """M‑1 de la revisión de C5: la fila `T-01` da encofrado y excavación (valorados) y concreto
    (pendiente, sin composición). Un hallazgo que no declara su partida y cita `T-01` (el de una
    regla ajena al núcleo) no tiene forma de saber cuál de las dos valoradas es: el informe no lo
    atribuye, y tampoco a la partida pendiente por ocupar el hueco del índice."""
    items = AdaptadorCivilTabular(codigos=CODIGOS_CIVIL).extraer(MUESTRA_CIVIL)
    concreto = linea_base.APU_CONCRETO.codigo_partida
    composiciones = composiciones_linea_base()
    del composiciones[concreto]

    resultado = elaborar(
        items,
        composiciones,
        linea_base.PARAMETROS_LINEA_BASE,
        codigo=linea_base.CODIGO_PRESUPUESTO,
        fecha=linea_base.FECHA_LINEA_BASE,
        plan=plan_secuencial,
    )

    (pendiente,) = resultado.pendientes
    assert pendiente.codigo_partida == concreto
    ajeno = Hallazgo(
        regla="X1",
        severidad=Severidad.INFO,
        descripcion="cita la fila del pendiente",
        origen_ids=(pendiente.origen_id,),
    )
    assert resultado.informe.partida_de(ajeno) is None
    (declarado,) = resultado.informe.por_regla()[CODIGO_PENDIENTE]
    assert resultado.informe.partida_de(declarado) == concreto
