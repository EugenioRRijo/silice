import openpyxl
import pytest

from adapters.bcv.informe import FormatoNoReconocido
from adapters.bcv.libro import FIRMA_OLE2, LIMITE_BYTES, abrir_libro
from tests.fixtures.bcv_cuadriculas import CARPETA, campo


def test_un_xlsx_renombrado_xls_se_rechaza(tmp_path):
    ruta = tmp_path / "x.xlsx"
    openpyxl.Workbook().save(ruta)
    falso = ruta.rename(tmp_path / "2_1_2c26_smc.xls")
    with pytest.raises(FormatoNoReconocido, match="renombrado"):
        abrir_libro(falso)


def test_mas_de_20_mb_se_rechaza(tmp_path):
    ruta = tmp_path / "grande.xls"
    ruta.write_bytes(FIRMA_OLE2 + b"\0" * LIMITE_BYTES)
    with pytest.raises(FormatoNoReconocido, match="20 MB"):
        abrir_libro(ruta)


def test_el_libro_real_abre_con_su_huella():
    libro = abrir_libro(CARPETA / campo("Tasa", "Archivo"))
    assert libro.huella == campo("Tasa", "SHA-256")
    assert len(libro.hojas) == int(campo("Tasa", "Hojas (N)"))


def test_un_xlsx_corrupto_o_sin_workbook_se_rechaza(tmp_path):
    """SEG-2: un .xlsx corrupto o sin estructura se rechaza con FormatoNoReconocido."""
    corrupto = tmp_path / "corrupto.xlsx"
    corrupto.write_bytes(b"PK\x03\x04" + b"\0" * 50)
    with pytest.raises(FormatoNoReconocido, match="danado|no es un libro"):
        abrir_libro(corrupto)


def _xlsx_con(tmp_path, nombres):
    libro = openpyxl.Workbook()
    libro.remove(libro.active)
    for nombre in nombres:
        libro.create_sheet(nombre).append(["x", 1.5])
    ruta = tmp_path / "t.xlsx"
    libro.save(ruta)
    return ruta


def test_un_xlsx_lee_solo_las_hojas_necesarias(tmp_path):
    libro = abrir_libro(_xlsx_con(tmp_path, ["01072026", "Notas", "02072026"]))
    assert [h.nombre for h in libro.hojas] == ["01072026", "02072026"]
    assert any("hojas ajenas no leidas: 1" in a for a in libro.avisos)


def test_un_xlsx_sin_hojas_necesarias_o_con_miles_se_rechaza(tmp_path):
    with pytest.raises(FormatoNoReconocido, match="ninguna hoja necesaria"):
        abrir_libro(_xlsx_con(tmp_path, ["Notas"]))
    with pytest.raises(FormatoNoReconocido, match="hojas; supera el maximo"):
        abrir_libro(_xlsx_con(tmp_path, [f"01{n:06d}" for n in range(300)]))
