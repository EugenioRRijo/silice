"""RF-110 (V-D3, sesión V7, RL-15): `scripts/anonimizar_libro.py` anonimiza un libro en el ZIP/XML.

Lo que se prueba, con el libro sintético de `tests/fixtures/libros_empresa.py` sembrado con datos
identificables FICTICIOS (cliente, contratista, ubicación, RIF con formato venezolano inventado,
autor, ruta de red de un vínculo externo, ruta de origen `absPath` y empresa en `docProps/app.xml`):

1. Los valores guardados (`<v>`) de TODAS las celdas de las hojas de datos (PRESUPUESTO, APU-01..03,
   INSUMOS, RESUMEN) quedan idénticos, texto a texto (tolerancia 0: el script los copia byte a
   byte, por eso se comparan las cadenas del XML y no floats), y las fórmulas siguen siendo las
   mismas fórmulas.
2. `procesar_libro` sobre el anonimizado da la memoria de cálculo a mano (`MEMORIA`: P.U.
   14,486725 / 30,2121875 / 4,056525; SUBSINIVA 4 358,20825) y los hallazgos son los del
   original salvo exactamente los dos L5 de vínculo externo (RL-15).
3. Ningún texto identificable aparece en NINGUNA pieza del ZIP anonimizado (búsqueda de bytes, sin
   distinguir mayúsculas), y cada uno SÍ estaba en el original (la búsqueda no es vacía).
4. Lo no identificable (descripciones, códigos, fuentes, montos) se conserva.
5. La CLI (`main(argv)`): código 0 y casos negativos (ruta inexistente, archivo que no es xlsx,
   SALIDA = ENTRADA, etiqueta vacía, PORTADA sin rótulos, DTD, verificación fallida) con su código.

Oráculos: los sustitutos son los literales documentados en la cabecera del script
(«CLIENTE-A», «CONTRATISTA-A», «UBICACIÓN ANONIMIZADA», «J-00000000-0», «anonimo»), los montos son
`MEMORIA` (memoria escrita a mano) y los hallazgos esperados del original son literales (L5 de
`nombre:MEDP1` y de `vinculo:mediciones.xlsx`). Ningún esperado sale del propio anonimizador.
"""

from __future__ import annotations

import io
import re
import zipfile
from collections import Counter
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from adapters.hoja_empresa.lector import leer_libro
from adapters.hoja_empresa.reglas import reglas_del_libro
from core.contracts.verificacion import Severidad
from fusion.proceso import procesar_libro
from scripts.anonimizar_libro import descubrir, descubrir_autores, main, verificar
from tests.fixtures.libros_empresa import (
    APUS_LIMPIOS,
    MEMORIA,
    D,
    Siembra,
    con,
    construir_libro,
)

# --- Datos identificables ficticios ---------------------------------------------------------------

CLIENTE = "Inversiones Ficticias Ejemplo C.A."
CONTRATISTA = "Constructora Inventada S.R.L."
UBICACION = "Av. Imaginaria 742, Ciudad Ficticia"
AUTOR = "Persona Inventada"
VINCULO = "file:///\\\\servidor-ficticio\\obras\\mediciones.xlsx"
# Un RIF inventado con formato venezolano. No es el neutro J-00000000-0: ese es justamente el
# sustituto del script y lo dejaría pasar sin cambiarlo.
RIF_PORTADA = "J-00000123-4"
RIF_MEDICIONES = "V-00000456-7"
OBRA = f"Sede de {CLIENTE} (ficticia)"
RUTA_ORIGEN = "C:\\Users\\PersonaInventada\\Obras Ficticias\\"
EMPRESA_APP = CONTRATISTA

#: Fórmula con el contratista dentro de un literal: solo el literal debe cambiar.
FORMULA_LITERAL = f'="Contrata: "&"{CONTRATISTA}"'

SIEMBRA = Siembra(
    cliente=CLIENTE,
    contratista=CONTRATISTA,
    ubicacion=UBICACION,
    autor=AUTOR,
    obra=OBRA,
    vinculo_externo=VINCULO,
    nombres_externos=("MEDP1",),
    parches={
        "PORTADA!A12": ("RIF:", None),
        "PORTADA!B12": (f"RIF {RIF_PORTADA}", None),
        "PORTADA!B13": (f"Contrata: {CONTRATISTA}", FORMULA_LITERAL),
        "MEDICIONES!A20": (f"Proveedor de bloques, RIF {RIF_MEDICIONES}", None),
    },
)

#: Textos que no pueden quedar en ninguna pieza del anonimizado (y que sí están en el original).
IDENTIFICABLES = (
    "Inversiones Ficticias",
    "Ficticias Ejemplo",
    "Constructora Inventada",
    "Inventada S.R.L",
    "Av. Imaginaria",
    "Imaginaria 742",
    "Ciudad Ficticia",
    RIF_PORTADA,
    "00000123",
    RIF_MEDICIONES,
    "00000456",
    AUTOR,
    "servidor-ficticio",
    "mediciones.xlsx",
    "PersonaInventada",
    "Obras Ficticias",
    "[1]Mediciones",
)

#: Hojas de datos cuyos valores guardados deben quedar idénticos.
HOJAS_DE_DATOS = ("PRESUPUESTO", "APU-01", "APU-02", "APU-03", "INSUMOS", "RESUMEN")

#: Sustitutos literales (cabecera de `scripts/anonimizar_libro.py`, etiqueta por omisión «A»). Las
#: variantes descubiertas se recortan de " .;-" en los extremos, así que el punto final de
#: «C.A.» y de «S.R.L.» queda tras el sustituto.
CAMBIOS_ESPERADOS: dict[str, str] = {
    "PORTADA!A2": "Sede de CLIENTE-A. (ficticia)",
    "PORTADA!B4": "CLIENTE-A.",
    "PORTADA!B5": "CONTRATISTA-A.",
    "PORTADA!B6": "UBICACIÓN ANONIMIZADA",
    "PORTADA!B12": "RIF J-00000000-0",
    "PORTADA!B13": "Contrata: CONTRATISTA-A.",
    "MEDICIONES!A20": "Proveedor de bloques, RIF J-00000000-0",
    "APU-01!F5": "CLIENTE-A.",
    "APU-02!F5": "CLIENTE-A.",
    "APU-03!F5": "CLIENTE-A.",
}

#: Los dos L5 de vínculo externo del original (RL-15): el anonimizado los pierde a propósito.
L5_DE_VINCULO = Counter(
    {
        ("L5", Severidad.ADVERTENCIA, ("nombre:MEDP1",)): 1,
        ("L5", Severidad.ADVERTENCIA, ("vinculo:mediciones.xlsx",)): 1,
    }
)

FECHA = date(2026, 10, 13)

# --- Ayudas ---------------------------------------------------------------------------------------


def _reescribir(datos: bytes, cambios: dict[str, Callable[[str], str]]) -> bytes:
    """Copia el ZIP aplicando `cambios[pieza](texto) -> texto` a las piezas nombradas."""
    entrada = io.BytesIO(datos)
    salida = io.BytesIO()
    with (
        zipfile.ZipFile(entrada) as zin,
        zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as zout,
    ):
        for info in zin.infolist():
            contenido = zin.read(info.filename)
            if info.filename in cambios:
                contenido = cambios[info.filename](contenido.decode("utf-8")).encode("utf-8")
            zout.writestr(info, contenido)
    return salida.getvalue()


def _con_ruta_y_empresa(datos: bytes) -> bytes:
    """Añade lo que Excel deja y XlsxWriter no escribe: la ruta de origen (`x15ac:absPath`) en
    `xl/workbook.xml` y la empresa en `docProps/app.xml`."""
    alternativo = (
        '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
        '<mc:Choice Requires="x15ac"><x15ac:absPath '
        'xmlns:x15ac="http://schemas.microsoft.com/office/spreadsheetml/2010/11/ac" '
        f'url="{RUTA_ORIGEN}"/></mc:Choice></mc:AlternateContent>'
    )

    def libro(xml: str) -> str:
        return re.sub(r"(<workbook\b[^>]*>)", lambda m: m.group(1) + alternativo, xml, count=1)

    def app(xml: str) -> str:
        assert "<Company></Company>" in xml
        return xml.replace("<Company></Company>", f"<Company>{EMPRESA_APP}</Company>")

    return _reescribir(datos, {"xl/workbook.xml": libro, "docProps/app.xml": app})


def _piezas(ruta: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(ruta) as z:
        return {n: z.read(n) for n in z.namelist()}


_CELDA = re.compile(r'<c r="(?P<ref>[A-Z]+\d+)"(?P<attrs>[^>]*?)(?:/>|>(?P<cuerpo>.*?)</c>)', re.S)


def _celdas_xml(
    piezas: dict[str, bytes], hoja: str
) -> dict[str, tuple[str, str | None, str | None]]:
    """Celda -> (atributos, texto de `<f>`, texto de `<v>`), leídos del XML crudo de la hoja."""
    libro = piezas["xl/workbook.xml"].decode("utf-8")
    rels = piezas["xl/_rels/workbook.xml.rels"].decode("utf-8")
    rid = re.search(rf'<sheet name="{re.escape(hoja)}" sheetId="\d+" r:id="(rId\d+)"', libro)
    assert rid is not None, hoja
    destino = re.search(rf'Id="{rid.group(1)}"[^>]*Target="([^"]+)"', rels)
    if destino is None:
        destino = re.search(rf'Target="([^"]+)"[^>]*Id="{rid.group(1)}"', rels)
    assert destino is not None, hoja
    xml = piezas["xl/" + destino.group(1).lstrip("/").removeprefix("xl/")].decode("utf-8")
    celdas = {}
    for m in _CELDA.finditer(xml):
        cuerpo = m.group("cuerpo") or ""
        f = re.search(r"<f(?:\s[^>]*)?>([^<]*)</f>", cuerpo)
        v = re.search(r"<v>([^<]*)</v>", cuerpo)
        celdas[m.group("ref")] = (
            m.group("attrs"),
            f.group(1) if f else None,
            v.group(1) if v else None,
        )
    return celdas


def _valores(ruta: Path, *, data_only: bool) -> dict[str, object]:
    """'HOJA!A1' -> valor de openpyxl, de todas las hojas."""
    wb = openpyxl.load_workbook(ruta, data_only=data_only)
    try:
        return {
            f"{ws.title}!{c.coordinate}": c.value
            for ws in wb.worksheets
            for fila in ws.iter_rows()
            for c in fila
            if c.value is not None
        }
    finally:
        wb.close()


def _firmas(hallazgos) -> Counter:
    return Counter((h.regla, h.severidad, tuple(h.origen_ids)) for h in hallazgos)


@dataclass(frozen=True)
class Par:
    entrada: Path
    salida: Path
    codigo: int


def _anonimizar(datos: bytes, carpeta: Path, *extra: str) -> Par:
    carpeta.mkdir(parents=True, exist_ok=True)
    entrada = carpeta / "entrada.xlsx"
    entrada.write_bytes(datos)
    salida = carpeta / "salida.xlsx"
    codigo = main([str(entrada), str(salida), *extra])
    return Par(entrada, salida, codigo)


@pytest.fixture(scope="module")
def par(tmp_path_factory: pytest.TempPathFactory) -> Par:
    datos = _con_ruta_y_empresa(construir_libro(siembra=SIEMBRA).datos)
    return _anonimizar(datos, tmp_path_factory.mktemp("anonimizar"))


# --- CLI: el camino feliz -------------------------------------------------------------------------


def test_la_cli_anonimiza_con_codigo_0_y_no_toca_la_entrada(tmp_path, capsys):
    datos = _con_ruta_y_empresa(construir_libro(siembra=SIEMBRA).datos)
    p = _anonimizar(datos, tmp_path)
    salida_estandar = capsys.readouterr().out
    assert p.codigo == 0
    assert p.salida.is_file()
    assert p.entrada.read_bytes() == datos, "la ENTRADA no se modifica"
    assert "Verificacion: ningun dato descubierto aparece en la SALIDA (salida.xlsx)." in (
        salida_estandar
    )


def test_anonimizar_dos_veces_da_los_mismos_bytes(par, tmp_path):
    otra = tmp_path / "otra.xlsx"
    assert main([str(par.entrada), str(otra)]) == 0
    assert otra.read_bytes() == par.salida.read_bytes()


# --- 1. Valores guardados y fórmulas --------------------------------------------------------------


@pytest.mark.parametrize("hoja", HOJAS_DE_DATOS)
def test_los_valores_guardados_y_las_formulas_quedan_identicos(par, hoja):
    """Tolerancia 0: cada `<v>` y cada `<f>` del XML crudo es la misma cadena en los dos libros."""
    original = _celdas_xml(_piezas(par.entrada), hoja)
    anonimo = _celdas_xml(_piezas(par.salida), hoja)
    assert anonimo.keys() == original.keys()
    con_formula = [ref for ref, (_a, f, _v) in original.items() if f is not None]
    con_valor = [ref for ref, (_a, _f, v) in original.items() if v is not None]
    # INSUMOS no tiene fórmulas (solo datos); las demás hojas de datos, decenas.
    minimo = 0 if hoja == "INSUMOS" else 5
    assert len(con_formula) >= minimo, f"{hoja}: la comparación no puede ser vacía"
    assert len(con_valor) >= 20, f"{hoja}: la comparación no puede ser vacía"
    for ref in con_formula:
        _a, f_orig, v_orig = original[ref]
        _a2, f_anon, v_anon = anonimo[ref]
        assert v_orig is not None, f"{hoja}!{ref}: el original trae valor guardado"
        assert f_anon == f_orig, f"{hoja}!{ref}: la fórmula cambió"
        assert v_anon == v_orig, f"{hoja}!{ref}: el valor guardado cambió"
    # Las celdas sin fórmula también (los textos van por índice de sharedStrings, que no cambia).
    assert anonimo == original


def test_openpyxl_lee_los_mismos_valores_salvo_los_textos_identificables(par):
    """data_only=True en TODAS las hojas: solo cambian las celdas con datos identificables, y a
    los sustitutos literales del script."""
    original = _valores(par.entrada, data_only=True)
    anonimo = _valores(par.salida, data_only=True)
    assert anonimo.keys() == original.keys()
    cambiadas = {k: anonimo[k] for k in original if anonimo[k] != original[k]}
    assert cambiadas == CAMBIOS_ESPERADOS
    for hoja in HOJAS_DE_DATOS:
        numericas = [
            k
            for k, v in original.items()
            if k.startswith(f"{hoja}!") and isinstance(v, (int, float))
        ]
        assert numericas, hoja
        for k in numericas:
            assert type(anonimo[k]) is type(original[k]), k
            assert repr(anonimo[k]) == repr(original[k]), f"{k}: tolerancia 0"


def test_las_formulas_siguen_siendo_formulas_y_solo_cambia_el_literal(par):
    original = _valores(par.entrada, data_only=False)
    anonimo = _valores(par.salida, data_only=False)
    formulas = {k for k, v in original.items() if isinstance(v, str) and v.startswith("=")}
    assert len(formulas) > 100
    cambiadas = {k: anonimo[k] for k in formulas if anonimo[k] != original[k]}
    assert cambiadas == {"PORTADA!B13": '="Contrata: "&"CONTRATISTA-A."'}
    assert all(isinstance(anonimo[k], str) and anonimo[k].startswith("=") for k in formulas)


# --- 2. Totales y hallazgos -----------------------------------------------------------------------


def test_procesar_libro_sobre_el_anonimizado_da_la_memoria_de_calculo(par):
    r = procesar_libro(par.salida.read_bytes(), par.salida.name, fecha_devolucion=FECHA)
    esperados = {5: D("14.486725"), 6: D("30.2121875"), 7: D("4.056525")}
    assert {k: r.precios_unitarios[k] for k in esperados} == esperados
    for fila, hoja in zip((5, 6, 7), ("APU-01", "APU-02", "APU-03"), strict=True):
        assert r.precios_unitarios[fila] == MEMORIA["componentes"][hoja]["PU"]
    assert r.totales_silice["SUBSINIVA"] == D("4358.20825")
    assert r.totales_silice["IVATOT"] == D("697.31332")
    assert r.totales_silice["TOTALGEN"] == D("5055.52157")


def test_reglas_del_libro_iguales_salvo_los_dos_l5_de_vinculo_externo(par):
    original = _firmas(reglas_del_libro(leer_libro(par.entrada.read_bytes(), "entrada.xlsx")))
    anonimo = _firmas(reglas_del_libro(leer_libro(par.salida.read_bytes(), "salida.xlsx")))
    assert original == L5_DE_VINCULO
    assert anonimo == Counter()
    assert original - anonimo == L5_DE_VINCULO
    assert anonimo - original == Counter()


def test_el_informe_de_procesar_libro_pierde_solo_los_l5_de_vinculo(par):
    original = procesar_libro(par.entrada.read_bytes(), "entrada.xlsx", fecha_devolucion=FECHA)
    anonimo = procesar_libro(par.salida.read_bytes(), "salida.xlsx", fecha_devolucion=FECHA)
    f_orig = _firmas(original.informe.hallazgos)
    f_anon = _firmas(anonimo.informe.hallazgos)
    assert f_orig - f_anon == L5_DE_VINCULO
    assert f_anon - f_orig == Counter()
    assert Counter((r, s) for r, s, _o in f_anon.elements()) == Counter(
        {("R2", Severidad.INFO): 1, ("R7", Severidad.INFO): 1}
    )


def test_escenario_con_varios_errores_conserva_todo_salvo_los_l5_de_vinculo(tmp_path):
    """El escenario del brief (F1, L1, L4, L5, L6, L7, L9): los errores que dependen de valores
    guardados (F1, L9) sobreviven a la anonimización; solo caen los dos L5 de vínculo."""
    apus = (
        APUS_LIMPIOS[0],
        con(APUS_LIMPIOS[1], cd_sin_equipos=True, utilidad_literal=D("0.02")),
        con(APUS_LIMPIOS[2], divisor_equipos=D("12")),
    )
    siembra = Siembra(
        mediciones_literales=("=9.73*2.75",),
        resumen_ref_rota=True,
        nombres_externos=("MEDP1", "MEDP2"),
        procura_dif=D("5"),
        cronograma_fila_rend=23,
        vinculo_externo=VINCULO,
    )
    p = _anonimizar(construir_libro(apus, siembra).datos, tmp_path)
    assert p.codigo == 0
    original = procesar_libro(p.entrada.read_bytes(), "entrada.xlsx", fecha_devolucion=FECHA)
    anonimo = procesar_libro(p.salida.read_bytes(), "salida.xlsx", fecha_devolucion=FECHA)
    f_orig = _firmas(original.informe.hallazgos)
    f_anon = _firmas(anonimo.informe.hallazgos)
    assert f_orig - f_anon == Counter(
        {
            ("L5", Severidad.ADVERTENCIA, ("nombre:MEDP1", "nombre:MEDP2")): 1,
            ("L5", Severidad.ADVERTENCIA, ("vinculo:mediciones.xlsx",)): 1,
        }
    )
    assert f_anon - f_orig == Counter()
    por_regla = Counter((r, s) for r, s, _o in f_anon.elements() if r[0] in "FL")
    assert por_regla == Counter(
        {
            ("F1", Severidad.ERROR): 1,
            ("L1", Severidad.ERROR): 3,
            ("L4", Severidad.ERROR): 1,
            ("L5", Severidad.ERROR): 1,
            ("L6", Severidad.ADVERTENCIA): 2,
            ("L7", Severidad.INFO): 1,
            ("L9", Severidad.ERROR): 3,
        }
    )
    # Los P.U. y totales de Silice tampoco cambian.
    assert dict(anonimo.precios_unitarios) == dict(original.precios_unitarios)
    assert dict(anonimo.totales_silice) == dict(original.totales_silice)


# --- 3. Nada identificable en ninguna pieza -------------------------------------------------------


@pytest.mark.parametrize("texto", IDENTIFICABLES)
def test_ningun_texto_identificable_queda_en_ninguna_pieza_del_zip(par, texto):
    aguja = texto.lower().encode("utf-8")  # las agujas son ASCII: bytes.lower basta
    en_original = [n for n, b in _piezas(par.entrada).items() if aguja in b.lower()]
    assert en_original, f"«{texto}» debía estar en el original (búsqueda no vacía)"
    en_anonimo = [n for n, b in _piezas(par.salida).items() if aguja in b.lower()]
    assert en_anonimo == []


def test_el_rif_cambia_al_neutro_tambien_fuera_de_la_portada(par):
    textos = _piezas(par.salida)["xl/sharedStrings.xml"].decode("utf-8")
    rifs = re.findall(r"[JVEGP]-?\d{7,9}-?\d?", textos)
    assert rifs == ["J-00000000-0", "J-00000000-0"]


def test_las_piezas_del_vinculo_externo_se_quitan(par):
    nombres_orig = set(_piezas(par.entrada))
    nombres_anon = set(_piezas(par.salida))
    assert nombres_orig - nombres_anon == {
        "xl/externalLinks/externalLink1.xml",
        "xl/externalLinks/_rels/externalLink1.xml.rels",
    }
    assert nombres_anon - nombres_orig == set()
    libro = _piezas(par.salida)["xl/workbook.xml"].decode("utf-8")
    assert "MEDP1" not in libro
    assert "absPath" not in libro
    # Los demás nombres definidos siguen (los usa el lector: FCAS, PRECIOS, SUBSINIVA...).
    for nombre in ("FCAS", "PRECIOS", "SUBSINIVA", "IVATOT", "TOTALGEN"):
        assert f'<definedName name="{nombre}">' in libro


def test_las_propiedades_del_documento_quedan_en_anonimo(par):
    piezas = _piezas(par.salida)
    core = piezas["docProps/core.xml"].decode("utf-8")
    assert re.search(r"<dc:creator>anonimo</dc:creator>", core)
    assert re.search(r"<cp:lastModifiedBy>anonimo</cp:lastModifiedBy>", core)
    assert re.search(r"<dc:title>anonimo</dc:title>", core)
    app = piezas["docProps/app.xml"].decode("utf-8")
    assert "<Company>anonimo</Company>" in app
    wb = openpyxl.load_workbook(par.salida)
    try:
        assert wb.properties.creator == "anonimo"
    finally:
        wb.close()


# --- 4. Lo no identificable se conserva -----------------------------------------------------------


def test_lo_no_identificable_se_conserva(par):
    anonimo = _valores(par.salida, data_only=True)
    esperados = {
        "INSUMOS!B5": "MT-01",
        "INSUMOS!C5": "Cemento gris tipo I (ficticio)",
        "INSUMOS!F5": "Cotización ficticia A",
        "INSUMOS!B7": "MT-03",
        "INSUMOS!F7": "REFERENCIAL — ajustar",
        "INSUMOS!C11": "Mezcladora de 1 saco (ficticia)",
        "PORTADA!A1": "PRESUPUESTO DE OBRA",
        "PORTADA!A4": "CLIENTE:",
        "PORTADA!A5": "CONTRATISTA:",
        "PORTADA!B7": "COVENIN 2000 (referencial)",
        "PORTADA!B8": "Dólares de los Estados Unidos de América (USD)",
        "PORTADA!B9": "Julio 2026",
        "PRESUPUESTO!C5": "E.311.000.001",
        "PRESUPUESTO!D6": "Pared de bloque de arcilla (ficticia)",
        "APU-02!C10": "Cemento gris tipo I (ficticio)",
    }
    assert {k: anonimo.get(k) for k in esperados} == esperados
    # Precios de INSUMOS (literales de la cabecera del generador), tolerancia 0.
    precios = {"INSUMOS!E5": D("10"), "INSUMOS!E6": D("25"), "INSUMOS!E7": D("0.8")}
    assert {k: D(repr(anonimo[k])) for k in precios} == precios


# --- Red de seguridad: `verificar` ve lo que hay en el original -----------------------------------


def test_verificar_encuentra_los_datos_en_el_original_y_nada_en_el_anonimizado(par):
    wb = openpyxl.load_workbook(par.entrada, keep_links=False)
    try:
        datos = descubrir(wb, "A")
    finally:
        wb.close()
    autores = descubrir_autores(_piezas(par.entrada))
    assert {d.tipo for d in datos} == {"obra", "cliente", "contratista", "ubicacion"}
    assert {d.texto for d in autores} >= {AUTOR, EMPRESA_APP}
    restos = verificar(par.entrada, datos, autores)
    assert {
        "cliente",
        "contratista",
        "ubicacion",
        "obra",
        "rif",
        "parte que debia quitarse",
        "nombre definido externo",
        "vinculo externo o ruta de origen",
        "propiedad del documento",
    } <= restos.keys()
    assert verificar(par.salida, datos, autores) == {}


# --- 5. Etiqueta y casos negativos de la CLI ----------------------------------------------------


def test_la_etiqueta_cambia_el_sustituto(tmp_path):
    p = _anonimizar(construir_libro(siembra=SIEMBRA).datos, tmp_path, "--etiqueta", "B")
    assert p.codigo == 0
    valores = _valores(p.salida, data_only=True)
    assert valores["PORTADA!B4"] == "CLIENTE-B."
    assert valores["PORTADA!B5"] == "CONTRATISTA-B."


def test_entrada_inexistente_da_2_y_no_escribe(tmp_path, capsys):
    salida = tmp_path / "salida.xlsx"
    assert main([str(tmp_path / "no_existe.xlsx"), str(salida)]) == 2
    assert "No existe la ENTRADA" in capsys.readouterr().err
    assert not salida.exists()


def test_archivo_que_no_es_xlsx_da_2_y_no_escribe(tmp_path, capsys):
    entrada = tmp_path / "texto.xlsx"
    entrada.write_text("esto no es un libro de Excel", encoding="utf-8")
    salida = tmp_path / "salida.xlsx"
    assert main([str(entrada), str(salida)]) == 2
    assert "no es un ZIP" in capsys.readouterr().err
    assert not salida.exists()


def test_salida_igual_a_la_entrada_da_2_y_no_toca_el_libro(tmp_path):
    datos = construir_libro(siembra=SIEMBRA).datos
    entrada = tmp_path / "libro.xlsx"
    entrada.write_bytes(datos)
    assert main([str(entrada), str(entrada)]) == 2
    assert entrada.read_bytes() == datos


def test_etiqueta_vacia_da_2(tmp_path):
    entrada = tmp_path / "libro.xlsx"
    entrada.write_bytes(construir_libro(siembra=SIEMBRA).datos)
    salida = tmp_path / "salida.xlsx"
    assert main([str(entrada), str(salida), "--etiqueta", "  "]) == 2
    assert not salida.exists()


def test_portada_sin_rotulos_da_1_y_no_escribe(tmp_path, capsys):
    wb = openpyxl.Workbook()
    wb.active.title = "PRESUPUESTO"
    wb.active["A1"] = "Sin portada"
    entrada = tmp_path / "sin_portada.xlsx"
    wb.save(entrada)
    salida = tmp_path / "salida.xlsx"
    assert main([str(entrada), str(salida)]) == 1
    assert "No se encontraron los rotulos" in capsys.readouterr().err
    assert not salida.exists()


def test_entrada_con_dtd_da_1_y_no_escribe(tmp_path, capsys):
    def con_dtd(xml: str) -> str:
        return xml.replace("?>", '?>\n<!DOCTYPE sst [<!ENTITY e "x">]>', 1)

    datos = _reescribir(construir_libro(siembra=SIEMBRA).datos, {"xl/sharedStrings.xml": con_dtd})
    entrada = tmp_path / "con_dtd.xlsx"
    entrada.write_bytes(datos)
    salida = tmp_path / "salida.xlsx"
    assert main([str(entrada), str(salida)]) == 1
    assert "No se anonimiza: xl/sharedStrings.xml" in capsys.readouterr().err
    assert not salida.exists()


def test_hoja_con_el_nombre_del_cliente_falla_la_verificacion_y_borra_la_salida(tmp_path, capsys):
    """El script no renombra hojas (cabecera): una hoja con el nombre del cliente deja el dato en
    `xl/workbook.xml` y en las fórmulas que la citan, la verificación lo ve y la SALIDA se borra."""
    cliente = "Grupo Ficticio Uno"
    apus = (APUS_LIMPIOS[0], APUS_LIMPIOS[1], con(APUS_LIMPIOS[2], hoja=cliente))
    p = _anonimizar(construir_libro(apus, Siembra(cliente=cliente)).datos, tmp_path)
    err = capsys.readouterr().err
    assert p.codigo == 1
    assert "VERIFICACION FALLIDA, SALIDA borrada" in err
    assert "cliente=" in err
    assert not p.salida.exists()
