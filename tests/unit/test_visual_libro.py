"""Réplica visual del libro de la empresa (Sílice 1.2.0, Task 2.2; RF-111, RS2-1).

`replicar_libro` extrae el **aspecto** del libro (relleno, fuente, bordes, alineación, formato
numérico, combinadas, anchos, altos, filas y columnas ocultas, paneles fijos) con coordenadas de
Excel: fila y columna base 1, `ref` sin `$`. La conversión a la instantánea de Univer (base 0)
ocurre en el frontend.

Los libros se construyen aquí con openpyxl (o XlsxWriter, cuando hace falta el valor guardado de una
fórmula) en memoria; ninguno es un libro real (NORMAS §1). Cada valor esperado está escrito a mano,
no sale de llamar al código bajo prueba.
"""

from __future__ import annotations

import hashlib
import io
from datetime import date
from decimal import Decimal
from typing import Any

import pytest
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side

from adapters.hoja_empresa import visual
from adapters.hoja_empresa.guardia import ArchivoRechazado
from adapters.hoja_empresa.visual import (
    LibroDemasiadoGrande,
    LibroVisual,
    replicar_libro,
)
from tests.fixtures.libros_empresa import MEMORIA, libro_limpio

NOMBRE = "libro_sintetico.xlsx"


def _bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _replica(wb: Workbook) -> LibroVisual:
    return replicar_libro(_bytes(wb), NOMBRE)


def _celda(libro: LibroVisual, hoja: str, ref: str) -> Any:
    h = next(h for h in libro.hojas if h.nombre == hoja)
    return next(c for c in h.celdas if c.ref == ref)


def _estilo(libro: LibroVisual, hoja: str, ref: str) -> Any:
    return libro.estilos[_celda(libro, hoja, ref).estilo]


# --- Hojas, combinadas, dimensiones --------------------------------------------------------------


def test_hojas_en_orden_y_visibilidad():
    wb = Workbook()
    wb.active.title = "PORTADA"
    wb.create_sheet("APU-01")
    wb.create_sheet("Oculta").sheet_state = "hidden"
    libro = _replica(wb)
    assert [h.nombre for h in libro.hojas] == ["PORTADA", "APU-01", "Oculta"]
    assert tuple(h.visible for h in libro.hojas) == (True, True, False)


def test_hoja_muy_oculta_no_es_visible():
    wb = Workbook()
    wb.active.title = "A"
    wb.create_sheet("B").sheet_state = "veryHidden"
    libro = _replica(wb)
    assert tuple(h.visible for h in libro.hojas) == (True, False)


def test_combinadas_anchos_altos_y_ocultas():
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    ws["A1"] = "Titulo"
    ws.merge_cells("A1:D1")
    ws.column_dimensions["B"].width = 20
    ws.row_dimensions[3].height = 30
    ws.column_dimensions["C"].hidden = True
    ws.row_dimensions[5].hidden = True
    h = _replica(wb).hojas[0]
    assert h.combinadas == ("A1:D1",)
    assert h.anchos[2] == "20"
    assert h.altos[3] == "30"
    assert h.ocultas_columnas == (3,)
    assert h.ocultas_filas == (5,)


def test_ancho_de_un_grupo_de_columnas_se_reparte_en_cada_una():
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    ws["A1"] = 1
    ws.column_dimensions["D"].width = 12.5
    ws.column_dimensions["D"].min = 4
    ws.column_dimensions["D"].max = 6
    h = _replica(wb).hojas[0]
    assert dict(h.anchos) == {4: "12.5", 5: "12.5", 6: "12.5"}


def test_el_ancho_y_el_alto_por_omision_no_se_escriben():
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    ws["A1"] = 1
    ws.column_dimensions["B"].width = 8.43
    ws.row_dimensions[2].height = 15
    ws.row_dimensions[4].height = 18.75
    h = _replica(wb).hojas[0]
    assert dict(h.anchos) == {}
    assert dict(h.altos) == {4: "18.75"}


def test_paneles_fijos_y_dimensiones_usadas():
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    ws["C7"] = "x"
    ws.freeze_panes = "B5"
    h = _replica(wb).hojas[0]
    assert h.fijas == "B5"
    assert (h.filas, h.columnas) == (7, 3)


def test_sin_paneles_fijos_ni_celdas():
    wb = Workbook()
    wb.active.title = "VACIA"
    h = _replica(wb).hojas[0]
    assert h.fijas is None
    assert (h.filas, h.columnas) == (0, 0)
    assert h.celdas == ()


# --- Celdas --------------------------------------------------------------------------------------


def test_coordenadas_base_1_y_ref_sin_dolar():
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    ws["F12"] = "dato"
    ws["AA3"] = "otro"
    h = _replica(wb).hojas[0]
    f12 = next(c for c in h.celdas if c.ref == "F12")
    assert (f12.ref, f12.fila, f12.columna, f12.mostrado) == ("F12", 12, 6, "dato")
    aa3 = next(c for c in h.celdas if c.ref == "AA3")
    assert (aa3.fila, aa3.columna) == (3, 27)
    assert all("$" not in c.ref for c in h.celdas)
    # Por filas y, dentro de cada fila, de izquierda a derecha.
    assert [c.ref for c in h.celdas] == ["AA3", "F12"]


def test_valor_mostrado_es_el_guardado_y_conserva_la_formula():
    import xlsxwriter

    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True})
    ws = wb.add_worksheet("H")
    ws.write_number("E12", 12.25)
    ws.write_formula("F12", "=E12*2", None, 24.5)
    ws.write_formula("G12", "=E12*3")  # nunca calculado: sin valor guardado (XlsxWriter pone 0)
    wb.close()
    libro = replicar_libro(buf.getvalue(), NOMBRE)
    f12 = _celda(libro, "H", "F12")
    assert f12.mostrado == "24.5"
    assert f12.formula == "=E12*2"
    e12 = _celda(libro, "H", "E12")
    assert e12.mostrado == "12.25"
    assert e12.formula is None


def test_numeros_como_texto_decimal_sin_float_ni_notacion_cientifica():
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    ws["A1"] = 0.1
    ws["A2"] = 100
    ws["A3"] = 1738.407
    ws["A4"] = 0.0000001
    ws["A5"] = 1.5e15
    h = _replica(wb).hojas[0]
    assert [c.mostrado for c in h.celdas] == [
        "0.1",
        "100",
        "1738.407",
        "0.0000001",
        "1500000000000000",
    ]


def test_fecha_texto_y_booleano_mostrados():
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    ws["A1"] = date(2026, 10, 18)
    ws["A2"] = "Rendimiento (u/dia)"
    ws["A3"] = True
    ws["A4"] = False
    h = _replica(wb).hojas[0]
    assert [c.mostrado for c in h.celdas] == [
        "2026-10-18",
        "Rendimiento (u/dia)",
        "TRUE",
        "FALSE",
    ]


def test_solo_van_las_celdas_con_valor_o_estilo_distinto_del_por_omision():
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    ws["A1"] = "con valor"
    ws["B2"].font = Font(bold=True)  # sin valor, con estilo
    ws["C3"].number_format = "General"  # tocada pero igual a la de omision
    ws["D4"].fill = PatternFill("solid", fgColor="FFFFFF00")
    h = _replica(wb).hojas[0]
    assert [c.ref for c in h.celdas] == ["A1", "B2", "D4"]
    assert _celda(_replica(wb), "H", "B2").mostrado == ""


# --- Estilos -------------------------------------------------------------------------------------


def test_estilos_deduplicados():
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    rojo = Font(bold=True, color="FFFF0000")
    for ref in ("A1", "B1", "C1"):
        ws[ref] = "x"
        ws[ref].font = rojo
    ws["A2"] = "normal"
    libro = _replica(wb)
    indices = {_celda(libro, "H", r).estilo for r in ("A1", "B1", "C1")}
    assert len(indices) == 1
    e = libro.estilos[indices.pop()]
    assert e.color == "#FF0000"
    assert e.negrita is True
    assert len(libro.estilos) == len(set(libro.estilos))  # ninguno repetido
    assert len(libro.estilos) == 2  # el rojo y el de A2


def test_atributos_de_fuente_relleno_y_formato_numerico():
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    ws["A1"] = 5
    ws["A1"].font = Font(name="Arial", sz=10.5, b=True, i=True, u="single", color="FF0000FF")
    ws["A1"].fill = PatternFill("solid", fgColor="FFFFFF00")
    ws["A1"].number_format = "#,##0.00"
    ws["A2"] = 6
    ws["A2"].font = Font(name="Arial", sz=10, u="double")
    e = _estilo(_replica(wb), "H", "A1")
    assert (e.fuente, e.tamano) == ("Arial", "10.5")
    assert (e.negrita, e.cursiva, e.subrayado) == (True, True, True)
    assert e.color == "#0000FF"
    assert e.relleno == "#FFFF00"
    assert e.formato_numero == "#,##0.00"
    e2 = _estilo(_replica(wb), "H", "A2")
    assert (e2.tamano, e2.subrayado, e2.negrita, e2.cursiva, e2.relleno) == (
        "10",
        True,
        False,
        False,
        None,
    )
    assert e2.formato_numero == "General"


def test_alineacion_y_ajuste_de_texto():
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    ws["A1"] = "a"
    ws["A1"].alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    ws["A2"] = "b"
    ws["A2"].alignment = Alignment(horizontal="right", vertical="center")
    ws["A3"] = "c"
    ws["A3"].alignment = Alignment(horizontal="general")
    ws["A4"] = "d"
    ws["A4"].alignment = Alignment(horizontal="centerContinuous", vertical="bottom")
    libro = _replica(wb)
    a1, a2, a3, a4 = (_estilo(libro, "H", f"A{i}") for i in (1, 2, 3, 4))
    assert (a1.alineacion_h, a1.alineacion_v, a1.ajustar_texto) == ("center", "top", True)
    assert (a2.alineacion_h, a2.alineacion_v, a2.ajustar_texto) == ("right", "center", False)
    assert (a3.alineacion_h, a3.alineacion_v) == (None, None)
    assert (a4.alineacion_h, a4.alineacion_v) == ("center", "bottom")


def test_bordes_arriba_derecha_abajo_izquierda():
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    ws["B2"] = 1
    ws["B2"].border = Border(
        top=Side(style="thin"),
        right=Side(style="medium"),
        bottom=Side(style="double"),
        left=Side(style=None),
    )
    ws["B3"] = 2
    ws["B3"].border = Border(top=Side(style="mediumDashDot"), bottom=Side(style="hair"))
    libro = _replica(wb)
    assert _estilo(libro, "H", "B2").bordes == ("thin", "medium", "double", None)
    assert _estilo(libro, "H", "B3").bordes == ("dashed", None, "hair", None)


# --- Colores -------------------------------------------------------------------------------------


def test_color_de_tema_con_tinte():
    """Tema 4 (accent1 de Office, #4472C4) con tinte 0.3999755851924192, a mano (ECMA-376 §18.8.19).

    RGB 68,114,196 → normalizados r=0.266667, g=0.447059, b=0.768627.
    max=0.768627, min=0.266667 → L = (max+min)/2 = 0.517647; S = (max-min)/(1-|2L-1|)
    = 0.501961/0.964706 = 0.520325; H = 218.4375° (el máximo es el azul).
    Tinte positivo (más claro): L' = L·(1-tinte) + tinte = 0.517647·0.600024 + 0.399976 = 0.710576.
    C = (1-|2L'-1|)·S = 0.578847·0.520325 = 0.301190; m = L' - C/2 = 0.559981;
    H/60 = 3.640625 está entre 3 y 4 → (r,g,b) = (m, X+m, C+m) con
    X = C·(1-|3.640625 mod 2 - 1|) = 0.301190·0.359375 = 0.108243:
    r = 0.559981·255 = 142.80 → 143 = 0x8F; g = 0.668224·255 = 170.40 → 170 = 0xAA;
    b = 0.861171·255 = 219.60 → 220 = 0xDC. Resultado: #8FAADC.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    ws["A1"] = "x"
    ws["A1"].font = Font(color=Color(theme=4, tint=0.3999755851924192))
    ws["A2"] = "y"
    ws["A2"].fill = PatternFill("solid", fgColor=Color(theme=4, tint=0.3999755851924192))
    libro = _replica(wb)
    assert _estilo(libro, "H", "A1").color == "#8FAADC"
    assert _estilo(libro, "H", "A2").relleno == "#8FAADC"


def test_color_de_tema_con_tinte_negativo_sobre_blanco():
    """Tema 0 (lt1, blanco #FFFFFF) con tinte -0.5, a mano: L = 1, S = 0; tinte negativo (más
    oscuro): L' = L·(1+tinte) = 0.5; sin saturación r = g = b = 0.5·255 = 127.5, que Excel
    muestra como 7F (empate hacia abajo: «Blanco, fondo 1, más oscuro 50 %» es #7F7F7F)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    ws["A1"] = "x"
    ws["A1"].fill = PatternFill("solid", fgColor=Color(theme=0, tint=-0.5))
    assert _estilo(_replica(wb), "H", "A1").relleno == "#7F7F7F"


def test_colores_de_tema_sin_tinte():
    # Indices de tema de SpreadsheetML: 0 = lt1 (blanco), 1 = dk1 (negro), 2 = lt2, 3 = dk2,
    # 4..9 = accent1..accent6 de la paleta de Office por omision.
    esperados = {
        0: "#FFFFFF",
        1: "#000000",
        2: "#E7E6E6",
        3: "#44546A",
        4: "#4472C4",
        5: "#ED7D31",
        6: "#A5A5A5",
        7: "#FFC000",
        8: "#5B9BD5",
        9: "#70AD47",
    }
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    for i, tema in enumerate(esperados, start=1):
        ws.cell(row=i, column=1, value=tema).fill = PatternFill("solid", fgColor=Color(theme=tema))
    libro = _replica(wb)
    for i, (tema, hexa) in enumerate(esperados.items(), start=1):
        assert _estilo(libro, "H", f"A{i}").relleno == hexa, f"tema {tema}"


def test_colores_indexados_de_la_paleta_de_excel():
    # ECMA-376 §18.8.27: 2 = rojo, 4 = azul, 5 = amarillo, 22 = plata C0C0C0, 64 = sistema (negro).
    esperados = {2: "#FF0000", 4: "#0000FF", 5: "#FFFF00", 22: "#C0C0C0", 64: "#000000"}
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    for i, indice in enumerate(esperados, start=1):
        ws.cell(row=i, column=1, value=indice).fill = PatternFill(
            "solid", fgColor=Color(indexed=indice)
        )
    libro = _replica(wb)
    for i, (indice, hexa) in enumerate(esperados.items(), start=1):
        assert _estilo(libro, "H", f"A{i}").relleno == hexa, f"indice {indice}"


def test_color_rgb_argb_pierde_el_alfa():
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    ws["A1"] = "x"
    ws["A1"].font = Font(color="80123ABC")
    assert _estilo(_replica(wb), "H", "A1").color == "#123ABC"


def test_relleno_que_no_es_solido_no_se_replica():
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    ws["A1"] = "x"
    ws["A1"].fill = PatternFill("lightGray", fgColor="FF00FF00")
    assert _estilo(_replica(wb), "H", "A1").relleno is None


# --- Libro sintético de la empresa ---------------------------------------------------------------


def test_libro_sintetico_de_la_empresa():
    c = libro_limpio()
    libro = replicar_libro(c.datos, NOMBRE)
    nombres = [h.nombre for h in libro.hojas]
    for esperada in ("PORTADA", "INSUMOS", "APU-01", "APU-02", "APU-03", "PRESUPUESTO"):
        assert esperada in nombres
    assert libro.nombre == NOMBRE
    assert libro.sha256 == hashlib.sha256(c.datos).hexdigest()
    ref = f"Q{c.filas_partida['APU-01']}"
    monto = _celda(libro, "PRESUPUESTO", ref)
    assert monto.mostrado == "1738.407"
    assert MEMORIA["montos"][1] == Decimal("1738.407")  # la memoria de cálculo del fixture
    assert monto.formula is not None and monto.formula.startswith("=")


def test_libro_sintetico_trae_estilos_en_las_celdas_de_titulo():
    c = libro_limpio()
    libro = replicar_libro(c.datos, NOMBRE)
    portada = next(h for h in libro.hojas if h.nombre == "PORTADA")
    assert any(libro.estilos[x.estilo].negrita for x in portada.celdas)
    assert any(libro.estilos[x.estilo].relleno is not None for x in portada.celdas)


# --- Guardia, límites e inmutabilidad ------------------------------------------------------------


def test_rechaza_lo_que_rechaza_la_guardia():
    with pytest.raises(ArchivoRechazado):
        replicar_libro(b"esto no es un zip", NOMBRE)
    with pytest.raises(ArchivoRechazado):
        replicar_libro(_bytes(Workbook()), "libro.xlsm")


def test_la_guardia_corre_antes_de_abrir_con_openpyxl(monkeypatch):
    import openpyxl

    def no_debe_abrirse(*a, **k):
        raise AssertionError("openpyxl abrió un archivo que la guardia rechaza")

    monkeypatch.setattr(openpyxl, "load_workbook", no_debe_abrirse)
    with pytest.raises(ArchivoRechazado):
        replicar_libro(b"esto no es un zip", NOMBRE)


def test_rechaza_mas_de_200000_celdas(monkeypatch):
    assert visual.MAX_CELDAS_VISUALES == 200_000
    monkeypatch.setattr(visual, "MAX_CELDAS_VISUALES", 10)
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    for i in range(1, 12):  # 11 celdas con valor
        ws.cell(row=i, column=1, value=i)
    with pytest.raises(LibroDemasiadoGrande):
        _replica(wb)


def test_libro_demasiado_grande_es_un_archivo_rechazado():
    assert issubclass(LibroDemasiadoGrande, ArchivoRechazado)


def test_justo_en_el_limite_no_se_rechaza(monkeypatch):
    monkeypatch.setattr(visual, "MAX_CELDAS_VISUALES", 10)
    wb = Workbook()
    ws = wb.active
    ws.title = "H"
    for i in range(1, 11):  # exactamente 10
        ws.cell(row=i, column=1, value=i)
    assert len(_replica(wb).hojas[0].celdas) == 10


def test_no_escribe_sobre_el_original():
    datos = libro_limpio().datos
    antes = hashlib.sha256(datos).hexdigest()
    copia = bytes(datos)
    replicar_libro(datos, NOMBRE)
    assert hashlib.sha256(datos).hexdigest() == antes
    assert datos == copia


def test_el_resultado_es_inmutable_y_con_slots():
    wb = Workbook()
    wb.active["A1"] = "x"
    libro = _replica(wb)
    with pytest.raises(AttributeError):
        libro.nombre = "otro"  # type: ignore[misc]
    assert not hasattr(libro, "__dict__")
    assert not hasattr(libro.hojas[0], "__dict__")
    assert not hasattr(libro.hojas[0].celdas[0], "__dict__")
    assert not hasattr(libro.estilos[0], "__dict__")
