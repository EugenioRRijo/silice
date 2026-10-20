"""«Reajuste de precios» por la API (ola 4 de la sesión 2b; UC-18, RF-71 y RF-72; RM-3).

Contrato de `api/rutas/reajuste.py` (paridad con `ui/paginas/reajuste.py`, que sigue siendo la
fuente de los mensajes). Todas con `CONVERTIR_Y_REAJUSTAR` (rol arma), sesión o clave de API:

- `GET /v1/reajuste/contexto?proyecto=` → versiones (con su motivo y si son revaloradas, las que
  sirven a la comprobación directa) y las opciones de los desplegables.
- `GET /v1/reajuste/parciales?proyecto=&version=` → estudios «parcial al corte» por reemplazar.
- `POST /v1/reajuste/calcular` → tabla lado a lado en texto, avisos y estado; no escribe.
- `POST /v1/reajuste/guardar` → guarda el estudio (confirma) y dice su número.
- `POST /v1/reajuste/excel` → el libro de la versión con las hojas Reajuste y Ejecutado.

Oráculos: los de la prueba de la lógica (`test_ui_reajuste.py`: escalatoria 15300.00,
comprobación 16000.00, diferencia en la otra moneda -0.82), el reconocido de noviembre de la
memoria OB-S2 (75,90) y los mensajes literales de la pantalla clásica, escritos a mano. La fecha
del cálculo se fija sustituyendo `ahora` de la ruta.
"""

from __future__ import annotations

import io
import sqlite3
from collections.abc import Iterator
from contextlib import closing
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import select

from api import caso_web, dependencias
from api.main import app
from api.rutas import reajuste as ruta_reajuste
from core import models
from core.catalog import abrir_sesion, crear_motor
from operacion import config
from operacion.base_datos import preparar_base
from tests.fixtures import control_ob_s2 as ob
from tests.fixtures import reajuste_ob_s2 as rz
from tests.fixtures.claves_api import CABECERA, crear_clave_de_prueba
from tests.fixtures.reajuste_pantalla import CAMPOS, COMPONENTES, EJECUTADO, PERIODOS
from tests.fixtures.reajuste_s1 import NOMBRE, sembrar
from tests.fixtures.sesion_web import (
    MENSAJE_PERMISO,
    cliente,
    con_cuenta,
    copiar_instalacion,
    plantilla_con_cuentas,
    protegida,
)

CIERRE = datetime(2027, 3, 5, 10, 0)  # después del último día de febrero: el estudio cierra
A_MITAD = datetime(2027, 2, 10, 10, 0)  # a mitad del periodo: «parcial al corte»
FALTA_INICIO = "Escriba «Inicio del contrato»: es obligatorio."
PERIODO_FANTASMA = "Período 7: no hay un período con ese número en la tabla."
SIN_VERSION = f"El proyecto {NOMBRE!r} no tiene una versión '009'."
PRIVADO = "texto-que-no-debe-volver-en-un-422"


@pytest.fixture(scope="module")
def _plantilla(tmp_path_factory) -> Path:
    return plantilla_con_cuentas(tmp_path_factory.mktemp("reajuste"))


@pytest.fixture(autouse=True)
def _entorno(_plantilla: Path, monkeypatch: pytest.MonkeyPatch) -> Iterator[Path]:
    """La instalación con cuentas y el caso por defecto con el contrato S1 (001 y 002 revalorada)"""
    monkeypatch.delenv(dependencias.VARIABLE_BASE, raising=False)
    monkeypatch.setattr(ruta_reajuste, "ahora", lambda: CIERRE)
    copiar_instalacion(_plantilla)
    ruta = config.ruta_caso_por_defecto()
    ruta.parent.mkdir(parents=True, exist_ok=True)
    preparar_base(ruta)
    motor = crear_motor(f"sqlite:///{ruta.as_posix()}")
    try:
        with abrir_sesion(motor) as sesion:
            sembrar(sesion)
            sesion.commit()
    finally:
        motor.dispose()
    caso_web.olvidar_todo()
    yield ruta
    caso_web.olvidar_todo()
    dependencias.cerrar_motores()


def _cuerpo(**cambios) -> dict:
    cuerpo = {
        "proyecto": NOMBRE,
        "version": "001",
        "clausula": dict(CAMPOS),
        "componentes": [dict(c) for c in COMPONENTES],
        "periodos": [dict(p) for p in PERIODOS],
        "ejecutado": [dict(e) for e in EJECUTADO],
    }
    cuerpo.update(cambios)
    return cuerpo


def _ana():
    return con_cuenta(app, "ana")


def _post(ruta: str, **cambios):
    navegador, ficha = _ana()
    return navegador.post(f"/v1/reajuste/{ruta}", json=_cuerpo(**cambios), headers=protegida(ficha))


def _filas_de(tabla: str, ruta: Path) -> int:
    with closing(sqlite3.connect(ruta)) as bd:
        return bd.execute(f'SELECT COUNT(*) FROM "{tabla}"').fetchone()[0]


# --- acceso (RM-3) --------------------------------------------------------------------------


def test_sin_sesion_es_401_y_con_clave_de_api_pasa() -> None:
    sin = cliente(app)
    assert sin.get("/v1/reajuste/contexto", params={"proyecto": NOMBRE}).status_code == 401
    assert sin.post("/v1/reajuste/calcular", json=_cuerpo()).status_code == 401
    clave = {CABECERA: crear_clave_de_prueba()}
    r = sin.get("/v1/reajuste/contexto", params={"proyecto": NOMBRE}, headers=clave)
    assert r.status_code == 200


@pytest.mark.parametrize("login", ["rita", "admin"])
def test_solo_quien_arma_reajusta(login: str) -> None:
    navegador, ficha = con_cuenta(app, login)
    pedidos = [
        navegador.get("/v1/reajuste/contexto", params={"proyecto": NOMBRE}),
        navegador.get("/v1/reajuste/parciales", params={"proyecto": NOMBRE, "version": "001"}),
        *(
            navegador.post(f"/v1/reajuste/{r}", json=_cuerpo(), headers=protegida(ficha))
            for r in ("calcular", "guardar", "excel")
        ),
    ]
    for r in pedidos:
        assert (r.status_code, r.json()) == (403, {"detalle": MENSAJE_PERMISO})


def test_el_post_exige_csrf() -> None:
    navegador, _ = _ana()
    r = navegador.post(
        "/v1/reajuste/calcular", json=_cuerpo(), headers={"Origin": "http://127.0.0.1:8000"}
    )
    assert r.status_code == 403


# --- contexto y parciales -------------------------------------------------------------------


def test_el_contexto_trae_las_versiones_y_las_opciones() -> None:
    navegador, _ = _ana()
    r = navegador.get("/v1/reajuste/contexto", params={"proyecto": NOMBRE})
    assert r.status_code == 200
    cuerpo = r.json()
    assert [(v["codigo"], v["revalorada"]) for v in cuerpo["versiones"]] == [
        ("001", False),
        ("002", True),
    ]
    assert cuerpo["versiones"][1]["motivo"] == "revaloracion"
    assert cuerpo["periodicidades"] == ["mensual", "trimestral"]
    assert cuerpo["mecanismos"] == ["escalatoria", "comprobacion", "ambos"]
    assert cuerpo["estructuras"] == ["por_partida", "declarada"]
    assert cuerpo["componentes"] == ["materiales", "equipos", "mano_obra"]


def test_un_proyecto_inexistente_es_404() -> None:
    navegador, _ = _ana()
    r = navegador.get("/v1/reajuste/contexto", params={"proyecto": "NO"})
    assert r.status_code == 404 and "NO" in r.json()["detalle"]


def test_sin_estudios_parciales_la_lista_esta_vacia() -> None:
    navegador, _ = _ana()
    r = navegador.get("/v1/reajuste/parciales", params={"proyecto": NOMBRE, "version": "001"})
    assert (r.status_code, r.json()) == (200, [])


# --- calcular -------------------------------------------------------------------------------


def test_calcular_da_la_tabla_lado_a_lado_y_no_escribe(_entorno: Path) -> None:
    antes = _filas_de("estudio_reajuste", _entorno)
    r = _post("calcular")
    assert r.status_code == 200, r.text
    cuerpo = r.json()
    primera = cuerpo["filas"][0]
    assert (primera["Escalatoria"], primera["Comprobación"], primera["Diferencia (otra)"]) == (
        "15300.00",
        "16000.00",
        "-0.82",
    )
    assert cuerpo["filas"][-1]["Período"] == "Total"
    assert cuerpo["columnas"][:3] == ["Período", "Escalatoria", "Comprobación"]
    assert cuerpo["estado"] == "cerrado" and cuerpo["mensaje_parcial"] == ""
    assert cuerpo["avisos"] == [] or all(isinstance(a, str) for a in cuerpo["avisos"])
    assert _filas_de("estudio_reajuste", _entorno) == antes == 0


def test_a_mitad_de_periodo_el_estudio_es_parcial_y_lo_dice(monkeypatch) -> None:
    monkeypatch.setattr(ruta_reajuste, "ahora", lambda: A_MITAD)
    cuerpo = _post("calcular").json()
    assert cuerpo["estado"] == "parcial al corte"
    # VB-5: con tildes y sin el código de la especificación; la fecha del corte en dd/mm/aaaa.
    assert cuerpo["mensaje_parcial"] == (
        "El último período aún no cierra: el estudio se guardará «parcial al corte» al "
        "10/02/2027 y se reemplaza al cerrar el mes."
    )


@pytest.mark.parametrize(
    ("cambios", "dice"),
    [
        ({"clausula": {**CAMPOS, "fecha_inicio": ""}}, FALTA_INICIO),
        (
            {"clausula": {**CAMPOS, "duracion_meses": "dos años"}},
            "Duración (meses): escriba un número entero de meses.",
        ),
        (
            {"ejecutado": [{"periodo": "7", "partida": "S1-R01", "cantidad": "1", "fuente": "x"}]},
            PERIODO_FANTASMA,
        ),
        ({"version": "009"}, SIN_VERSION),
    ],
)
def test_un_dato_malo_es_422_con_el_mensaje_de_la_pantalla(cambios: dict, dice: str) -> None:
    r = _post("calcular", **cambios)
    assert (r.status_code, r.json()) == (422, {"detalle": dice})


def test_un_cuerpo_invalido_es_422_y_no_devuelve_lo_enviado() -> None:
    navegador, ficha = _ana()
    r = navegador.post(
        "/v1/reajuste/calcular",
        json={"proyecto": 5, "version": PRIVADO},
        headers=protegida(ficha),
    )
    assert r.status_code == 422 and PRIVADO not in r.text


def test_la_responsable_en_blanco_es_la_persona_de_la_sesion() -> None:
    sin_responsable = {**CAMPOS, "responsable": ""}
    r = _post("calcular", clausula=sin_responsable)
    assert r.status_code == 200, r.text  # la sesión de ana pone el nombre


# --- guardar --------------------------------------------------------------------------------


def test_guardar_confirma_el_estudio_con_su_fuente_y_responsable(_entorno: Path) -> None:
    r = _post("guardar", fuente="memoria del estudio", responsable="Ing. Prueba")
    assert r.status_code == 200, r.text
    cuerpo = r.json()
    assert cuerpo["estado"] == "cerrado" and cuerpo["estudio"] > 0
    motor = crear_motor(f"sqlite:///{_entorno.as_posix()}")
    try:
        with abrir_sesion(motor) as sesion:
            estudio = sesion.scalars(select(models.EstudioReajuste)).one()
            assert estudio.id == cuerpo["estudio"]
            assert estudio.fuente == "memoria del estudio"
            assert estudio.responsable == "Ing. Prueba"
    finally:
        motor.dispose()


def test_guardar_sin_responsable_usa_la_sesion(_entorno: Path) -> None:
    assert _post("guardar", fuente="f").status_code == 200
    motor = crear_motor(f"sqlite:///{_entorno.as_posix()}")
    try:
        with abrir_sesion(motor) as sesion:
            assert "ana" in sesion.scalars(select(models.EstudioReajuste.responsable)).one()
    finally:
        motor.dispose()


def test_guardar_un_dato_malo_no_escribe(_entorno: Path) -> None:
    r = _post("guardar", fuente="f", clausula={**CAMPOS, "fecha_inicio": ""})
    assert (r.status_code, r.json()) == (422, {"detalle": FALTA_INICIO})
    assert _filas_de("estudio_reajuste", _entorno) == 0


def test_un_parcial_se_ofrece_y_el_cerrado_lo_reemplaza(monkeypatch, _entorno: Path) -> None:
    monkeypatch.setattr(ruta_reajuste, "ahora", lambda: A_MITAD)
    parcial = _post("guardar", fuente="parcial").json()
    assert parcial["estado"] == "parcial al corte"
    navegador, _ = _ana()
    lista = navegador.get(
        "/v1/reajuste/parciales", params={"proyecto": NOMBRE, "version": "001"}
    ).json()
    assert lista == [{"id": parcial["estudio"], "fecha_calculo": "2027-02-10"}]
    monkeypatch.setattr(ruta_reajuste, "ahora", lambda: CIERRE)
    cerrado = _post("guardar", fuente="cerrado", reemplaza_a=parcial["estudio"]).json()
    assert cerrado["estado"] == "cerrado"
    nuevos = navegador.get(
        "/v1/reajuste/parciales", params={"proyecto": NOMBRE, "version": "001"}
    ).json()
    assert nuevos == []  # el parcial ya tiene quien lo reemplace


# --- Excel ----------------------------------------------------------------------------------


def test_el_excel_trae_las_hojas_reajuste_y_ejecutado() -> None:
    r = _post("excel")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert 'filename="reajuste_001.xlsx"' in r.headers["content-disposition"]
    libro = openpyxl.load_workbook(io.BytesIO(r.content))
    assert {"Reajuste", "Ejecutado"} <= set(libro.sheetnames)
    assert Decimal(str(libro["Reajuste"].cell(2, 2).value)) == Decimal(15300)


# --- lo ejecutado desde las valuaciones (RF-72) ---------------------------------------------


def test_con_valuaciones_el_reconocido_es_el_de_la_memoria_ob_s2(_entorno: Path) -> None:
    motor = crear_motor(f"sqlite:///{_entorno.as_posix()}")
    try:
        with abrir_sesion(motor) as sesion:
            ob.sembrar(sesion)
            ob.declarar_linea_base(sesion)
            rz.sembrar_revaloracion(sesion)
            ob.registrar_v01_y_v02(sesion)
            sesion.commit()
    finally:
        motor.dispose()
    clausula = {
        "fecha_inicio": "2026-11-02",
        "duracion_meses": "6",
        "fecha_referencia_primer_anio": "2026-11-02",
        "periodicidad": "mensual",
        "mecanismo": "comprobacion",
        "estructura": "declarada",
        "fecha_base_indices": "2026-10-01",
        "r": "0",
        "fuente_r": "memoria OB-S2: la clausula ficticia no declara R",
        "fuente": "contrato ficticio de OB-S2",
        "responsable": ob.INSPECTOR,
    }
    componentes = [
        {"componente": c, "serie": s, "peso": str(rz.PESOS[c])} for c, s in rz.SERIES.items()
    ]
    periodos = [
        {
            "desde": "2026-11-01",
            "hasta": "2026-11-30",
            "version_comprobacion": rz.REVALORADA,
            "excepcion": "",
        }
    ]
    r = _post(
        "calcular",
        proyecto=ob.NOMBRE,
        version=ob.CODIGO,
        clausula=clausula,
        componentes=componentes,
        periodos=periodos,
        ejecutado=[],
        con_valuaciones=True,
    )
    assert r.status_code == 200, r.text
    assert r.json()["filas"][0]["Reconocido"] == "75.90"
    guardado = _post(
        "guardar",
        proyecto=ob.NOMBRE,
        version=ob.CODIGO,
        clausula=clausula,
        componentes=componentes,
        periodos=periodos,
        ejecutado=[],
        con_valuaciones=True,
        fuente="valuaciones",
    )
    assert guardado.status_code == 200, guardado.text
    assert _filas_de("cita_valuacion", _entorno) > 0


def test_un_numero_enorme_es_422_y_no_500() -> None:
    enorme = [{**c, "peso": "1E999999999"} for c in COMPONENTES]
    r = _post("calcular", componentes=enorme)
    assert r.status_code == 422, r.text
    assert isinstance(r.json()["detalle"], str)
    g = _post("calcular", clausula={**CAMPOS, "duracion_meses": "1" + "0" * 40})
    assert g.status_code == 422, g.text


def test_el_nombre_del_excel_solo_lleva_caracteres_seguros() -> None:
    assert ruta_reajuste.nombre_excel("001") == "reajuste_001.xlsx"
    peligroso = 'a"b\r\n\u0416/..'
    assert ruta_reajuste.nombre_excel(peligroso) == "reajuste_a_b____...xlsx"
