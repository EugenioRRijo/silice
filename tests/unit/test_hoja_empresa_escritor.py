"""Escritor del libro calcado (Silice 1.5, RF-107 y RF-108; spec F-D3 a F-D5, F-D9, F-D10, §1).

La prueba arma `DatosDevolucion` A MANO (los `ApuCalculado` salen de la memoria de cálculo del
fixture, no de `fusion`), llama a `escribir_devolucion` y reabre el resultado con openpyxl y con
`zipfile`. Así el oráculo es el dato de entrada: lo que el escritor debe escribir es exactamente lo
que se le dio como «valor de Silice».

Paleta de la plantilla (perfil `arenaza`): título 501C1C, encabezado 7F2D2D, subtotal FBEFD6,
editable (ámbar) FFF6DD, total E9C1C1, texto 1F2328, fuente Arial.
"""

from __future__ import annotations

import io
import re
import zipfile

import openpyxl
import pytest

from adapters.hoja_empresa.escritor import escribir_devolucion
from adapters.hoja_empresa.escritor.devolucion import TITULO_DEVUELTO
from adapters.hoja_empresa.lector import leer_libro
from adapters.hoja_empresa.modelo import (
    ApuCalculado,
    DatosDevolucion,
    FilaAlerta,
    FilaHallazgo,
    FilaTraza,
    Rotulo,
)
from tests.fixtures.libros_empresa import (
    APUS_LIMPIOS,
    MEMORIA,
    D,
    Siembra,
    con,
    construir_libro,
    libro_limpio,
)

A1, A2, A3 = APUS_LIMPIOS
TOL = 1e-9
#: Lo que la memoria no lista: base de MO/día, equipos/día y herramientas por unidad.
EXTRA = {
    "APU-01": (D("35"), D("15"), D("0.3325")),
    "APU-02": (D("50"), D("8"), D("0.59375")),
    "APU-03": (D("30"), D("0"), D("0.1425")),
}
HOJAS_ORIGINALES = (
    "PORTADA",
    "MEDICIONES",
    "INSUMOS",
    "APU-01",
    "APU-02",
    "APU-03",
    "PRESUPUESTO",
    "RESUMEN",
    "CRONOGRAMA",
    "PROCURA",
    "PLAN DE PAGOS",
)


def _calculado(hoja: str) -> ApuCalculado:
    m = MEMORIA["componentes"][hoja]
    base, eq_dia, herr = EXTRA[hoja]
    return ApuCalculado(
        hoja=hoja,
        materiales=m["MAT"],
        mano_obra=m["MO"],
        equipos=m["EQ"],
        transporte=m["TR"],
        costo_directo=m["CD"],
        administracion=m["ADM"],
        subtotal=m["SUB"],
        utilidad=m["UTI"],
        precio_unitario=m["PU"],
        base_mo_dia=base,
        equipos_dia=eq_dia,
        herramientas=herr,
    )


def _datos(origen: bytes, **extra) -> DatosDevolucion:
    libro = leer_libro(origen, "recibido.xlsx")
    base = {
        "libro": libro,
        "apus": {h: _calculado(h) for h in ("APU-01", "APU-02", "APU-03")},
        "precios_unitarios": {
            5: MEMORIA["componentes"]["APU-01"]["PU"],
            6: MEMORIA["componentes"]["APU-02"]["PU"],
            7: MEMORIA["componentes"]["APU-03"]["PU"],
        },
        "totales": {k: MEMORIA[k] for k in ("SUBSINIVA", "IVATOT", "TOTALGEN")},
        "metadatos": {
            "sha256_origen": libro.sha256,
            "version": "9.9.9-prueba",
            "fecha": "2026-10-13",
            "producido_por": "Silice",
        },
    }
    base.update(extra)
    return DatosDevolucion(**base)


def _abrir(datos: bytes):
    return (
        openpyxl.load_workbook(io.BytesIO(datos)),
        openpyxl.load_workbook(io.BytesIO(datos), data_only=True),
    )


@pytest.fixture(scope="module")
def devuelto():
    c = libro_limpio()
    datos = _datos(
        c.datos,
        hallazgos=(
            FilaHallazgo(
                "L4", "ERROR", "CD sin equipos (sembrado)", "APU-02!G36", D("154.28"), "corregido"
            ),
            FilaHallazgo("L6", "ADVERTENCIA", "Utilidad a mano", "APU-02!F39", None, "a confirmar"),
            FilaHallazgo("R2", "INFO", "no aplicable", "", None, "informado"),
        ),
        alertas=(
            FilaAlerta(
                "EQ-01",
                "Mezcladora",
                "sin fuente",
                "USD",
                D("30"),
                "2026-07-31",
                "",
                None,
                None,
                None,
                "la FUENTE está vacía",
            ),
        ),
        trazabilidad=(FilaTraza("Parámetro FCAS", "1.9", "INSUMOS!E16", "factor"),),
        reglas_evaluadas=(("L4", "Componente fuera del total"), ("L6", "Parámetro global")),
        reglas_no_evaluadas=(("R11", "Precio fuera de rango", "sin mapeo"),),
    )
    salida = escribir_devolucion(c.datos, datos)
    return c, salida, _abrir(salida)


# --- Hojas y nombres -----------------------------------------------------------------------------


def test_mismas_hojas_sin_aux_y_las_tres_nuevas_al_final(devuelto):
    _, _, (wf, _) = devuelto
    assert tuple(wf.sheetnames) == (
        *HOJAS_ORIGINALES,
        "AUDITORÍA",
        "TRAZABILIDAD",
        "ALERTAS USD-Bs",
    )


def test_cada_apu_publica_sus_nombres_definidos(devuelto):
    _, _, (wf, wv) = devuelto
    nombres = {n: d.attr_text for n, d in wf.defined_names.items()}
    for i in (1, 2, 3):
        for rot in ("MAT", "MO", "EQ", "TR", "CD", "ADM", "SUB", "UTI", "PU", "REND", "ESC"):
            assert f"APU_0{i}_{rot}" in nombres
    # La maqueta es uniforme: el mismo rótulo en la misma celda en los tres APU.
    assert {nombres[f"APU_0{i}_PU"] for i in (1, 2, 3)} == {
        "'APU-01'!$G$52",
        "'APU-02'!$G$52",
        "'APU-03'!$G$52",
    }
    for i in (1, 2, 3):
        hoja = f"APU-0{i}"
        assert abs(wv[hoja]["G52"].value - float(MEMORIA["componentes"][hoja]["PU"])) <= TOL
    # Nombres de la plantilla que se conservan
    assert nombres["FCAS"] == "INSUMOS!$E$16"
    assert nombres["SUBSINIVA"] == "PRESUPUESTO!$Q$10"


def test_presupuesto_lee_por_nombre_con_formula_y_valor_guardado(devuelto):
    _, _, (wf, wv) = devuelto
    assert wf["PRESUPUESTO"]["P5"].value == "=APU_01_PU"
    assert wf["PRESUPUESTO"]["F6"].value == "=APU_02_REND"
    assert wf["CRONOGRAMA"]["D4"].value == "=APU_01_REND"  # el INDIRECT se volvió nombre
    for fila, hoja in ((5, "APU-01"), (6, "APU-02"), (7, "APU-03")):
        assert (
            abs(wv["PRESUPUESTO"][f"P{fila}"].value - float(MEMORIA["componentes"][hoja]["PU"]))
            <= TOL
        )
    assert abs(wv["PRESUPUESTO"]["Q10"].value - float(MEMORIA["SUBSINIVA"])) <= TOL
    assert abs(wv["PRESUPUESTO"]["Q12"].value - float(MEMORIA["TOTALGEN"])) <= TOL


def test_ninguna_formula_sin_valor_guardado_ni_ref_ni_indirect_ni_vinculos(devuelto):
    _, salida, (wf, _) = devuelto
    with zipfile.ZipFile(io.BytesIO(salida)) as z:
        assert not [n for n in z.namelist() if "externalLink" in n]
        sin_valor = []
        for nombre in z.namelist():
            if not nombre.startswith("xl/worksheets/sheet"):
                continue
            xml = z.read(nombre).decode("utf-8")
            for celda in re.finditer(r"<c ([^>]*?)(?:/>|>(.*?)</c>)", xml, re.S):
                atributos, cuerpo = celda.group(1), celda.group(2) or ""
                f = re.search(r"<f>(.*?)</f>", cuerpo, re.S)
                if f is None:
                    continue
                v = re.search(r"<v>(.*?)</v>", cuerpo, re.S)
                # Un <v> vacío solo vale si la fórmula puede dar texto vacío (la guarda
                # `IF($B12="","",…)` de una línea sin código), como exige la aceptación local.
                puede_dar_vacio = "&quot;&quot;" in f.group(1) or '""' in f.group(1)
                if v is None or (v.group(1) == "" and not puede_dar_vacio):
                    sin_valor.append((nombre, atributos))
        assert sin_valor == []
    formulas = [
        c.value
        for ws in wf.worksheets
        for fila in ws.iter_rows()
        for c in fila
        if isinstance(c.value, str) and c.value.startswith("=")
    ]
    assert len(formulas) > 100
    assert not [f for f in formulas if "#REF!" in f.upper() or "INDIRECT(" in f.upper()]
    assert not [n for n, d in wf.defined_names.items() if "[" in (d.attr_text or "")]


def test_propiedades_del_documento_sin_el_autor_original(devuelto):
    c, salida, (wf, _) = devuelto
    assert wf.properties.title == TITULO_DEVUELTO
    assert wf.properties.creator == "Silice"
    with zipfile.ZipFile(io.BytesIO(salida)) as z:
        todo = b"".join(z.read(n) for n in z.namelist() if n.startswith("docProps/"))
    assert b"Autor Ficticio" not in todo
    assert leer_libro(c.datos, "x.xlsx").sha256.encode() in todo
    assert b"9.9.9-prueba" in todo


# --- Paleta, fuente y celdas editables -----------------------------------------------------------


def _relleno(c) -> str | None:
    return c.fill.fgColor.rgb[-6:] if c.fill and c.fill.fill_type else None


def test_paleta_y_fuente_del_apu_regenerado(devuelto):
    _, _, (wf, _) = devuelto
    ws = wf["APU-01"]
    assert (_relleno(ws["A1"]), ws["A1"].font.color.rgb[-6:]) == ("501C1C", "FFFFFF")
    assert _relleno(ws["A9"]) == "7F2D2D"  # encabezado de materiales
    assert _relleno(ws["G14"]) == "FBEFD6"  # subtotal de materiales
    assert _relleno(ws["E10"]) == "FFF6DD"  # cantidad editable (ámbar)
    assert _relleno(ws["F23"]) == "FFF6DD"  # rendimiento editable
    assert _relleno(ws["G52"]) == "E9C1C1"  # precio unitario total
    assert ws["A10"].font.color.rgb[-6:] == "1F2328"
    for ref in ("A1", "A9", "C10", "G52"):
        assert ws[ref].font.name == "Arial"
    # Hojas copiadas: el ámbar de la cantidad sigue en PRESUPUESTO.
    assert _relleno(wf["PRESUPUESTO"]["G5"]) == "FFF6DD"
    # Hojas nuevas con la paleta del libro.
    assert _relleno(wf["AUDITORÍA"]["A1"]) == "501C1C"
    assert wf["AUDITORÍA"]["A1"].font.name == "Arial"


def test_apu_uniforme_con_formulas_de_linea_y_valores(devuelto):
    _, _, (wf, wv) = devuelto
    assert wf["APU-01"]["F10"].value == '=IF($B10="","",IFERROR(VLOOKUP($B10,PRECIOS,4,0),0))'
    assert wv["APU-01"]["F10"].value == 10
    assert wf["APU-01"]["G10"].value == '=IF($B10="",0,N($E10)*N($F10))'
    assert abs(wv["APU-01"]["G10"].value - 2.0) <= TOL
    assert wf["APU-02"]["B12"].value == "MT-03"
    assert abs(wv["APU-02"]["G48"].value - float(MEMORIA["componentes"]["APU-02"]["CD"])) <= TOL
    # El título del APU sigue enlazado a su fila de PRESUPUESTO.
    assert wf["APU-03"]["A2"].value == "=+PRESUPUESTO!D7"
    assert wv["APU-03"]["A2"].value == "Demolición de friso (ficticia)"


# --- Las tres hojas nuevas -----------------------------------------------------------------------


def test_auditoria_cuenta_por_severidad_y_lista_hallazgos_reglas_y_cambios(devuelto):
    _, _, (wf, wv) = devuelto
    ws = wv["AUDITORÍA"]
    resumen = {ws[f"A{f}"].value: ws[f"B{f}"].value for f in range(6, 11)}
    assert resumen == {"CRITICO": 0, "ERROR": 1, "ADVERTENCIA": 1, "INFO": 1, "Total": 3}
    assert wf["AUDITORÍA"]["B7"].value.startswith("=COUNTIF(")
    filas = [tuple(ws.cell(f, c).value for c in range(1, 7)) for f in range(13, 16)]
    assert filas == [
        ("L4", "ERROR", "CD sin equipos (sembrado)", "APU-02!G36", 154.28, "corregido"),
        ("L6", "ADVERTENCIA", "Utilidad a mano", "APU-02!F39", None, "a confirmar"),
        ("R2", "INFO", "no aplicable", None, None, "informado"),
    ]
    textos = [c.value for fila in ws.iter_rows() for c in fila if isinstance(c.value, str)]
    assert any(t.startswith("Los impactos no se suman") for t in textos)
    assert "sin mapeo" in textos
    assert "HOJA-OMITIDA" in textos  # AUX no sale (cambio del calcado)


def test_trazabilidad_y_alertas(devuelto):
    _, _, (_, wv) = devuelto
    traza = [c.value for fila in wv["TRAZABILIDAD"].iter_rows() for c in fila if c.value]
    assert "Parámetro FCAS" in traza and "INSUMOS!E16" in traza
    assert "Hoja que no sale en la devolución" in traza
    assert "AUX" in traza
    al = wv["ALERTAS USD-Bs"]
    assert [al.cell(5, c).value for c in (1, 3, 8, 9, 11)] == [
        "EQ-01",
        "sin fuente",
        "no contrastable",
        "no contrastable",
        "la FUENTE está vacía",
    ]


def test_sin_hallazgos_ni_alertas():
    c = libro_limpio()
    salida = escribir_devolucion(c.datos, _datos(c.datos))
    _, wv = _abrir(salida)
    assert wv["AUDITORÍA"]["A13"].value == "Sin hallazgos."
    assert wv["AUDITORÍA"]["B10"].value == 0
    assert wv["ALERTAS USD-Bs"]["A5"].value == "Sin alertas."


# --- Valores de Silice, no los del libro ---------------------------------------------------------


def test_los_valores_guardados_son_los_de_silice_no_los_recibidos():
    """El libro recibido tiene H1 en APU-02 (P guardado 28,28375); Silice dice 30,2121875."""
    c = construir_libro((A1, con(A2, cd_sin_equipos=True), A3))
    salida = escribir_devolucion(c.datos, _datos(c.datos))
    wf, wv = _abrir(salida)
    assert abs(wv["PRESUPUESTO"]["P6"].value - 30.2121875) <= TOL
    assert abs(wv["APU-02"]["G48"].value - 24.96875) <= TOL  # CD con los equipos
    assert wf["APU-02"]["G48"].value == "=G44+G45+G46+G47"


def test_plantilla_sin_partida_no_se_copia_y_ref_rota_pasa_a_valor():
    apus = (*APUS_LIMPIOS, con(A1, numero=9, hoja="APU-x", en_presupuesto=False))
    c = construir_libro(
        apus,
        Siembra(resumen_ref_rota=True, nombres_externos=("MEDP1",), vinculo_externo="m.xlsx"),
    )
    salida = escribir_devolucion(c.datos, _datos(c.datos))
    wf, _ = _abrir(salida)
    assert "APU-x" not in wf.sheetnames
    assert "MEDP1" not in dict(wf.defined_names.items())
    assert not isinstance(wf["RESUMEN"]["B13"].value, str) or not wf["RESUMEN"][
        "B13"
    ].value.startswith("=")
    textos = [
        c.value for fila in wf["AUDITORÍA"].iter_rows() for c in fila if isinstance(c.value, str)
    ]
    assert "FORMULA-A-VALOR" in textos


def test_misma_entrada_mismos_bytes():
    c = libro_limpio()
    assert escribir_devolucion(c.datos, _datos(c.datos)) == escribir_devolucion(
        c.datos, _datos(c.datos)
    )


# --- Releer la devolución ------------------------------------------------------------------------


def test_la_devolucion_se_relee_con_el_mismo_lector(devuelto):
    _, salida, _ = devuelto
    libro = leer_libro(salida, "devuelto.xlsx")
    for a in libro.apus:
        assert a.guardados[Rotulo.PU] == MEMORIA["componentes"][a.hoja]["PU"]
    assert [(a.codigo_partida, a.unidad) for a in libro.apus] == [
        ("E.311.000.001", "m2"),
        ("E.321.000.002", "m2"),
        ("E.041.000.003", "m2"),
    ]


def test_cabecera_del_apu_en_otras_filas_se_relee_bien():
    """Hallazgo de la auditoría final: el kit tomaba los rótulos de la cabecera de las filas fijas
    4-6 del APU original. Con la cabecera en las filas 3-4, la devolución escribía «Unidad:» junto
    al código de partida y, al releerla, la unidad del APU salía «E.311.000.001». La cabecera debe
    salir por rótulo (código de partida, unidad), como el resto del lector."""
    c = construir_libro(siembra=Siembra(cabecera_compacta=True))
    salida = escribir_devolucion(c.datos, _datos(c.datos))
    libro = leer_libro(salida, "devuelto.xlsx")
    assert [(a.codigo_partida, a.unidad) for a in libro.apus] == [
        ("E.311.000.001", "m2"),
        ("E.321.000.002", "m2"),
        ("E.041.000.003", "m2"),
    ]
