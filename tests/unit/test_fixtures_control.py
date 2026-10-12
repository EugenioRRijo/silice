"""Los datos de entrada de S2-0: el XLSX de OB-S2 y las variantes declaradas de otros dominios.

El oráculo es la memoria `control_ob_s2.md` (§2): V-01 con A 40 y V-02 con A 50 y B 60, inspector
«Ing. Ficticio», costos reales de 1 280,00 USD. Aquí se leen las celdas del archivo versionado y
se comparan con esos literales; nada se calcula con código del sistema.
"""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import openpyxl

from core.contracts.apu import ModalidadManoObra
from tests.fixtures import caso_ti001
from tests.fixtures.control_fronteras import (
    PARTIDA_SIS001,
    PARTIDA_TI001,
    sis001_sin_mano_de_obra,
    ti001_a_destajo,
)

XLSX = Path(__file__).resolve().parents[1] / "fixtures" / "control" / "ob_s2.xlsx"


def _hoja(libro, nombre: str) -> list[list[object]]:
    return [[celda.value for celda in fila] for fila in libro[nombre].iter_rows()]


def test_el_xlsx_de_ob_s2_trae_las_tres_hojas_con_las_columnas_del_spec():
    libro = openpyxl.load_workbook(XLSX)
    assert libro.sheetnames == ["Valuaciones", "Lineas", "Costos"]
    assert _hoja(libro, "Valuaciones")[0] == ["numero", "desde", "hasta", "fecha", "inspector"]
    assert _hoja(libro, "Lineas")[0] == [
        "numero",
        "codigo_partida",
        "cantidad",
        "fecha_medicion",
        "observacion",
    ]
    assert _hoja(libro, "Costos")[0] == [
        "numero",
        "codigo_partida",
        "monto",
        "moneda",
        "fecha",
        "fuente",
    ]


def test_el_xlsx_de_ob_s2_repite_las_valuaciones_y_el_costo_de_la_memoria():
    libro = openpyxl.load_workbook(XLSX)
    assert _hoja(libro, "Valuaciones")[1:] == [
        ["V-01", "2026-11-02", "2026-11-06", "2026-11-09", "Ing. Ficticio"],
        ["V-02", "2026-11-09", "2026-11-13", "2026-11-16", "Ing. Ficticio"],
    ]
    lineas = [fila[:3] for fila in _hoja(libro, "Lineas")[1:]]
    assert lineas == [["V-01", "OB-A", "40"], ["V-02", "OB-A", "50"], ["V-02", "OB-B", "60"]]
    costos = _hoja(libro, "Costos")[1:]
    assert len(costos) == 1 and costos[0][2:5] == ["1280.00", "USD", "2026-11-13"]
    # Todo número viaja como texto (spec §6.2): se convierte a Decimal sin pasar por float.
    assert all(isinstance(fila[2], str) for fila in _hoja(libro, "Lineas")[1:])


def test_la_cantidad_acumulada_del_xlsx_es_la_de_la_memoria_a_90_y_60():
    libro = openpyxl.load_workbook(XLSX)
    acumulado: dict[str, Decimal] = {}
    for _, codigo, cantidad, *_ in _hoja(libro, "Lineas")[1:]:
        acumulado[codigo] = acumulado.get(codigo, Decimal(0)) + Decimal(cantidad)
    assert acumulado == {"OB-A": Decimal(90), "OB-B": Decimal(60)}


def test_ti001_a_destajo_no_deja_jornal_que_escalar_y_no_toca_el_caso():
    original = caso_ti001.COMPOSICIONES[PARTIDA_TI001]
    variante = ti001_a_destajo()
    assert variante.codigo_partida == PARTIDA_TI001
    assert [linea.modalidad for linea in variante.mano_obra] == [ModalidadManoObra.DESTAJO]
    assert [linea.sueldo for linea in variante.mano_obra] == [Decimal(50)]
    # El caso real sigue a jornal: la variante es una copia, no una edición.
    assert all(linea.modalidad is ModalidadManoObra.JORNAL for linea in original.mano_obra)
    assert variante.materiales == original.materiales
    assert variante.rendimiento == original.rendimiento


def test_sis001_sin_mano_de_obra_conserva_el_resto_del_apu():
    variante = sis001_sin_mano_de_obra()
    assert variante.codigo_partida == PARTIDA_SIS001
    assert variante.mano_obra == ()
    assert variante.rendimiento > 0
