"""SF-E4 (RF-59, spec 9.5): el Excel de la hoja con formulas vivas.

El oraculo es un evaluador de formulas de Excel (paquete `formulas`, grupo `dev`): calcula cada
celda del libro escrito y se compara con el `Decimal` del sistema con tolerancia de +-0,005 por
celda
(Excel trabaja en coma flotante). Oraculo adicional: la linea base sigue sumando 1 586,61 USD.
Casos: la linea base (5 partidas), TI-001 (10 partidas, bono por oficio) y un presupuesto sintetico
con mano de obra a destajo y equipos. Nada escribe fuera de `tmp_path`.
"""

from __future__ import annotations

from datetime import date
from decimal import ROUND_HALF_UP, Decimal

import pytest
from openpyxl import load_workbook

from core.budget import elaborar, exportar_excel, plan_secuencial
from core.budget.equilibrio import punto_equilibrio
from core.budget.excel import TablaExportable, exportar_hoja
from core.contracts.apu import (
    ComposicionAPU,
    LineaEquipo,
    LineaManoObra,
    LineaMaterial,
    ModalidadManoObra,
    ParametrosCosto,
)
from core.contracts.dominio import Dominio
from core.contracts.item_computo import ItemComputo, OrigenTipo
from tests.fixtures import apu_linea_base as linea_base
from tests.fixtures import caso_ti001 as ti001
from tests.fixtures.computo_auditado import items_auditados
from tests.fixtures.evaluador_excel import _celda, evaluar

D = Decimal
TOLERANCIA = D("0.005")

HOJAS_SIN_EXTRAS = [
    "APU",
    "Presupuesto",
    "Plan y curva",
    "Punto de equilibrio",
    "Parametros",
    "Auditoria",
    "Trazabilidad",
]


# ---------------------------------------------------------------------------------------------
# El evaluador
# ---------------------------------------------------------------------------------------------


def _cerca(calculado: float, exacto: Decimal) -> bool:
    return abs(D(str(calculado)) - exacto) <= TOLERANCIA


# ---------------------------------------------------------------------------------------------
# Presupuestos de prueba
# ---------------------------------------------------------------------------------------------


@pytest.fixture(scope="module")
def base():
    items = items_auditados()
    composiciones = {apu.codigo_partida: apu for apu in linea_base.APUS_LINEA_BASE}
    return elaborar(
        items,
        composiciones,
        linea_base.PARAMETROS_LINEA_BASE,
        codigo="001",
        fecha=linea_base.FECHA_LINEA_BASE,
        plan=plan_secuencial,
    )


@pytest.fixture(scope="module")
def libro_base(base, tmp_path_factory):
    ruta = tmp_path_factory.mktemp("hoja") / "base.xlsx"
    exportar_hoja(base.presupuesto, base.informe, ruta, linea_base.PARAMETROS_LINEA_BASE)
    return ruta


def _item(codigo: str, cantidad: str) -> ItemComputo:
    return ItemComputo(
        codigo_partida=codigo,
        descripcion=f"Partida {codigo}",
        unidad="m2",
        cantidad=D(cantidad),
        origen_id=f"{codigo}-1",
        origen_tipo=OrigenTipo.MANUAL,
        dominio=Dominio.CIVIL,
    )


def _sintetico():
    """Material, equipo, jornal y destajo con un parametro de cada tipo distinto del defecto."""
    apu = ComposicionAPU(
        codigo_partida="S-01",
        descripcion="Partida sintetica",
        unidad="m2",
        rendimiento=D("12.5"),
        materiales=(LineaMaterial("Tornillo", "pieza", D("3.25"), D("0.4075")),),
        equipos=(LineaEquipo("Taladro", D("2"), D("85.50"), D("0.03")),),
        mano_obra=(
            LineaManoObra("Ayudante a jornal", D("2"), D("11.3")),
            LineaManoObra("Instalador a destajo", D("1"), D("1.15"), ModalidadManoObra.DESTAJO),
        ),
    )
    parametros = ParametrosCosto(
        fcas=D("4.5"), bono_alimentacion=D("2.75"), administracion=D("0.12"), utilidad=D("0.08")
    )
    return elaborar(
        [_item("S-01", "7.333")],
        {"S-01": apu},
        parametros,
        codigo="SIN",
        fecha=date(2026, 10, 5),
        plan=plan_secuencial,
    ), parametros


# ---------------------------------------------------------------------------------------------
# Estructura
# ---------------------------------------------------------------------------------------------


def test_una_hoja_por_pestana_mas_trazabilidad_y_auditoria(libro_base):
    assert load_workbook(libro_base).sheetnames == HOJAS_SIN_EXTRAS


def test_las_hojas_que_core_no_conoce_entran_como_tablas_exportables(base, tmp_path):
    espacio = TablaExportable(
        titulo="Espacio y metrajes",
        columnas=("Ambiente", "Largo (m)"),
        filas=(("Oficina", D("4.250")), ("Bano", D("2.1"))),
    )
    precios = TablaExportable(
        titulo="Precios", columnas=("Insumo", "Ref"), filas=(("Cemento", "CEM019"),)
    )

    ruta = exportar_hoja(
        base.presupuesto,
        base.informe,
        tmp_path / "extras.xlsx",
        linea_base.PARAMETROS_LINEA_BASE,
        hojas_extra=(precios, espacio),
    )

    libro = load_workbook(ruta)
    assert libro.sheetnames == [
        "Espacio y metrajes",
        "APU",
        "Presupuesto",
        "Plan y curva",
        "Punto de equilibrio",
        "Precios",
        "Parametros",
        "Auditoria",
        "Trazabilidad",
    ]
    assert [c.value for c in libro["Espacio y metrajes"][1]] == ["Ambiente", "Largo (m)"]
    assert libro["Espacio y metrajes"]["B2"].value == pytest.approx(4.25)
    assert libro["Precios"]["B2"].value == "CEM019"


@pytest.mark.parametrize("titulo", ["", "x" * 32, "a:b", "a/b", "[a]", "APU", "Presupuesto"])
def test_un_titulo_invalido_o_repetido_se_rechaza(titulo, base, tmp_path):
    tabla = TablaExportable(titulo=titulo, columnas=("a",), filas=())
    with pytest.raises(ValueError, match="titulo"):
        exportar_hoja(
            base.presupuesto,
            base.informe,
            tmp_path / "x.xlsx",
            linea_base.PARAMETROS_LINEA_BASE,
            hojas_extra=(tabla,),
        )


def test_la_trazabilidad_se_escribe_desde_su_tabla(base, tmp_path):
    traza = TablaExportable(
        titulo="Trazabilidad",
        columnas=("Version", "Objeto", "Fuente"),
        filas=(("002", "apu", "memoria B"),),
    )

    ruta = exportar_hoja(
        base.presupuesto,
        base.informe,
        tmp_path / "t.xlsx",
        linea_base.PARAMETROS_LINEA_BASE,
        trazabilidad=traza,
    )

    hoja = load_workbook(ruta)["Trazabilidad"]
    assert [c.value for c in hoja[2]] == ["002", "apu", "memoria B"]


def test_cantidades_y_precios_van_sin_redondear_con_formato_de_dos_decimales(base, tmp_path):
    ruta = exportar_hoja(
        base.presupuesto, base.informe, tmp_path / "f.xlsx", linea_base.PARAMETROS_LINEA_BASE
    )
    hoja = load_workbook(ruta)["Presupuesto"]
    primera = base.presupuesto.partidas[0]

    cantidad = hoja["E2"]
    assert D(str(cantidad.value)) == primera.item.cantidad  # no 2 decimales
    assert cantidad.number_format == "#,##0.00"
    assert hoja["F2"].number_format == "#,##0.00"
    assert hoja["G2"].value.startswith("=")  # total de la fila: formula


# ---------------------------------------------------------------------------------------------
# Las formulas evaluadas coinciden con el sistema
# ---------------------------------------------------------------------------------------------


def test_la_linea_base_sigue_sumando_1586_61(base, libro_base):
    celdas = evaluar(libro_base)
    libro = load_workbook(libro_base)
    total = celdas[("PRESUPUESTO", _celda(libro, "Presupuesto", "TOTAL", "G"))]

    assert _cerca(total, linea_base.TOTAL_PRESUPUESTO_AUDITADO)
    assert _cerca(total, base.presupuesto.total)


def _comprobar_libro(presupuesto, parametros, ruta):
    """Cada celda con formula del libro evaluado contra el `Decimal` del sistema."""
    celdas = evaluar(ruta)
    libro = load_workbook(ruta)
    presupuesto_hoja = libro["Presupuesto"]

    # Presupuesto: PU y total de cada renglon, y el TOTAL
    for fila, partida in enumerate(presupuesto.partidas, start=2):
        assert _cerca(celdas[("PRESUPUESTO", f"F{fila}")], partida.resultado.precio_unitario), fila
        assert _cerca(celdas[("PRESUPUESTO", f"G{fila}")], partida.total), fila
    total = _celda(libro, "Presupuesto", "TOTAL", "G")
    assert _cerca(celdas[("PRESUPUESTO", total)], presupuesto.total)

    # APU: cada resumen de cada bloque
    etiquetas = {
        "Materiales": "materiales",
        "Equipos": "equipos",
        "Mano de obra": "mano_obra",
        "Costo directo": "costo_directo",
        "Con administración": "con_administracion",
        "Precio unitario": "precio_unitario",
    }
    hoja_apu = libro["APU"]
    bloque = -1
    comprobadas = 0
    for fila in hoja_apu.iter_rows(min_col=1, max_col=2):
        a, b = fila[0].value, fila[1].value
        if isinstance(a, str) and " · " in a:
            bloque += 1
        if a in etiquetas and isinstance(b, str) and b.startswith("=") and a != "Materiales":
            esperado = getattr(presupuesto.partidas[bloque].resultado, etiquetas[a])
            assert _cerca(celdas[("APU", f"B{fila[0].row}")], esperado), (bloque, a)
            comprobadas += 1
        elif a == "Materiales" and isinstance(b, str) and b.startswith("="):
            esperado = presupuesto.partidas[bloque].resultado.materiales
            assert _cerca(celdas[("APU", f"B{fila[0].row}")], esperado), (bloque, a)
            comprobadas += 1
    assert comprobadas == 6 * len(presupuesto.partidas)

    # Curva
    for fila, punto in enumerate(presupuesto.curva, start=2):
        assert _cerca(celdas[("PLAN Y CURVA", f"C{fila}")], punto.acumulado), fila
        if presupuesto.total:
            esperado = punto.acumulado / presupuesto.total * 100
            assert abs(D(str(celdas[("PLAN Y CURVA", f"D{fila}")])) - esperado) <= TOLERANCIA

    # Punto de equilibrio
    equilibrio = punto_equilibrio(presupuesto, parametros)
    for etiqueta, esperado in (
        ("Ventas (V)", equilibrio.venta),
        ("Costo variable (CV)", equilibrio.costo_variable),
        ("Costo fijo (CF)", equilibrio.costo_fijo),
        ("Punto de equilibrio (PE)", equilibrio.punto),
    ):
        celda = _celda(libro, "Punto de equilibrio", etiqueta)
        assert _cerca(celdas[("PUNTO DE EQUILIBRIO", celda)], esperado), etiqueta
    assert presupuesto_hoja.max_row > 1


def test_cada_formula_de_la_linea_base_coincide_con_el_sistema(base, libro_base):
    _comprobar_libro(base.presupuesto, linea_base.PARAMETROS_LINEA_BASE, libro_base)


def test_cada_formula_de_ti001_coincide_con_el_sistema(tmp_path):
    resultado = ti001.elaborar_caso()
    ruta = exportar_hoja(
        resultado.presupuesto, resultado.informe, tmp_path / "ti.xlsx", ti001.PARAMETROS
    )

    _comprobar_libro(resultado.presupuesto, ti001.PARAMETROS, ruta)

    celdas = evaluar(ruta)
    total = _celda(load_workbook(ruta), "Presupuesto", "TOTAL", "G")
    assert _cerca(celdas[("PRESUPUESTO", total)], D("50503.27"))


def test_cada_formula_con_destajo_equipo_y_parametros_propios_coincide(tmp_path):
    resultado, parametros = _sintetico()
    ruta = exportar_hoja(resultado.presupuesto, resultado.informe, tmp_path / "s.xlsx", parametros)

    _comprobar_libro(resultado.presupuesto, parametros, ruta)


def test_los_parametros_del_libro_son_los_de_la_version(tmp_path):
    resultado, parametros = _sintetico()
    ruta = exportar_hoja(resultado.presupuesto, resultado.informe, tmp_path / "p.xlsx", parametros)

    hoja = load_workbook(ruta)["Parametros"]
    valores = {fila[0].value: fila[1].value for fila in hoja.iter_rows(min_row=2)}
    assert {k: D(str(v)) for k, v in valores.items()} == {
        "FCAS": parametros.fcas,
        "Bono de alimentación": parametros.bono_alimentacion,
        "Administración": parametros.administracion,
        "Utilidad": parametros.utilidad,
    }


def test_sin_partidas_el_libro_se_escribe_y_el_pe_no_es_aplicable(tmp_path):
    resultado = elaborar(
        [], {}, ParametrosCosto(), codigo="VAC", fecha=date(2026, 10, 5), plan=None
    )

    ruta = exportar_hoja(
        resultado.presupuesto, resultado.informe, tmp_path / "v.xlsx", ParametrosCosto()
    )

    libro = load_workbook(ruta)
    assert libro.sheetnames == HOJAS_SIN_EXTRAS
    mensajes = [c.value for fila in libro["Punto de equilibrio"].iter_rows() for c in fila]
    assert any(isinstance(m, str) and "no aplica" in m for m in mensajes)


# ---------------------------------------------------------------------------------------------
# M-2 (RS-8): el total exacto es el contractual; la nota declara el redondeo de presentacion
# ---------------------------------------------------------------------------------------------


def _tres_de_1_004():
    parametros = ParametrosCosto(administracion=D(0), utilidad=D(0))
    composiciones = {
        c: ComposicionAPU(
            codigo_partida=c,
            descripcion=c,
            unidad="m2",
            rendimiento=D(1),
            materiales=(LineaMaterial("Pieza", "pieza", D(1), D("1.004")),),
        )
        for c in ("A", "B", "C")
    }
    resultado = elaborar(
        [_item(c, "1") for c in "ABC"],
        composiciones,
        parametros,
        codigo="M2",
        fecha=date(2026, 10, 5),
        plan=plan_secuencial,
    )
    return resultado, parametros


def test_m2_el_total_es_el_exacto_y_la_nota_declara_la_diferencia_de_presentacion(tmp_path):
    resultado, parametros = _tres_de_1_004()
    ruta = exportar_hoja(resultado.presupuesto, resultado.informe, tmp_path / "m2.xlsx", parametros)
    libro = load_workbook(ruta)
    celdas = evaluar(ruta)

    exacto = resultado.presupuesto.total  # 3,012
    filas = sum(
        (p.total.quantize(D("0.01"), ROUND_HALF_UP) for p in resultado.presupuesto.partidas), D(0)
    )
    mostrado = exacto.quantize(D("0.01"), ROUND_HALF_UP)
    assert (filas, mostrado) == (D("3.00"), D("3.01"))  # el caso de M-2: no cuadran

    assert _cerca(celdas[("PRESUPUESTO", _celda(libro, "Presupuesto", "TOTAL", "G"))], exacto)
    suma = celdas[("PRESUPUESTO", _celda(libro, "Presupuesto", "Suma de filas mostradas", "G"))]
    total_mostrado = celdas[("PRESUPUESTO", _celda(libro, "Presupuesto", "Total mostrado", "G"))]
    diferencia = celdas[
        ("PRESUPUESTO", _celda(libro, "Presupuesto", "Diferencia de redondeo", "G"))
    ]
    assert (round(suma, 2), round(total_mostrado, 2), round(diferencia, 2)) == (3.00, 3.01, 0.01)
    textos = [c.value for fila in libro["Presupuesto"].iter_rows() for c in fila if c.value]
    assert any(isinstance(t, str) and "total exacto" in t and "redondeo" in t for t in textos)


# ---------------------------------------------------------------------------------------------
# Lo que no cambia: exportar_excel y sus cuatro hojas
# ---------------------------------------------------------------------------------------------


def test_exportar_excel_sigue_con_sus_cuatro_hojas_y_sin_formulas(base, tmp_path):
    ruta = exportar_excel(base.presupuesto, base.informe, tmp_path / "cuatro.xlsx")

    libro = load_workbook(ruta)
    assert libro.sheetnames == ["Presupuesto", "APU", "Curva", "Auditoria"]
    formulas = [
        c.value
        for hoja in libro
        for fila in hoja.iter_rows()
        for c in fila
        if isinstance(c.value, str) and c.value.startswith("=")
    ]
    assert formulas == []
