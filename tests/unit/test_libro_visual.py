"""El libro visual ficticio (`tests/fixtures/libro_visual.py`) y su dibujo con `vista_calcada`.

Por qué existe: las pruebas del visor y de la réplica de estilos necesitan un libro con la **forma
visual** de los de la empresa: títulos en celdas combinadas, encabezados con relleno oscuro y letra
blanca, columnas y filas ocultas, filas separadoras bajas, paneles inmovilizados, color de pestaña,
comentarios y formatos numéricos con paréntesis y raya. Los libros reales son confidenciales y no
se versionan ni derivados de ellos (NORMAS §1; regla de Eugenio, 2026-10-18): este libro es
**inventado** y copia solo el *tipo* de rasgos, no sus datos.

Oráculo a mano (todo literal, sin leer nada del generador):

    PRESUPUESTO  filas 5-7 → montos 120 × 14,5 = 1 740 ; 80 × 30 = 2 400 ; 50 × 4 = 200
                 fila 8 oculta (no suma) ; total sin IVA = 1 740 + 2 400 + 200 = 4 340
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from tests.fixtures.libro_visual import libro_visual
from ui import libros


@pytest.fixture(scope="module")
def datos() -> bytes:
    return libro_visual()


@pytest.fixture(scope="module")
def libro(datos):
    return load_workbook(io.BytesIO(datos))


def test_las_hojas_y_los_rasgos_de_la_portada(libro) -> None:
    assert libro.sheetnames == ["PORTADA", "PRESUPUESTO"]
    portada = libro["PORTADA"]
    assert {str(r) for r in portada.merged_cells.ranges} == {"B2:E2", "C5:D5", "C6:D6"}
    assert portada.row_dimensions[1].height == 6
    assert portada.row_dimensions[2].height == 30
    assert portada["B2"].value == "PRESUPUESTO DE OBRA (EJEMPLO FICTICIO)"
    assert portada["B2"].fill.fgColor.rgb == "FF1F4E79"
    assert portada["B2"].font.b and portada["B2"].font.color.rgb == "FFFFFFFF"
    assert portada["C5"].value == "Cliente Ficticio S.A."


def test_los_rasgos_de_presupuesto(libro) -> None:
    hoja = libro["PRESUPUESTO"]
    assert hoja.freeze_panes == "A5"
    assert hoja.sheet_properties.tabColor.rgb == "FF1F4E79"
    assert [c for c, d in hoja.column_dimensions.items() if d.hidden] == ["H"]
    assert [n for n, d in hoja.row_dimensions.items() if d.hidden] == [8]
    assert hoja.row_dimensions[9].height == 4
    assert {str(r) for r in hoja.merged_cells.ranges} == {"A10:F10"}
    assert hoja["E6"].comment is not None and "ficticio" in hoja["E6"].comment.text
    assert hoja["G5"].number_format == "#,##0.00;(#,##0.00);—"


def test_los_valores_guardados_siguen_la_memoria_a_mano(datos) -> None:
    hoja = load_workbook(io.BytesIO(datos), data_only=True)["PRESUPUESTO"]
    assert [hoja[f"G{n}"].value for n in (5, 6, 7)] == [1740, 2400, 200]
    assert hoja["G10"].value == 4340  # la fila 8 oculta no suma
    formulas = load_workbook(io.BytesIO(datos))["PRESUPUESTO"]
    assert formulas["G5"].value == "=E5*F5" and formulas["G10"].value == "=SUM(G5:G7)"


def test_la_vista_dibuja_combinadas_rellenos_y_formatos(datos) -> None:
    html = libros.vista_calcada(datos, "PRESUPUESTO").html
    assert 'colspan="6"' in html  # A10:F10, el total
    minusculas = html.lower()
    assert "background:#1f4e79" in minusculas  # encabezado oscuro
    assert "color:#ffffff" in minusculas  # letra blanca sobre él
    assert ">4340.00<" in html and ">1740.00<" in html  # dos decimales por el formato


def test_la_vista_muestra_lo_oculto_o_lo_deja_fuera(datos) -> None:
    con = libros.vista_calcada(datos, "PRESUPUESTO", ocultas=True).html
    sin = libros.vista_calcada(datos, "PRESUPUESTO", ocultas=False).html
    assert con.count('title="Columna oculta en el libro"') == 1
    assert con.count('title="Fila oculta en el libro"') == 1
    assert "FILA AUXILIAR OCULTA" in con
    assert "oculta en el libro" not in sin and "FILA AUXILIAR OCULTA" not in sin


def test_la_vista_escapa_el_texto_del_libro(datos) -> None:
    html = libros.vista_calcada(datos, "PRESUPUESTO").html
    assert "&lt;b&gt;m2&lt;/b&gt; &amp; más" in html
    assert "<b>m2</b>" not in html
