"""SF-E4, hallazgo M-1 y RNF-12: ningun texto exportado se interpreta como formula.

Toda celda de texto cuyo primer caracter sea `=`, `+`, `-`, `@`, tabulador o retorno de carro se
escribe como texto (tipo cadena y `quotePrefix`), en las DOS exportaciones (`exportar_excel` y
`exportar_hoja`) y en las tablas genericas que arma la interfaz. Un numero negativo legitimo sigue
siendo numero. Oraculo: el libro releido con `openpyxl` (tipo de cada celda) y el XML crudo de las
hojas (ninguna `<f>` con la carga util).
"""

from __future__ import annotations

import re
import zipfile
from dataclasses import replace
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.budget import elaborar, exportar_excel
from core.budget.excel import TablaExportable, exportar_hoja
from core.contracts.apu import ComposicionAPU, LineaMaterial, ParametrosCosto
from core.contracts.dominio import Dominio
from core.contracts.item_computo import ItemComputo, OrigenTipo
from core.contracts.verificacion import Hallazgo, Severidad

D = Decimal
PREFIJOS = ["=", "+", "-", "@", "\t", "\r"]
CARGA = "cmd|' /C calc'!A0"


def _presupuesto(prefijo: str):
    """Un presupuesto cuyo texto de usuario (insumo, periodo, hallazgo) lleva `prefijo`."""
    malo = prefijo + CARGA
    item = ItemComputo(
        codigo_partida="X-01",
        descripcion=malo,
        unidad="m2",
        cantidad=D(2),
        origen_id="X-1",
        origen_tipo=OrigenTipo.MANUAL,
        dominio=Dominio.CIVIL,
    )
    apu = ComposicionAPU(
        codigo_partida="X-01",
        descripcion=malo,
        unidad="m2",
        rendimiento=D(1),
        materiales=(LineaMaterial(malo, "pieza", D(1), D(3)),),
    )
    resultado = elaborar(
        [item],
        {"X-01": apu},
        ParametrosCosto(),
        codigo="M1",
        fecha=date(2026, 10, 5),
        plan=[(malo, (("X-01", D(1)),))],
    )
    hallazgo = Hallazgo(
        regla="R1", severidad=Severidad.ERROR, descripcion=malo, origen_ids=("X-1",)
    )
    return resultado.presupuesto, replace(resultado.informe, hallazgos=(hallazgo,))


def _celdas_de_texto(ruta: Path, prefijo: str):
    libro = load_workbook(ruta)
    return [
        (hoja.title, c)
        for hoja in libro
        for fila in hoja.iter_rows()
        for c in fila
        if isinstance(c.value, str) and c.value.startswith(prefijo)
    ]


def _formulas_con_carga(ruta: Path) -> list[str]:
    """Las formulas (`<f>`) de las hojas que contienen la carga util: debe no haber ninguna."""
    with zipfile.ZipFile(ruta) as zip_:
        xml = "".join(
            zip_.read(n).decode("utf-8") for n in zip_.namelist() if n.startswith("xl/worksheets/")
        )
    return [f for f in re.findall(r"<f>([^<]*)</f>", xml) if "cmd" in f]


@pytest.mark.parametrize("prefijo", PREFIJOS)
def test_exportar_excel_escribe_como_texto_cada_prefijo_peligroso(prefijo, tmp_path):
    presupuesto, informe = _presupuesto(prefijo)

    ruta = exportar_excel(presupuesto, informe, tmp_path / "m1.xlsx")

    celdas = _celdas_de_texto(ruta, prefijo)
    hojas = {titulo for titulo, _ in celdas}
    assert {"Presupuesto", "APU", "Curva", "Auditoria"} <= hojas
    for titulo, celda in celdas:
        assert celda.data_type == "s", (titulo, celda.coordinate)
        assert celda.quotePrefix, (titulo, celda.coordinate)
    assert _formulas_con_carga(ruta) == []


@pytest.mark.parametrize("prefijo", PREFIJOS)
def test_exportar_hoja_escribe_como_texto_cada_prefijo_peligroso(prefijo, tmp_path):
    presupuesto, informe = _presupuesto(prefijo)
    tabla = TablaExportable(
        titulo="Espacio y metrajes",
        columnas=(prefijo + "col",),
        filas=((prefijo + CARGA,),),
    )

    ruta = exportar_hoja(
        presupuesto, informe, tmp_path / "m1h.xlsx", ParametrosCosto(), hojas_extra=(tabla,)
    )

    celdas = _celdas_de_texto(ruta, prefijo)
    hojas = {titulo for titulo, _ in celdas}
    assert {"Presupuesto", "APU", "Plan y curva", "Auditoria", "Espacio y metrajes"} <= hojas
    for titulo, celda in celdas:
        assert celda.data_type == "s", (titulo, celda.coordinate)
        assert celda.quotePrefix, (titulo, celda.coordinate)
    assert _formulas_con_carga(ruta) == []


def test_un_numero_negativo_legitimo_sigue_siendo_numero_y_un_texto_negativo_sigue_siendo_texto(
    tmp_path,
):
    presupuesto, informe = _presupuesto("")
    tabla = TablaExportable(
        titulo="Cifras",
        columnas=("texto", "decimal", "entero", "formula"),
        filas=(("-5", D("-5.25"), -3, "=1+1"),),
    )

    ruta = exportar_hoja(
        presupuesto, informe, tmp_path / "neg.xlsx", ParametrosCosto(), hojas_extra=(tabla,)
    )

    hoja = load_workbook(ruta)["Cifras"]
    assert (hoja["A2"].data_type, hoja["A2"].value) == ("s", "-5")
    assert (hoja["B2"].data_type, hoja["B2"].value) == ("n", -5.25)
    assert (hoja["C2"].data_type, hoja["C2"].value) == ("n", -3)
    assert (hoja["D2"].data_type, hoja["D2"].value) == ("s", "=1+1")


def test_un_texto_sin_prefijo_peligroso_no_cambia(tmp_path):
    presupuesto, informe = _presupuesto("")  # el texto empieza por "c": normal

    ruta = exportar_excel(presupuesto, informe, tmp_path / "ok.xlsx")

    celda = load_workbook(ruta)["Presupuesto"]["C2"]
    assert celda.value == CARGA
    assert not celda.quotePrefix
