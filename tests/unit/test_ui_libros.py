"""Página «Libro de la empresa» (RF-109, UC-25; spec 1.5 §1 y §5) y su apoyo sin Streamlit.

Dos partes:

1. `ui.libros` (funciones puras): tablas de totales, hallazgos, alertas y trazabilidad; filtros;
   celdas citadas por hoja; nombre del libro devuelto; la vista calcada y su modo ligero.
2. `ui.paginas.libros` con `AppTest`: sin libro, con el libro limpio (resumen, R2 y R7 en INFO,
   descarga), con el libro de H1 (F1 de 154,28 y la celda P6 resaltada en la vista), con un archivo
   rechazado por la guardia, con un fallo inesperado (mensaje fijo, sin el texto del error), con un
   PDF simple sin Excel (L10) y sin permiso.

Subida de archivos: `AppTest` de Streamlit 1.62 maneja `st.file_uploader`
(`at.file_uploader(key=...).set_value((nombre, octetos, mime))`), como en
`test_ui_actualizacion_alta.py`; no hace falta sustituirlo (la página usa además `a.file_uploader`
dentro de columnas, que una sustitución de `st.file_uploader` no alcanzaría).

Oráculos (independientes de lo que se prueba):

* Libro limpio: la memoria de cálculo a mano de `tests/fixtures/libros_empresa.py` (`MEMORIA`):
  SUBSINIVA 4 358,20825, IVATOT 697,31332, TOTALGEN 5 055,52157 -> 4358.21 / 697.31 / 5055.52.
* Libro con H1 (`cd_sin_equipos=True` en APU-02), a mano: el CD guardado omite EQ = 1,59375, así
  que CD = 24,96875 − 1,59375 = 23,375; ADM = 2,3375; SUB = 25,7125; UTI = 2,57125;
  P.U. guardado = 28,28375 frente a 30,2121875 del motor. Diferencia 1,9284375 × 80 = 154,275
  -> F1 impacto 154,28 (ROUND_HALF_UP). Subtotal del libro = 4 358,20825 − 154,275 = 4 203,93325
  (4203.93); IVA del libro 0,16 × 4 203,93325 = 672,62932 (672.63); diferencia de IVA
  0,16 × 154,275 = 24,684 (24.68); diferencia total 154,275 + 24,684 = 178,959 (178.96).
* PDF simple: dos renglones de 50,00 (10 × 5,00 y 20 × 2,50) con TOTAL PRESUPUESTO impreso de
  110,00 -> L10 ERROR, impacto +10,00; con 100,00 no hay hallazgos.
* Tamaños de hoja: lo que openpyxl lee del libro (no `ui.libros`); topes 20 000 filas y 200
  columnas del docstring del módulo (frontera tope / tope + 1).
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook
from streamlit.testing.v1 import AppTest

from adapters.hoja_empresa.modelo import FilaAlerta, FilaHallazgo, FilaTraza
from core.contracts.verificacion import Hallazgo, Severidad
from fusion.proceso import procesar_libro
from tests.fixtures.libros_empresa import (
    APUS_LIMPIOS,
    MEMORIA,
    Siembra,
    con,
    construir_libro,
    libro_limpio,
)
from tests.fixtures.paginas_s3 import con_ficha, preparar
from ui import libros
from ui import sesion as ui_sesion
from ui.sesion import MENSAJE_SIN_PERMISO

DEVOLUCION = date(2026, 10, 13)
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TOL = Decimal("0.005")  # media unidad del segundo decimal: tolerancia de un monto presentado


def _apus_h1():
    return (APUS_LIMPIOS[0], con(APUS_LIMPIOS[1], cd_sin_equipos=True), APUS_LIMPIOS[2])


@pytest.fixture(scope="module")
def limpio():
    c = libro_limpio()
    return c, procesar_libro(c.datos, "libro.xlsx", fecha_devolucion=DEVOLUCION)


@pytest.fixture(scope="module")
def con_h1():
    c = construir_libro(_apus_h1())
    return c, procesar_libro(c.datos, "libro.xlsx", fecha_devolucion=DEVOLUCION)


def _h(regla: str, severidad: Severidad, *origenes: str) -> Hallazgo:
    return Hallazgo(regla=regla, severidad=severidad, descripcion="x", origen_ids=origenes)


def _fila(codigo: str, severidad: str, celda: str = "") -> FilaHallazgo:
    return FilaHallazgo(codigo, severidad, "d", celda, None, "informado")


def _alerta(**cambios) -> FilaAlerta:
    base = {
        "codigo": "MT-01",
        "descripcion": "Cemento (ficticio)",
        "tipo": "USD-Bs",
        "moneda": "USD",
        "precio": Decimal("10"),
        "fecha": "2026-07-31",
        "fuente": "Cotización ficticia",
        "tasa": None,
        "precio_bs": None,
        "variacion": None,
        "mensaje": "m",
    }
    base.update(cambios)
    return FilaAlerta(**base)


# =================================================================================================
# 1. Funciones puras de ui.libros
# =================================================================================================


def test_totales_del_libro_limpio_coinciden_con_la_memoria_a_mano(limpio) -> None:
    _, r = limpio
    filas = libros.filas_totales(r.libro.totales, r.totales_silice)
    assert [f["Concepto"] for f in filas] == ["Subtotal sin IVA", "IVA", "Total general"]
    esperados = (MEMORIA["SUBSINIVA"], MEMORIA["IVATOT"], MEMORIA["TOTALGEN"])
    for fila, memoria in zip(filas, esperados, strict=True):
        assert set(fila) == {"Concepto", "Libro entregado", "Según Silice", "Diferencia"}
        for columna in ("Libro entregado", "Según Silice"):
            assert abs(Decimal(fila[columna]) - memoria) <= TOL
        assert fila["Diferencia"] == "0.00"
    assert [f["Según Silice"] for f in filas] == ["4358.21", "697.31", "5055.52"]


def test_totales_del_libro_con_h1_dan_la_diferencia_calculada_a_mano(con_h1) -> None:
    _, r = con_h1
    filas = libros.filas_totales(r.libro.totales, r.totales_silice)
    assert [(f["Libro entregado"], f["Según Silice"], f["Diferencia"]) for f in filas] == [
        ("4203.93", "4358.21", "154.28"),
        ("672.63", "697.31", "24.68"),
        ("4876.56", "5055.52", "178.96"),
    ]


def test_totales_sin_dato_dicen_no_disponible_y_no_inventan_la_diferencia() -> None:
    filas = libros.filas_totales(
        {"SUBSINIVA": Decimal("10"), "IVATOT": True, "TOTALGEN": "no es número"},
        {"SUBSINIVA": None, "IVATOT": Decimal("1.6")},
    )
    assert filas == [
        {
            "Concepto": "Subtotal sin IVA",
            "Libro entregado": "10.00",
            "Según Silice": "no disponible",
            "Diferencia": "",
        },
        {
            "Concepto": "IVA",
            "Libro entregado": "no disponible",  # un booleano no es un monto
            "Según Silice": "1.60",
            "Diferencia": "",
        },
        {
            "Concepto": "Total general",
            "Libro entregado": "no disponible",
            "Según Silice": "no disponible",
            "Diferencia": "",
        },
    ]


@pytest.mark.parametrize(
    ("libro", "silice", "esperado"),
    [
        ("10.004", "10", "0.00"),  # −0,004 redondea a cero: sin «-0.00»
        ("10.005", "10", "-0.01"),  # −0,005 ya es un céntimo (ROUND_HALF_UP)
        ("10", "10.005", "0.01"),
    ],
)
def test_diferencia_de_totales_en_la_frontera_del_centimo(libro, silice, esperado) -> None:
    [fila, *_] = libros.filas_totales({"SUBSINIVA": libro}, {"SUBSINIVA": silice})
    assert fila["Diferencia"] == esperado


def test_totales_con_pdf_agregan_la_columna_de_lo_impreso() -> None:
    @dataclass
    class Impreso:
        subtotal: Decimal | None = Decimal("4358.2")
        iva: Decimal | None = None
        total: Decimal | None = Decimal("5055.52")

    filas = libros.filas_totales({}, {}, Impreso())
    assert [f["PDF entregado"] for f in filas] == ["4358.20", "no disponible", "5055.52"]
    assert all("PDF entregado" not in f for f in libros.filas_totales({}, {}))


def test_tabla_de_hallazgos_del_libro_limpio_es_r2_y_r7_en_info(limpio) -> None:
    _, r = limpio
    tabla = libros.filas_hallazgos(r.hallazgos_tabla)
    assert [list(f) for f in tabla] == [list(libros.COLUMNAS_HALLAZGOS)] * 2
    assert list(libros.COLUMNAS_HALLAZGOS) == [
        "Código",
        "Severidad",
        "Descripción",
        "Celda",
        "Impacto",
        "Acción",
    ]
    assert [
        (f["Código"], f["Severidad"], f["Celda"], f["Impacto"], f["Acción"]) for f in tabla
    ] == [
        ("R2", "INFO", "", "", "informado"),
        ("R7", "INFO", "", "", "informado"),
    ]


def test_tabla_de_hallazgos_con_h1_trae_f1_con_su_impacto_y_su_celda(con_h1) -> None:
    _, r = con_h1
    tabla = libros.filas_hallazgos(r.hallazgos_tabla)
    resumen = [(f["Código"], f["Severidad"], f["Celda"], f["Impacto"], f["Acción"]) for f in tabla]
    assert resumen == [
        ("F1", "ERROR", "PRESUPUESTO!P6", "154.28", "corregido"),
        ("L4", "ERROR", "APU-02!G36", "154.28", "corregido"),
        ("L9", "ERROR", "PRESUPUESTO!H6", "127.50", "corregido"),  # 1,59375 × 80
        ("L9", "ERROR", "RESUMEN!B10", "127.50", "corregido"),
        ("R2", "INFO", "", "", "informado"),
        ("R7", "INFO", "", "", "informado"),
    ]


def test_opciones_de_filtro_ordenan_por_gravedad_y_dejan_lo_desconocido_al_final(con_h1) -> None:
    _, r = con_h1
    sev, reglas, hojas = libros.opciones_de_filtro(r.hallazgos_tabla)
    assert sev == ["ERROR", "INFO"]
    assert reglas == ["F1", "L4", "L9", "R2", "R7"]
    assert hojas == sorted(["(sin celda)", "APU-02", "PRESUPUESTO", "RESUMEN"])
    filas = [_fila("A", "ZETA"), _fila("B", "INFO"), _fila("C", "CRITICO"), _fila("D", "ALFA")]
    assert libros.opciones_de_filtro(filas)[0] == ["CRITICO", "INFO", "ALFA", "ZETA"]


def test_filtrar_hallazgos_combina_los_tres_filtros_y_vacio_deja_pasar_todo(con_h1) -> None:
    _, r = con_h1
    filas = r.hallazgos_tabla
    assert libros.filtrar_hallazgos(filas) == list(filas)
    assert [f.codigo for f in libros.filtrar_hallazgos(filas, severidades=["INFO"])] == [
        "R2",
        "R7",
    ]
    solo = libros.filtrar_hallazgos(filas, severidades=["ERROR"], hojas=["PRESUPUESTO"])
    assert [(f.codigo, f.celda) for f in solo] == [
        ("F1", "PRESUPUESTO!P6"),
        ("L9", "PRESUPUESTO!H6"),
    ]
    assert libros.filtrar_hallazgos(filas, reglas=["L9"], hojas=["RESUMEN"])[0].celda == (
        "RESUMEN!B10"
    )
    assert libros.filtrar_hallazgos(filas, reglas=["F1"], hojas=["RESUMEN"]) == []
    assert [f.codigo for f in libros.filtrar_hallazgos(filas, hojas=[libros.SIN_HOJA])] == [
        "R2",
        "R7",
    ]


@pytest.mark.parametrize(
    ("celda", "hoja"),
    [
        ("PRESUPUESTO!P6", "PRESUPUESTO"),
        ("'APU 1'!$G$48", "APU 1"),
        ("'Hoja de O''Neil'!A1:B2", "Hoja de O'Neil"),
        ("APU-02!PU", "APU-02"),  # no es una celda: se toma lo de antes del «!»
        ("", libros.SIN_HOJA),
        ("!A1", libros.SIN_HOJA),
        ("PDF:presupuesto.pdf:p1:total", libros.SIN_HOJA),
    ],
)
def test_hoja_de_una_celda(celda, hoja) -> None:
    assert libros.hoja_de(celda) == hoja


def test_marcas_por_hoja_del_libro_con_h1(con_h1) -> None:
    _, r = con_h1
    assert libros.marcas_por_hoja(r.informe.hallazgos) == {
        "PRESUPUESTO": {
            "P6": ("F1",),
            "H6": ("L9",),
            "I6": ("L9",),
            "J6": ("L9",),
            "K6": ("L9",),
            "L6": ("L9",),
            "Q10": ("L9",),
        },
        "APU-02": {"G40": ("F1",), "G36": ("L4",), "G34": ("L4",)},
        "RESUMEN": {"B10": ("L9",)},
    }


def test_marcas_expanden_rangos_por_filas_sin_repetir_codigos_y_con_tope() -> None:
    hallazgos = [
        _h("L1", Severidad.ERROR, "'Hoja X'!$B$2:A1", "sin hoja", "Hoja X!A1"),
        _h("L1", Severidad.ERROR, "'Hoja X'!A1"),
        _h("L9", Severidad.ERROR, "'Hoja X'!a1"),
    ]
    assert libros.marcas_por_hoja(hallazgos) == {
        "Hoja X": {"A1": ("L1", "L9"), "B1": ("L1",), "A2": ("L1",), "B2": ("L1",)}
    }
    assert libros.MAX_CELDAS_RANGO == 500
    for ultima, cuantas in ((499, 499), (500, 500), (501, 500)):
        marcas = libros.marcas_por_hoja([_h("L1", Severidad.ERROR, f"H!A1:A{ultima}")])
        assert len(marcas["H"]) == cuantas
    assert libros.marcas_por_hoja([_h("R2", Severidad.INFO)]) == {}


def test_hay_tasas_y_filas_de_alertas_sin_tasa_dicen_no_contrastable(limpio) -> None:
    _, r = limpio
    assert r.alertas and not libros.hay_tasas(r.alertas)
    filas = libros.filas_alertas(r.alertas)
    assert list(filas[0]) == [
        "Código",
        "Descripción",
        "Tipo",
        "Moneda",
        "Precio",
        "Fecha",
        "Fuente",
        "Tasa BCV",
        "Precio en Bs",
        "Variación",
        "Mensaje",
    ]
    assert all(f["Tasa BCV"] == f["Precio en Bs"] == "no contrastable" for f in filas)
    sin_fuente = {
        (f["Código"], f["Precio"], f["Fuente"]) for f in filas if f["Tipo"] == "sin fuente"
    }
    assert sin_fuente == {("EQ-01", "30.00", ""), ("MT-03", "0.80", "REFERENCIAL — ajustar")}


@pytest.mark.parametrize(
    ("variacion", "texto"),
    [
        (Decimal("0.05"), "+5.00 %"),
        (Decimal("0.05001"), "+5.00 %"),
        (Decimal("-0.035"), "-3.50 %"),
        (Decimal("0"), "+0.00 %"),
        (None, ""),
    ],
)
def test_filas_de_alertas_con_tasa_formatean_tasa_bs_y_variacion(variacion, texto) -> None:
    alerta = _alerta(tasa=Decimal("36.123456"), precio_bs=Decimal("361.23456"), variacion=variacion)
    assert libros.hay_tasas([_alerta(), alerta])
    [fila] = libros.filas_alertas([alerta])
    assert (fila["Tasa BCV"], fila["Precio en Bs"], fila["Variación"]) == (
        "36.1235",
        "361.23",
        texto,
    )


def test_filas_de_trazabilidad() -> None:
    traza = [FilaTraza("Fecha del libro", "2026-07-31", "PORTADA: «Julio 2026»")]
    assert libros.filas_traza(traza) == [
        {
            "Dato": "Fecha del libro",
            "Valor": "2026-07-31",
            "Origen": "PORTADA: «Julio 2026»",
            "Nota": "",
        }
    ]


@pytest.mark.parametrize(
    ("nombre", "esperado"),
    [
        ("Presupuesto obra.xlsx", "Presupuesto obra - devuelto por Silice.xlsx"),
        ("obra", "obra - devuelto por Silice.xlsx"),
        ("c:/carpeta/sub/obra.v2.xlsx", "obra.v2 - devuelto por Silice.xlsx"),
        ("..\\..\\obra.xlsx", "obra - devuelto por Silice.xlsx"),
        ("CON.xlsx", "_CON - devuelto por Silice.xlsx"),
        (".xlsx", "libro - devuelto por Silice.xlsx"),
        ("", "libro - devuelto por Silice.xlsx"),
        ("...", "libro - devuelto por Silice.xlsx"),
    ],
)
def test_nombre_del_libro_devuelto(nombre, esperado) -> None:
    assert libros.nombre_devuelto(nombre) == esperado


def test_hojas_del_libro_limpio_en_su_orden(limpio) -> None:
    c, _ = limpio
    assert libros.hojas_del_libro(c.datos) == (
        "PORTADA",
        "MEDICIONES",
        "INSUMOS",
        "APU-01",
        "APU-02",
        "APU-03",
        "PRESUPUESTO",
        "RESUMEN",
        "CRONOGRAMA",
        "PROCURA",
        "PLAN DE PAGOS",
        "AUX",
    )


def test_vista_calcada_resalta_lo_citado_y_declara_lo_que_queda_fuera(con_h1) -> None:
    c, r = con_h1
    hoja = load_workbook(io.BytesIO(c.datos), data_only=True)["PRESUPUESTO"]
    filas, columnas = hoja.max_row, hoja.max_column
    assert filas < libros.FILAS_VISTA and columnas < libros.COLUMNAS_VISTA
    ultima = libros._numero_a_letras(columnas)
    siguiente = libros._numero_a_letras(columnas + 1)
    marcas = dict(libros.marcas_por_hoja(r.informe.hallazgos)["PRESUPUESTO"])
    marcas |= {
        f"{ultima}{filas}": ("Z9",),  # la esquina usada: dentro
        f"{siguiente}{filas}": ("Z9",),  # una columna más: fuera
        f"A{filas + 1}": ("Z9",),  # una fila más: fuera
        "AA1": ("Z9",),  # más allá de la columna Z: fuera
    }
    vista = libros.vista_calcada(c.datos, "PRESUPUESTO", marcas)
    assert (vista.filas, vista.columnas, vista.desde, vista.hasta) == (filas, columnas, 1, filas)
    assert vista.aviso == ""
    assert vista.fuera == tuple(
        sorted(
            [
                ("AA1", ("Z9",)),
                (f"A{filas + 1}", ("Z9",)),
                (f"{siguiente}{filas}", ("Z9",)),
            ]
        )
    )
    assert 'title="P6: F1"' in vista.html
    assert 'title="Q10: L9"' in vista.html
    assert f'title="{ultima}{filas}: Z9"' in vista.html
    assert vista.html.count(f"box-shadow:inset 0 0 0 2px {libros.COLOR_MARCA}") == 8
    with pytest.raises(KeyError):
        libros.vista_calcada(c.datos, "NO EXISTE")


def test_vista_calcada_escapa_el_texto_ajeno() -> None:
    c = construir_libro(siembra=Siembra(obra='<script>alert("x")</script> & Cía'))
    vista = libros.vista_calcada(c.datos, "PORTADA")
    assert "<script>" not in vista.html
    assert "&lt;script&gt;alert(&quot;x&quot;)&lt;/script&gt; &amp; Cía" in vista.html


def _libro(celdas: dict[str, object], combinada: str | None = None) -> bytes:
    wb = Workbook()
    hoja = wb.active
    hoja.title = "H"
    for ref, valor in celdas.items():
        hoja[ref] = valor
    if combinada:
        hoja.merge_cells(combinada)
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _con_combinadas(*refs: str) -> bytes:
    """Un libro de dos celdas cuya combinada `A1:A2` se reescribe en el XML como `refs`, sin tocar
    `<dimension>` (que sigue en `A1:A2`): lo único grande es lo que declaran las combinadas."""
    original = _libro({"A1": 1, "A2": 2}, "A1:A2")
    bloque = re.compile(rb"<mergeCells[ >].*?</mergeCells>", re.DOTALL)
    nuevas = f'<mergeCells count="{len(refs)}">'
    nuevas += "".join(f'<mergeCell ref="{ref}"/>' for ref in refs) + "</mergeCells>"
    salida = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as zin, zipfile.ZipFile(salida, "w") as zout:
        for info in zin.infolist():
            datos = zin.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                assert len(bloque.findall(datos)) == 1 and b'<dimension ref="A1:A2"' in datos
                datos = bloque.sub(nuevas.encode(), datos)
            zout.writestr(info, datos)
    return salida.getvalue()


@pytest.mark.parametrize(
    ("celdas", "ligero"),
    [
        ({"A1": 1, "A20000": 2}, False),  # 20 000 filas: el tope
        ({"A1": 1, "A20001": 2}, True),  # una más: modo ligero
        ({"A1": 1, "GR1": 2}, False),  # 200 columnas (GR): el tope
        ({"A1": 1, "GS1": 2}, True),  # 201
    ],
)
def test_abrir_libro_pasa_a_modo_ligero_justo_despues_del_tope(celdas, ligero) -> None:
    assert (libros.TOPE_FILAS, libros.TOPE_COLUMNAS) == (20_000, 200)
    libro = libros.abrir_libro(_libro(celdas))
    try:
        assert bool(getattr(libro, "read_only", False)) is ligero
        vista = libros.vista_calcada(libro, "H", {"A1": ("L1",)})
    finally:
        libro.close()
    assert bool(vista.aviso) is ligero
    assert 'title="A1: L1"' in vista.html and vista.fuera == ()
    if ligero:
        filas = 20_001 if "A20001" in celdas else 1
        columnas = 201 if "GS1" in celdas else 1
        assert f"(esta hoja: {libros._miles(filas)} × {columnas})" in vista.aviso
        assert "20.000 filas" in vista.aviso


@pytest.mark.parametrize(
    ("ref", "ligero"),
    [
        ("A1:A20000", False),  # la combinada llega al tope de filas
        ("A1:A20001", True),  # lo pasa por una
        ("A1:GS1", True),  # pasa el tope de columnas
        ("A1:A2", False),
    ],
)
def test_abrir_libro_mide_las_combinadas_sin_cargar_la_hoja(ref, ligero) -> None:
    libro = libros.abrir_libro(_con_combinadas(ref))
    try:
        assert bool(getattr(libro, "read_only", False)) is ligero
    finally:
        libro.close()


def test_abrir_libro_suma_el_area_de_las_combinadas() -> None:
    """Tres combinadas de 200 × 10 000 = 2 000 000 celdas: cada una dentro del tope, juntas
    6 000 000 > 4 000 000 (`TOPE_AREA_COMBINADA`). (La frontera exacta, 4 000 000, no se prueba:
    si la guardia acierta, openpyxl expande cuatro millones de celdas y la prueba dura minutos.)"""
    assert libros.TOPE_AREA_COMBINADA == 4_000_000
    libro = libros.abrir_libro(_con_combinadas("A1:GR10000", "A1:GR10000", "A1:GR10000"))
    try:
        assert getattr(libro, "read_only", False) is True
    finally:
        libro.close()


# =================================================================================================
# 2. La página con AppTest
# =================================================================================================


def _pagina() -> None:
    from ui.paginas import libros as pagina

    pagina.render()


@pytest.fixture(autouse=True)
def _cerrar_motores():
    yield
    ui_sesion.limpiar_motores()


@pytest.fixture(autouse=True)
def _sin_buscar_componentes(monkeypatch):
    """`AppTest` recorre, en la primera corrida de cada instancia, los metadatos de todos los
    paquetes instalados buscando componentes v2 de Streamlit (unos 2 s en Windows). La página no
    usa ninguno: se omite esa búsqueda para que el archivo quede en segundos."""
    from streamlit.components.v2.component_manager import BidiComponentManager

    monkeypatch.setattr(
        BidiComponentManager, "discover_and_register_components", lambda *_, **__: None
    )


def _abrir(monkeypatch, tmp_path, rol: str = "arma") -> AppTest:
    monkeypatch.setenv("SILICE_CASO", str(tmp_path / "caso.db"))
    ficha = preparar(monkeypatch, tmp_path, rol)
    at = con_ficha(AppTest.from_function(_pagina, default_timeout=120), ficha)
    return at.run()


def _subir(at: AppTest, clave: str, nombre: str, datos: bytes, mime: str) -> AppTest:
    at.file_uploader(key=clave).set_value((nombre, datos, mime))
    at = at.run()
    assert not at.exception, [e.value for e in at.exception]
    return at


def _pulsar(at: AppTest, clave: str) -> AppTest:
    at.button(key=clave).click()
    at = at.run()
    assert not at.exception, [e.value for e in at.exception]
    return at


def _tabla(at: AppTest, columna: str):
    return next(d.value for d in at.dataframe if columna in d.value.columns)


def _metricas(at: AppTest) -> dict[str, str]:
    return {m.label: m.value for m in at.metric}


def _textos(at: AppTest) -> str:
    partes = [e.value for grupo in (at.error, at.warning, at.info, at.success) for e in grupo]
    partes += [c.value for c in at.caption] + [m.value for m in at.markdown]
    partes += [c.value for c in at.code]
    return "\n".join(partes)


def test_sin_libro_invita_a_subir_y_un_archivo_que_no_es_xlsx_se_rechaza(
    monkeypatch, tmp_path
) -> None:
    """Quien revisa también tiene `LEER_PROYECTO`: ve la página entera."""
    at = _abrir(monkeypatch, tmp_path, rol="revisa")
    assert not at.exception, [e.value for e in at.exception]
    assert len(at.file_uploader) == 3
    [aviso] = [i.value for i in at.info]
    assert aviso.startswith("Suba el libro de Excel de la empresa para empezar.")
    assert not at.button and not at.error and not at.metric

    at = _subir(at, "libros_xlsx", "libro.xlsx", b"esto no es un ZIP", MIME_XLSX)
    # El mensaje de la guardia, pasado por `texto_seguro` (los signos de Markdown escapados).
    assert [e.value for e in at.error] == [
        r"El archivo «libro\.xlsx» no es un libro \.xlsx válido \(está dañado o no es de "
        r"Excel\)\."
    ]
    assert not at.button and not at.metric and "libros_verificado" not in at.session_state


def test_libro_limpio_resumen_hallazgos_exactos_y_descarga(monkeypatch, tmp_path, limpio) -> None:
    c, _ = limpio
    at = _abrir(monkeypatch, tmp_path)
    at = _subir(at, "libros_xlsx", "libro de prueba.xlsx", c.datos, MIME_XLSX)
    assert any("Pulse «Auditar»" in x.value for x in at.caption)
    at = _pulsar(at, "libros_auditar")
    metricas = _metricas(at)
    assert metricas["Sin IVA, libro entregado"] == "4358.21"
    assert metricas["Sin IVA, según Silice"] == "4358.21"
    assert metricas["Diferencia (Silice − libro)"] == "0.00"
    assert (metricas["Critico"], metricas["Error"], metricas["Advertencia"], metricas["Info"]) == (
        "0",
        "0",
        "0",
        "2",
    )
    assert [s.value for s in at.success][0] == "Cumple: ningún hallazgo de error ni crítico."
    totales = _tabla(at, "Según Silice")
    assert list(totales["Según Silice"]) == ["4358.21", "697.31", "5055.52"]
    hallazgos = _tabla(at, "Acción")
    assert list(hallazgos.columns) == list(libros.COLUMNAS_HALLAZGOS)
    assert list(zip(hallazgos["Código"], hallazgos["Severidad"], strict=True)) == [
        ("R2", "INFO"),
        ("R7", "INFO"),
    ]
    assert any(x.value.startswith("2 de 2 hallazgos.") for x in at.caption)
    assert any("fecha del libro 31/07/2026" in x.value for x in at.caption)  # «Julio 2026»
    assert not at.get("download_button")

    at = _pulsar(at, "libros_preparar")
    [boton] = at.get("download_button")
    assert boton.proto.label == "Descargar «libro de prueba - devuelto por Silice.xlsx»"
    _, nombre, octetos = at.session_state["libros_devuelto"]
    assert nombre == "libro de prueba - devuelto por Silice.xlsx"
    devuelto = load_workbook(io.BytesIO(octetos), read_only=True)
    try:
        assert {"AUDITORÍA", "TRAZABILIDAD", "ALERTAS USD-Bs"} <= set(devuelto.sheetnames)
    finally:
        devuelto.close()


def test_libro_con_h1_no_cumple_y_la_vista_resalta_la_celda_de_f1(monkeypatch, tmp_path) -> None:
    datos = construir_libro(_apus_h1()).datos
    at = _abrir(monkeypatch, tmp_path)
    at = _subir(at, "libros_xlsx", "obra.xlsx", datos, MIME_XLSX)
    at = _pulsar(at, "libros_auditar")
    metricas = _metricas(at)
    assert metricas["Sin IVA, libro entregado"] == "4203.93"
    assert metricas["Diferencia (Silice − libro)"] == "154.28"
    assert (metricas["Error"], metricas["Info"]) == ("4", "2")
    assert [e.value for e in at.error] == [
        "No cumple: 4 hallazgo(s) de error o crítico(s). Véalos en «Hallazgos» y en la hoja "
        "AUDITORÍA del libro devuelto."
    ]
    f1 = _tabla(at, "Acción").iloc[0]
    assert (f1["Código"], f1["Celda"], f1["Impacto"]) == ("F1", "PRESUPUESTO!P6", "154.28")

    at.segmented_control(key="libros_seccion").set_value("Vista de la hoja")
    at = at.run()
    assert not at.exception, [e.value for e in at.exception]
    assert at.selectbox(key="libros_hoja").value == "PRESUPUESTO"  # la del primer hallazgo
    [html] = [h.proto.body for h in at.get("html")]
    assert 'title="P6: F1"' in html
    propios = _tabla(at, "Acción")
    assert list(propios["Celda"]) == ["PRESUPUESTO!P6", "PRESUPUESTO!H6"]
    assert not at.warning

    # Si dibujar la hoja falla, la página lo dice (sin el texto del error) y sigue entera.
    def rota(*_, **__):
        raise RuntimeError("detalle interno /ruta/secreta")

    monkeypatch.setattr(libros, "vista_calcada", rota)
    at.toggle(key="libros_ocultas").set_value(False)  # otra vista: no está guardada
    at = at.run()
    assert not at.exception, [e.value for e in at.exception]
    assert [w.value for w in at.warning] == [
        "No se pudo dibujar esta hoja; los hallazgos siguen en la tabla."
    ]
    assert [x.value for x in at.code] == ["RuntimeError"]
    assert "/ruta/secreta" not in _textos(at)
    assert at.button(key="libros_preparar")  # el resto de la página sigue


def test_fallo_inesperado_al_auditar_da_el_mensaje_fijo_sin_el_texto_del_error(
    monkeypatch, tmp_path, limpio
) -> None:
    import fusion.proceso
    from ui.paginas import libros as pagina

    def falla(*_, **__):
        raise ZeroDivisionError("detalle interno /ruta/secreta")

    monkeypatch.setattr(fusion.proceso, "procesar_libro", falla)
    c, _ = limpio
    at = _abrir(monkeypatch, tmp_path)
    at = _subir(at, "libros_xlsx", "libro.xlsx", c.datos, MIME_XLSX)
    at = _pulsar(at, "libros_auditar")
    assert [e.value for e in at.error] == [pagina.MENSAJE_FALLO]
    assert [x.value for x in at.code] == ["ZeroDivisionError"]
    assert "/ruta/secreta" not in _textos(at)
    assert not at.metric and "libros_resultado" not in at.session_state


def _pdf_simple(total: str) -> bytes:
    """Una página «Presupuesto» con la cabecera Items/Descripción/Unidad/Cantidad/Precio
    Unitario/Total, dos renglones de 50,00 y la línea TOTAL PRESUPUESTO (`adapters/hoja_empresa/
    pdf.py`)."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    salida = io.BytesIO()
    hoja = canvas.Canvas(salida, pagesize=letter)
    hoja.setFont("Helvetica", 9)
    columnas = (40, 80, 300, 350, 410, 490)
    cabecera = ("Items", "Descripción", "Unidad", "Cantidad", "Precio Unitario", "Total ($)")
    renglones = (
        ("1", "Friso liso ficticio", "m2", "10.00", "5.00", "50.00"),
        ("2", "Pintura ficticia", "m2", "20.00", "2.50", "50.00"),
    )
    hoja.drawString(40, 740, "Presupuesto")
    y = 700
    for fila in (cabecera, *renglones):
        for x, texto in zip(columnas, fila, strict=True):
            hoja.drawString(x, y, texto)
        y -= 18
    hoja.drawString(40, y, "TOTAL PRESUPUESTO")
    hoja.drawString(490, y, total)
    hoja.showPage()
    hoja.save()
    return salida.getvalue()


def test_un_pdf_simple_sin_excel_se_revisa_con_l10(monkeypatch, tmp_path) -> None:
    at = _abrir(monkeypatch, tmp_path)
    at = _subir(at, "libros_pdf", "presupuesto.pdf", _pdf_simple("110.00"), "application/pdf")
    assert any("Pulse «Revisar el PDF»" in x.value for x in at.caption)
    at = _pulsar(at, "libros_auditar_pdf")
    metricas = _metricas(at)
    assert metricas["Renglones del presupuesto"] == "2"
    assert metricas["Renglones de cómputos"] == "0"
    assert metricas["Total impreso"] == "110.00"
    assert (metricas["Error"], metricas["Info"]) == ("1", "0")
    [fila] = _tabla(at, "Acción").to_dict("records")
    assert (fila["Código"], fila["Severidad"], fila["Impacto"], fila["Acción"]) == (
        "L10",
        "ERROR",
        "10.00",  # 110,00 impreso − (50,00 + 50,00)
        "informado",
    )
    assert fila["Descripción"].startswith("suma de renglones ≠ total impreso")
    assert fila["Celda"] == "PDF:presupuesto.pdf:p1:total"
    assert not at.get("download_button")  # un PDF no se devuelve

    # Otro PDF cuyo total cuadra (100,00): el resultado anterior no se muestra y no hay hallazgos.
    at = _subir(at, "libros_pdf", "presupuesto.pdf", _pdf_simple("100.00"), "application/pdf")
    assert any("Pulse «Revisar el PDF»" in x.value for x in at.caption) and not at.metric
    at = _pulsar(at, "libros_auditar_pdf")
    assert "No hay hallazgos." in [x.value for x in at.success]
    assert _metricas(at)["Total impreso"] == "100.00" and not at.error


def test_sin_leer_proyecto_no_ve_la_pagina(monkeypatch, tmp_path) -> None:
    """El administrador no tiene `LEER_PROYECTO` (`seguridad.permisos`, S3 §5.1)."""
    at = _abrir(monkeypatch, tmp_path, rol="administrador")
    assert [e.value for e in at.error] == [MENSAJE_SIN_PERMISO]
    assert not at.file_uploader and not at.info


def test_el_libro_devuelto_se_previsualiza_con_sus_hojas_nuevas(
    monkeypatch, tmp_path, limpio
) -> None:
    """Tras «Preparar el libro devuelto» la página dibuja ese libro (no el recibido): el selector
    de hojas trae las tres nuevas y la vista de AUDITORÍA muestra las reglas evaluadas."""
    c, _ = limpio
    at = _abrir(monkeypatch, tmp_path)
    at = _subir(at, "libros_xlsx", "libro de prueba.xlsx", c.datos, MIME_XLSX)
    at = _pulsar(at, "libros_auditar")
    assert not at.get("html")  # antes de preparar no hay libro devuelto que dibujar
    assert not [s for s in at.selectbox if s.key == "libros_dev_hoja"]

    at = _pulsar(at, "libros_preparar")
    hojas = at.selectbox(key="libros_dev_hoja").options
    assert hojas[-3:] == ["AUDITORÍA", "TRAZABILIDAD", "ALERTAS USD-Bs"]  # al final, tras las suyas
    assert hojas[0] == "PORTADA"
    at.selectbox(key="libros_dev_hoja").set_value("AUDITORÍA")
    at = at.run()
    assert not at.exception, [e.value for e in at.exception]
    [html] = [h.proto.body for h in at.get("html")]
    assert "RESUMEN POR SEVERIDAD" in html and "REGLAS EVALUADAS" in html
    assert "R2 no aplicable" in html  # los dos hallazgos INFO del libro limpio
    assert any(x.value.startswith("Libro devuelto, hoja AUDITORÍA") for x in at.caption)


def test_si_no_se_puede_dibujar_el_libro_devuelto_la_descarga_sigue(
    monkeypatch, tmp_path, limpio
) -> None:
    c, _ = limpio
    at = _abrir(monkeypatch, tmp_path)
    at = _subir(at, "libros_xlsx", "libro de prueba.xlsx", c.datos, MIME_XLSX)
    at = _pulsar(at, "libros_auditar")

    def rota(*_, **__):
        raise RuntimeError("detalle interno /ruta/secreta")

    monkeypatch.setattr(libros, "vista_calcada", rota)
    at = _pulsar(at, "libros_preparar")
    assert [w.value for w in at.warning] == [
        "No se pudo dibujar esta hoja; el libro se descarga igual."
    ]
    assert "/ruta/secreta" not in _textos(at)
    [boton] = at.get("download_button")
    assert boton.proto.label.startswith("Descargar «")
