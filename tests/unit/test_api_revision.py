"""«Revisar un paquete» por la API (T7 de la migración final; UC-23, RF-86, RF-87; S3 §8.3-§10).

Contrato de `api/rutas/revision.py`, la página `ui/paginas/revisar.py` hecha servicio sobre
`aplicacion.revision` (la misma lógica que usa Streamlit):

- `POST /v1/revision/paquete` (multipart `paquete` y, opcional, `lista`): verifica el `.silice`
  (firma, integridad y cadena), registra lo recibido una vez por SHA-256, lo reaudita en memoria a
  la fecha de emisión y responde todo lo que la pantalla clásica muestra: banner de estado y de
  versión, avisos, hallazgos por origen, diferencias de precio, contra la entrega anterior,
  contra una lista del caso, quién cambió qué y los dictámenes anteriores.
- `POST /v1/revision/contra-lista` (`paquete`, `lista`): el contraste con otra lista del caso.
- `POST /v1/revision/dictamen` (`paquete`, `resultado`, `notas`): el dictamen, atado al SHA-256
  del archivo; solo con la sesión de una persona.
- `POST /v1/revision/excel` (`paquete`, `tipo` = presupuesto | revision): el Excel regenerado aquí.
- `GET /v1/revision/plantilla` y `POST /v1/revision/excel-externo` (`presupuesto`): la pestaña
  «Presupuesto en Excel».

Todas exigen `REVISAR_PAQUETE` (rol revisa) con sesión y CSRF, o la clave de API; ninguna se
guarda en el navegador (`no-store`). Oráculos: los mensajes literales de la pantalla clásica y de
`paquete.verificar` / `paquete.formato` (copiados a mano), la memoria del ejemplo mínimo
(`tests/fixtures/canonico/ejemplo_v1.py`: total 836,7975), las filas leídas con `sqlite3` directo
y el SHA-256 con `hashlib`. El router se monta en una aplicación de prueba con lo que pone
`api.main` alrededor de cada ruta (la línea de `api/main.py` la aplica el orquestador).
"""

from __future__ import annotations

import copy
import hashlib
import io
import sqlite3
from collections.abc import Iterator
from contextlib import closing
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from sqlalchemy import select

from api import caso_web, dependencias, main
from api.proteccion import proteger
from api.rutas import auditorias
from api.rutas import revision as ruta_revision
from api.rutas import sesion as rutas_sesion
from auditoria.reauditoria import Servicios
from core import models
from core.budget.canonico import json_canonico
from core.catalog import abrir_sesion, crear_motor
from operacion import config
from operacion.base_datos import preparar_base
from operacion.instalacion import sesion_de_instalacion
from seguridad import firma
from seguridad.modelo import EmpresaRegistrada
from tests.fixtures.canonico.ejemplo_v1 import EJEMPLO_V1
from tests.fixtures.canonico_minimo import sembrar_minimo
from tests.fixtures.claves_api import CABECERA, crear_clave_de_prueba
from tests.fixtures.instalacion import AHORA
from tests.fixtures.paquetes import EMPRESA, con_manifiesto, paquete
from tests.fixtures.sesion_web import (
    MENSAJE_PERMISO,
    cliente,
    con_cuenta,
    copiar_instalacion,
    plantilla_con_cuentas,
    protegida,
)

PAR = firma.generar_par()
TOTAL = "836.7975"  # 20 m3 x 41,839875 USD/m3, la memoria del ejemplo mínimo
VERIFICADO = (
    "Verificado: emitido por Contratista Demo, C.A. el 2026-11-10. No se alteró desde que se firmó."
)
NO_VERIFICADO_SIN_FIRMA = (
    "No verificado: dice venir de Contratista Demo, C.A., pero no se pudo comprobar. El paquete "
    "no está firmado: las huellas solo prueban que no se dañó por accidente, no que nadie lo "
    "cambió. Se audita igual."
)
RECHAZADO_ZIP = (
    "Rechazado: El archivo no es un paquete de Silice válido (el ZIP está dañado). Pida uno nuevo."
)
PRIMERA = "Versión 005 del 2026-11-10; es la primera que recibe de este proyecto."
SIN_ANTERIOR = "Esta instalación no recibió una entrega anterior de este proyecto."
PRECIOS_OK = "Los precios y totales que declara el paquete coinciden con los recalculados."
REGLAS = "Reglas corridas a la fecha de emisión (2026-11-10): R1-R7, R11-R12."
NOTAS_FALTAN = "Un dictamen con observaciones o de rechazo necesita notas."
GUARDADO = "Dictamen guardado, atado al SHA-256 de este archivo."
EXCEL_EXTERNO = "No verificado: un Excel no trae cadena ni firma. Se audita igual."
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SIN_DICTAMEN = (
    "Este paquete no quedó registrado (está alterado o no trae los datos mínimos): no admite "
    "dictamen. Pida al emisor uno nuevo."
)
MENSAJE_TOPE = (
    "El archivo pesa más de 20 MiB, el máximo de un paquete. Pida al emisor un paquete más pequeño."
)


def _app() -> FastAPI:
    """La sesión (`/v1/sesion`), el router de revisión, los manejadores `{"detalle": ...}` de
    `api.main` y `proteger`; los servicios de la reauditoría, sin referencia ni tasas."""
    prueba = FastAPI()
    prueba.add_exception_handler(HTTPException, main.manejar_http_exception)
    prueba.add_exception_handler(LookupError, main.manejar_lookup_error)
    prueba.add_exception_handler(ValueError, main.manejar_value_error)
    prueba.include_router(rutas_sesion.router)
    prueba.include_router(ruta_revision.router)
    prueba.dependency_overrides[auditorias.servicios_de_la_api] = Servicios
    proteger(prueba)
    return prueba


APP = _app()


@pytest.fixture(scope="module")
def _plantilla(tmp_path_factory) -> Path:
    return plantilla_con_cuentas(tmp_path_factory.mktemp("revision"))


@pytest.fixture(autouse=True)
def _entorno(_plantilla: Path, monkeypatch: pytest.MonkeyPatch) -> Iterator[Path]:
    """La instalación con cuentas, la clave de «Contratista Demo, C.A.» registrada y el caso por
    defecto con MIN-001 sembrado (su «Lista de prueba»)."""
    monkeypatch.delenv(dependencias.VARIABLE_BASE, raising=False)
    copiar_instalacion(_plantilla)
    with sesion_de_instalacion() as bd:
        bd.add(
            EmpresaRegistrada(
                empresa=EMPRESA,
                rif="J-00000000-0",
                publica=PAR.publica,
                huella=firma.huella_completa(PAR.publica),
                estado="activa",
                comprobada_con="Ana, por teléfono",
                metodo="telefono",
                fecha_comprobacion=AHORA,
                registrada_por="admin",
            )
        )
        bd.commit()
    ruta = config.ruta_caso_por_defecto()
    ruta.parent.mkdir(parents=True, exist_ok=True)
    preparar_base(ruta)
    motor = crear_motor(f"sqlite:///{ruta.as_posix()}")
    try:
        with abrir_sesion(motor) as caso:
            sembrar_minimo(caso)
            caso.commit()
    finally:
        motor.dispose()
    caso_web.olvidar_todo()
    yield ruta
    caso_web.olvidar_todo()
    dependencias.cerrar_motores()


def _documento() -> bytes:
    return json_canonico(copy.deepcopy(EJEMPLO_V1))


def _firmado(**mas) -> bytes:
    return paquete(par=PAR, presupuesto=_documento(), anexos={}, **mas)


def _rita() -> tuple[TestClient, str]:
    return con_cuenta(APP, "rita")


def _subir(navegador, ficha, datos: bytes, ruta: str = "/v1/revision/paquete", **campos):
    return navegador.post(
        ruta,
        files={"paquete": ("recibido.silice", datos, "application/zip")},
        data=campos,
        headers=protegida(ficha),
    )


def _filas(tabla: str) -> list[sqlite3.Row]:
    with closing(sqlite3.connect(config.ruta_caso_por_defecto())) as bd:
        bd.row_factory = sqlite3.Row
        return list(bd.execute(f"SELECT * FROM {tabla} ORDER BY id"))  # noqa: S608 (nombre fijo)


# --- acceso ---------------------------------------------------------------------------------


def test_sin_sesion_es_401_y_con_clave_de_api_pasa() -> None:
    sin = cliente(APP)
    archivos = {"paquete": ("recibido.silice", _firmado(), "application/zip")}
    assert sin.post("/v1/revision/paquete", files=archivos).status_code == 401
    assert sin.get("/v1/revision/plantilla").status_code == 401
    r = sin.post(
        "/v1/revision/paquete", files=archivos, headers={CABECERA: crear_clave_de_prueba()}
    )
    assert r.status_code == 200, r.text
    assert _filas("paq_recibido")[0]["recibido_por"] == "api (sin clave)"


@pytest.mark.parametrize("login", ["ana", "admin"])
def test_solo_quien_revisa_revisa(login: str) -> None:
    """REVISAR_PAQUETE es del rol revisa: arma y el administrador reciben el 403 del rol."""
    navegador, ficha = con_cuenta(APP, login)
    pedidos = [
        _subir(navegador, ficha, _firmado()),
        _subir(navegador, ficha, _firmado(), "/v1/revision/contra-lista", lista="1"),
        _subir(navegador, ficha, _firmado(), "/v1/revision/dictamen", resultado="conforme"),
        _subir(navegador, ficha, _firmado(), "/v1/revision/excel", tipo="presupuesto"),
        navegador.get("/v1/revision/plantilla"),
    ]
    for r in pedidos:
        assert (r.status_code, r.json()) == (403, {"detalle": MENSAJE_PERMISO})
    assert _filas("paq_recibido") == []


def test_los_post_exigen_csrf() -> None:
    navegador, _ = _rita()
    r = navegador.post(
        "/v1/revision/paquete",
        files={"paquete": ("recibido.silice", _firmado(), "application/zip")},
        headers={"Origin": "http://127.0.0.1:8000"},
    )
    assert r.status_code == 403
    assert _filas("paq_recibido") == []


# --- analizar un paquete --------------------------------------------------------------------


def test_un_paquete_firmado_se_verifica_se_registra_y_se_reaudita() -> None:
    navegador, ficha = _rita()
    datos = _firmado()
    r = _subir(navegador, ficha, datos)
    assert r.status_code == 200, r.text
    assert r.headers["cache-control"] == "no-store"
    cuerpo = r.json()
    sha = hashlib.sha256(datos).hexdigest()
    assert cuerpo["sha256"] == sha
    assert (cuerpo["estado"], cuerpo["color"], cuerpo["legible"]) == ("verificado", "verde", True)
    assert cuerpo["banner"] == VERIFICADO
    assert cuerpo["banner_version"] == PRIMERA
    assert cuerpo["avisos"] == [] and cuerpo["aviso_armo"] is None
    assert cuerpo["ayuda_registro"] is None and cuerpo["retroactiva"] is None
    re = cuerpo["reauditoria"]
    assert re["termino"] is True and re["avisos"] == []
    assert Decimal(re["total_recalculado"]) == Decimal(TOTAL)  # exacto, como texto
    assert re["total_recalculado_texto"] == "836.80"  # como la clásica: dos decimales
    assert re["reglas"] == REGLAS
    assert re["precios_coinciden"] is True and re["mensaje_precios"] == PRECIOS_OK
    assert re["diferencias"] == []
    # Sin referencia MaPreX ni tasas (servicios vacíos), R11 da un único INFO con los tres
    # insumos, que el emisor no declaró (como `test_ui_revision`, conjunto exacto)
    assert [lote["titulo"] for lote in re["hallazgos"]] == [
        "El emisor no los declaró",
        "Los declaró y ya no aparecen",
        "En común",
    ]
    [r11] = re["hallazgos"][0]["hallazgos"]
    assert (r11["regla"], r11["severidad"]) == ("R11", "INFO")
    assert r11["descripcion"].startswith("3 insumos no contrastables")
    assert re["hallazgos"][1]["hallazgos"] == [] and re["hallazgos"][2]["hallazgos"] == []
    assert cuerpo["entrega_anterior"] is None and cuerpo["sin_entrega_anterior"] == SIN_ANTERIOR
    assert cuerpo["cambios_titulo"] == "Quién cambió qué"
    assert [c["version"] for c in cuerpo["cambios"]] == ["004", "005"]
    assert cuerpo["cambios"][0]["responsable"] == "Ana <ana>"
    assert cuerpo["dictamenes"] == []
    assert cuerpo["resultados"] == ["conforme", "con observaciones", "rechazado"]
    assert cuerpo["admite_dictamen"] is True and cuerpo["sin_dictamen"] is None
    [fila] = _filas("paq_recibido")
    assert (fila["huella_archivo"], fila["recibido_por"], fila["estado"]) == (
        sha,
        "Rita <rita>",
        "verificado",
    )
    guardado = Path(fila["ruta"])
    assert guardado.parent == config.carpeta_datos() / "paquetes"
    assert guardado.read_bytes() == datos and guardado.name == f"{sha}-recibido.silice"


def test_el_mismo_paquete_otra_vez_no_se_duplica_y_el_banner_lo_dice() -> None:
    navegador, ficha = _rita()
    datos = _firmado(destinatario="Obra Norte")
    _subir(navegador, ficha, datos)
    cuerpo = _subir(navegador, ficha, datos).json()
    hoy = date.today().isoformat()  # el reloj de la sesión, como `ui_sesion.ahora`
    assert cuerpo["banner_version"] == (
        f"Versión 005 del 2026-11-10; ya recibió este mismo paquete el {hoy}."
    )
    # «Ya recibió…» ya lo dice el banner de versión: no se repite entre los avisos
    assert cuerpo["avisos"] == ["Este paquete se emitió para Obra Norte."]
    assert len(_filas("paq_recibido")) == 1


def test_un_paquete_sin_firma_se_audita_igual() -> None:
    navegador, ficha = _rita()
    cuerpo = _subir(navegador, ficha, paquete(par=None, presupuesto=_documento(), anexos={})).json()
    assert (cuerpo["estado"], cuerpo["color"]) == ("no_verificado", "ambar")
    assert cuerpo["banner"] == NO_VERIFICADO_SIN_FIRMA
    assert Decimal(cuerpo["reauditoria"]["total_recalculado"]) == Decimal(TOTAL)


def test_una_empresa_no_registrada_lleva_la_ayuda_para_registrarla() -> None:
    navegador, ficha = _rita()
    ajeno = paquete(par=firma.generar_par(), presupuesto=_documento(), anexos={})
    cuerpo = _subir(navegador, ficha, ajeno).json()
    assert cuerpo["estado"] == "no_verificado"
    assert cuerpo["banner"].endswith(
        "Su clave no está registrada en esta instalación. Se audita igual."
    )
    assert cuerpo["ayuda_registro"] == (
        "Para que sus paquetes salgan verificados, el administrador debe registrar la clave "
        "de esa empresa (Administración), comprobando su huella por otro canal."
    )


def test_un_archivo_que_no_es_un_paquete_se_rechaza_con_el_mensaje_de_la_clasica() -> None:
    navegador, ficha = _rita()
    r = _subir(navegador, ficha, b"esto no es un zip")
    assert r.status_code == 200
    cuerpo = r.json()
    assert (cuerpo["estado"], cuerpo["color"], cuerpo["legible"]) == ("rechazado", "rojo", False)
    assert cuerpo["banner"] == RECHAZADO_ZIP
    assert cuerpo["banner_version"] is None and cuerpo["reauditoria"] is None
    assert _filas("paq_recibido") == []
    assert not (config.carpeta_datos() / "paquetes").exists()


def test_un_paquete_alterado_dice_que_archivo_cambio() -> None:
    from tests.fixtures.paquetes import reempaquetar

    navegador, ficha = _rita()
    alterado = reempaquetar(_firmado(), {"informe.json": b'{"otro":1}'})
    cuerpo = _subir(navegador, ficha, alterado).json()
    assert (cuerpo["estado"], cuerpo["color"]) == ("alterado", "rojo")
    assert cuerpo["banner"] == (
        "Alterado: informe.json cambió después de firmarse. No confíe en este paquete; pida uno "
        "nuevo."
    )


def test_un_documento_con_la_forma_equivocada_es_422_con_el_mensaje_y_sin_traza() -> None:
    """La clásica lo pinta con `st.error(str(error))`: aquí es el `detalle` del 422."""
    navegador, ficha = _rita()
    r = _subir(navegador, ficha, paquete(par=PAR, anexos={}))  # presupuesto.json de juguete
    assert r.status_code == 422
    assert "Traceback" not in r.text and str(config.carpeta_datos()) not in r.text


def test_quien_revisa_y_figura_como_emisor_ve_el_aviso() -> None:
    from seguridad import empresas

    with sesion_de_instalacion() as bd:
        empresas.guardar_empresa_propia(bd, EMPRESA, "J-00000000-0", "", "admin", AHORA)
        bd.commit()
    sin_firma = paquete(par=None, presupuesto=_documento(), anexos={})
    navegador, ficha = _rita()
    cuerpo = _subir(navegador, ficha, con_manifiesto(sin_firma, usuario_emisor="Rita <rita>"))
    assert cuerpo.json()["aviso_armo"] == (
        "Usted armó o cambió este presupuesto: lo está revisando la misma persona que lo armó "
        "(figura en la versión 005). Su dictamen lo dejará registrado; no se le impide darlo."
    )


def test_un_archivo_de_mas_del_tope_es_413_sin_guardar_nada(monkeypatch) -> None:
    monkeypatch.setattr(ruta_revision, "TOPE_PAQUETE", 10)
    navegador, ficha = _rita()
    r = _subir(navegador, ficha, b"x" * 11)
    assert r.status_code == 413
    assert r.json()["detalle"] == MENSAJE_TOPE
    assert _filas("paq_recibido") == []


def test_sin_archivo_es_422_con_un_mensaje_fijo() -> None:
    navegador, ficha = _rita()
    r = navegador.post(
        "/v1/revision/paquete", data={"lista": "abc-que-no-vuelve"}, headers=protegida(ficha)
    )
    assert r.status_code == 422
    assert "abc-que-no-vuelve" not in r.text


# --- contra la entrega anterior y contra una lista ------------------------------------------


def test_la_segunda_entrega_se_compara_con_la_anterior() -> None:
    """Memoria a mano: la v004 con 25 m3 en vez de 20 suma 25 x 41,839875 = 1 045,996875; la
    v005 (20 m3) baja 209,199375, todo por cantidad."""
    from tests.fixtures.paquetes import EMITIDO

    navegador, ficha = _rita()
    anterior = copy.deepcopy(EJEMPLO_V1)
    anterior["renglones"][0]["cantidad"] = "25"
    anterior["renglones"][0]["item"]["cantidad"] = "25"
    anterior["renglones"][0]["total"] = "1045.99687500"
    antes = paquete(
        par=PAR,
        presupuesto=json_canonico(anterior),
        anexos={},
        codigo="004",
        codigos=("004",),
        paquete_id="22222222-2222-4222-8222-222222222222",
        emitido=EMITIDO.replace(day=3),
    )
    assert _subir(navegador, ficha, antes).status_code == 200
    cuerpo = _subir(navegador, ficha, _firmado()).json()
    previa = cuerpo["entrega_anterior"]
    assert cuerpo["sin_entrega_anterior"] is None
    assert previa["titulo"] == "Contra la entrega anterior (004)"
    assert Decimal(previa["total_antes"]) == Decimal("1045.996875")
    assert Decimal(previa["total_despues"]) == Decimal(TOTAL)
    assert previa["diferencia_texto"] == "-209.20"
    # el reparto viaja como montos (exacto y texto), no como una frase con punto decimal
    assert "por_cantidad_y_precio" not in previa
    assert Decimal(previa["desvio_cantidad"]) == Decimal("-209.199375")
    assert previa["desvio_cantidad_texto"] == "-209.20"
    assert Decimal(previa["desvio_precio"]) == 0 and previa["desvio_precio_texto"] == "0.00"
    [partida] = previa["partidas"]
    assert (partida["codigo_partida"], partida["estado"]) == ("MIN-001", "cambiada")
    assert cuerpo["cambios_titulo"] == "Quién cambió qué desde la entrega anterior"


def _lista_del_revisor(caso_ruta: Path) -> int:
    motor = crear_motor(f"sqlite:///{caso_ruta.as_posix()}")
    try:
        with abrir_sesion(motor) as caso:
            arena = caso.scalars(
                select(models.Insumo).where(models.Insumo.codigo == "MAT-001")
            ).one()
            mia = models.ListaPrecios(
                nombre="Lista del revisor", moneda="USD", fecha_vigencia=date(2026, 11, 5)
            )
            mia.precios.append(models.PrecioInsumo(insumo=arena, precio=Decimal("16.00")))
            caso.add(mia)
            caso.commit()
            return mia.id
    finally:
        motor.dispose()


def test_el_contraste_con_mi_lista_usa_la_mas_nueva_y_se_puede_cambiar(_entorno: Path) -> None:
    """Memoria a mano: con la arena a 16,00 el total es 20 x 44,369875 = 887,3975 (Δ 50,60)."""
    mia = _lista_del_revisor(_entorno)
    navegador, ficha = _rita()
    cuerpo = _subir(navegador, ficha, _firmado()).json()
    assert [x["nombre"] for x in cuerpo["listas"]] == ["Lista del revisor", "Lista de prueba"]
    contra = cuerpo["contra_lista"]
    assert contra["lista_id"] == mia
    assert contra["diferencia"]["titulo"] == "El paquete con los precios de mi lista"
    assert Decimal(contra["diferencia"]["total_despues"]) == Decimal("887.3975")
    assert contra["diferencia"]["diferencia_texto"] == "50.60"
    assert contra["diferencia"]["desvio_precio_texto"] == "50.60"
    assert contra["diferencia"]["desvio_cantidad_texto"] == "0.00"
    assert contra["no_contrastables"] == (
        "No contrastables (no están en mi lista; no suman): EQU-001, MO-001"
    )
    otra = next(x["id"] for x in cuerpo["listas"] if x["nombre"] == "Lista de prueba")
    r = _subir(navegador, ficha, _firmado(), "/v1/revision/contra-lista", lista=str(otra))
    assert r.status_code == 200, r.text
    assert r.headers["cache-control"] == "no-store"
    assert r.json()["diferencia"]["diferencia_texto"] == "0.00"  # sus mismos precios
    assert r.json()["no_contrastables"] is None


def test_una_lista_que_no_existe_es_404() -> None:
    navegador, ficha = _rita()
    r = _subir(navegador, ficha, _firmado(), "/v1/revision/contra-lista", lista="999")
    assert r.status_code == 404


# --- dictamen -------------------------------------------------------------------------------


def test_el_dictamen_pide_notas_si_hay_observaciones_y_se_guarda_atado_al_archivo() -> None:
    navegador, ficha = _rita()
    datos = _firmado()
    _subir(navegador, ficha, datos)
    r = _subir(navegador, ficha, datos, "/v1/revision/dictamen", resultado="con observaciones")
    assert (r.status_code, r.json()) == (422, {"detalle": NOTAS_FALTAN})
    r = _subir(
        navegador,
        ficha,
        datos,
        "/v1/revision/dictamen",
        resultado="con observaciones",
        notas="Falta la memoria de cálculo de R4",
    )
    assert r.status_code == 200, r.text
    assert r.json()["mensaje"] == GUARDADO
    [fila] = _filas("paq_dictamen")
    [recibido] = _filas("paq_recibido")
    assert (fila["resultado"], fila["usuario"], fila["estado_verificacion"]) == (
        "con observaciones",
        "Rita <rita>",
        "verificado",
    )
    assert fila["huella_archivo"] == recibido["huella_archivo"] == hashlib.sha256(datos).hexdigest()
    assert fila["notas"] == "Falta la memoria de cálculo de R4" and not fila["revisor_armo"]
    [anterior] = r.json()["dictamenes"]
    assert anterior["texto"].startswith(
        "Dictamen anterior: con observaciones · Rita <rita> · "
    ) and anterior["texto"].endswith(" · verificación entonces: verificado")
    # y el próximo análisis del mismo archivo lo muestra
    assert _subir(navegador, ficha, datos).json()["dictamenes"] == r.json()["dictamenes"]


def test_un_dictamen_desconocido_es_422_sin_repetir_lo_enviado() -> None:
    navegador, ficha = _rita()
    r = _subir(navegador, ficha, _firmado(), "/v1/revision/dictamen", resultado="<script>x")
    assert r.status_code == 422
    assert "<script>" not in r.text
    assert _filas("paq_dictamen") == []


def test_un_paquete_alterado_no_admite_dictamen_y_no_es_un_500() -> None:
    """Un archivo cambiado tras firmarse se audita igual pero no queda registrado (recibir no lo
    guarda): su dictamen no tendría a qué atarse. Decisión del orquestador (ronda 1 de T7)."""
    from tests.fixtures.paquetes import reempaquetar

    navegador, ficha = _rita()
    alterado = reempaquetar(_firmado(), {"informe.json": b'{"otro":1}'})
    cuerpo = _subir(navegador, ficha, alterado).json()
    assert cuerpo["estado"] == "alterado"
    assert cuerpo["admite_dictamen"] is False and cuerpo["sin_dictamen"] == SIN_DICTAMEN
    r = _subir(
        navegador, ficha, alterado, "/v1/revision/dictamen", resultado="rechazado", notas="x"
    )
    assert (r.status_code, r.json()) == (422, {"detalle": SIN_DICTAMEN})
    assert _filas("paq_dictamen") == []


def test_el_mismo_archivo_dos_veces_a_la_vez_no_es_un_500(monkeypatch) -> None:
    """Carrera: la otra petición registró el archivo entre la búsqueda por huella y la inserción.
    Se simula haciendo que la búsqueda no lo vea una vez: la recepción relee lo ya registrado."""
    from paquete import modelo

    navegador, ficha = _rita()
    datos = _firmado()
    assert _subir(navegador, ficha, datos).status_code == 200
    real = modelo.buscar_por_huella
    vistas: list[str] = []

    def ciega_una_vez(sesion, huella):
        vistas.append(huella)
        return None if len(vistas) == 1 else real(sesion, huella)

    monkeypatch.setattr(modelo, "buscar_por_huella", ciega_una_vez)
    r = _subir(navegador, ficha, datos)
    assert r.status_code == 200, r.text
    assert r.json()["banner_version"].startswith("Versión 005 del 2026-11-10; ya recibió")
    assert len(_filas("paq_recibido")) == 1


def test_el_dictamen_de_un_paquete_rechazado_no_se_guarda() -> None:
    navegador, ficha = _rita()
    r = _subir(navegador, ficha, b"no es zip", "/v1/revision/dictamen", resultado="conforme")
    assert (r.status_code, r.json()) == (422, {"detalle": RECHAZADO_ZIP})
    assert _filas("paq_dictamen") == []


def test_el_dictamen_lo_da_una_persona_no_la_clave_de_api() -> None:
    sin = cliente(APP)
    r = sin.post(
        "/v1/revision/dictamen",
        files={"paquete": ("recibido.silice", _firmado(), "application/zip")},
        data={"resultado": "conforme"},
        headers={CABECERA: crear_clave_de_prueba()},
    )
    assert (r.status_code, r.json()) == (403, {"detalle": ruta_revision.MENSAJE_SOLO_PERSONA})
    assert _filas("paq_dictamen") == []


# --- Excel ----------------------------------------------------------------------------------


@pytest.mark.parametrize("tipo", ["presupuesto", "revision"])
def test_el_excel_se_regenera_aqui_con_su_nombre(tipo: str) -> None:
    navegador, ficha = _rita()
    datos = _firmado()
    r = _subir(navegador, ficha, datos, "/v1/revision/excel", tipo=tipo)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == XLSX and r.headers["cache-control"] == "no-store"
    sha = hashlib.sha256(datos).hexdigest()
    assert f'filename="{tipo}-{sha[:12]}.xlsx"' in r.headers["content-disposition"]
    hojas = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True).sheetnames
    revision_hojas = {"Verificación", "Hallazgos", "Diferencias", "Cambios"}
    assert revision_hojas <= set(hojas) if tipo == "revision" else not revision_hojas & set(hojas)


def test_un_tipo_de_excel_desconocido_es_422() -> None:
    navegador, ficha = _rita()
    r = _subir(navegador, ficha, _firmado(), "/v1/revision/excel", tipo="del-emisor")
    assert r.status_code == 422 and "del-emisor" not in r.text


# --- presupuesto en Excel (UC-05) -----------------------------------------------------------


def test_la_plantilla_se_descarga_con_las_cinco_hojas_del_formato() -> None:
    navegador, _ = _rita()
    r = navegador.get("/v1/revision/plantilla")
    assert r.status_code == 200 and r.headers["content-type"] == XLSX
    assert 'filename="plantilla-presupuesto.xlsx"' in r.headers["content-disposition"]
    hojas = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True).sheetnames
    assert hojas == ["Datos", "Presupuesto", "APU", "Insumos", "Curva"]


def test_la_plantilla_se_audita_como_excel_externo_y_se_dice_no_verificado() -> None:
    navegador, ficha = _rita()
    plantilla = navegador.get("/v1/revision/plantilla").content
    r = navegador.post(
        "/v1/revision/excel-externo",
        files={"presupuesto": ("externo.xlsx", plantilla, XLSX)},
        headers=protegida(ficha),
    )
    assert r.status_code == 200, r.text
    cuerpo = r.json()
    assert cuerpo["aviso"] == EXCEL_EXTERNO
    assert Decimal(cuerpo["total"]) == Decimal(TOTAL) and cuerpo["total_texto"] == "836.80"
    # La plantilla: 20 m3 sin regla (R1) y un plan «Dia 1» que no nombra la partida (R7, dos)
    assert [(h["regla"], h["severidad"]) for h in cuerpo["hallazgos"]] == [
        ("R1", "ADVERTENCIA"),
        ("R7", "ADVERTENCIA"),
        ("R7", "ADVERTENCIA"),
    ]
    r1, r7_periodo, r7_partida = (h["descripcion"] for h in cuerpo["hallazgos"])
    assert r1.startswith("la cantidad de EJ-P1 (20 m3) es manual")
    assert "'Dia 1'" in r7_periodo and "EJ-P1 no aparece" in r7_partida
    assert _filas("paq_recibido") == []  # en memoria: nada se guarda


def test_un_excel_externo_que_no_es_xlsx_lo_rechaza_la_guardia() -> None:
    navegador, ficha = _rita()
    r = navegador.post(
        "/v1/revision/excel-externo",
        files={"presupuesto": ("macro.xlsm", b"PK\x03\x04", XLSX)},
        headers=protegida(ficha),
    )
    assert r.status_code == 422
    assert r.json()["detalle"].startswith("El archivo «macro.xlsm»")


def test_un_excel_externo_de_mas_del_tope_es_413(monkeypatch) -> None:
    monkeypatch.setattr(ruta_revision, "TOPE_EXCEL", 10)
    navegador, ficha = _rita()
    r = navegador.post(
        "/v1/revision/excel-externo",
        files={"presupuesto": ("externo.xlsx", b"x" * 11, XLSX)},
        headers=protegida(ficha),
    )
    assert (r.status_code, r.json()) == (
        413,
        {"detalle": "El presupuesto en Excel supera el máximo de 20 MiB."},
    )
