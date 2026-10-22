"""Pruebas de los fixtures de báscula (ejemplo construido).

Los valores esperados están escritos aquí a mano (los de `tests/fixtures/bascula/ORACULO.md`), no se
importan de `construir.py`: así la prueba no es circular. Tampoco usa código de `produccion/`.
"""

from __future__ import annotations

import csv
import io
import zipfile
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from tests.fixtures.bascula import construir

ENCABEZADOS = [
    "Ticket",
    "Fecha",
    "Placa",
    "Cliente",
    "Material",
    "Bruto (kg)",
    "Tara (kg)",
    "Neto (kg)",
]

# Oráculo a mano (ORACULO.md §1), como texto, igual que sale del CSV.
ESPERADO_BASICO = [
    ["TCK-001", "2026-10-15", "EJ-001", "Cliente Ejemplo A", "ARENA", "42500", "14200", "28300"],
    ["TCK-002", "2026-10-15", "EJ-002", "Cliente Ejemplo B", "PIEDRA34", "38000", "13500", "24500"],
    ["TCK-003", "2026-10-16", "EJ-003", "Cliente Ejemplo A", "ARENA", "42000", "45000", "-3000"],
    ["TCK-004", "2026-10-16", "EJ-004", "Cliente Ejemplo C", "ARENA", "40000", "15000", "20000"],
    ["TCK-005", "2026-12-31", "EJ-005", "Cliente Ejemplo B", "PIEDRA34", "40000", "15000", "25000"],
    ["TCK-006", "2026-10-17", "EJ-006", "Cliente Ejemplo C", "GRAVA9", "40000", "15000", "25000"],
]

GENERADORES_SIN_ARGUMENTOS = [
    construir.csv_basico,
    construir.xlsx_basico,
    construir.xlsx_sin_unidad,
    construir.csv_coma_decimal,
    construir.xlsx_con_macros_renombrado,
]


def _leer_csv(datos: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(datos.decode("utf-8")), delimiter=";"))


def _leer_xlsx(datos: bytes) -> list[list[object]]:
    libro = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    try:
        return [list(fila) for fila in libro.worksheets[0].iter_rows(values_only=True)]
    finally:
        libro.close()


@pytest.mark.parametrize("generador", GENERADORES_SIN_ARGUMENTOS)
def test_cada_generador_devuelve_bytes(generador) -> None:
    datos = generador()
    assert isinstance(datos, bytes)
    assert len(datos) > 0


def test_csv_desmedido_devuelve_bytes() -> None:
    assert isinstance(construir.csv_desmedido(3), bytes)


@pytest.mark.parametrize("generador", GENERADORES_SIN_ARGUMENTOS)
def test_los_generadores_son_deterministas(generador) -> None:
    assert generador() == generador()


def test_csv_basico_tiene_encabezado_y_seis_filas_de_datos() -> None:
    filas = _leer_csv(construir.csv_basico())
    assert filas[0] == ENCABEZADOS
    assert len(filas) - 1 == 6
    assert filas[1:] == ESPERADO_BASICO


def test_xlsx_basico_tiene_encabezado_y_seis_filas_de_datos() -> None:
    filas = _leer_xlsx(construir.xlsx_basico())
    assert list(filas[0]) == ENCABEZADOS
    assert len(filas) - 1 == 6


def test_xlsx_basico_guarda_los_pesos_como_numeros_nativos_y_la_fecha_como_texto() -> None:
    for fila in _leer_xlsx(construir.xlsx_basico())[1:]:
        assert isinstance(fila[1], str)  # fecha ISO en texto
        for peso in fila[5:8]:
            assert isinstance(peso, int) and not isinstance(peso, bool)


def test_csv_y_xlsx_basicos_coinciden_celda_a_celda() -> None:
    filas_csv = _leer_csv(construir.csv_basico())
    filas_xlsx = _leer_xlsx(construir.xlsx_basico())
    assert len(filas_csv) == len(filas_xlsx) == 7
    for i, (f_csv, f_xlsx) in enumerate(zip(filas_csv, filas_xlsx, strict=True)):
        assert len(f_csv) == len(f_xlsx) == 8, f"fila {i}"
        for j, (celda_csv, celda_xlsx) in enumerate(zip(f_csv, f_xlsx, strict=True)):
            # el CSV es texto; el xlsx trae enteros nativos en los pesos
            assert celda_csv == str(celda_xlsx), f"celda ({i}, {j})"


def test_las_cuentas_del_oraculo_cuadran_a_mano() -> None:
    # fila 1: 42 500 - 14 200 = 28 300 kg = 28,300 t; / 1,60 = 17,6875 m3
    assert 42500 - 14200 == 28300
    assert Decimal(28300) / 1000 == Decimal("28.3")
    assert Decimal("28.3") / Decimal("1.60") == Decimal("17.6875")
    assert Decimal("17.6875") * Decimal("1.60") == Decimal("28.3000")
    # fila 2: 38 000 - 13 500 = 24 500 kg = 24,5 t
    assert 38000 - 13500 == 24500
    assert Decimal(24500) / 1000 == Decimal("24.5")
    # fila 3: tara > bruto y neto = bruto - tara = -3 000
    assert 45000 > 42000
    assert 42000 - 45000 == -3000
    # fila 4: bruto - tara = 25 000, no 20 000
    assert 40000 - 15000 == 25000 != 20000
    # fila 5: 2026-10-18 a 2026-12-31 son 13 + 30 + 31 = 74 dias
    from datetime import date

    assert (date(2026, 12, 31) - date(2026, 10, 18)).days == 13 + 30 + 31 == 74
    # coma decimal: 1.234,56 - 234,56 = 1.000,00
    assert Decimal("1234.56") - Decimal("234.56") == Decimal("1000.00")


def test_xlsx_sin_unidad_no_trae_unidad_en_los_pesos() -> None:
    filas = _leer_xlsx(construir.xlsx_sin_unidad())
    assert list(filas[0]) == [
        "Ticket",
        "Fecha",
        "Placa",
        "Cliente",
        "Material",
        "Bruto",
        "Tara",
        "Neto",
    ]
    assert len(filas) - 1 == 6
    # los datos son los mismos que los del basico
    assert [[str(c) for c in f] for f in filas[1:]] == ESPERADO_BASICO


def test_csv_coma_decimal_trae_los_tres_formatos() -> None:
    filas = _leer_csv(construir.csv_coma_decimal())
    assert filas[0] == ENCABEZADOS
    assert [f[0] for f in filas[1:]] == ["TCK-101", "TCK-102", "TCK-103"]
    assert [f[5:8] for f in filas[1:]] == [
        ["1.234,56", "234,56", "1.000,00"],
        ["1,234.56", "234.56", "1,000.00"],
        ["42500", "14200", "28,300"],
    ]


def test_xlsx_con_macros_renombrado_contiene_vbaproject() -> None:
    datos = construir.xlsx_con_macros_renombrado()
    with zipfile.ZipFile(io.BytesIO(datos)) as zf:
        nombres = zf.namelist()
        assert "xl/vbaProject.bin" in nombres
        assert "xl/workbook.xml" in nombres  # por lo demas parece un .xlsx
        assert zf.testzip() is None
    # y el basico NO lo trae
    with zipfile.ZipFile(io.BytesIO(construir.xlsx_basico())) as zf:
        assert "xl/vbaProject.bin" not in zf.namelist()


def test_csv_desmedido_200000_filas_tiene_200001_lineas() -> None:
    datos = construir.csv_desmedido(200000)
    lineas = datos.decode("utf-8").splitlines()
    assert len(lineas) == 200001
    assert datos.count(b"\n") == 200001
    assert lineas[0] == ";".join(ENCABEZADOS)
    assert lineas[1] == "TCK-000001;2026-10-01;EJ-000001;Cliente Ejemplo A;ARENA;40000;15000;25000"
    assert lineas[-1] == "TCK-200000;2026-10-01;EJ-200000;Cliente Ejemplo A;ARENA;40000;15000;25000"


@pytest.mark.parametrize("filas", [0, 1, 5])
def test_csv_desmedido_cuenta_encabezado_mas_filas(filas: int) -> None:
    assert len(construir.csv_desmedido(filas).decode("utf-8").splitlines()) == filas + 1


def test_csv_desmedido_rechaza_cantidad_negativa() -> None:
    with pytest.raises(ValueError):
        construir.csv_desmedido(-1)
