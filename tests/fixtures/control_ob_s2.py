# tests/fixtures/control_ob_s2.py
"""La obra OB-S2 (spec S2 §7.1): los datos de `tests/fixtures/memorias/control_ob_s2.md` en código.

Una excavación (A), una tubería (B) y un relleno (C) en USD, con administración 0,15, utilidad 0,10,
FCAS 6,00 y bono 1,00: el PU es el costo directo por 1,265. El plan son cuatro semanas de cinco días
hábiles desde el lunes 2026-11-02. **Datos ficticios** (NORMAS.md §1): el inspector, las actas y los
costos reales los inventa la memoria.

Los códigos de partida son `OB-A`, `OB-B`, `OB-C` y `OB-D` (esta última, para la valuación sembrada,
no está presupuestada). Este módulo es la única copia de los datos: lo usan las pruebas de
valuaciones, valor ganado, recuperación y la interfaz de control.
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from core import models
from core.budget import crear_proyecto, elaborar, guardar_presupuesto
from core.budget.curva import PeriodoPlan
from core.budget.valuaciones import (
    Calendario,
    CostoReal,
    InstantaneaControl,
    LineaBase,
    registrar_costo,
    registrar_linea_base,
    registrar_valuacion,
)
from core.catalog import Catalogo
from core.contracts.apu import (
    ComposicionAPU,
    LineaEquipo,
    LineaManoObra,
    LineaMaterial,
    ParametrosCosto,
)
from core.contracts.dominio import Dominio
from core.contracts.item_computo import ItemComputo, OrigenTipo
from core.verification.valuaciones import LineaValuacion, Valuacion

D = Decimal
NOMBRE = "OB-S2"
CODIGO = "OB-001"
FECHA = date(2026, 11, 2)
MONEDA = "USD"
PARAMETROS = ParametrosCosto(D("6.00"), D("1.00"), D("0.15"), D("0.10"))
INICIO = date(2026, 11, 2)
CALENDARIO = Calendario(INICIO, 5)
CORTE = date(2026, 11, 13)
INSPECTOR = "Ing. Ficticio"
RESPONSABLE = "Residente Ficticio"
FUENTE = "acta ficticia de OB-S2 (NORMAS.md seccion 1)"
CARPETA = Path(__file__).resolve().parent / "control"
RUTA_XLSX = CARPETA / "ob_s2.xlsx"

#: (código, descripción, unidad, rendimiento, obreros, sueldo por obrero, equipo, material)
_DATOS = {
    "OB-A": ("Excavacion de zanja", "m3", D(10), (D(4), D("3.00")), (D(120), D("0.10")), None),
    "OB-B": ("Tuberia de drenaje", "m", D(20), (D(2), D("2.50")), (D(30), D("0.10")), D("3.00")),
    "OB-C": ("Relleno compactado", "m3", D(10), (D(4), D("2.50")), (D(60), D("0.10")), D("12.00")),
}
CANTIDADES = {"OB-A": D(100), "OB-B": D(200), "OB-C": D(100)}
#: Plan de la línea base: fracción de cada partida por semana (S2 §7.1).
PLAN: list[PeriodoPlan] = [
    ("Semana 1", (("OB-A", D("0.5")),)),
    ("Semana 2", (("OB-A", D("0.5")), ("OB-B", D("0.5")))),
    ("Semana 3", (("OB-B", D("0.5")), ("OB-C", D("0.5")))),
    ("Semana 4", (("OB-C", D("0.5")),)),
]
#: Unitarios del corte `u'` (costo directo): el material de B pasa de 3,00 a 4,00 y el de C de 12,00
#: a 14,10.
UNITARIOS_BASE = {"OB-A": D("10.00"), "OB-B": D("5.00"), "OB-C": D("20.00")}
UNITARIOS_DEL_CORTE = {"OB-A": D("10.00"), "OB-B": D("6.00"), "OB-C": D("22.10")}
COSTO_REAL = D("1280.00")


def precios_del_corte():
    """Los precios del corte de la memoria, por comprobación directa (importa S2-2 al llamarse)."""
    from core.budget.valor_ganado import PreciosDelCorte

    return PreciosDelCorte("comprobacion", "version revalorada ficticia", UNITARIOS_DEL_CORTE)


def composiciones(materiales: dict[str, D] | None = None) -> dict[str, ComposicionAPU]:
    """Los tres APU. `materiales` cambia el precio del material de una partida (la revaloración)."""
    precios = {"OB-B": D("3.00"), "OB-C": D("12.00"), **(materiales or {})}
    apus: dict[str, ComposicionAPU] = {}
    for codigo, (descripcion, unidad, rendimiento, obreros, equipo, material) in _DATOS.items():
        apus[codigo] = ComposicionAPU(
            codigo,
            descripcion,
            unidad,
            rendimiento,
            materiales=()
            if material is None
            else (LineaMaterial(f"Material de {descripcion}", unidad, D(1), precios[codigo]),),
            equipos=(LineaEquipo(f"Equipo de {descripcion}", D(1), equipo[0], equipo[1]),),
            mano_obra=(LineaManoObra(f"Cuadrilla de {descripcion}", obreros[0], obreros[1]),),
        )
    return apus


def items() -> list[ItemComputo]:
    return [
        ItemComputo(
            codigo,
            _DATOS[codigo][0],
            _DATOS[codigo][1],
            cantidad,
            f"memoria-ob-s2:{codigo}",
            OrigenTipo.MANUAL,
            Dominio.CIVIL,
        )
        for codigo, cantidad in CANTIDADES.items()
    ]


def elaborar_ob_s2(codigo: str = CODIGO, plan=PLAN):
    """El presupuesto OB-001 con su curva y su informe (puro: no toca la base)."""
    return elaborar(
        items(), composiciones(), PARAMETROS, codigo=codigo, fecha=FECHA, moneda=MONEDA, plan=plan
    )


def sembrar(sesion) -> models.Presupuesto:
    """El proyecto OB-S2 con la versión OB-001 guardada y su plan. No confirma."""
    lista = models.ListaPrecios(nombre="OB-S2 USD", moneda=MONEDA, fecha_vigencia=FECHA)
    sesion.add(lista)
    sesion.flush()
    catalogo = Catalogo(sesion)
    for apu in composiciones().values():
        catalogo.cargar_composicion(apu, lista, Dominio.CIVIL, FECHA, "memoria OB-S2")
    resultado = elaborar_ob_s2()
    proyecto = crear_proyecto(sesion, NOMBRE)
    return guardar_presupuesto(
        sesion,
        resultado.presupuesto,
        resultado.informe,
        proyecto=proyecto,
        lista=lista,
        parametros=PARAMETROS,
        plan=PLAN,
    )


def declarar_linea_base(sesion, perspectiva: str = "costo", desde: date = INICIO):
    return registrar_linea_base(
        sesion,
        NOMBRE,
        CODIGO,
        desde=desde,
        perspectiva=perspectiva,
        calendario=CALENDARIO,
        fuente=FUENTE,
        responsable=RESPONSABLE,
    )


#: Las valuaciones de la memoria: V-01 y V-02 (registradas) y la sembrada V-S.
V01 = dict(
    numero="V-01",
    desde=date(2026, 11, 2),
    hasta=date(2026, 11, 6),
    fecha=date(2026, 11, 9),
    inspector=INSPECTOR,
    lineas=[LineaValuacion("OB-A", D(40))],
)
V02 = dict(
    numero="V-02",
    desde=date(2026, 11, 9),
    hasta=date(2026, 11, 13),
    fecha=date(2026, 11, 16),
    inspector=INSPECTOR,
    lineas=[LineaValuacion("OB-A", D(50)), LineaValuacion("OB-B", D(60))],
)
VS = dict(
    numero="V-03",
    desde=date(2026, 11, 12),
    hasta=date(2026, 11, 20),
    fecha=date(2026, 11, 19),
    inspector="",
    lineas=[
        LineaValuacion("OB-B", D(150)),
        LineaValuacion("OB-D", D(5)),
        LineaValuacion("OB-C", D(10), date(2026, 11, 23)),
    ],
)
CORTE_DE_VS = date(2026, 11, 20)


def registrar(sesion, datos: dict, **extra):
    return registrar_valuacion(
        sesion, NOMBRE, fuente=FUENTE, responsable=RESPONSABLE, **datos, **extra
    )


def registrar_v01_y_v02(sesion):
    """V-01 y V-02, registradas, y el costo real de la memoria (1 280,00 USD)."""
    primera = registrar(sesion, V01, ahora=datetime(2026, 11, 9, 9, 0))
    segunda = registrar(sesion, V02, ahora=datetime(2026, 11, 16, 9, 0))
    registrar_costo(
        sesion,
        NOMBRE,
        monto=COSTO_REAL,
        moneda=MONEDA,
        fecha=date(2026, 11, 13),
        fuente="libro de obra ficticio",
        responsable=RESPONSABLE,
    )
    return primera, segunda


def valuacion(datos: dict, id: int) -> Valuacion:
    """Una valuación de la memoria como objeto puro (sin base)."""
    return Valuacion(id=id, **datos)


def instantanea(
    valuaciones, *, corte: date = CORTE, perspectiva: str = "costo", costos=None
) -> InstantaneaControl:
    """La instantánea de control de OB-S2 sin base de datos (para las pruebas puras)."""
    resultado = elaborar_ob_s2()
    return InstantaneaControl(
        proyecto=NOMBRE,
        corte=corte,
        linea_base=LineaBase(
            1, NOMBRE, CODIGO, INICIO, perspectiva, FUENTE, RESPONSABLE, CALENDARIO
        ),
        presupuesto=resultado.presupuesto,
        plan=tuple(PLAN),
        parametros_costo=PARAMETROS,
        moneda=MONEDA,
        valuaciones=tuple(valuaciones),
        costos=tuple(
            costos
            if costos is not None
            else [CostoReal(1, COSTO_REAL, MONEDA, date(2026, 11, 13), "libro de obra ficticio")]
        ),
    )


def v01_v02() -> list[Valuacion]:
    return [valuacion(V01, 1), valuacion(V02, 2)]


def escribir_xlsx(ruta: Path, *, valuaciones=None, costos=True) -> Path:
    """Escribe el XLSX de valuaciones de §6.2 con V-01, V-02 (y el costo real de la memoria).

    Los números van como texto, que es lo exacto por construcción (`leer_tabla`); las fechas, en
    ISO. `valuaciones` permite escribir otras (por ejemplo V-S)."""
    import openpyxl

    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Valuaciones"
    hoja.append(["numero", "desde", "hasta", "fecha", "inspector"])
    lineas = libro.create_sheet("Lineas")
    lineas.append(["numero", "codigo_partida", "cantidad", "fecha_medicion", "observacion"])
    for datos in valuaciones or (V01, V02):
        hoja.append(
            [
                datos["numero"],
                datos["desde"].isoformat(),
                datos["hasta"].isoformat(),
                "" if datos["fecha"] is None else datos["fecha"].isoformat(),
                datos["inspector"],
            ]
        )
        for linea in datos["lineas"]:
            medicion = linea.fecha_medicion
            lineas.append(
                [
                    datos["numero"],
                    linea.codigo_partida,
                    str(linea.cantidad),
                    "" if medicion is None else medicion.isoformat(),
                    linea.observacion,
                ]
            )
    if costos:
        hoja_costos = libro.create_sheet("Costos")
        hoja_costos.append(["numero", "codigo_partida", "monto", "moneda", "fecha", "fuente"])
        hoja_costos.append(
            ["", "", str(COSTO_REAL), MONEDA, "2026-11-13", "libro de obra ficticio"]
        )
    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta)
    return ruta
