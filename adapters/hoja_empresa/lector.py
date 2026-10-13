"""Lector del libro de Excel de la empresa: `.xlsx` → `LibroLeido` (spec 1.5, F-D3, F-D6, F-D10).

Principios:

- **Por rótulo, nunca por fila fija.** Los APU de la empresa tienen más de diez formas (de 41 a 81
  filas, con y sin bloque de transporte); el lector busca cada valor por el rótulo de su fila y la
  columna que dice el `Perfil` (causa C1 del análisis de causas raíz).
- **Dos lecturas del mismo archivo.** openpyxl abre el libro una vez con fórmulas y otra con los
  valores guardados (`data_only`). Se guardan las dos cosas, celda por celda, para que las reglas
  L1-L9 comparen lo que el libro *dice* con lo que *calcula*.
- **Todo dato lleva su celda** (`Celda`), para que cada hallazgo apunte a `HOJA!CELDA`.
- **`Decimal` siempre** (NORMAS §2.3): los números del libro se convierten con `Decimal(str(v))`.
- **El lector no juzga.** Una celda rara (texto donde va un número, `#REF!`, vacía) no hace fallar
  la lectura: el dato se omite o vale lo que Excel usaría con `N()` (0), y las reglas lo informan.
  Solo se lanza `ArchivoRechazado`: por la guardia (F-D10), si openpyxl no puede abrir el libro o si
  falta la hoja PRESUPUESTO.

Límites declarados:

- El rendimiento de un APU sale del rótulo REND; si no está o no tiene número, de la primera fila
  cuyo rótulo contiene "Rendimiento" y tiene número en la columna de REND; si tampoco, vale 1 (no
  divide) y `guardados` no tiene REND, de modo que las reglas pueden señalarlo. Si el respaldo
  encuentra el número, `rotulos[REND]` apunta a la celda de la que se leyó.
- Una línea de insumo con cantidad o precio vacío o no numérico vale 0 en ese campo (lo que haría la
  plantilla con `N()`); la celda queda en `valores` para las reglas.
- Solo se reconocen como números los valores numéricos guardados; un número escrito como texto vale
  como texto (igual que `N()` en Excel).
"""

from __future__ import annotations

import hashlib
import io
import re
import warnings
import zipfile
from collections.abc import Iterator, Mapping
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from typing import Any

from adapters.hoja_empresa.guardia import ArchivoRechazado, verificar_archivo
from adapters.hoja_empresa.modelo import (
    ApuLeido,
    Celda,
    FilaEspecial,
    InsumoLeido,
    LibroLeido,
    LineaApu,
    ParametrosLibro,
    PartidaLeida,
    Portada,
    Rotulo,
    TipoFilaEspecial,
    TipoInsumo,
    clave,
)
from adapters.hoja_empresa.perfil import Perfil, cargar_perfil, coincide, normalizar

__all__ = ["leer_libro"]

_CERO = Decimal(0)
_UNO = Decimal(1)
_TOLERANCIA_PU = Decimal("0.000001")
_TOLERANCIA_REASIGNAR = Decimal("0.005")  # medio centavo: pasada final de `_unir` (F-D6)
_REF_CELDA = re.compile(r"^(?:'((?:[^']|'')+)'|([^'!\[\]]+))!\$?([A-Za-z]{1,3})\$?(\d+)$")
_COORD = re.compile(r"^([A-Z]{1,3})(\d+)$")


# --- Conversión de valores -----------------------------------------------------------------------


def _convertir(v: Any) -> object:
    """Valor guardado de una celda → lo que va en `LibroLeido.valores`."""
    if v is None or isinstance(v, (bool, str, Decimal)):
        return v
    if isinstance(v, (int, float)):
        return Decimal(str(v))
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == time(0) else v.isoformat()
    if isinstance(v, (date, time)):
        return v.isoformat()
    return str(v)


def _num(v: object) -> Decimal | None:
    """Número de un valor ya convertido; `None` si no es un número finito (texto, error, vacío)."""
    if isinstance(v, Decimal) and v.is_finite():
        return v
    return None


def _entero(v: object) -> int | None:
    n = _num(v)
    if n is None or n != n.to_integral_value():
        return None
    return int(n)


def _texto(v: object) -> str:
    """Texto para mostrar de un valor convertido: '' si está vacío."""
    if v is None:
        return ""
    return v.strip() if isinstance(v, str) else str(v)


def _col(letra: str) -> int:
    n = 0
    for ch in letra:
        n = n * 26 + (ord(ch) - 64)
    return n


def _letra(col: int) -> str:
    letras = ""
    while col:
        col, r = divmod(col - 1, 26)
        letras = chr(65 + r) + letras
    return letras


def _ref(fila: int, col: int) -> str:
    return f"{_letra(col)}{fila}"


# --- Una hoja en memoria -------------------------------------------------------------------------


class _Hoja:
    """Las celdas no vacías de una hoja: fórmula (si la hay) y valor guardado convertido."""

    __slots__ = ("formulas", "nombre", "por_fila", "valores")

    def __init__(self, nombre: str) -> None:
        self.nombre = nombre
        self.valores: dict[tuple[int, int], object] = {}
        self.formulas: dict[tuple[int, int], str] = {}
        self.por_fila: dict[int, list[int]] = {}

    def _indexar(self) -> None:
        por_fila: dict[int, list[int]] = {}
        for f, c in sorted(self.valores):
            por_fila.setdefault(f, []).append(c)
        self.por_fila = por_fila

    @property
    def filas(self) -> list[int]:
        return list(self.por_fila)

    def celdas(self) -> Iterator[tuple[int, int, object]]:
        """Recorre las celdas no vacías por filas y, dentro de cada fila, de izquierda a
        derecha."""
        for f, cols in self.por_fila.items():
            for c in cols:
                yield f, c, self.valores[(f, c)]

    def valor(self, fila: int, col: int) -> object:
        return self.valores.get((fila, col))

    def formula(self, fila: int, col: int) -> str | None:
        return self.formulas.get((fila, col))

    def numero(self, fila: int, col: int) -> Decimal | None:
        return _num(self.valores.get((fila, col)))

    def a_la_derecha(self, fila: int, col: int) -> object:
        """El primer valor no vacío a la derecha de (fila, col) en la misma fila, o `None`."""
        for c in self.por_fila.get(fila, ()):
            if c > col:
                v = self.valores[(fila, c)]
                if v is not None and v != "":
                    return v
        return None

    def celda(self, fila: int, col: int) -> Celda:
        return Celda(self.nombre, _ref(fila, col))


def _texto_formula(celda: Any) -> str | None:
    """El texto de la fórmula de una celda openpyxl (incluye fórmulas matriciales), o `None`."""
    if getattr(celda, "data_type", None) != "f":
        return None
    v = celda.value
    texto = getattr(v, "text", None) or (v if isinstance(v, str) else None)
    if not texto:
        return None
    return texto if texto.startswith("=") else f"={texto}"


def _abrir(datos: bytes, nombre: str, data_only: bool) -> Any:
    """Abre el libro en modo de solo lectura (sin VBA ni vínculos externos).

    El modo normal de openpyxl aplica el estilo de cada celda combinada al cargar: en el libro B
    eso es el 80 % del tiempo (unos 7 s por libro). El lector no necesita estilos, y en solo
    lectura las dos pasadas de un libro de 35 hojas tardan alrededor de 1,5 s. El llamador cierra
    el libro (`close`), porque en este modo openpyxl deja abierto el ZIP.
    """
    import openpyxl

    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            return openpyxl.load_workbook(
                io.BytesIO(datos),
                read_only=True,
                data_only=data_only,
                keep_vba=False,
                keep_links=False,
            )
    except Exception as exc:  # openpyxl lanza tipos muy variados ante un XML dañado
        raise ArchivoRechazado(
            f"No se pudo abrir el libro «{nombre}»: el archivo .xlsx está dañado o no es de Excel."
        ) from exc


def _celdas_reales(ws: Any) -> Iterator[Any]:
    """Las celdas no vacías de una hoja de solo lectura, por filas.

    `reset_dimensions` hace que openpyxl no rellene hasta la dimensión declarada en el XML (que
    puede mentir y decir XFD1048576): cada fila llega solo hasta su última celda real.
    """
    ws.reset_dimensions()
    for fila in ws.iter_rows():
        for celda in fila:
            if celda.value is not None:
                yield celda


def _cargar_hojas(wf: Any, wv: Any) -> tuple[dict[str, _Hoja], dict[str, str], dict[str, object]]:
    """Lee todas las celdas no vacías de las dos pasadas: fórmulas (`wf`) y valores guardados
    (`wv`). Una celda con fórmula sin valor guardado (libro nunca calculado) queda con `None`."""
    hojas: dict[str, _Hoja] = {}
    formulas: dict[str, str] = {}
    valores: dict[str, object] = {}
    nombres_v = set(wv.sheetnames)
    for ws in wf.worksheets:
        guardados_v: dict[tuple[int, int], object] = {}
        if ws.title in nombres_v:
            for cv in _celdas_reales(wv[ws.title]):
                guardados_v[(cv.row, cv.column)] = cv.value
        h = _Hoja(ws.title)
        for celda in _celdas_reales(ws):
            f, c = celda.row, celda.column
            tf = _texto_formula(celda)
            if tf is not None:
                guardado = _convertir(guardados_v.get((f, c)))
                h.formulas[(f, c)] = tf
                formulas[clave(ws.title, _ref(f, c))] = tf
            else:
                guardado = _convertir(celda.value)
            h.valores[(f, c)] = guardado
            valores[clave(ws.title, _ref(f, c))] = guardado
        h._indexar()
        hojas[ws.title] = h
    return hojas, formulas, valores


def _buscar_hoja(hojas: Mapping[str, _Hoja], nombre: str) -> _Hoja | None:
    if nombre in hojas:
        return hojas[nombre]
    n = normalizar(nombre)
    return next((h for t, h in hojas.items() if normalizar(t) == n), None)


# --- Nombres definidos y vínculos externos -------------------------------------------------------


def _nombres(wf: Any) -> dict[str, str]:
    """Nombres del libro (`NOMBRE`) y de cada hoja (`HOJA!NOMBRE`) → su texto (`attr_text`)."""
    out: dict[str, str] = {}
    for n, dn in wf.defined_names.items():
        out[n] = dn.attr_text or ""
    for ws in wf.worksheets:
        for n, dn in ws.defined_names.items():
            out[f"{ws.title}!{n}"] = dn.attr_text or ""
    return out


def _resolver(texto: str | None) -> Celda | None:
    """`INSUMOS!$E$112` → `Celda('INSUMOS', 'E112')`; rangos, vínculos externos y `#REF!` → None."""
    if not texto:
        return None
    m = _REF_CELDA.match(texto.strip().lstrip("="))
    if m is None:
        return None
    hoja = m.group(1).replace("''", "'") if m.group(1) is not None else m.group(2)
    return Celda(hoja, f"{m.group(3).upper()}{m.group(4)}")


def _vinculos_externos(datos: bytes) -> tuple[str, ...]:
    """Los `Target` de `xl/externalLinks/_rels/*.rels` (el archivo al que apunta cada vínculo)."""
    from defusedxml import ElementTree

    destinos: list[str] = []
    try:
        with zipfile.ZipFile(io.BytesIO(datos)) as zf:
            for n in sorted(zf.namelist()):
                if not (n.startswith("xl/externalLinks/_rels/") and n.endswith(".rels")):
                    continue
                try:
                    raiz = ElementTree.fromstring(zf.read(n))
                except Exception:  # XML dañado o con entidades: se omite, no se adivina
                    continue
                for el in raiz.iter():
                    if el.tag.rsplit("}", 1)[-1] == "Relationship" and el.get("Target"):
                        destinos.append(el.get("Target", ""))
    except (zipfile.BadZipFile, OSError, ValueError):
        return ()
    return tuple(destinos)


# --- PORTADA, parámetros e INSUMOS ---------------------------------------------------------------


def _portada(h: _Hoja | None, perfil: Perfil) -> Portada:
    if h is None:
        return Portada()
    campos: dict[str, str] = {}
    exactos = {
        campo: {normalizar(v).rstrip(":").strip() for v in variantes}
        for campo, variantes in perfil.rotulos_portada.items()
    }
    ancla = normalizar(perfil.ancla_obra)
    for f, c, v in h.celdas():
        if not isinstance(v, str):
            continue
        n = normalizar(v).rstrip(":").strip()
        if "obra" not in campos and n == ancla:
            campos["obra"] = _texto(h.valor(f + 1, c))
            continue
        for campo, textos in exactos.items():
            if campo not in campos and n in textos:
                campos[campo] = _texto(h.a_la_derecha(f, c))
                break
    return Portada(**campos)


def _parametros(
    perfil: Perfil, nombres: Mapping[str, str], valores: Mapping[str, object]
) -> ParametrosLibro:
    importes: dict[str, Decimal] = {}
    celdas: dict[str, Celda] = {}
    for campo, nombre in perfil.nombres_parametros.items():
        celda = _resolver(nombres.get(nombre))
        importes[campo] = _CERO
        if celda is None:
            continue
        celdas[nombre] = celda
        importes[campo] = _num(valores.get(str(celda))) or _CERO
    return ParametrosLibro(**importes, celdas=celdas)


def _insumos(h: _Hoja | None, perfil: Perfil) -> tuple[InsumoLeido, ...]:
    if h is None:
        return ()
    cols = {k: _col(v) for k, v in perfil.columnas_insumos.items()}
    out: list[InsumoLeido] = []
    for f in h.filas:
        codigo = h.valor(f, cols["codigo"])
        if not perfil.es_codigo_insumo(codigo):
            continue
        codigo = str(codigo).strip()
        try:
            tipo = TipoInsumo(codigo[:2].upper())
        except ValueError:
            continue
        out.append(
            InsumoLeido(
                codigo=codigo,
                tipo=tipo,
                descripcion=_texto(h.valor(f, cols["descripcion"])),
                unidad=_texto(h.valor(f, cols["unidad"])),
                precio=h.numero(f, cols["precio"]) or _CERO,
                fuente=_texto(h.valor(f, cols["fuente"])),
                celda=h.celda(f, cols["precio"]),
            )
        )
    return tuple(out)


# --- APU -----------------------------------------------------------------------------------------


def _fila_referida(formula: str | None, hoja_presupuesto: str) -> int | None:
    """`=+PRESUPUESTO!D23` o `=PRESUPUESTO!$D$23` → 23. Texto, `#REF!` u otra hoja → None.

    Se acepta cualquier columna: lo que une el APU con su partida es la fila (spec F-D6)."""
    if not formula:
        return None
    m = re.fullmatch(r"=\s*\+?\s*(.+?)\s*", formula)
    celda = _resolver(m.group(1)) if m else None
    if celda is None or normalizar(celda.hoja) != normalizar(hoja_presupuesto):
        return None
    return int(_COORD.match(celda.ref).group(2))  # type: ignore[union-attr]


def _leer_apu(h: _Hoja, perfil: Perfil, params: ParametrosLibro, hoja_presupuesto: str) -> ApuLeido:
    col_rot = _col(perfil.columna_rotulos_apu)
    col_cab = _col(perfil.columna_cabecera_apu)
    col_pct = _col(perfil.columna_porcentaje_apu)
    cl = {k: _col(v) for k, v in perfil.columnas_linea.items()}
    m = _COORD.match(perfil.celda_titulo_apu)
    tf, tc = int(m.group(2)), _col(m.group(1))  # type: ignore[union-attr]

    codigo_partida: str | None = None
    unidad: str | None = None
    filas_rotulo: dict[Rotulo, int] = {}
    lineas: list[LineaApu] = []
    for f in h.filas:
        txt = h.valor(f, col_rot)
        if isinstance(txt, str):
            if codigo_partida is None and coincide(txt, perfil.rotulos_codigo_partida):
                codigo_partida = _texto(h.valor(f, col_cab))
            elif unidad is None and coincide(txt, perfil.rotulos_unidad):
                unidad = _texto(h.valor(f, col_cab))
            else:
                for r, ancla in perfil.rotulos_apu.items():
                    if r not in filas_rotulo and coincide(txt, ancla.variantes):
                        filas_rotulo[r] = f
                        break
        codigo = h.valor(f, cl["codigo"])
        if perfil.es_codigo_insumo(codigo):
            codigo = str(codigo).strip()
            try:
                tipo = TipoInsumo(codigo[:2].upper())
            except ValueError:
                continue
            lineas.append(
                LineaApu(
                    tipo=tipo,
                    codigo=codigo,
                    descripcion=_texto(h.valor(f, cl["descripcion"])),
                    unidad=_texto(h.valor(f, cl["unidad"])),
                    cantidad=h.numero(f, cl["cantidad"]) or _CERO,
                    precio=h.numero(f, cl["precio"]) or _CERO,
                    celda_codigo=h.celda(f, cl["codigo"]),
                    celda_cantidad=h.celda(f, cl["cantidad"]),
                )
            )

    rotulos: dict[Rotulo, Celda] = {}
    guardados: dict[Rotulo, Decimal] = {}
    for r, f in filas_rotulo.items():
        col = _col(perfil.rotulos_apu[r].columna)
        rotulos[r] = h.celda(f, col)
        n = h.numero(f, col)
        if n is not None:
            guardados[r] = n

    rendimiento = guardados.get(Rotulo.REND)
    if rendimiento is None:
        col_rend = _col(perfil.rotulos_apu[Rotulo.REND].columna)
        respaldo = normalizar(perfil.respaldo_rendimiento)
        for f in h.filas:
            if respaldo in normalizar(h.valor(f, col_rot)):
                n = h.numero(f, col_rend)
                if n is not None:
                    rendimiento = n
                    rotulos[Rotulo.REND] = h.celda(f, col_rend)
                    guardados[Rotulo.REND] = n
                    break
    if rendimiento is None:
        rendimiento = _UNO

    def _pct(r: Rotulo, global_: Decimal) -> Decimal:
        f = filas_rotulo.get(r)
        n = h.numero(f, col_pct) if f is not None else None
        return global_ if n is None else n

    return ApuLeido(
        hoja=h.nombre,
        titulo=_texto(h.valor(tf, tc)),
        fila_presupuesto=_fila_referida(h.formula(tf, tc), hoja_presupuesto),
        codigo_partida=codigo_partida or "",
        unidad=unidad or "",
        rendimiento=rendimiento,
        lineas=tuple(lineas),
        escombro=guardados.get(Rotulo.ESC, _CERO),
        administracion=_pct(Rotulo.ADM, params.administracion),
        utilidad=_pct(Rotulo.UTI, params.utilidad),
        herramientas=params.herramientas if Rotulo.HERR in filas_rotulo else _CERO,
        rotulos=rotulos,
        guardados=guardados,
    )


# --- PRESUPUESTO ---------------------------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class _Renglon:
    fila: int
    numero: int
    codigo: str
    descripcion: str
    unidad: str
    cantidad: Decimal
    capitulo: str | None
    guardados: Mapping[str, Decimal]


def _fila_encabezado(h: _Hoja, perfil: Perfil) -> int | None:
    enc = [(_col(c), v) for c, v in perfil.encabezado_presupuesto.items()]
    for f in h.filas:
        if all(coincide(h.valor(f, c), variantes) for c, variantes in enc):
            return f
    return None


def _presupuesto(h: _Hoja, perfil: Perfil) -> tuple[list[_Renglon], tuple[FilaEspecial, ...]]:
    cp = {k: _col(v) for k, v in perfil.columnas_presupuesto.items()}
    guardar = [(letra, _col(letra)) for letra in perfil.columnas_guardadas]
    inicio = _fila_encabezado(h, perfil)
    renglones: list[_Renglon] = []
    especiales: list[FilaEspecial] = []
    capitulo: str | None = None
    for f in h.filas:
        if inicio is not None and f <= inicio:
            continue
        b = h.valor(f, cp["numero"])
        numero = _entero(b)
        descripcion = h.valor(f, cp["descripcion"])
        if numero is not None and numero > 0 and isinstance(descripcion, str):
            if not descripcion.strip():
                continue
            renglones.append(
                _Renglon(
                    fila=f,
                    numero=numero,
                    codigo=_texto(h.valor(f, cp["codigo"])),
                    descripcion=descripcion.strip(),
                    unidad=_texto(h.valor(f, cp["unidad"])),
                    cantidad=h.numero(f, cp["cantidad"]) or _CERO,
                    capitulo=capitulo,
                    guardados={
                        letra: n for letra, c in guardar if (n := h.numero(f, c)) is not None
                    },
                )
            )
            continue
        monto = h.numero(f, cp["monto"])
        if isinstance(b, str) and b.strip():
            tipo = next(
                (t for t, variantes in perfil.filas_especiales.items() if coincide(b, variantes)),
                TipoFilaEspecial.CAPITULO,
            )
            especiales.append(FilaEspecial(fila=f, texto=b.strip(), tipo=tipo, valor=monto))
            if tipo is TipoFilaEspecial.CAPITULO:
                capitulo = b.strip()
            elif tipo is TipoFilaEspecial.SUBTOTAL:
                capitulo = None
            elif tipo is TipoFilaEspecial.TOTAL:
                break
            continue
        formula_q = (h.formula(f, cp["monto"]) or "").upper()
        if (numero is None or numero <= 0) and "SUM(" in formula_q:
            especiales.append(
                FilaEspecial(fila=f, texto="", tipo=TipoFilaEspecial.SUMA, valor=monto)
            )
    return renglones, tuple(especiales)


def _unir(
    renglones: list[_Renglon], apus: tuple[ApuLeido, ...], prefijo: str, col_pu: str
) -> dict[int, str | None]:
    """Fila de PRESUPUESTO → hoja APU (spec F-D6).

    1. Por la fórmula del título de cada APU (`fila_presupuesto`).
    2. Si varios APU reclaman la misma fila, gana el que coincide por número (`APU-nn`); si ninguno,
       el de P.U. guardado igual al P guardado de la fila; si tampoco, el primero en el orden del
       libro (determinista). El conflicto lo informa L2, no el lector.
    3. Una partida sin APU se une por número a `APU-nn` si esa hoja existe y no está tomada.
    4. Pasada final: una partida que sigue sin APU toma un APU todavía libre cuyo P.U. guardado
       coincide con el P guardado de la fila (tolerancia `_TOLERANCIA_REASIGNAR`); si hay varios,
       el primero en el orden del libro.

    Por qué la pasada 4: un APU que pierde un conflicto en el paso 2 queda libre, pero su título
    apunta a la fila equivocada (causa C2: parche atado al nombre de la hoja), así que ningún paso
    anterior lo vuelve a mirar. En el libro C-1, APU-08 y APU-09 reclaman la fila 23; gana APU-08 y
    la fila 24 (alucobond, el ~40 % de la obra) quedaba sin APU aunque el P.U. de APU-09 es
    exactamente su P guardado. Sin esta pasada la partida de mayor monto se recalcularía sin
    análisis. La coincidencia de precio es una evidencia fuerte (dos APU distintos con el mismo
    P.U. a 0,005 son raros) y el conflicto sigue informado por L2, así que la reasignación queda a
    la vista del usuario. La tolerancia de 0,005 es medio centavo: el redondeo de presentación.
    """
    reclamos: dict[int, list[ApuLeido]] = {}
    for a in apus:
        if a.fila_presupuesto is not None:
            reclamos.setdefault(a.fila_presupuesto, []).append(a)
    por_hoja = {a.hoja: a for a in apus}
    union: dict[int, str | None] = {}
    tomadas: set[str] = set()
    for r in renglones:
        candidatos = reclamos.get(r.fila, [])
        if not candidatos:
            continue
        elegido = candidatos[0]
        if len(candidatos) > 1:
            por_numero = f"{prefijo}-{r.numero:02d}"
            p_fila = r.guardados.get(col_pu)
            elegido = next(
                (a for a in candidatos if a.hoja == por_numero),
                next(
                    (
                        a
                        for a in candidatos
                        if p_fila is not None
                        and Rotulo.PU in a.guardados
                        and abs(a.guardados[Rotulo.PU] - p_fila) <= _TOLERANCIA_PU
                    ),
                    candidatos[0],
                ),
            )
        union[r.fila] = elegido.hoja
        tomadas.add(elegido.hoja)
    for r in renglones:
        if r.fila in union:
            continue
        hoja = f"{prefijo}-{r.numero:02d}"
        if hoja in por_hoja and hoja not in tomadas:
            union[r.fila] = hoja
            tomadas.add(hoja)
        else:
            union[r.fila] = None
    for r in renglones:
        p_fila = r.guardados.get(col_pu)
        if union.get(r.fila) is not None or p_fila is None:
            continue
        libre = next(
            (
                a
                for a in apus
                if a.hoja not in tomadas
                and Rotulo.PU in a.guardados
                and abs(a.guardados[Rotulo.PU] - p_fila) <= _TOLERANCIA_REASIGNAR
            ),
            None,
        )
        if libre is not None:
            union[r.fila] = libre.hoja
            tomadas.add(libre.hoja)
    return union


# --- Programación, avance y gastos ---------------------------------------------------------------


def _programacion(hojas: Mapping[str, _Hoja], perfil: Perfil) -> dict[str, Decimal]:
    out: dict[str, Decimal] = {}
    for k, (logica, variantes) in perfil.programacion.items():
        h = _buscar_hoja(hojas, perfil.hoja(logica))
        if h is None:
            continue
        for f, c, v in h.celdas():
            if isinstance(v, str) and coincide(v, variantes):
                n = _num(h.a_la_derecha(f, c))
                if n is not None:
                    out[k] = n
                    break
    return out


def _seccion(
    h: _Hoja | None,
    inicio: tuple[str, ...],
    fin: tuple[str, ...],
    col_num: int,
    col_cant: int,
) -> dict[int, Decimal]:
    """Número de partida → cantidad, entre el rótulo de inicio de una sección y el de la
    siguiente."""
    out: dict[int, Decimal] = {}
    if h is None:
        return out
    dentro = False
    for f in h.filas:
        marca = h.valor(f, col_num)
        if isinstance(marca, str):
            if not dentro and coincide(marca, inicio):
                dentro = True
            elif dentro and coincide(marca, fin):
                break
            continue
        if not dentro:
            continue
        n = _entero(marca)
        q = h.numero(f, col_cant)
        if n is not None and n > 0 and q is not None:
            out.setdefault(n, q)
    return out


def _gastado(h: _Hoja | None, perfil: Perfil) -> dict[int, Decimal]:
    out: dict[int, Decimal] = {}
    if h is None:
        return out
    col_num = _col(perfil.columnas_gastos["columna_numero"])
    col_g = _col(perfil.columnas_gastos["columna_gastado"])
    for f in h.filas:
        n = _entero(h.valor(f, col_num))
        q = h.numero(f, col_g)
        if n is not None and n > 0 and q is not None:
            out.setdefault(n, q)
    return out


# --- Entrada pública -----------------------------------------------------------------------------


def leer_libro(datos: bytes, nombre: str, perfil: Perfil | None = None) -> LibroLeido:
    """Lee el libro `.xlsx` de la empresa y lo devuelve como `LibroLeido`.

    Lanza `ArchivoRechazado` si la guardia lo rechaza, si openpyxl no puede abrirlo o si no tiene la
    hoja PRESUPUESTO. Cualquier otra rareza del libro se tolera (ver la cabecera del módulo).
    """
    verificar_archivo(datos, nombre)
    perfil = perfil or cargar_perfil()
    wf = _abrir(datos, nombre, data_only=False)
    try:
        wv = _abrir(datos, nombre, data_only=True)
        try:
            # En solo lectura el XML de cada hoja se analiza al recorrerla: un XML dañado aparece
            # aquí y no al abrir.
            hojas, formulas, valores = _cargar_hojas(wf, wv)
        except Exception as exc:
            raise ArchivoRechazado(
                f"No se pudo leer el libro «{nombre}»: alguna hoja está dañada."
            ) from exc
        finally:
            wv.close()
        nombres = _nombres(wf)
        nombres_hojas = tuple(wf.sheetnames)
    finally:
        wf.close()

    h_pres = _buscar_hoja(hojas, perfil.hoja("PRESUPUESTO"))
    if h_pres is None:
        raise ArchivoRechazado(
            f"El libro «{nombre}» no tiene la hoja {perfil.hoja('PRESUPUESTO')}: no parece un "
            "presupuesto de la plantilla de la empresa."
        )
    params = _parametros(perfil, nombres, valores)

    prefijo = normalizar(perfil.prefijo_apu)
    apus = tuple(
        _leer_apu(h, perfil, params, h_pres.nombre)
        for t, h in hojas.items()
        if normalizar(t).startswith(prefijo)
    )
    renglones, especiales = _presupuesto(h_pres, perfil)
    union = _unir(
        renglones, apus, perfil.prefijo_apu, perfil.columnas_presupuesto["precio_unitario"]
    )
    partidas = tuple(
        PartidaLeida(
            fila=r.fila,
            numero=r.numero,
            codigo=r.codigo,
            descripcion=r.descripcion,
            unidad=r.unidad,
            cantidad=r.cantidad,
            capitulo=r.capitulo,
            hoja_apu=union.get(r.fila),
            guardados=r.guardados,
        )
        for r in renglones
    )

    totales: dict[str, Decimal] = {}
    for n in perfil.nombres_totales:
        celda = _resolver(nombres.get(n))
        v = _num(valores.get(str(celda))) if celda is not None else None
        if v is not None:
            totales[n] = v

    h_control = _buscar_hoja(hojas, perfil.hoja("CONTROL DE OBRA"))
    sec = perfil.secciones_control_obra
    col_num = _col(perfil.columnas_control_obra["columna_numero"])
    col_cant = _col(perfil.columnas_control_obra["columna_cantidad"])

    return LibroLeido(
        nombre=nombre,
        sha256=hashlib.sha256(datos).hexdigest(),
        hojas=nombres_hojas,
        portada=_portada(_buscar_hoja(hojas, perfil.hoja("PORTADA")), perfil),
        parametros=params,
        insumos=_insumos(_buscar_hoja(hojas, perfil.hoja("INSUMOS")), perfil),
        apus=apus,
        partidas=partidas,
        filas_especiales=especiales,
        totales=totales,
        nombres=nombres,
        vinculos_externos=_vinculos_externos(datos),
        formulas=formulas,
        valores=valores,
        programacion=_programacion(hojas, perfil),
        ejecutado=_seccion(
            h_control, sec["inicio_ejecutado"], sec["fin_ejecutado"], col_num, col_cant
        ),
        final_previsto=_seccion(
            h_control, sec["inicio_final"], sec["fin_final"], col_num, col_cant
        ),
        gastado=_gastado(_buscar_hoja(hojas, perfil.hoja("GASTOS")), perfil),
        perfil=perfil.nombre,
    )
