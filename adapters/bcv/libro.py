"""Abre un libro del BCV (.xls con xlrd, .xlsx con openpyxl; ambos importados en la función) y lo
entrega como cuadrícula neutra. Un número que llega como float se lleva a Decimal con repr y se
declara en `Libro.avisos` (S1 §7.2)."""

from __future__ import annotations

import hashlib
import io
import zipfile
import zlib
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from adapters.bcv.informe import FormatoNoReconocido

Celda = str | Decimal | date | None
LIMITE_BYTES = 20 * 1024 * 1024
FIRMA_OLE2 = bytes.fromhex("D0CF11E0A1B11AE1")
FIRMA_ZIP = b"PK\x03\x04"
MAX_ENTRADAS = 5000
MAX_DESCOMPRIMIDO = 60 * 1024 * 1024
MAX_PIEZA = 30 * 1024 * 1024
MAX_RAZON = 200
MAX_CELDAS_XLS = 2_000_000  # filas x columnas declaradas por hoja de un .xls
MAX_HOJAS_XLS = 50


def _verificar_xlsx_seguro(contenido: bytes, nombre: str) -> None:
    """Verifica que el contenedor sea un ZIP válido y seguro antes de pasarlo a openpyxl (SEG-2).
    Protege contra zip corrupto, zip bombs, compresión anómala y macros."""
    try:
        zf = zipfile.ZipFile(io.BytesIO(contenido))
    except (zipfile.BadZipFile, zipfile.LargeZipFile, ValueError, OSError, EOFError) as exc:
        raise FormatoNoReconocido(
            f"{nombre}: no es un libro .xlsx valido (esta danado o no es de Excel)"
        ) from exc
    with zf:
        entradas = zf.infolist()
        if len(entradas) > MAX_ENTRADAS:
            raise FormatoNoReconocido(
                f"{nombre}: tiene {len(entradas)} piezas internas; "
                f"supera el maximo de {MAX_ENTRADAS}"
            )
        total = 0
        comprimido = 0
        for e in entradas:
            total += e.file_size
            comprimido += e.compress_size
            if e.file_size and (e.compress_size == 0 or e.file_size > MAX_RAZON * e.compress_size):
                raise FormatoNoReconocido(
                    f"{nombre}: la pieza {e.filename} tiene compresion anomala "
                    f"(mas de {MAX_RAZON} a 1)"
                )
            if e.file_size > MAX_PIEZA:
                raise FormatoNoReconocido(
                    f"{nombre}: la pieza {e.filename} supera el limite de {MAX_PIEZA // 2**20} MB"
                )
        if total > MAX_DESCOMPRIMIDO:
            raise FormatoNoReconocido(
                f"{nombre}: tamano descomprimido supera el limite de "
                f"{MAX_DESCOMPRIMIDO // 2**20} MB"
            )
        if total and (comprimido == 0 or total > MAX_RAZON * comprimido):
            raise FormatoNoReconocido(f"{nombre}: razon de compresion global anomala")
        nombres = {e.filename.replace("\\", "/") for e in entradas}
        if any(n.rsplit("/", 1)[-1].lower() == "vbaproject.bin" for n in nombres):
            raise FormatoNoReconocido(f"{nombre}: contiene macros (vbaProject.bin)")
        if "xl/workbook.xml" not in nombres:
            raise FormatoNoReconocido(f"{nombre}: no tiene la estructura de un libro de Excel")
        try:
            fallo = zf.testzip()
            if fallo:
                raise FormatoNoReconocido(f"{nombre}: archivo corrupto en la pieza {fallo}")
        except (zlib.error, zipfile.BadZipFile, RuntimeError, ValueError) as exc:
            raise FormatoNoReconocido(f"{nombre}: error de descompresion ({exc})") from exc


@dataclass(frozen=True, slots=True)
class Hoja:
    nombre: str
    filas: tuple[tuple[Celda, ...], ...]


@dataclass(frozen=True, slots=True)
class Libro:
    archivo: str
    huella: str
    hojas: tuple[Hoja, ...]
    avisos: tuple[str, ...] = ()


def abrir_libro(ruta: Path) -> Libro:
    tamano = ruta.stat().st_size
    if tamano > LIMITE_BYTES:
        raise FormatoNoReconocido(f"{ruta.name}: pesa {tamano / 2**20:.1f} MB; el limite es 20 MB")
    contenido = ruta.read_bytes()
    sufijo = ruta.suffix.lower()
    if sufijo == ".xls":
        if not contenido.startswith(FIRMA_OLE2):
            extra = " (es un .xlsx renombrado)" if contenido.startswith(FIRMA_ZIP) else ""
            raise FormatoNoReconocido(
                f"{ruta.name}: la extension .xls no coincide con el contenido{extra}"
            )
        hojas, floats = _xls(contenido, ruta.name)
    elif sufijo == ".xlsx":
        if not contenido.startswith(FIRMA_ZIP):
            raise FormatoNoReconocido(
                f"{ruta.name}: la extension .xlsx no coincide con el contenido"
            )
        _verificar_xlsx_seguro(contenido, ruta.name)
        hojas, floats = _xlsx(ruta)
    else:
        raise FormatoNoReconocido(f"{ruta.name}: se leen .xls y .xlsx")
    avisos = (
        (f"{floats} celdas numericas leidas como float y llevadas a Decimal con repr",)
        if floats
        else ()
    )
    return Libro(ruta.name, hashlib.sha256(contenido).hexdigest(), hojas, avisos)


def _xls(contenido: bytes, nombre: str) -> tuple[tuple[Hoja, ...], int]:
    import xlrd  # extra `bcv` (RV‑24), importación diferida

    try:
        libro = xlrd.open_workbook(file_contents=contenido, on_demand=True)
    except Exception as error:  # xlrd lanza varias clases ante un archivo dañado
        raise FormatoNoReconocido(f"{nombre}: no se pudo abrir ({error})") from None
    try:
        if libro.nsheets > MAX_HOJAS_XLS:
            raise FormatoNoReconocido(
                f"{nombre}: tiene {libro.nsheets} hojas; supera el maximo de {MAX_HOJAS_XLS}"
            )
        floats, hojas = 0, []
        for indice in range(libro.nsheets):
            hoja = libro.sheet_by_index(indice)
            if hoja.nrows * hoja.ncols > MAX_CELDAS_XLS:
                raise FormatoNoReconocido(
                    f"{nombre}: la hoja {hoja.name!r} declara {hoja.nrows} filas x {hoja.ncols} "
                    f"columnas; supera el maximo de {MAX_CELDAS_XLS} celdas"
                )
            filas = []
            for r in range(hoja.nrows):
                fila: list[Celda] = []
                for c in range(hoja.ncols):
                    celda = hoja.cell(r, c)
                    if celda.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        fila.append(None)
                    elif celda.ctype == xlrd.XL_CELL_TEXT:
                        fila.append(celda.value.strip() or None)
                    elif celda.ctype == xlrd.XL_CELL_NUMBER:
                        floats += 1
                        fila.append(Decimal(repr(celda.value)))
                    elif celda.ctype == xlrd.XL_CELL_DATE:
                        fila.append(xlrd.xldate_as_datetime(celda.value, libro.datemode).date())
                    else:
                        fila.append(str(celda.value))
                filas.append(tuple(fila))
            hojas.append(Hoja(hoja.name, tuple(filas)))
            libro.unload_sheet(indice)
    finally:
        libro.release_resources()
    return tuple(hojas), floats


def _xlsx(ruta: Path) -> tuple[tuple[Hoja, ...], int]:
    import openpyxl

    try:
        libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    except Exception as error:
        raise FormatoNoReconocido(f"{ruta.name}: no se pudo abrir ({error})") from error
    floats, hojas = 0, []
    try:
        for hoja in libro.worksheets:
            filas = []
            for valores in hoja.iter_rows(values_only=True):
                fila: list[Celda] = []
                for v in valores:
                    if isinstance(v, float):
                        floats += 1
                        fila.append(Decimal(repr(v)))
                    elif isinstance(v, int) and not isinstance(v, bool):
                        fila.append(Decimal(v))
                    elif isinstance(v, datetime):
                        fila.append(v.date())
                    elif isinstance(v, str):
                        fila.append(v.strip() or None)
                    else:
                        fila.append(None if v is None else str(v))
                filas.append(tuple(fila))
            hojas.append(Hoja(hoja.title, tuple(filas)))
    finally:
        libro.close()
    return tuple(hojas), floats
