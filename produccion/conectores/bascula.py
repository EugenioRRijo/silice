"""Lector de reportes de bascula `.csv` y `.xlsx` (Tarea 4.6; RT-S4-5, RT-S4-6, RNF-25).

Codigo puro: recibe bytes y devuelve `LecturaBascula`; no toca la base, no usa el reloj (la fecha
de referencia la manda el cliente) y no conoce el catalogo de materiales. Por eso un pesaje cuyo
material no existe en el catalogo (`GRAVA9` del oraculo) o que necesita una densidad que no hay sale
aqui como LEIDO y la importacion (Tarea 4.7) lo manda a la bandeja con «material desconocido» o
«sin densidad».

Archivo y formato
-----------------
- La extension (`.csv` o `.xlsx`, sin distinguir mayusculas) decide el lector.
- CSV: UTF-8 con o sin BOM, o Latin-1 si no decodifica. El delimitador (`;`, `,` o tabulador) se
  detecta con `csv.Sniffer` sobre los primeros 4 KiB, con `;` si no se puede decidir.
- XLSX: `openpyxl.load_workbook(read_only=True, data_only=True)`, primera hoja (no la activa), y el
  libro siempre se cierra. Las filas se piden con tope: no se itera sin limite.
- Numeracion: la fila 1 del archivo es la primera fila; `fila` de cada pesaje es la del archivo. Las
  filas en blanco se saltan pero cuentan en la numeracion.
- Mas de `MAX_FILAS` filas de datos o de `MAX_COLUMNAS` columnas: `ReporteRechazado`.

Unidad del peso (RT-S4-6)
-------------------------
Sale del parametro `unidad_peso` (`"kg"` o `"t"`; si se da, manda) o del encabezado de las
columnas de peso. Sin ninguna: `ReporteRechazado("Indique si los pesos están en kg o en t.")`.
Los pesos se devuelven siempre en toneladas (`kg / 1000`) y con `Decimal`.

Numeros de texto
----------------
`1.234,56` y `1,234.56` se leen exactos (el ultimo separador es el decimal si hay de los dos tipos).
Un solo tipo de separador repetido es de miles (`1.234.567`). Un solo separador con exactamente tres
cifras despues y de 1 a 3 cifras antes (sin empezar en 0) es AMBIGUO (`28,300`): no se adivina y la
fila es dudosa. Las celdas numericas nativas de XLSX no son ambiguas. No se admiten espacios
internos.

Orden de reglas por fila (la PRIMERA que salta es el unico motivo de un pesaje dudoso)
--------------------------------------------------------------------------------------
 1. «falta el ticket»                       ticket vacio.
 2. «ticket repetido en el archivo»         el primero se lee o se evalua; los demas son repetidos
                                            aunque el primero haya salido dudoso.
 3. «fecha ilegible»                        no es fecha (ISO `AAAA-MM-DD` o `DD/MM/AAAA`, con hora
                                            opcional, o celda de fecha nativa).
 4. «fecha futura»                          fecha > `fecha_referencia`.
 5. «falta el material»                     material vacio.
 6. «peso ilegible» / «número ambiguo»      primer peso con problema, en el orden bruto, tara, neto.
                                            Un peso vacio es ilegible, salvo el neto cuando existen
                                            bruto y tara (entonces se calcula).
 7. «peso negativo»                         bruto o tara negativos (o el neto sin bruto y tara).
 8. «tara mayor o igual que el bruto»       con bruto y tara; incluye el neto negativo coherente.
 9. «neto distinto de bruto − tara»         solo si el archivo trae neto; tolerancia
                                            `TOLERANCIA_NETO_T` (0: exacto). El signo es U+2212.
Si falta la columna `neto` (o su celda esta vacia) se calcula `bruto − tara` y no hay regla 9. Sin
columna `bruto` o `tara`, ese peso vale 0 y las reglas 8 y 9 no aplican.

Un pesaje bueno nunca es duplicado de otro del mismo archivo; el cruce con la base (reimportar) es
de la importacion, que usa `huella` y el ticket.
"""

from __future__ import annotations

import csv
import hashlib
import io
import math
import re
from collections.abc import Iterator, Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from itertools import islice
from pathlib import PurePath

import openpyxl

from produccion.conectores.detector_columnas import (
    CAMPOS,
    MapeoNecesario,
    faltantes,
    mejor_encabezado,
    unidad_de_los_pesos,
)
from produccion.fuentes import (
    FILAS_DE_ENCABEZADO,
    MAX_COLUMNAS,
    MAX_FILAS,
    TOLERANCIA_NETO_T,
)

__all__ = [
    "LecturaBascula",
    "PesajeDudoso",
    "PesajeLeido",
    "ReporteRechazado",
    "leer",
]

_MENSAJE_SIN_UNIDAD = "Indique si los pesos están en kg o en t."
_MOTIVO_NETO = "neto distinto de bruto − tara"
_MILES = Decimal(1000)
_PESO_MAXIMO = Decimal(10) ** 15  # cota de cordura del valor absoluto (kg o t)
_MAX_CIFRAS = 18


@dataclass(frozen=True, slots=True)
class PesajeLeido:
    fila: int
    ticket: str
    fecha: date
    placa: str
    cliente: str
    material: str
    bruto_t: Decimal
    tara_t: Decimal
    neto_t: Decimal


@dataclass(frozen=True, slots=True)
class PesajeDudoso:
    fila: int
    ticket: str
    fecha_texto: str
    placa: str
    material: str
    bruto_texto: str
    tara_texto: str
    neto_texto: str
    motivo: str


@dataclass(frozen=True, slots=True)
class LecturaBascula:
    huella: str
    pesajes: tuple[PesajeLeido, ...]
    dudosos: tuple[PesajeDudoso, ...]


class ReporteRechazado(ValueError):  # noqa: N818 (nombre fijado por el plan)
    """El archivo no se puede importar; el mensaje esta escrito para el usuario."""


class _NumeroIlegibleError(Exception):
    pass


class _NumeroAmbiguoError(Exception):
    pass


# --------------------------------------------------------------------------------------- celdas


def _vacia(celda: object) -> bool:
    return celda is None or (isinstance(celda, str) and not celda.strip())


def _texto(celda: object) -> str:
    """Texto de una celda tal como la ve el usuario (para la bandeja)."""
    if celda is None:
        return ""
    if isinstance(celda, datetime):
        return celda.isoformat(sep=" ")
    if isinstance(celda, date):
        return celda.isoformat()
    if isinstance(celda, float) and celda.is_integer():
        return str(int(celda))
    return str(celda).strip()


_ISO = re.compile(r"(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T].*)?")
_DIA_MES_ANIO = re.compile(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})(?:\s.*)?")


def _fecha(celda: object) -> date | None:
    if isinstance(celda, datetime):
        return celda.date()
    if isinstance(celda, date):
        return celda
    if not isinstance(celda, str):
        return None
    texto = celda.strip()
    try:
        if coincide := _ISO.fullmatch(texto):
            return date(int(coincide[1]), int(coincide[2]), int(coincide[3]))
        if coincide := _DIA_MES_ANIO.fullmatch(texto):
            return date(int(coincide[3]), int(coincide[2]), int(coincide[1]))
    except ValueError:
        return None
    return None


def _grupos_de_miles(grupos: Sequence[str]) -> bool:
    primero, *resto = grupos
    return (
        primero.isdigit() and len(primero) <= 3 and all(g.isdigit() and len(g) == 3 for g in resto)
    )


def _numero_de_texto(texto: str) -> Decimal:
    s = texto.strip()
    signo = ""
    if s[:1] in ("+", "-"):
        signo, s = s[0], s[1:]
    if not re.fullmatch(r"[0-9.,]+", s) or s[0] in ".," or s[-1] in ".,":
        raise _NumeroIlegibleError
    puntos, comas = s.count("."), s.count(",")
    if puntos and comas:
        decimal = "." if s.rfind(".") > s.rfind(",") else ","
        miles = "," if decimal == "." else "."
        if s.count(decimal) != 1:
            raise _NumeroIlegibleError
        entero, fraccion = s.split(decimal)
        grupos = entero.split(miles)
        if not _grupos_de_miles(grupos):
            raise _NumeroIlegibleError
        valor = "".join(grupos) + "." + fraccion
    elif puntos or comas:
        separador = "." if puntos else ","
        if s.count(separador) > 1:
            grupos = s.split(separador)
            if not _grupos_de_miles(grupos):
                raise _NumeroIlegibleError
            valor = "".join(grupos)
        else:
            entero, fraccion = s.split(separador)
            if len(fraccion) == 3 and 1 <= len(entero) <= 3 and not entero.startswith("0"):
                raise _NumeroAmbiguoError
            valor = entero + "." + fraccion
    else:
        valor = s
    if sum(c.isdigit() for c in valor) > _MAX_CIFRAS:
        raise _NumeroIlegibleError
    return Decimal(signo + valor)


def _numero(celda: object) -> Decimal:
    """Peso de una celda sin convertir de unidad; si no se puede, lanza un error de numero."""
    if isinstance(celda, bool):
        raise _NumeroIlegibleError
    if isinstance(celda, int):
        valor = Decimal(celda)
    elif isinstance(celda, float):
        if not math.isfinite(celda):
            raise _NumeroIlegibleError
        valor = Decimal(repr(celda))
    elif isinstance(celda, Decimal):
        valor = celda
    elif isinstance(celda, str):
        valor = _numero_de_texto(celda)
    else:
        raise _NumeroIlegibleError
    if not valor.is_finite() or abs(valor) >= _PESO_MAXIMO:
        raise _NumeroIlegibleError
    return valor


# ----------------------------------------------------------------------------------- archivos


def _acotada(fila: Sequence[object]) -> list[object]:
    """Fila sin celdas vacias al final; rechaza la que pasa de `MAX_COLUMNAS` columnas."""
    celdas = list(fila)
    while celdas and _vacia(celdas[-1]):
        celdas.pop()
    if len(celdas) > MAX_COLUMNAS:
        raise ReporteRechazado(
            f"El archivo tiene más de {MAX_COLUMNAS} columnas; deje solo las del reporte."
        )
    return celdas


def _delimitador(texto: str) -> str:
    muestra = texto[:4096]
    if len(texto) > 4096 and "\n" in muestra:
        muestra = muestra[: muestra.rfind("\n")]
    try:
        return csv.Sniffer().sniff(muestra, delimiters=";,\t").delimiter
    except csv.Error:
        return ";"


def _filas_de_csv(datos: bytes, tope: int) -> list[list[object]]:
    if datos.startswith(b"PK\x03\x04"):
        raise ReporteRechazado(
            "El archivo parece un libro de Excel; guárdelo con extensión .xlsx o expórtelo a CSV."
        )
    try:
        texto = datos.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = datos.decode("latin-1")
    lector = csv.reader(io.StringIO(texto, newline=""), delimiter=_delimitador(texto))
    try:
        return [_acotada(fila) for fila in islice(lector, tope)]
    except csv.Error as error:
        raise ReporteRechazado(
            "No se pudo leer el archivo CSV: tiene un formato inválido."
        ) from error


def _filas_de_xlsx(datos: bytes, tope: int) -> Iterator[list[object]]:
    """Primera hoja del libro, hasta `tope` filas; el libro se cierra al terminar o al fallar."""
    try:
        libro = openpyxl.load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    except Exception as error:  # archivo ajeno: cualquier fallo del formato es un rechazo
        raise ReporteRechazado("El archivo no es un libro de Excel (.xlsx) válido.") from error
    try:
        if libro.worksheets:
            hoja = libro.worksheets[0]
            for fila in islice(hoja.iter_rows(values_only=True), tope):
                yield _acotada(fila)
    except ReporteRechazado:
        raise
    except Exception as error:  # hoja corrupta
        raise ReporteRechazado("No se pudo leer la hoja del libro de Excel.") from error
    finally:
        libro.close()


# ------------------------------------------------------------------------------------- mapeo


def _validar_mapeo(mapeo: Mapping[str, int]) -> None:
    for campo, indice in mapeo.items():
        if campo not in CAMPOS:
            raise ReporteRechazado(
                f"Campo desconocido en el mapeo: «{campo}». Los campos son: {', '.join(CAMPOS)}."
            )
        if (
            isinstance(indice, bool)
            or not isinstance(indice, int)
            or not 0 <= indice < MAX_COLUMNAS
        ):
            raise ReporteRechazado(f"La columna indicada para «{campo}» no es válida.")


# --------------------------------------------------------------------------------------- filas


def _celda(fila: Sequence[object], columnas: Mapping[str, int], campo: str) -> object:
    indice = columnas.get(campo)
    if indice is None or indice >= len(fila):
        return None
    return fila[indice]


def _evaluar(
    fila: Sequence[object],
    columnas: Mapping[str, int],
    *,
    unidad: str,
    fecha_referencia: date,
    vistos: set[str],
    numero_de_fila: int,
) -> PesajeLeido | PesajeDudoso:
    """Aplica las reglas en el orden del docstring del modulo; la primera que salta decide."""
    textos = {campo: _texto(_celda(fila, columnas, campo)) for campo in CAMPOS}
    ticket = textos["ticket"]

    def dudoso(motivo: str) -> PesajeDudoso:
        return PesajeDudoso(
            fila=numero_de_fila,
            ticket=ticket,
            fecha_texto=textos["fecha"],
            placa=textos["placa"],
            material=textos["material"],
            bruto_texto=textos["bruto"],
            tara_texto=textos["tara"],
            neto_texto=textos["neto"],
            motivo=motivo,
        )

    if not ticket:
        return dudoso("falta el ticket")
    if ticket in vistos:
        return dudoso("ticket repetido en el archivo")
    vistos.add(ticket)

    fecha = _fecha(_celda(fila, columnas, "fecha"))
    if fecha is None:
        return dudoso("fecha ilegible")
    if fecha > fecha_referencia:
        return dudoso("fecha futura")
    if not textos["material"]:
        return dudoso("falta el material")

    con_bruto_y_tara = "bruto" in columnas and "tara" in columnas
    pesos: dict[str, Decimal] = {}
    for campo in ("bruto", "tara", "neto"):
        if campo not in columnas:
            continue
        celda = _celda(fila, columnas, campo)
        if _vacia(celda):
            if campo == "neto" and con_bruto_y_tara:
                continue
            return dudoso("peso ilegible")
        try:
            pesos[campo] = _numero(celda)
        except _NumeroAmbiguoError:
            return dudoso("número ambiguo")
        except _NumeroIlegibleError:
            return dudoso("peso ilegible")

    divisor = _MILES if unidad == "kg" else Decimal(1)
    bruto = pesos.get("bruto", Decimal(0)) / divisor
    tara = pesos.get("tara", Decimal(0)) / divisor
    neto_archivo = pesos["neto"] / divisor if "neto" in pesos else None

    if (
        bruto < 0
        or tara < 0
        or (not con_bruto_y_tara and neto_archivo is not None and neto_archivo < 0)
    ):
        return dudoso("peso negativo")
    if con_bruto_y_tara:
        if tara >= bruto:
            return dudoso("tara mayor o igual que el bruto")
        if neto_archivo is not None and abs(neto_archivo - (bruto - tara)) > TOLERANCIA_NETO_T:
            return dudoso(_MOTIVO_NETO)
    neto = neto_archivo if neto_archivo is not None else bruto - tara
    return PesajeLeido(
        fila=numero_de_fila,
        ticket=ticket,
        fecha=fecha,
        placa=textos["placa"],
        cliente=textos["cliente"],
        material=textos["material"],
        bruto_t=bruto,
        tara_t=tara,
        neto_t=neto,
    )


# ----------------------------------------------------------------------------------------- API


def leer(
    datos: bytes,
    nombre: str,
    *,
    fecha_referencia: date,
    unidad_peso: str | None = None,
    mapeo: Mapping[str, int] | None = None,
) -> LecturaBascula:
    """Lee un reporte de bascula; ver el docstring del modulo para formato, unidad y reglas.

    `mapeo` (campo -> indice 0-based) manda sobre la deteccion campo por campo; lo que no mapea se
    sigue detectando. Si faltan columnas indispensables se propaga `MapeoNecesario` (no es rechazo:
    es la peticion de un mapeo). Todo lo demas que impide importar es `ReporteRechazado`.
    """
    extension = PurePath(nombre).suffix.lower()
    if extension not in (".csv", ".xlsx"):
        raise ReporteRechazado("El archivo debe ser .csv o .xlsx.")
    if unidad_peso is not None and unidad_peso not in ("kg", "t"):
        raise ReporteRechazado(_MENSAJE_SIN_UNIDAD)
    if mapeo is not None:
        _validar_mapeo(mapeo)

    huella = hashlib.sha256(datos).hexdigest()
    tope = (
        FILAS_DE_ENCABEZADO + MAX_FILAS + 1
    )  # lo justo para saber si hay una fila de datos de mas
    filas = _filas_de_csv(datos, tope) if extension == ".csv" else list(_filas_de_xlsx(datos, tope))
    if all(not any(not _vacia(c) for c in fila) for fila in filas):
        raise ReporteRechazado("El archivo no tiene filas.")

    fila_encabezado, detectadas = mejor_encabezado(filas)
    if len(filas) - fila_encabezado - 1 > MAX_FILAS:
        raise ReporteRechazado(
            f"El archivo tiene más de {MAX_FILAS} filas de datos; divídalo en reportes menores."
        )
    columnas = {**detectadas, **(mapeo or {})}
    faltan = faltantes(columnas)
    if faltan:
        raise MapeoNecesario(faltan)

    unidad = unidad_peso or unidad_de_los_pesos(filas[fila_encabezado], columnas)
    if unidad is None:
        raise ReporteRechazado(_MENSAJE_SIN_UNIDAD)

    pesajes: list[PesajeLeido] = []
    dudosos: list[PesajeDudoso] = []
    vistos: set[str] = set()
    for indice in range(fila_encabezado + 1, len(filas)):
        fila = filas[indice]
        if all(_vacia(c) for c in fila):
            continue
        resultado = _evaluar(
            fila,
            columnas,
            unidad=unidad,
            fecha_referencia=fecha_referencia,
            vistos=vistos,
            numero_de_fila=indice + 1,
        )
        if isinstance(resultado, PesajeLeido):
            pesajes.append(resultado)
        else:
            dudosos.append(resultado)
    return LecturaBascula(huella=huella, pesajes=tuple(pesajes), dudosos=tuple(dudosos))
