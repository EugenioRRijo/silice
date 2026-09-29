"""Listas de precios nuevas y el historial de cambios que producen (UC‑02, Sesión I1).

Tres responsabilidades, en el orden en que las usa el caso de uso:

1. **Leer** un archivo CSV o XLSX con las columnas `tipo`, `insumo`, `unidad` y `precio`. Todo
   importe se construye con `Decimal` **desde el texto** del archivo, nunca desde un `float`
   (CLAUDE.md §2.3). El CSV ya es texto en el archivo: `pandas.read_csv(dtype=str)` lo lee tal
   cual. El XLSX es distinto: `pandas.read_excel(dtype=str)` **no** basta, porque el motor de
   Excel (`openpyxl`) analiza toda celda numérica como `float` de Python *antes* de que `pandas`
   pueda aplicar `dtype`, que solo convierte ese `float` a texto *después*, delegando el formateo a
   un arreglo intermedio de `numpy` que este módulo no controla ni puede garantizar entre
   versiones. Por eso el XLSX se lee celda por celda con `openpyxl` (`_leer_tabla_excel`), en
   Python puro y sin arreglos intermedios.
2. **Crear** la lista nueva a partir de la vigente: copia todos sus precios y sobrescribe los que
   trae el archivo. Así la lista nueva vale por sí sola para reconstruir cualquier APU, que es lo
   que exige `Catalogo.composicion`. Los insumos que el catálogo no conoce **no se crean** salvo
   que el usuario los elija: `insumos_desconocidos` los lista sin escribir nada y
   `crear_lista_desde_archivo(..., dar_de_alta=...)` da de alta solo los elegidos, con precio solo
   en la lista nueva; los demás se devuelven en `ResumenLista.desconocidos` (UC‑02, flujo 2a).
3. **Detectar y registrar** qué insumos cambiaron de precio entre dos listas. Ese historial es el
   que alimentará al módulo predictivo (Sesión I6), de modo que se persiste siempre, aunque el
   presupuesto nuevo se descarte (UC‑02, flujo 7a).

Un insumo se identifica por (tipo, descripción, unidad), no por su código: el archivo lo escribe un
proveedor que no conoce los códigos internos. La descripción se compara normalizada con
`core.verification.texto.normalizar_texto`, la única tabla de normalización de texto del sistema
(principio DRY), y la unidad con `normalizar_unidad`, la única tabla de alias.

**Homónimos.** La línea base carga como variantes los insumos con la misma descripción y distinto
precio (`Agua`, `Pala`, `Cinta métrica`, `Nivel de mano`; docs/modelo_datos.md §6). Una fila del
archivo actualiza **todas** las variantes de esa descripción y unidad: el archivo declara un precio
por insumo, y distinguir variantes exigiría que el proveedor escribiera los códigos internos.
"""

from __future__ import annotations

from collections import defaultdict
from collections.abc import Collection, Iterable
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
import pandas as pd
from sqlalchemy import select
from sqlalchemy.orm import Session

from core import models
from core.catalog.repositorio import Catalogo
from core.contracts.unidades import normalizar_unidad
from core.verification.texto import normalizar_texto

__all__ = [
    "COLUMNAS_ARCHIVO",
    "PRIMERA_FILA_DE_DATOS",
    "SUFIJOS_EXCEL",
    "CambioDetectado",
    "PrecioLeido",
    "ResumenLista",
    "cambios_precio",
    "crear_lista_desde_archivo",
    "insumos_afectados",
    "insumos_desconocidos",
    "leer_lista_precios",
    "leer_tabla",
    "registrar_cambios",
    "tipo_insumo",
    "variacion_relativa",
]

#: Columnas obligatorias del archivo de precios.
COLUMNAS_ARCHIVO = ("tipo", "insumo", "unidad", "precio")

#: Extensiones que se saben leer. `.xls` exige `xlrd`, que no es dependencia del proyecto.
SUFIJO_CSV = ".csv"
SUFIJOS_EXCEL = (".xlsx", ".xlsm")

#: Primera fila con datos del archivo (y de cada hoja de un libro): la 1 es la de encabezados.
PRIMERA_FILA_DE_DATOS = 2

#: Grafías de la columna `tipo` que no son el valor canónico de `models.TipoInsumo`.
_ALIAS_TIPO = {
    "mano de obra": models.TipoInsumo.MANO_OBRA,
    "mano_de_obra": models.TipoInsumo.MANO_OBRA,
}

#: Un insumo del catálogo visto desde el archivo: (tipo, descripción normalizada, unidad).
_ClaveInsumo = tuple[str, str, str | None]


@dataclass(frozen=True, slots=True)
class PrecioLeido:
    """Una fila del archivo ya validada: qué insumo describe y a qué precio."""

    tipo: models.TipoInsumo
    descripcion: str
    unidad: str | None
    precio: Decimal


@dataclass(frozen=True, slots=True)
class ResumenLista:
    """Qué produjo la carga de un archivo de precios.

    `desconocidos` son las filas cuyo (tipo, descripción, unidad) no está en el catálogo y que no
    se dieron de alta: los ignorados del flujo 2a y los rechazados del 2b (unidad no equivalente a
    la del catálogo). Se reportan y no intervienen en el recálculo, porque dar de alta un insumo es
    una decisión del administrador, no un efecto secundario de leer un archivo (UC‑02, flujo 2a).
    `dados_de_alta` son los insumos que el administrador eligió dar de alta (expediente C6.2), con
    precio solo en `lista`. Sigue el patrón de `ResumenCarga`: la carga devuelve lo que hizo, no
    solo lo que creó.
    """

    lista: models.ListaPrecios
    desconocidos: tuple[PrecioLeido, ...] = ()
    dados_de_alta: tuple[models.Insumo, ...] = ()


_CERO = Decimal(0)


def variacion_relativa(anterior: Decimal, nuevo: Decimal) -> Decimal:
    """Fracción (nuevo − anterior) / anterior: la fórmula de docs/modelo_datos.md §6.

    Única definición del sistema (INV‑004): la usan `CambioDetectado.variacion` y la incidencia por
    partida de `core.budget.actualizacion`.

    **Base nula.** Con `anterior == 0` la variación relativa no está definida. Política declarada:
    devuelve 0, y el registro conserva `precio_anterior = 0`, que es lo que distingue ese caso de
    un precio que no cambió. Un precio 0 es válido (`leer_lista_precios` acepta cualquier no
    negativo), así que la fórmula no puede lanzar `DivisionByZero` en medio de UC‑02.
    """
    return (nuevo - anterior) / anterior if anterior else _CERO


@dataclass(frozen=True, slots=True)
class CambioDetectado:
    """Un insumo cuyo precio difiere entre dos listas, antes de persistirlo."""

    insumo: models.Insumo
    precio_anterior: Decimal
    precio_nuevo: Decimal

    @property
    def variacion(self) -> Decimal:
        """Fracción (nuevo − anterior) / anterior, con la política de base nula declarada."""
        return variacion_relativa(self.precio_anterior, self.precio_nuevo)


# ---------------------------------------------------------------------------------------------
# Lectura del archivo
# ---------------------------------------------------------------------------------------------


def leer_lista_precios(ruta: Path) -> list[PrecioLeido]:
    """Lee un archivo de precios CSV o XLSX y devuelve sus filas validadas, en su mismo orden.

    Lanza `ValueError` citando el archivo y el número de fila si falta una columna, si el tipo de
    insumo no existe o si el precio no es un número no negativo: una lista de precios a medias
    valoraría el presupuesto en silencio (UC‑02, flujo 2a de docs/ERS.md).

    El separador decimal es el punto. Aceptar también la coma exigiría decidir qué significa el
    separador de miles, y ninguna fuente del proyecto lo necesita todavía.
    """
    tabla = leer_tabla(ruta)
    faltantes = [columna for columna in COLUMNAS_ARCHIVO if columna not in tabla.columns]
    if faltantes:
        raise ValueError(
            f"{ruta.name}: faltan las columnas {', '.join(faltantes)}; "
            f"se esperan {', '.join(COLUMNAS_ARCHIVO)}"
        )
    return [
        _a_precio_leido(fila, numero, ruta)
        for numero, fila in enumerate(tabla.to_dict("records"), start=PRIMERA_FILA_DE_DATOS)
    ]


def leer_tabla(ruta: Path, hoja: str | None = None) -> pd.DataFrame:
    """El archivo como tabla de texto: ningún importe pasa por `float` en ningún paso.

    Única lectura de CSV y XLSX sin `float` del sistema: la usan las listas de precios y la
    importación del presupuesto de un tercero (`core.budget.importacion`, expediente C6.4).
    `hoja` elige una hoja del XLSX por su nombre (sin ella, la activa).
    """
    sufijo = ruta.suffix.lower()
    if sufijo == SUFIJO_CSV:
        if hoja is not None:
            raise ValueError(f"{ruta.name}: un CSV no tiene hojas; se esperaba la hoja {hoja!r}")
        tabla = pd.read_csv(ruta, dtype=str, keep_default_na=False, encoding="utf-8")
    elif sufijo in SUFIJOS_EXCEL:
        tabla = _leer_tabla_excel(ruta, hoja)
    else:
        raise ValueError(
            f"{ruta.name}: formato no soportado; se leen {', '.join((SUFIJO_CSV, *SUFIJOS_EXCEL))}"
        )
    tabla.columns = [str(columna).strip().lower() for columna in tabla.columns]
    return tabla.fillna("")


def _leer_tabla_excel(ruta: Path, hoja: str | None = None) -> pd.DataFrame:
    """Lee un XLSX celda por celda con `openpyxl`, sin la ruta de `pandas.read_excel(dtype=str)`.

    Esa ruta no evita el `float`: el motor `openpyxl` analiza toda celda numérica del XML como
    `float` de Python al leer el archivo, y el parámetro `dtype` de `pandas.read_excel` solo
    convierte ese `float` a texto *después*, pasando por un arreglo intermedio de `numpy` cuyo
    formateo no depende de este módulo. Aquí se lee celda por celda y se convierte cada valor a
    texto de una vez, en Python puro y sin arreglos intermedios: si la celda ya es texto se usa tal
    cual (el caso exacto por construcción, el recomendado para la columna `precio`); si Excel la
    guardó como número, se usa `str()` de Python sobre ese valor, que para cualquier precio real
    (con la precisión que un ser humano puede escribir en una celda) reconstruye el mismo texto que
    se escribió: CPython garantiza que su algoritmo de la representación más corta cumple
    `float(str(x)) == x`.

    Sin `hoja` lee la hoja activa; con ella, la de ese nombre (`ValueError` si el libro no la
    tiene).
    """
    libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        if hoja is not None and hoja not in libro.sheetnames:
            raise ValueError(f"{ruta.name}: falta la hoja {hoja!r}")
        filas = (libro[hoja] if hoja is not None else libro.active).iter_rows(values_only=True)
        try:
            encabezados = [_texto_celda(valor) for valor in next(filas)]
        except StopIteration:
            return pd.DataFrame()
        registros = [
            dict(zip(encabezados, (_texto_celda(valor) for valor in fila), strict=False))
            for fila in filas
            if any(valor is not None for valor in fila)
        ]
    finally:
        libro.close()
    return pd.DataFrame.from_records(registros, columns=encabezados)


def _texto_celda(valor: object) -> str:
    """Una celda de Excel como texto exacto, sin pasar por ningún arreglo de `numpy`."""
    if valor is None:
        return ""
    if isinstance(valor, str):
        return valor.strip()
    return str(valor).strip()


def _a_precio_leido(fila: dict[str, object], numero: int, ruta: Path) -> PrecioLeido:
    descripcion = _texto(fila["insumo"])
    if not descripcion:
        raise ValueError(f"{ruta.name}, fila {numero}: la columna insumo está vacía")
    unidad = _texto(fila["unidad"])
    return PrecioLeido(
        tipo=_a_tipo(_texto(fila["tipo"]), numero, ruta),
        descripcion=descripcion,
        unidad=normalizar_unidad(unidad) if unidad else None,
        precio=_a_precio(_texto(fila["precio"]), numero, ruta),
    )


def _texto(valor: object) -> str:
    return str(valor).strip()


def tipo_insumo(texto: str) -> models.TipoInsumo:
    """El tipo de insumo de la columna `tipo`, con sus grafías (`mano de obra`, `mano_de_obra`).

    Única traducción de esa columna: la usan las listas de precios y la importación del
    presupuesto de un tercero. `ValueError` sin posición si no es un tipo conocido; quien la
    llama añade el archivo, la hoja y la fila.
    """
    clave = texto.strip().lower()
    if clave in _ALIAS_TIPO:
        return _ALIAS_TIPO[clave]
    try:
        return models.TipoInsumo(clave)
    except ValueError:
        esperados = ", ".join(tipo.value for tipo in models.TipoInsumo)
        raise ValueError(
            f"tipo de insumo {texto!r} desconocido; se espera uno de {esperados}"
        ) from None


def _a_tipo(valor: str, numero: int, ruta: Path) -> models.TipoInsumo:
    try:
        return tipo_insumo(valor)
    except ValueError as error:
        raise ValueError(f"{ruta.name}, fila {numero}: {error}") from None


def _a_precio(valor: str, numero: int, ruta: Path) -> Decimal:
    try:
        precio = Decimal(valor)
    except InvalidOperation as error:
        raise ValueError(
            f"{ruta.name}, fila {numero}: el precio {valor!r} no es un número decimal"
        ) from error
    # `is_finite()` va primero y en la misma condicion: `Decimal("nan") < 0` lanzaria
    # `InvalidOperation` (un `ArithmeticError`, no el `ValueError` que promete el docstring de
    # `leer_lista_precios`) y `Decimal("Infinity") < 0` es `False`, asi que un precio infinito
    # entraria a la lista en silencio.
    if not precio.is_finite() or precio < 0:
        raise ValueError(
            f"{ruta.name}, fila {numero}: el precio {valor} no es un numero no negativo"
        )
    return precio


# ---------------------------------------------------------------------------------------------
# Creación de la lista
# ---------------------------------------------------------------------------------------------


def crear_lista_desde_archivo(
    session: Session,
    ruta: Path,
    nombre: str,
    moneda: str,
    fecha_vigencia: date,
    origen: str = "",
    dar_de_alta: Collection[PrecioLeido] = (),
) -> ResumenLista:
    """Crea la lista de precios que declara el archivo, partiendo de la vigente a esa fecha.

    La lista nueva nace con **todos** los precios de la anterior y sobrescribe los del archivo, de
    modo que valora el catálogo completo: una lista incompleta haría fallar la reconstrucción de
    cualquier APU con un insumo sin precio (`CatalogoIncompleto`).

    Deja la transacción abierta —confirmarla o deshacerla es de quien llama, como en
    `core.budget.guardar_presupuesto`— y no modifica la lista anterior: es inmutable desde que un
    presupuesto la referencia (docs/modelo_datos.md §2.2).

    Sin `origen` queda el nombre del archivo, que es la procedencia del dato (CLAUDE.md §2.6).

    `dar_de_alta` son las filas de `insumos_desconocidos` que el administrador eligió (UC‑02,
    flujo 2a; expediente C6.2): se crean con código nuevo de su tipo y su precio entra solo en la
    lista nueva. Las filas elegidas con la misma identidad de catálogo (tipo, descripción
    normalizada y, solo en materiales, unidad normalizada) dan **una** alta. Se validan antes de
    escribir nada: `ValueError` si alguna no es un insumo desconocido de este archivo, o si dos
    con la misma identidad traen precios distintos. Sin elegidos, el catálogo no gana insumos.
    """
    leidos = leer_lista_precios(ruta)
    anterior = Catalogo(session).lista_vigente(fecha_vigencia)
    catalogados = _insumos_por_clave(session)
    elegidos = _elegidos(leidos, catalogados, dar_de_alta, ruta)

    lista = models.ListaPrecios(
        nombre=nombre,
        moneda=moneda,
        fecha_vigencia=fecha_vigencia,
        origen=origen or ruta.name,
    )
    session.add(lista)
    session.flush()

    catalogo = Catalogo(session)
    dados_de_alta: list[models.Insumo] = []
    por_identidad: dict[_ClaveInsumo, models.Insumo] = {}
    for leido in elegidos:
        identidad = _identidad_alta(leido)
        insumo = por_identidad.get(identidad)
        if insumo is None:  # varias filas con la misma identidad: una sola alta
            insumo = catalogo.dar_de_alta_insumo(leido.tipo, leido.descripcion, leido.unidad)
            por_identidad[identidad] = insumo
            dados_de_alta.append(insumo)
        # Indexada por la clave de la fila (con su unidad, p. ej. «dia» en un equipo), para que el
        # recorrido de abajo le fije el precio en la lista nueva.
        catalogados[_clave(leido.tipo.value, leido.descripcion, leido.unidad)] = [insumo]

    precios = {precio.insumo_id: precio.precio for precio in anterior.precios}
    desconocidos: list[PrecioLeido] = []
    for leido in leidos:
        variantes = catalogados.get(_clave(leido.tipo.value, leido.descripcion, leido.unidad))
        if not variantes:
            desconocidos.append(leido)
            continue
        for insumo in variantes:
            precios[insumo.id] = leido.precio

    session.add_all(
        models.PrecioInsumo(lista_id=lista.id, insumo_id=insumo_id, precio=precio)
        for insumo_id, precio in precios.items()
    )
    session.flush()
    return ResumenLista(
        lista=lista, desconocidos=tuple(desconocidos), dados_de_alta=tuple(dados_de_alta)
    )


def insumos_desconocidos(session: Session, ruta: Path) -> tuple[PrecioLeido, ...]:
    """La vista previa del flujo 2a de UC‑02: las filas del archivo que se pueden dar de alta.

    Son las filas cuyo (tipo, descripción) normalizado no está en el catálogo con **ninguna**
    unidad, en el orden del archivo. Una fila cuya descripción sí está, con una unidad no
    equivalente, es el flujo 2b («la fila se rechaza y se reporta»): no se ofrece, porque darla de
    alta crearía un homónimo con otra unidad del insumo que ese flujo rechaza. No escribe nada.
    """
    return _candidatos_de_alta(leer_lista_precios(ruta), _insumos_por_clave(session))


def _candidatos_de_alta(
    leidos: Iterable[PrecioLeido], catalogados: dict[_ClaveInsumo, list[models.Insumo]]
) -> tuple[PrecioLeido, ...]:
    conocidas = {(tipo, descripcion) for tipo, descripcion, _unidad in catalogados}
    return tuple(
        leido
        for leido in leidos
        if (leido.tipo.value, normalizar_texto(leido.descripcion)) not in conocidas
    )


def _elegidos(
    leidos: list[PrecioLeido],
    catalogados: dict[_ClaveInsumo, list[models.Insumo]],
    dar_de_alta: Collection[PrecioLeido],
    ruta: Path,
) -> list[PrecioLeido]:
    """Los elegidos para el alta, en el orden del archivo; `ValueError` si alguno no se ofrece."""
    candidatos = _candidatos_de_alta(leidos, catalogados)
    ajenos = [leido for leido in dar_de_alta if leido not in candidatos]
    if ajenos:
        raise ValueError(
            f"{ruta.name}: {', '.join(repr(a.descripcion) for a in ajenos)} no es un insumo "
            "desconocido de este archivo; solo se dan de alta los que lista "
            "insumos_desconocidos (UC-02, flujo 2a)"
        )
    elegidos = set(dar_de_alta)
    filas = {id(leido): numero for numero, leido in enumerate(leidos, start=PRIMERA_FILA_DE_DATOS)}
    en_orden = [leido for leido in candidatos if leido in elegidos]
    por_identidad: dict[_ClaveInsumo, list[PrecioLeido]] = defaultdict(list)
    for leido in en_orden:
        por_identidad[_identidad_alta(leido)].append(leido)
    for mismas in por_identidad.values():
        if len({leido.precio for leido in mismas}) > 1:
            raise ValueError(
                f"{ruta.name}, filas {', '.join(str(filas[id(leido)]) for leido in mismas)}: "
                f"{mismas[0].descripcion!r} se dara de alta una sola vez y esas filas le dan "
                f"precios distintos ({', '.join(str(leido.precio) for leido in mismas)})"
            )
    return en_orden


def _identidad_alta(leido: PrecioLeido) -> _ClaveInsumo:
    """La identidad de catálogo del insumo que daría de alta la fila: tipo, descripción
    normalizada y unidad normalizada solo en materiales (equipos y mano de obra van sin unidad,
    como en `Catalogo.dar_de_alta_insumo`)."""
    unidad = leido.unidad if leido.tipo is models.TipoInsumo.MATERIAL else None
    return _clave(leido.tipo.value, leido.descripcion, unidad)


def _insumos_por_clave(session: Session) -> dict[_ClaveInsumo, list[models.Insumo]]:
    """Índice del catálogo por (tipo, descripción, unidad); una clave puede traer variantes."""
    indice: dict[_ClaveInsumo, list[models.Insumo]] = defaultdict(list)
    for insumo in session.scalars(select(models.Insumo).order_by(models.Insumo.codigo)):
        indice[_clave(insumo.tipo, insumo.descripcion, insumo.unidad)].append(insumo)
    return indice


def _clave(tipo: str, descripcion: str, unidad: str | None) -> _ClaveInsumo:
    return (tipo, normalizar_texto(descripcion), normalizar_unidad(unidad) if unidad else None)


# ---------------------------------------------------------------------------------------------
# Historial de cambios
# ---------------------------------------------------------------------------------------------


def insumos_afectados(
    session: Session, lista_anterior: models.ListaPrecios, lista_nueva: models.ListaPrecios
) -> list[CambioDetectado]:
    """Insumos con precio distinto entre las dos listas, ordenados por código de insumo.

    Un insumo que solo tiene precio en una de las dos no es un cambio de precio: no hay variación
    que calcular, y `CambioPrecio` exige los dos importes.
    """
    anteriores = _precios_de(session, lista_anterior)
    nuevos = session.execute(
        select(models.Insumo, models.PrecioInsumo.precio)
        .join(models.PrecioInsumo, models.PrecioInsumo.insumo_id == models.Insumo.id)
        .where(models.PrecioInsumo.lista_id == lista_nueva.id)
        .order_by(models.Insumo.codigo)
    ).all()
    return [
        CambioDetectado(insumo=insumo, precio_anterior=anteriores[insumo.id], precio_nuevo=precio)
        for insumo, precio in nuevos
        if insumo.id in anteriores and anteriores[insumo.id] != precio
    ]


def registrar_cambios(
    session: Session, lista_anterior: models.ListaPrecios, lista_nueva: models.ListaPrecios
) -> list[models.CambioPrecio]:
    """Persiste un `CambioPrecio` por insumo afectado, fechado en la vigencia de la lista nueva.

    Idempotente: si ese par de listas ya tiene su historial registrado, devuelve el existente en
    vez de duplicarlo. Como `crear_lista_desde_archivo`, no confirma la transacción.
    """
    registrados = list(
        session.scalars(
            select(models.CambioPrecio)
            .where(
                models.CambioPrecio.lista_anterior_id == lista_anterior.id,
                models.CambioPrecio.lista_nueva_id == lista_nueva.id,
            )
            .order_by(models.CambioPrecio.id)
        )
    )
    if registrados:
        return registrados

    cambios = [
        models.CambioPrecio(
            lista_anterior_id=lista_anterior.id,
            lista_nueva_id=lista_nueva.id,
            insumo_id=detectado.insumo.id,
            precio_anterior=detectado.precio_anterior,
            precio_nuevo=detectado.precio_nuevo,
            variacion=detectado.variacion,
            fecha=lista_nueva.fecha_vigencia,
        )
        for detectado in insumos_afectados(session, lista_anterior, lista_nueva)
    ]
    session.add_all(cambios)
    session.flush()
    return cambios


def cambios_precio(
    session: Session,
    insumo: str | None = None,
    desde: date | None = None,
    hasta: date | None = None,
) -> list[models.CambioPrecio]:
    """El historial de `CambioPrecio`, en orden estable (fecha, id), con filtros opcionales.

    Única consulta del historial del sistema: la usan `api/rutas/listas.py` y
    `ui/paginas/historico.py`, que antes la reproducían por separado (hallazgo menor diferido de
    la bitácora F.1; principio DRY). `insumo` filtra por descripción exacta; `desde` y `hasta`
    acotan la fecha del cambio, ambos inclusivos.
    """
    consulta = (
        select(models.CambioPrecio)
        .join(models.Insumo)
        .order_by(models.CambioPrecio.fecha, models.CambioPrecio.id)
    )
    if insumo is not None:
        consulta = consulta.where(models.Insumo.descripcion == insumo)
    if desde is not None:
        consulta = consulta.where(models.CambioPrecio.fecha >= desde)
    if hasta is not None:
        consulta = consulta.where(models.CambioPrecio.fecha <= hasta)
    return list(session.scalars(consulta))


def _precios_de(session: Session, lista: models.ListaPrecios) -> dict[int, Decimal]:
    return {
        insumo_id: precio
        for insumo_id, precio in session.execute(
            select(models.PrecioInsumo.insumo_id, models.PrecioInsumo.precio).where(
                models.PrecioInsumo.lista_id == lista.id
            )
        ).all()
    }
