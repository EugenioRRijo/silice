"""Exportación del presupuesto a Excel: el libro de Elaborar (`exportar_excel`) y el de la hoja del
proyecto (`exportar_hoja`, SF‑E4).

`exportar_excel` presenta el presupuesto, APU, curva de inversión y auditoría. Es el único lugar del
núcleo donde se redondea: sus celdas se escriben con dos decimales porque un presupuesto se lee en
centavos, mientras que los `Decimal` del modelo conservan todos sus dígitos (NORMAS.md §2.3). La
cantidad de decimales no se declara aquí: se toma de
`core.verification.informe.DECIMALES_PRESENTACION`, que ya la fija para el informe de auditoría. Su
libro de **cuatro hojas** y su estructura no cambian (CP‑05, CP‑07, indicador 3 y RNF‑04 lo miden).

`exportar_hoja` es el libro de la hoja del proyecto (spec §9.5): una hoja por pestaña y las de
Trazabilidad y Auditoría. Cantidades y precios van **sin redondear**, con formato de presentación a
dos decimales, y los totales son **fórmulas** vivas (`SUM`, `ROUND`, `SUMIF`, `SUMPRODUCT`: el
archivo guarda los nombres en inglés, que Excel muestra traducidos), de modo que el total del Excel
coincide con `Presupuesto.total` y el principio 3 de NORMAS (redondear solo al presentar) se
respeta. M‑2 (RS‑8): el total exacto es el contractual; una nota declara la diferencia de redondeo
de presentación entre la suma de las filas mostradas y el total mostrado. Las hojas que `core` no
conoce (Espacio y metrajes, origen de precios) llegan como `TablaExportable`.

**M‑1 y RNF‑12, en las dos exportaciones:** toda celda de texto cuyo primer carácter sea `=`, `+`,
`-`, `@`, tabulador o retorno de carro se escribe como texto (tipo cadena y `quotePrefix`), nunca
como fórmula; un número negativo sigue siendo número.

El libro lleva siempre la hoja de auditoría, aunque el presupuesto no tenga hallazgos: el informe se
entrega siempre (principio 7 de NORMAS.md §2), y la hoja enumera las reglas evaluadas (UC‑05 5a).

La tabla de mano de obra de la hoja APU incluye la columna Modalidad: bajo DESTAJO, `sueldo` no es
un sueldo diario sino el precio por unidad de partida del artículo 114 de la LOTTT (decisión D9), y
un Excel que no distingue las dos modalidades invita al malentendido que D9 existe para evitar.
"""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from core.budget.equilibrio import EstadoEquilibrio, punto_equilibrio
from core.contracts.apu import ParametrosCosto
from core.contracts.presupuesto import PartidaPresupuestada, Presupuesto
from core.verification.informe import DECIMALES_PRESENTACION, InformeAuditoria

__all__ = ["TablaExportable", "exportar_excel", "exportar_hoja"]

HOJA_PRESUPUESTO = "Presupuesto"
HOJA_APU = "APU"
HOJA_CURVA = "Curva"
HOJA_AUDITORIA = "Auditoria"

# Hojas de `exportar_hoja` (una por pestaña de la hoja del proyecto, spec §9.1).
PESTANA_ESPACIO = "Espacio y metrajes"
PESTANA_PLAN = "Plan y curva"
PESTANA_EQUILIBRIO = "Punto de equilibrio"
PESTANA_PRECIOS = "Precios"
PESTANA_PARAMETROS = "Parametros"
PESTANA_TRAZABILIDAD = "Trazabilidad"

#: Orden de las pestañas; una tabla extra con uno de estos títulos ocupa su lugar y las demás van al
#: final, en el orden recibido.
ORDEN_PESTANAS = (
    PESTANA_ESPACIO,
    HOJA_APU,
    HOJA_PRESUPUESTO,
    PESTANA_PLAN,
    PESTANA_EQUILIBRIO,
    PESTANA_PRECIOS,
    PESTANA_PARAMETROS,
    HOJA_AUDITORIA,
    PESTANA_TRAZABILIDAD,
)
_HOJAS_PROPIAS = {
    HOJA_APU,
    HOJA_PRESUPUESTO,
    PESTANA_PLAN,
    PESTANA_EQUILIBRIO,
    PESTANA_PARAMETROS,
    HOJA_AUDITORIA,
    PESTANA_TRAZABILIDAD,
}
_CARACTERES_PROHIBIDOS = set("[]:*?/\\")
COLUMNAS_TRAZABILIDAD = (
    "Versión de origen",
    "Versión resultante",
    "Objeto",
    "Clave",
    "Campo",
    "Valor anterior",
    "Valor nuevo",
    "Fuente del cambio",
    "Responsable",
    "Fecha y hora",
)

PASO_REDONDEO = Decimal(1).scaleb(-DECIMALES_PRESENTACION)
FORMATO_PRESENTACION = "#,##0." + "0" * DECIMALES_PRESENTACION
CIEN = Decimal(100)
CERO = Decimal(0)

_NEGRITA = Font(bold=True)

#: Bloque de las reglas evaluadas en la hoja de auditoría: columna G (la 7), tras la F vacía.
_COLUMNA_REGLAS = 7
_TITULOS_REGLAS = ("Regla evaluada", "Nombre", "Hallazgos")

#: Primer carácter de un texto que Excel podría interpretar como fórmula (M‑1, RNF‑12).
_PELIGROSOS = ("=", "+", "-", "@", "\t", "\r")


class _Formula:
    """Una fórmula escrita por el sistema (nunca texto del usuario): no se escapa."""

    __slots__ = ("texto",)

    def __init__(self, texto: str) -> None:
        self.texto = texto


@dataclass(frozen=True, slots=True)
class TablaExportable:
    """Una hoja genérica que `core` no conoce: título, columnas y filas (spec §5.2, §9.5).

    La arma quien conoce el contenido (`ui.hoja`: Espacio y metrajes, origen de precios). Los
    textos pasan por el escape de M‑1; `Decimal`, enteros y fechas van como números y fechas.
    """

    titulo: str
    columnas: tuple[str, ...]
    filas: tuple[tuple[object, ...], ...] = field(default=())

    def __post_init__(self) -> None:
        object.__setattr__(self, "columnas", tuple(self.columnas))
        object.__setattr__(self, "filas", tuple(tuple(fila) for fila in self.filas))


# ---------------------------------------------------------------------------------------------
# Escritura con el escape de M‑1
# ---------------------------------------------------------------------------------------------


def _asegurar_texto(celda: Cell) -> None:
    """M‑1: un texto con prefijo peligroso queda como cadena con `quotePrefix`, no como fórmula."""
    valor = celda.value
    if isinstance(valor, str) and valor.startswith(_PELIGROSOS):
        celda.data_type = "s"
        celda.quotePrefix = True


def _agregar(hoja: Worksheet, valores: Sequence[object], formato: Sequence[int] = ()) -> None:
    """`append` con el escape de M‑1; `formato` son los índices (desde 0) con formato de dos
    decimales."""
    formulas = {i for i, valor in enumerate(valores) if isinstance(valor, _Formula)}
    hoja.append([v.texto if isinstance(v, _Formula) else v for v in valores])
    if not valores:
        return
    fila = hoja[hoja.max_row]
    for indice, celda in enumerate(fila):
        if indice not in formulas:
            _asegurar_texto(celda)
    for indice in formato:
        if indice < len(fila):
            fila[indice].number_format = FORMATO_PRESENTACION


def _poner(hoja: Worksheet, fila: int, columna: int, valor: object) -> Cell:
    celda = hoja.cell(row=fila, column=columna, value=valor)
    _asegurar_texto(celda)
    return celda


def _redondear(valor: Decimal) -> Decimal:
    return valor.quantize(PASO_REDONDEO, rounding=ROUND_HALF_UP)


def _encabezado(hoja: Worksheet, titulos: list[str]) -> None:
    _agregar(hoja, titulos)
    for celda in hoja[hoja.max_row]:
        celda.font = _NEGRITA


# ---------------------------------------------------------------------------------------------
# El libro de Elaborar: cuatro hojas, con dos decimales
# ---------------------------------------------------------------------------------------------


def exportar_excel(presupuesto: Presupuesto, informe: InformeAuditoria, ruta: Path) -> Path:
    """Escribe el libro con las cuatro hojas y devuelve la ruta escrita."""
    libro = Workbook()
    _hoja_presupuesto(libro.active, presupuesto)
    _hoja_apu(libro.create_sheet(HOJA_APU), presupuesto)
    _hoja_curva(libro.create_sheet(HOJA_CURVA), presupuesto)
    _hoja_auditoria(libro.create_sheet(HOJA_AUDITORIA), informe)
    libro.save(ruta)
    return ruta


def _hoja_presupuesto(hoja: Worksheet, presupuesto: Presupuesto) -> None:
    hoja.title = HOJA_PRESUPUESTO
    _encabezado(
        hoja,
        ["N.º", "Código", "Descripción", "Unidad", "Cantidad", "Precio unitario", "Total"],
    )
    for numero, partida in enumerate(presupuesto.partidas, start=1):
        _agregar(
            hoja,
            [
                numero,
                partida.item.codigo_partida,
                partida.item.descripcion,
                partida.item.unidad,
                _redondear(partida.item.cantidad),
                _redondear(partida.resultado.precio_unitario),
                _redondear(partida.total),
            ],
        )
    fila = hoja.max_row + 1
    hoja.cell(row=fila, column=6, value="TOTAL").font = _NEGRITA
    hoja.cell(row=fila, column=7, value=_redondear(presupuesto.total)).font = _NEGRITA


def _hoja_apu(hoja: Worksheet, presupuesto: Presupuesto) -> None:
    for partida in presupuesto.partidas:
        _bloque_apu(hoja, partida)
        _agregar(hoja, [])


def _bloque_apu(hoja: Worksheet, partida: PartidaPresupuestada) -> None:
    apu = partida.apu
    _agregar(hoja, [f"{apu.codigo_partida} · {apu.descripcion}"])
    hoja[hoja.max_row][0].font = _NEGRITA
    _agregar(hoja, ["Unidad", apu.unidad, "Rendimiento", _redondear(apu.rendimiento)])

    _encabezado(hoja, ["Materiales", "Unidad", "Cantidad", "Precio", "Total"])
    for linea in apu.materiales:
        _agregar(
            hoja,
            [
                linea.descripcion,
                linea.unidad,
                _redondear(linea.cantidad),
                _redondear(linea.precio),
                _redondear(linea.total),
            ],
        )

    _encabezado(hoja, ["Equipos", "Cantidad", "Precio", "Depreciación", "Total"])
    for equipo in apu.equipos:
        _agregar(
            hoja,
            [
                equipo.descripcion,
                _redondear(equipo.cantidad),
                _redondear(equipo.precio),
                equipo.depreciacion,
                _redondear(equipo.total),
            ],
        )

    _encabezado(hoja, ["Mano de obra", "Obreros", "Sueldo", "Modalidad", "Total"])
    for obrero in apu.mano_obra:
        _agregar(
            hoja,
            [
                obrero.descripcion,
                _redondear(obrero.cantidad),
                _redondear(obrero.sueldo),
                obrero.modalidad.value,
                _redondear(obrero.total),
            ],
        )

    resultado = partida.resultado
    for etiqueta, valor in (
        ("Materiales", resultado.materiales),
        ("Equipos", resultado.equipos),
        ("Mano de obra", resultado.mano_obra),
        ("Costo directo", resultado.costo_directo),
        ("Con administración", resultado.con_administracion),
        ("Precio unitario", resultado.precio_unitario),
    ):
        _agregar(hoja, [etiqueta, _redondear(valor)])


def _hoja_curva(hoja: Worksheet, presupuesto: Presupuesto) -> None:
    _encabezado(hoja, ["Período", "Monto", "Acumulado", "% acumulado"])
    total = presupuesto.total
    for punto in presupuesto.curva:
        porcentaje = CERO if total == CERO else punto.acumulado / total * CIEN
        _agregar(
            hoja,
            [
                punto.periodo,
                _redondear(punto.monto),
                _redondear(punto.acumulado),
                _redondear(porcentaje),
            ],
        )


def _hoja_auditoria(hoja: Worksheet, informe: InformeAuditoria) -> None:
    """Una fila por hallazgo (columnas A a E) y, a la derecha, las reglas evaluadas.

    Las reglas van en un bloque aparte (columnas G a I, tras la F vacía) para que el libro declare
    lo que verificó aunque no haya hallazgos (UC‑05 5a, INV‑092) sin mezclar filas de otra clase en
    la tabla de hallazgos.
    """
    _encabezado(hoja, ["Regla", "Severidad", "Descripción", "Impacto", "Orígenes"])
    for hallazgo in informe.hallazgos:
        _agregar(
            hoja,
            [
                hallazgo.regla,
                hallazgo.severidad.name,
                hallazgo.descripcion,
                None if hallazgo.impacto is None else _redondear(hallazgo.impacto),
                ", ".join(hallazgo.origen_ids),
            ],
        )
    reglas = informe.cuenta_por_regla_evaluada()
    if not reglas:
        return
    for columna, titulo in enumerate(_TITULOS_REGLAS, start=_COLUMNA_REGLAS):
        _poner(hoja, 1, columna, titulo).font = _NEGRITA
    for fila, valores in enumerate(reglas, start=2):
        for columna, valor in enumerate(valores, start=_COLUMNA_REGLAS):
            _poner(hoja, fila, columna, valor)


# ---------------------------------------------------------------------------------------------
# El libro de la hoja del proyecto (SF‑E4)
# ---------------------------------------------------------------------------------------------


def exportar_hoja(
    presupuesto: Presupuesto,
    informe: InformeAuditoria,
    ruta: Path,
    parametros: ParametrosCosto,
    *,
    trazabilidad: TablaExportable | None = None,
    hojas_extra: Sequence[TablaExportable] = (),
) -> Path:
    """Escribe el libro de la hoja del proyecto con fórmulas vivas y devuelve la ruta (spec §9.5).

    `parametros` son los de la versión: se escriben en la hoja Parametros y las fórmulas de los APU
    los leen. `trazabilidad` son los `RegistroCambio` de la versión como tabla
    (`core.budget.edicion.tabla_trazabilidad`); sin ella la hoja lleva solo sus encabezados.
    `hojas_extra` son las hojas que `core` no conoce (Espacio y metrajes, Precios…): una tabla con
    título de una pestaña ocupa su lugar y las demás van al final. `ValueError` si un título no es
    un nombre de hoja válido o repite el de una hoja propia.
    """
    extras = _validar_titulos(hojas_extra)
    libro = Workbook()
    libro.remove(libro.active)
    hojas = {titulo: libro.create_sheet(titulo) for titulo in _orden(extras)}

    celdas_apu = _hoja_apu_con_formulas(hojas[HOJA_APU], presupuesto)
    total = _hoja_presupuesto_con_formulas(hojas[HOJA_PRESUPUESTO], presupuesto, celdas_apu)
    _hoja_plan(hojas[PESTANA_PLAN], presupuesto, total)
    _hoja_equilibrio(hojas[PESTANA_EQUILIBRIO], presupuesto, parametros, total)
    _hoja_parametros(hojas[PESTANA_PARAMETROS], parametros)
    _hoja_auditoria(hojas[HOJA_AUDITORIA], informe)
    _escribir_tabla(
        hojas[PESTANA_TRAZABILIDAD],
        trazabilidad or TablaExportable(PESTANA_TRAZABILIDAD, COLUMNAS_TRAZABILIDAD),
    )
    for tabla in extras:
        _escribir_tabla(hojas[tabla.titulo], tabla)
    libro.save(ruta)
    return ruta


def _validar_titulos(tablas: Sequence[TablaExportable]) -> list[TablaExportable]:
    vistos: set[str] = set()
    for tabla in tablas:
        titulo = tabla.titulo
        if (
            not titulo
            or len(titulo) > 31
            or _CARACTERES_PROHIBIDOS & set(titulo)
            or titulo.startswith("'")
            or titulo.endswith("'")
        ):
            raise ValueError(
                f"titulo de hoja invalido: {titulo!r} (1 a 31 caracteres, sin []:*?/\\)"
            )
        if titulo in _HOJAS_PROPIAS or titulo.lower() in vistos:
            raise ValueError(f"titulo de hoja repetido: {titulo!r}")
        vistos.add(titulo.lower())
    return list(tablas)


def _orden(extras: Sequence[TablaExportable]) -> list[str]:
    titulos = {tabla.titulo for tabla in extras}
    previstos = [t for t in ORDEN_PESTANAS if t in _HOJAS_PROPIAS or t in titulos]
    return previstos + [tabla.titulo for tabla in extras if tabla.titulo not in ORDEN_PESTANAS]


def _escribir_tabla(hoja: Worksheet, tabla: TablaExportable) -> None:
    _encabezado(hoja, list(tabla.columnas))
    for fila in tabla.filas:
        valores = [_valor_de_tabla(valor) for valor in fila]
        _agregar(
            hoja,
            valores,
            formato=[i for i, v in enumerate(valores) if isinstance(v, Decimal)],
        )


def _valor_de_tabla(valor: object) -> object:
    if valor is None or isinstance(valor, (str, int, float, Decimal, date, datetime)):
        return valor
    return str(valor)


def _ref(hoja: str, celda: str) -> str:
    return f"'{hoja}'!{celda}"


def _hoja_parametros(hoja: Worksheet, parametros: ParametrosCosto) -> None:
    _encabezado(hoja, ["Parámetro", "Valor"])
    for etiqueta, valor in (
        ("FCAS", parametros.fcas),
        ("Bono de alimentación", parametros.bono_alimentacion),
        ("Administración", parametros.administracion),
        ("Utilidad", parametros.utilidad),
    ):
        _agregar(hoja, [etiqueta, valor])


#: Dónde lee cada fórmula los parámetros: filas 2 a 5 de la hoja Parametros, columna B.
_FCAS = _ref(PESTANA_PARAMETROS, "$B$2")
_BONO = _ref(PESTANA_PARAMETROS, "$B$3")
_ADMINISTRACION = _ref(PESTANA_PARAMETROS, "$B$4")
_UTILIDAD = _ref(PESTANA_PARAMETROS, "$B$5")


def _hoja_apu_con_formulas(hoja: Worksheet, presupuesto: Presupuesto) -> list[tuple[str, str]]:
    """Un bloque por renglón; devuelve, por renglón, las celdas de su costo directo y su PU."""
    celdas: list[tuple[str, str]] = []
    for partida in presupuesto.partidas:
        celdas.append(_bloque_apu_con_formulas(hoja, partida))
        _agregar(hoja, [])
    return celdas


def _bloque_apu_con_formulas(hoja: Worksheet, partida: PartidaPresupuestada) -> tuple[str, str]:
    """El bloque APU con las fórmulas de NORMAS.md §4: materiales sin dividir; equipos entre el
    rendimiento; mano de obra a jornal con FCAS y bono entre el rendimiento y a destajo completa;
    administración y utilidad en cascada."""
    apu = partida.apu
    _agregar(hoja, [f"{apu.codigo_partida} · {apu.descripcion}"])
    hoja[hoja.max_row][0].font = _NEGRITA
    _agregar(hoja, ["Unidad", apu.unidad, "Rendimiento", apu.rendimiento], formato=[3])
    rendimiento = f"D{hoja.max_row}"

    def fila_siguiente() -> int:
        return hoja.max_row + 1

    _encabezado(hoja, ["Materiales", "Unidad", "Cantidad", "Precio", "Total"])
    primera = fila_siguiente()
    for linea in apu.materiales:
        f = fila_siguiente()
        _agregar(
            hoja,
            [
                linea.descripcion,
                linea.unidad,
                linea.cantidad,
                linea.precio,
                _Formula(f"=C{f}*D{f}"),
            ],
            formato=[2, 3, 4],
        )
    materiales = (primera, hoja.max_row) if apu.materiales else None

    _encabezado(hoja, ["Equipos", "Cantidad", "Precio", "Depreciación", "Total"])
    primera = fila_siguiente()
    for equipo in apu.equipos:
        f = fila_siguiente()
        _agregar(
            hoja,
            [
                equipo.descripcion,
                equipo.cantidad,
                equipo.precio,
                equipo.depreciacion,
                _Formula(f"=B{f}*C{f}*D{f}"),
            ],
            formato=[1, 2, 4],
        )
    equipos = (primera, hoja.max_row) if apu.equipos else None

    _encabezado(hoja, ["Mano de obra", "Obreros", "Sueldo", "Modalidad", "Total"])
    primera = fila_siguiente()
    for obrero in apu.mano_obra:
        f = fila_siguiente()
        _agregar(
            hoja,
            [
                obrero.descripcion,
                obrero.cantidad,
                obrero.sueldo,
                obrero.modalidad.value,
                _Formula(f"=B{f}*C{f}"),
            ],
            formato=[1, 2, 4],
        )
    mano_obra = (primera, hoja.max_row) if apu.mano_obra else None

    def sumar(rango: tuple[int, int] | None, formula: str) -> _Formula:
        return _Formula(formula.format(a=rango[0], b=rango[1]) if rango else "=0")

    def fila_de(etiqueta: str, formula: _Formula) -> int:
        _agregar(hoja, [etiqueta, formula], formato=[1])
        return hoja.max_row

    fila_mat = fila_de("Materiales", sumar(materiales, "=SUM(E{a}:E{b})"))
    fila_equ = fila_de("Equipos", sumar(equipos, f"=SUM(E{{a}}:E{{b}})/{rendimiento}"))
    fila_mo = fila_de(
        "Mano de obra",
        sumar(
            mano_obra,
            '=(SUMIF(D{a}:D{b},"jornal",E{a}:E{b})*(1+'
            + _FCAS
            + ")+"
            + _BONO
            + '*SUMIF(D{a}:D{b},"jornal",B{a}:B{b}))/'
            + rendimiento
            + '+SUMIF(D{a}:D{b},"destajo",E{a}:E{b})',
        ),
    )
    fila_cd = fila_de("Costo directo", _Formula(f"=B{fila_mat}+B{fila_equ}+B{fila_mo}"))
    fila_adm = fila_de("Con administración", _Formula(f"=B{fila_cd}*(1+{_ADMINISTRACION})"))
    fila_pu = fila_de("Precio unitario", _Formula(f"=B{fila_adm}*(1+{_UTILIDAD})"))
    return f"B{fila_cd}", f"B{fila_pu}"


def _hoja_presupuesto_con_formulas(
    hoja: Worksheet, presupuesto: Presupuesto, celdas_apu: list[tuple[str, str]]
) -> int:
    """La hoja Presupuesto; devuelve la fila de su TOTAL (las demás hojas lo citan)."""
    _encabezado(
        hoja,
        [
            "N.º",
            "Código",
            "Descripción",
            "Unidad",
            "Cantidad",
            "Precio unitario",
            "Total",
            "Costo directo",
        ],
    )
    for numero, (partida, (cd, pu)) in enumerate(
        zip(presupuesto.partidas, celdas_apu, strict=True), start=1
    ):
        f = hoja.max_row + 1
        _agregar(
            hoja,
            [
                numero,
                partida.item.codigo_partida,
                partida.item.descripcion,
                partida.item.unidad,
                partida.item.cantidad,
                _Formula(f"={_ref(HOJA_APU, pu)}"),
                _Formula(f"=E{f}*F{f}"),
                _Formula(f"={_ref(HOJA_APU, cd)}"),
            ],
            formato=[4, 5, 6, 7],
        )
    ultima = hoja.max_row
    total = ultima + 1
    rango = f"G2:G{ultima}"
    _agregar(
        hoja,
        ["TOTAL", None, None, None, None, None, _Formula(f"=SUM({rango})" if ultima > 1 else "=0")],
        formato=[6],
    )
    for celda in hoja[total]:
        celda.font = _NEGRITA
    _agregar(
        hoja,
        [
            "Nota",
            "El total exacto es el contractual. Las filas se muestran redondeadas a "
            f"{DECIMALES_PRESENTACION} decimales: la suma de las filas mostradas puede diferir del "
            "total mostrado por redondeo de presentación.",
        ],
    )
    redondeo = DECIMALES_PRESENTACION
    suma_filas = hoja.max_row + 1
    _agregar(
        hoja,
        [
            "Suma de filas mostradas",
            None,
            None,
            None,
            None,
            None,
            _Formula(f"=SUMPRODUCT(ROUND({rango},{redondeo}))" if ultima > 1 else "=0"),
        ],
        formato=[6],
    )
    mostrado = hoja.max_row + 1
    _agregar(
        hoja,
        ["Total mostrado", None, None, None, None, None, _Formula(f"=ROUND(G{total},{redondeo})")],
        formato=[6],
    )
    _agregar(
        hoja,
        [
            "Diferencia de redondeo",
            None,
            None,
            None,
            None,
            None,
            _Formula(f"=G{mostrado}-G{suma_filas}"),
        ],
        formato=[6],
    )
    return total


def _hoja_plan(hoja: Worksheet, presupuesto: Presupuesto, total: int) -> None:
    _encabezado(hoja, ["Período", "Monto", "Acumulado", "% acumulado"])
    ventas = _ref(HOJA_PRESUPUESTO, f"$G${total}")
    for punto in presupuesto.curva:
        f = hoja.max_row + 1
        _agregar(
            hoja,
            [
                punto.periodo,
                punto.monto,
                _Formula(f"=SUM(B$2:B{f})"),
                _Formula(f"=IF({ventas}=0,0,C{f}/{ventas}*100)"),
            ],
            formato=[1, 2, 3],
        )


def _hoja_equilibrio(
    hoja: Worksheet, presupuesto: Presupuesto, parametros: ParametrosCosto, total: int
) -> None:
    """Costo variable, costo fijo y PE como fórmulas sobre el Presupuesto y los Parametros; la
    ubicación en la curva y el estado son los de `punto_equilibrio`. Si el estado no admite las
    fórmulas (V = 0, margen ≤ 0) los valores van fijos y el mensaje lo dice."""
    r = punto_equilibrio(presupuesto, parametros)
    ultima = total - 1
    _encabezado(hoja, ["Concepto", "Valor"])
    _agregar(hoja, ["Ventas (V)", _Formula(f"={_ref(HOJA_PRESUPUESTO, f'G{total}')}")], [1])
    v = hoja.max_row
    cv_formula = (
        _Formula(
            f"=SUMPRODUCT({_ref(HOJA_PRESUPUESTO, f'E2:E{ultima}')},"
            f"{_ref(HOJA_PRESUPUESTO, f'H2:H{ultima}')})"
        )
        if ultima > 1
        else _Formula("=0")
    )
    _agregar(hoja, ["Costo variable (CV)", cv_formula], [1])
    cv = hoja.max_row
    _agregar(hoja, ["Costo fijo (CF)", _Formula(f"=B{cv}*{_ADMINISTRACION}")], [1])
    cf = hoja.max_row
    _agregar(hoja, ["Utilidad", _Formula(f"=B{v}-B{cv}-B{cf}")], [1])
    con_formulas = r.estado in (EstadoEquilibrio.CALCULADO, EstadoEquilibrio.NO_SE_ALCANZA)
    if con_formulas:
        _agregar(hoja, ["Margen de contribución (m)", _Formula(f"=(B{v}-B{cv})/B{v}")])
        _agregar(hoja, ["Punto de equilibrio (PE)", _Formula(f"=B{cf}*B{v}/(B{v}-B{cv})")], [1])
        pe = hoja.max_row
        _agregar(hoja, ["PE como % de las ventas", _Formula(f"=B{pe}/B{v}*100")], [1])
    else:
        _agregar(hoja, ["Margen de contribución (m)", r.margen])
        _agregar(hoja, ["Punto de equilibrio (PE)", r.punto], [1])
        _agregar(hoja, ["PE como % de las ventas", r.porcentaje], [1])
    _agregar(hoja, ["Período en que se alcanza", r.periodo])
    _agregar(hoja, ["Fracción del período", r.fraccion_periodo])
    _agregar(hoja, ["Estado", r.estado.value])
    _agregar(hoja, ["Mensaje", r.mensaje])
