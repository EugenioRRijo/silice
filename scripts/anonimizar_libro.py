"""Anonimiza un libro de presupuesto real (.xlsx) para que pueda entrar al repositorio.

Por que existe: la decision V-D3 de Silice 1.5
(`docs/especificaciones/2026-10-13-silice-1.5-fusion-design.md`) deja los libros de Excel reales
fuera del repositorio, porque son confidenciales. Lo que si entra son casos *derivados*: el mismo
libro, con sus formulas, hojas y montos, pero sin nada que identifique al cliente, al contratista,
la obra, la direccion ni a las personas que lo firman.

Por que el script descubre los datos en vez de traerlos escritos: una lista de nombres dentro del
repositorio seria, ella misma, la fuga que se quiere evitar. Los datos identificables se leen de la
PORTADA de la plantilla (rotulos CLIENTE, CONTRATISTA, UBICACION, firmas y nombre de la obra) y del
patron del RIF; despues se sustituyen en todas las hojas, en los encabezados y pies de impresion, en
los comentarios, en los mensajes de las validaciones de datos y en las propiedades del documento.

Lo que se QUITA en vez de sustituirse, porque no se puede revisar texto por texto:

- Las imagenes embebidas de cada hoja (los logos del membrete del cliente o del contratista): se
  eliminan todas. El caso anonimizado no lleva ninguna imagen.
- Los hipervinculos de celda (destino, texto mostrado e informacion en pantalla): una direccion web
  o de correo identifica tanto como un nombre. El texto de la celda queda, ya sustituido.
- Las validaciones de datos de tipo lista con los valores escritos en la propia regla
  (`"opcion1,opcion2"`): suelen enumerar personas o empresas. Las que apuntan a un rango se
  conservan, porque ese rango ya se anonimizo como cualquier celda.

Por que verifica al final: sustituir es facil de hacer a medias (un espacio doble, una mayuscula, un
literal dentro de una formula). Por eso el script relee la SALIDA y busca de nuevo cada dato
descubierto (tambien en hipervinculos y validaciones) y comprueba que no quede ninguna imagen,
ningun hipervinculo ni ninguna lista literal de validacion; si algo sobrevive, borra la SALIDA y
termina con codigo 1. Un caso anonimizado que no paso esa verificacion no existe.

Lo que no hace: no toca montos, cantidades ni fechas, y no anonimiza nombres que no esten en la
PORTADA (por ejemplo, un proveedor citado solo en una hoja de insumos). Esa revision sigue siendo de
quien sube el caso.

Uso::

    uv run python scripts/anonimizar_libro.py ENTRADA.xlsx SALIDA.xlsx [--etiqueta A]
"""

from __future__ import annotations

import argparse
import re
import sys
import unicodedata
import zipfile
from collections import defaultdict
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.packaging.custom import CustomPropertyList
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.formula import ArrayFormula

#: Texto que sustituye a cada clase de dato. `{e}` es la etiqueta del caso.
SUSTITUTOS: dict[str, str] = {
    "cliente": "CLIENTE-{e}",
    "contratista": "CONTRATISTA-{e}",
    "ubicacion": "UBICACIÓN ANONIMIZADA",
    "firma": "PERSONA-{e}",
    "rif": "J-00000000-0",
}

#: Rotulos de la PORTADA (normalizados: sin acentos, en minusculas, sin ':').
ROTULOS: dict[str, str] = {
    "cliente": "cliente",
    "contratista": "contratista",
    "ubicacion": "ubicacion",
}
ROTULOS_FIRMA: tuple[str, ...] = ("elaborado por", "aprobado por")
ROTULO_OBRA = "presupuesto de obra"

RIF = re.compile(r"(?<![A-Za-z0-9])[JVEGP]-?\d{7,9}-?\d?(?![0-9])")
RIF_NEUTRO = SUSTITUTOS["rif"]

#: Un fragmento mas corto que esto no se usa como variante (seria demasiado generico).
LARGO_MINIMO_VARIANTE = 5

PROPIEDADES = ("creator", "lastModifiedBy", "title", "subject", "keywords", "description")
PROPIEDADES_EXTENDIDAS = ("company", "manager")
ANONIMO = "anonimo"

LITERAL_FORMULA = re.compile(r'"(?:[^"]|"")*"')


@dataclass(frozen=True)
class Dato:
    """Un texto identificable descubierto en el libro y lo que lo sustituye."""

    tipo: str
    texto: str
    sustituto: str

    @property
    def patron(self) -> re.Pattern[str]:
        palabras = [re.escape(p) for p in self.texto.split()]
        return re.compile(r"\s+".join(palabras), re.IGNORECASE)


def _normalizar(texto: str) -> str:
    sin_acentos = unicodedata.normalize("NFKD", texto)
    sin_acentos = "".join(c for c in sin_acentos if not unicodedata.combining(c))
    return " ".join(sin_acentos.casefold().strip().rstrip(":").split())


def _es_texto(valor: object) -> bool:
    return isinstance(valor, str) and bool(valor.strip()) and not valor.startswith("=")


def _variantes(tipo: str, valor: str) -> list[str]:
    """El valor completo, la parte antes de " (" y la parte antes de ",".

    En la ubicacion, ademas, cada tramo entre comas: una direccion es una lista de lugares y
    cualquiera de ellos (la calle, el edificio) puede aparecer suelto en otra hoja.
    """
    completo = " ".join(valor.split())
    candidatos = [completo]
    antes_parentesis = completo.split(" (")[0].strip()
    candidatos.append(antes_parentesis)
    candidatos.append(antes_parentesis.split(",")[0].strip())
    if tipo == "ubicacion":
        candidatos.extend(t.strip() for t in antes_parentesis.split(","))
    vistos: list[str] = []
    for c in candidatos:
        c = c.strip(" .;-")
        if len(c) >= LARGO_MINIMO_VARIANTE and any(ch.isalpha() for ch in c):
            if c.casefold() not in {v.casefold() for v in vistos}:
                vistos.append(c)
    return vistos


def _hoja_portada(libro: Workbook):
    for hoja in libro.worksheets:
        if _normalizar(hoja.title) == "portada":
            return hoja
    return None


def _a_la_derecha(hoja, fila: int, columna: int) -> str | None:
    for c in range(columna + 1, columna + 4):
        valor = hoja.cell(row=fila, column=c).value
        if valor is not None and str(valor).strip():
            return valor if _es_texto(valor) else None
    return None


def descubrir(libro: Workbook, etiqueta: str) -> list[Dato]:
    """Lee la PORTADA y devuelve los datos identificables, del mas largo al mas corto."""
    portada = _hoja_portada(libro)
    if portada is None:
        return []
    crudos: dict[str, list[str]] = defaultdict(list)
    obras: list[str] = []
    for fila in portada.iter_rows():
        for celda in fila:
            if not isinstance(celda.value, str):
                continue
            rotulo = _normalizar(celda.value)
            for tipo, esperado in ROTULOS.items():
                if rotulo == esperado:
                    valor = _a_la_derecha(portada, celda.row, celda.column)
                    if valor:
                        crudos[tipo].append(valor)
            for firma in ROTULOS_FIRMA:
                if rotulo.startswith(firma):
                    resto = celda.value.split(":", 1)[1] if ":" in celda.value else ""
                    debajo = portada.cell(row=celda.row + 1, column=celda.column).value
                    for valor in (resto, debajo):
                        if _es_texto(valor) and set(str(valor).strip()) - {"_", " "}:
                            crudos["firma"].append(str(valor))
            if rotulo == ROTULO_OBRA:
                debajo = portada.cell(row=celda.row + 1, column=celda.column).value
                if _es_texto(debajo):
                    obras.append(str(debajo))

    datos: list[Dato] = []
    conocidos: set[str] = set()
    for tipo in ("cliente", "contratista", "ubicacion", "firma"):
        sustituto = SUSTITUTOS[tipo].format(e=etiqueta)
        for valor in crudos.get(tipo, []):
            for variante in _variantes(tipo, valor):
                if variante.casefold() in conocidos:
                    continue
                conocidos.add(variante.casefold())
                datos.append(Dato(tipo, variante, sustituto))
    datos.sort(key=lambda d: len(d.texto), reverse=True)

    # El nombre de la obra solo es identificable si nombra al cliente; entonces se sustituye
    # por el mismo nombre con el cliente ya anonimizado.
    clientes = [d for d in datos if d.tipo == "cliente"]
    for obra in obras:
        if any(d.patron.search(obra) for d in clientes):
            anonima = _sustituir_texto(obra, clientes)[0]
            completo = " ".join(obra.split())
            if completo.casefold() not in conocidos:
                conocidos.add(completo.casefold())
                datos.insert(0, Dato("obra", completo, anonima))
    return datos


def _sustituir_texto(texto: str, datos: list[Dato]) -> tuple[str, set[str]]:
    """Sustituye los datos y el RIF en un texto; devuelve el texto y los tipos que cambiaron."""
    tipos: set[str] = set()
    for dato in datos:
        nuevo = dato.patron.sub(lambda _m, s=dato.sustituto: s, texto)
        if nuevo != texto:
            tipos.add(dato.tipo)
            texto = nuevo
    nuevo = RIF.sub(lambda m: m.group(0) if m.group(0) == RIF_NEUTRO else RIF_NEUTRO, texto)
    if nuevo != texto:
        tipos.add("rif")
        texto = nuevo
    return texto, tipos


def _sustituir_formula(formula: str, datos: list[Dato]) -> tuple[str, set[str]]:
    """Solo toca los literales entre comillas; las referencias y funciones quedan igual."""
    tipos: set[str] = set()

    def literal(m: re.Match[str]) -> str:
        nuevo, t = _sustituir_texto(m.group(0), datos)
        tipos.update(t)
        return nuevo

    return LITERAL_FORMULA.sub(literal, formula), tipos


def _sustituir_valor(valor: object, datos: list[Dato]) -> tuple[object, set[str]]:
    if isinstance(valor, ArrayFormula):
        texto, tipos = _sustituir_formula(valor.text or "", datos)
        if tipos:
            valor.text = texto
        return valor, tipos
    if not isinstance(valor, str):
        return valor, set()
    if valor.startswith("="):
        return _sustituir_formula(valor, datos)
    return _sustituir_texto(valor, datos)


def _partes_encabezado(hoja) -> Iterator[object]:
    hf = hoja.HeaderFooter
    for nombre in (
        "oddHeader",
        "oddFooter",
        "evenHeader",
        "evenFooter",
        "firstHeader",
        "firstFooter",
    ):
        bloque = getattr(hf, nombre, None)
        if bloque is None:
            continue
        for lado in ("left", "center", "right"):
            parte = getattr(bloque, lado, None)
            if parte is not None:
                yield parte


def anonimizar(libro: Workbook, datos: list[Dato]) -> dict[str, dict[str, int]]:
    """Aplica la sustitucion a todo el libro. Devuelve {tipo: {hoja: celdas cambiadas}}."""
    informe: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))

    def anotar(tipos: set[str], lugar: str) -> None:
        for t in tipos:
            informe[t][lugar] += 1

    for hoja in libro.worksheets:
        for fila in hoja.iter_rows():
            for celda in fila:
                if celda.value is not None:
                    nuevo, tipos = _sustituir_valor(celda.value, datos)
                    if tipos:
                        celda.value = nuevo
                        anotar(tipos, hoja.title)
                if celda.comment is not None:
                    texto, tipos = _sustituir_texto(celda.comment.text or "", datos)
                    celda.comment.text = texto
                    celda.comment.author = ANONIMO
                    anotar(tipos, f"{hoja.title} (comentarios)")
        for parte in _partes_encabezado(hoja):
            if parte.text:
                texto, tipos = _sustituir_texto(parte.text, datos)
                if tipos:
                    parte.text = texto
                    anotar(tipos, f"{hoja.title} (encabezado/pie)")
    return {t: dict(h) for t, h in informe.items()}


def _es_lista_literal(validacion) -> bool:
    """Validacion de tipo lista cuyos valores van escritos en la regla (`"a,b,c"`)."""
    formula = (validacion.formula1 or "").strip()
    return validacion.type == "list" and formula.startswith('"')


_TEXTOS_VALIDACION = ("prompt", "promptTitle", "error", "errorTitle")


def quitar_imagenes_vinculos_y_validaciones(
    libro: Workbook, datos: list[Dato]
) -> tuple[int, int, int]:
    """Quita las imagenes (logos), los hipervinculos de celda y las listas literales de
    validacion; sustituye los datos en los mensajes de las validaciones que quedan. Devuelve los
    tres conteos."""
    imagenes = vinculos = listas = 0
    for hoja in libro.worksheets:
        imagenes += len(hoja._images)
        hoja._images = []
        for fila in hoja.iter_rows():
            for celda in fila:
                if celda.hyperlink is not None:
                    celda.hyperlink = None
                    vinculos += 1
        hoja._hyperlinks = []
        conservadas = []
        for validacion in hoja.data_validations.dataValidation:
            if _es_lista_literal(validacion):
                listas += 1
                continue
            for campo in _TEXTOS_VALIDACION:
                texto = getattr(validacion, campo, None)
                if texto:
                    setattr(validacion, campo, _sustituir_texto(texto, datos)[0])
            conservadas.append(validacion)
        hoja.data_validations.dataValidation = conservadas
    return imagenes, vinculos, listas


def _es_externo(destino: str | None) -> bool:
    return bool(destino) and re.search(r"\[\d+\]", destino) is not None


def limpiar_vinculos_y_propiedades(libro: Workbook) -> tuple[int, int]:
    """Quita vinculos externos y nombres que apuntan a ellos; deja las propiedades en 'anonimo'."""
    vinculos = len(libro._external_links)
    libro._external_links = []
    nombres = 0
    for nombre in [n for n, d in libro.defined_names.items() if _es_externo(d.attr_text)]:
        del libro.defined_names[nombre]
        nombres += 1
    for hoja in libro.worksheets:
        for nombre in [n for n, d in hoja.defined_names.items() if _es_externo(d.attr_text)]:
            del hoja.defined_names[nombre]
            nombres += 1

    props = libro.properties
    for campo in PROPIEDADES:
        setattr(props, campo, ANONIMO)
    for campo in PROPIEDADES_EXTENDIDAS:
        if hasattr(props, campo):
            setattr(props, campo, ANONIMO)
    libro.custom_doc_props = CustomPropertyList()
    return vinculos, nombres


def _textos_del_libro(libro: Workbook) -> Iterator[str]:
    """Todo texto de la SALIDA en el que podria quedar un dato: celdas, encabezados, nombres..."""
    props = libro.properties
    for campo in (*PROPIEDADES, "category", "identifier", "contentStatus"):
        valor = getattr(props, campo, None)
        if isinstance(valor, str):
            yield valor
    for prop in libro.custom_doc_props:
        yield str(getattr(prop, "value", ""))
    for nombre, definido in libro.defined_names.items():
        yield nombre
        yield definido.attr_text or ""
    for hoja in libro.worksheets:
        yield hoja.title
        for nombre, definido in hoja.defined_names.items():
            yield nombre
            yield definido.attr_text or ""
        for parte in _partes_encabezado(hoja):
            if parte.text:
                yield parte.text
        for validacion in hoja.data_validations.dataValidation:
            for campo in ("formula1", "formula2", *_TEXTOS_VALIDACION):
                texto = getattr(validacion, campo, None)
                if texto:
                    yield str(texto)
        for fila in hoja.iter_rows():
            for celda in fila:
                valor = celda.value
                if isinstance(valor, ArrayFormula):
                    yield valor.text or ""
                elif isinstance(valor, str):
                    yield valor
                if celda.comment is not None:
                    yield celda.comment.text or ""
                    yield celda.comment.author or ""
                vinculo = celda.hyperlink
                if vinculo is not None:
                    for campo in ("target", "location", "display", "tooltip"):
                        yield str(getattr(vinculo, campo, None) or "")


def verificar(salida: Path, datos: list[Dato]) -> dict[str, int]:
    """Relee la SALIDA y cuenta los datos descubiertos que sobreviven, por tipo."""
    libro = openpyxl.load_workbook(salida, keep_links=False)
    restos: dict[str, int] = defaultdict(int)
    if libro._external_links:
        restos["vinculo externo"] += len(libro._external_links)
    # Las imagenes se cuentan en el paquete y no solo en openpyxl, que sin Pillow no las carga.
    with zipfile.ZipFile(salida) as paquete:
        restos["imagen"] += sum(1 for n in paquete.namelist() if n.startswith("xl/media/"))
    for hoja in libro.worksheets:
        restos["imagen"] += len(hoja._images)
        restos["lista de validacion literal"] += sum(
            1 for v in hoja.data_validations.dataValidation if _es_lista_literal(v)
        )
        for fila in hoja.iter_rows():
            for celda in fila:
                if celda.hyperlink is not None:
                    restos["hipervinculo"] += 1
    for texto in _textos_del_libro(libro):
        for dato in datos:
            if dato.patron.search(texto):
                restos[dato.tipo] += 1
        if any(m.group(0) != RIF_NEUTRO for m in RIF.finditer(texto)):
            restos["rif"] += 1
    return {tipo: n for tipo, n in restos.items() if n}


def _imprimir_informe(
    datos: list[Dato],
    informe: dict[str, dict[str, int]],
    vinculos: int,
    nombres: int,
    quitados: tuple[int, int, int],
) -> None:
    por_tipo: dict[str, int] = defaultdict(int)
    for dato in datos:
        por_tipo[dato.tipo] += 1
    print("Datos descubiertos (variantes por tipo):")
    for tipo, n in sorted(por_tipo.items()):
        print(f"  {tipo}: {n}")
    print("Celdas cambiadas por tipo de dato:")
    if not informe:
        print("  (ninguna)")
    for tipo in sorted(informe):
        hojas = informe[tipo]
        total = sum(hojas.values())
        detalle = ", ".join(f"{h}={n}" for h, n in sorted(hojas.items()))
        print(f"  {tipo}: {total} [{detalle}]")
    print(f"Vinculos externos quitados: {vinculos}")
    print(f"Nombres definidos externos quitados: {nombres}")
    imagenes, hipervinculos, listas = quitados
    print(f"Imagenes (logos) quitadas: {imagenes}")
    print(f"Hipervinculos de celda quitados: {hipervinculos}")
    print(f"Listas de validacion literales quitadas: {listas}")
    print(f"Propiedades del documento: '{ANONIMO}'")


def main(argv: list[str] | None = None) -> int:
    analizador = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    analizador.add_argument("entrada", type=Path, help="libro real (.xlsx); no se modifica")
    analizador.add_argument("salida", type=Path, help="libro anonimizado (.xlsx)")
    analizador.add_argument(
        "--etiqueta", default="A", help="letra del caso: CLIENTE-<etiqueta>, CONTRATISTA-..."
    )
    args = analizador.parse_args(argv)

    entrada: Path = args.entrada
    salida: Path = args.salida
    if not entrada.is_file():
        print(f"No existe la ENTRADA: {entrada}", file=sys.stderr)
        return 2
    if salida.resolve() == entrada.resolve() or (salida.exists() and salida.samefile(entrada)):
        print(
            "La SALIDA no puede ser la ENTRADA: el libro real no se sobrescribe.", file=sys.stderr
        )
        return 2
    etiqueta = args.etiqueta.strip()
    if not etiqueta:
        print("La etiqueta no puede estar vacia.", file=sys.stderr)
        return 2

    libro = openpyxl.load_workbook(entrada, keep_links=False)
    datos = descubrir(libro, etiqueta)
    if not any(d.tipo in ("cliente", "contratista") for d in datos):
        print(
            "No se encontraron los rotulos CLIENTE/CONTRATISTA en la PORTADA: "
            "no se puede garantizar el anonimato y no se escribe nada.",
            file=sys.stderr,
        )
        return 1

    informe = anonimizar(libro, datos)
    quitados = quitar_imagenes_vinculos_y_validaciones(libro, datos)
    vinculos, nombres = limpiar_vinculos_y_propiedades(libro)
    salida.parent.mkdir(parents=True, exist_ok=True)
    libro.save(salida)
    _imprimir_informe(datos, informe, vinculos, nombres, quitados)

    restos = verificar(salida, datos)
    if restos:
        salida.unlink(missing_ok=True)
        detalle = ", ".join(f"{t}={n}" for t, n in sorted(restos.items()))
        print(f"VERIFICACION FALLIDA, SALIDA borrada. Restos por tipo: {detalle}", file=sys.stderr)
        return 1
    print(f"Verificacion: ningun dato descubierto aparece en la SALIDA ({salida.name}).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
