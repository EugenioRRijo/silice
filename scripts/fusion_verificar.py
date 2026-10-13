"""Aceptación de Silice 1.5 sobre los libros reales de la empresa (spec 1.5 §7).

Por qué existe: la 1.5 se construye sin pruebas nuevas (regla del 2026-10-12) y los libros reales
son confidenciales, así que no pueden ser casos de la suite. Este script es la herramienta de
aceptación que se corre A MANO, en el equipo que tiene los libros, contra la carpeta local donde
están (por ejemplo `docs/auditoria/`, ignorada por Git). No es una prueba de pytest y no corre en
la CI.

Qué comprueba, como metas OK/FALLA con su evidencia:

1. **Hallazgos esperados** (tabla del spec §7 y cifras de las causas raíz H1-H18): para cada libro,
   `fusion.proceso.procesar_libro` debe producir la regla esperada en la celda esperada, con la
   severidad y el impacto documentados.
2. **Valores de Silice**: en A, B-1 y B-2 (cuyo PRESUPUESTO cuadra con sus APU) el P.U. de Silice
   es el P guardado de cada partida; los totales de Silice son coherentes y, donde hay cifra
   documentada, iguales a ella (en C-2, el subtotal entregado más los 205,92 de H1).
3. **Devolución**: el libro devuelto (`devolver`) se escribe en `--salida` y se REABRE con openpyxl:
   sin `#REF!`, sin `INDIRECT`, sin vínculos ni nombres rotos, con un valor guardado en cada
   fórmula, con P y totales iguales a los de Silice, con las hojas AUDITORÍA, TRAZABILIDAD y
   ALERTAS USD-Bs, con el rótulo de P.U. en la misma celda en todos los APU y con la huella
   SHA-256 del libro recibido en sus propiedades.
4. **PDF simples** (001 y 002), por `fusion.proceso.auditar_pdf_simple`: la regla L10 con los
   subtipos esperados.
5. **PDF entregados** (L11): cada PDF de la carpeta se identifica con un literal documentado
   (`TOTALES_DOCUMENTADOS`) y la lectura del producto (`leer_totales_impresos`) se compara con ese
   literal (metas `PDF-*.LECTURA`); en cada libro, L11 debe dar ERROR en exactamente los campos en
   que el literal del PDF emparejado difiere del total guardado del libro (openpyxl), con impacto
   = PDF − libro, con su signo.
6. **Conjunto exacto de hallazgos** (meta `CONJUNTO-EXACTO`, control de falsos positivos): el
   multiconjunto (regla, severidad, celdas de origen) de cada libro, B-1 incluido, contra el
   congelado en `expectativas.json`. Sin ese archivo la meta queda PENDIENTE.
7. **Valores conservados** (meta `VALORES-CONSERVADOS`): las filas "Valor guardado conservado" de
   TRAZABILIDAD de la devolución deben ser 0 o estar en la lista blanca local.

Por qué los oráculos no salen de Silice: lo esperado viene de un literal documentado (las cifras
de `docs/auditoria/2026-10-13-causas-raiz-libros-arenaza.md` y del spec 1.5; los totales con IVA
de cada PDF, del encargo V7 de la revisión de calidad), de una memoria aritmética escrita aquí
(H1: 1,82565 × 1,10 × 1,02 por u) o de leer el libro RECIBIDO con openpyxl, aparte del lector del
adaptador (celdas con `#REF!`, nombres con vínculo externo, P guardado, números de partida,
fórmulas de MEDICIONES y de los APU). Comparar Silice contra Silice no demostraría nada: donde una
meta compara la devolución con `ResultadoFusion` (D-P, D-TOTALES), compara además con el oráculo.
La única lectura del producto que se usa para decidir qué esperar es la identificación del libro
por contenido (`FIRMAS`), y solo cuando `expectativas.json` no lo identifica por su SHA-256.

Cómo identifica cada libro sin escribir nombres de clientes en el código: si hay
`expectativas.json`, por la SHA-256 del `.xlsx` (`libros_sha256`) o por su nombre (`libros`); si
no, por su contenido: número de APU y de partidas y subtotal guardado (`FIRMAS`). C-1 y C-2 tienen
el mismo subtotal; C-1 es el que tiene dos APU que reclaman la misma fila (conflicto L2). Cada PDF
se identifica por su SHA-256 (`pdf_sha256`) o, si no figura, por el literal documentado cuyo total
con IVA aparece en su texto (pdfplumber directo, sin el lector del adaptador). El PDF de cada
libro es el que comparte nombre base con el `.xlsx`; para C-2 se cotejan además (L11) los otros
PDF identificados como "C". Nada de lo que imprime lleva nombres de archivo ni de clientes: los
libros se llaman "Libro A", "Libro B-1"… y toda evidencia pasa por un filtro que borra los nombres
reales.

`expectativas.json` (opcional, LOCAL, nunca al repositorio). Por defecto es
`CARPETA/expectativas.json`; `--expectativas RUTA` lo lee (y `--congelar` lo escribe) en otra
ruta. **Aviso:** `.gitignore` cubre las hojas de cálculo y los PDF de `docs/auditoria/**`, no los
`.json`: un `expectativas.json` dentro del repositorio quedaría versionable. Por eso `--congelar`
se niega a escribir dentro de un repositorio de Git en una ruta que Git no ignore, y una corrida
que lee un archivo así lo avisa. Lleva las celdas de los hallazgos de libros confidenciales:
guárdalo fuera del repositorio (o en una ruta ignorada). Formato::

    {"libros_sha256": {"<sha256 del xlsx>": "A", ...},
     "libros": {"<archivo>.xlsx": "A",
                "<otro>.xlsx": {"etiqueta": "C-2", "pdf": ["<impreso>.pdf"]}},
     "pdf_sha256": {"<sha256 del pdf>": "C", ...},
     "hallazgos": {"A": [["L1", "ERROR", "CRONOGRAMA!F10|APU-01!F27"], ...], ...},
     "valores_conservados": {"C-2": ["HOJA!CELDA", ...]},
     "valores_conservados_vistos": {...}}

Las claves de `pdf_sha256` son las de `TOTALES_DOCUMENTADOS` (A, B-1, B-2, C, QA, RC). En
`hallazgos`, el tercer elemento son las celdas de origen normalizadas `HOJA!CELDA` unidas por `|`;
un origen que no es celda (un PDF, un nombre definido) va como `~` + 10 cifras de su SHA-256, para
no escribir rutas ni nombres. `valores_conservados_vistos` lo escribe `--congelar` como ayuda para
la revisión: la lista blanca que cuenta es solo `valores_conservados`, que escribe una persona.

`--congelar` escribe (fusionando con lo que ya tenga el archivo) `libros_sha256`, `pdf_sha256`,
`hallazgos` y `valores_conservados_vistos` con lo que produce la corrida. La primera vez lo
revisa una persona hallazgo por hallazgo: congelar no demuestra nada, por eso en esa corrida
`CONJUNTO-EXACTO` queda PENDIENTE. Desde entonces la meta compara exacto (OK si idéntico, FALLA
con lo que sobra y lo que falta).

Uso::

    uv run python scripts/fusion_verificar.py CARPETA [--salida DIR] [--json]
        [--fecha-libro 2026-07-31] [--fecha-devolucion AAAA-MM-DD] [--pdf-simple RUTA ...]
        [--expectativas RUTA] [--congelar]

Código de salida: 0 si ninguna meta falla (las PENDIENTE no fallan), 1 si alguna falla, 2 si no
se pudo correr (falta `fusion.proceso`, la carpeta no existe o `--congelar` escribiría en una
ruta versionable).
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import sys
import traceback
import unicodedata
import zipfile
from collections import Counter
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import openpyxl
from defusedxml.ElementTree import fromstring as leer_xml

RAIZ = Path(__file__).resolve().parents[1]
if str(RAIZ) not in sys.path:
    sys.path.insert(0, str(RAIZ))

from scripts.meta_alpha import Estado, Fila, _filas_a_json, render_tabla  # noqa: E402

#: Fecha de los libros (PORTADA "Julio 2026" -> último día del mes, spec F-D8). Se pasa explícita
#: para que la corrida no dependa de cómo el lector interprete la portada.
FECHA_LIBRO = date(2026, 7, 31)
SALIDA_POR_DEFECTO = RAIZ / "data" / "confidencial" / "devoluciones"
PDF_SIMPLES = (
    RAIZ / "data" / "samples" / "telecom" / "Presupuesto_1_ARENAZA.pdf",
    RAIZ / "data" / "samples" / "telecom" / "Presupuesto_2_ARENAZA.pdf",
)

#: Tolerancias explícitas: medio céntimo por celda (la unidad impresa es el céntimo) y un céntimo
#: para totales que suman redondeos y para impactos documentados con dos decimales (el literal
#: está redondeado al céntimo: H1 es 205,9236 y se documenta 205,92). Los impactos se comparan
#: CON SU SIGNO: un impacto con el signo opuesto falla.
TOL_CELDA = Decimal("0.005")
TOL_CENTIMO = Decimal("0.01")
TOL_IMPACTO = Decimal("0.01")

#: IVA documentado (causas raíz, cifras de H1: "con IVA (x1,16)").
IVA_DOCUMENTADO = Decimal("0.16")

#: H1 (causas raíz): el CD de la partida del alucobond (PRESUPUESTO fila 24 en C-1 y C-2) omite
#: los equipos, 1,82565 USD/u. Memoria: Silice suma el equipo y conserva la administración del
#: 10 % y la utilidad del 2 % (H2) => P.U. Silice = P guardado + 1,82565 x 1,10 x 1,02
#: = P + 2,0483793 por u; con la cantidad 100,53, impacto 205,9236 (documentado 205,92).
H1_FILA = 24
H1_DELTA_PU = Decimal("1.82565") * Decimal("1.10") * Decimal("1.02")
IMPACTO_H1 = Decimal("205.92")
#: H1 deja fuera de `CD = G12+G21` el equipo: la fórmula literal del CD en el APU de H1.
FORMULA_CD_H1 = "=G12+G21"
HOJA_H1: dict[str, str] = {"C-1": "APU-09", "C-2": "APU-07"}
#: H2: (10 % - 2 %) x 94,7004 x 100,53 = 761,62 (global - local > 0: el libro cobra de menos).
IMPACTO_H2 = Decimal("761.62")
#: H3: RESUMEN, composición del costo (24 495,46) - subtotal (24 311,92) = +183,53; L9 informa
#: impacto = valor de la hoja - SUBSINIVA (`reglas._l9_resumen`), así que el signo es +.
IMPACTO_H3 = Decimal("183.53")


@dataclass(frozen=True, slots=True)
class Totales:
    """Subtotal sin IVA, IVA y total con IVA, tal como se documentan (dos decimales)."""

    subtotal: Decimal
    iva: Decimal
    total: Decimal

    def campos(self) -> tuple[Decimal, Decimal, Decimal]:
        return self.subtotal, self.iva, self.total


def _t(subtotal: str, iva: str, total: str) -> Totales:
    return Totales(Decimal(subtotal), Decimal(iva), Decimal(total))


#: Totales impresos documentados de cada PDF entregado (causas raíz, tabla de alcance y párrafo
#: "PDF entregados"; IVA y total del encargo V7). "C" es el de C-1, C-2 e "IR" (valores
#: idénticos); "QA" la versión de julio; "RC" el ÚLTIMO bloque ("Opción definitiva"), que es el que
#: debe tomar `leer_totales_impresos`.
TOTALES_DOCUMENTADOS: dict[str, Totales] = {
    "A": _t("9955.65", "1592.90", "11548.55"),
    "B-1": _t("12917.54", "2066.81", "14984.34"),
    "B-2": _t("16491.80", "2638.69", "19130.49"),
    "C": _t("24311.92", "3889.91", "28201.83"),
    "QA": _t("14542.09", "2326.73", "16868.82"),
    "RC": _t("15451.69", "2472.27", "17923.97"),
}
#: Totales documentados de lo que guarda cada libro (los de su PDF propio, sin cruzar).
TOTALES_DEL_LIBRO: dict[str, str] = {"A": "A", "B-1": "B-1", "B-2": "B-2", "C-1": "C", "C-2": "C"}

HOJAS_NUEVAS = ("AUDITORÍA", "TRAZABILIDAD", "ALERTAS USD-Bs")
ROTULO_PU = "PRECIO UNITARIO TOTAL"
PROPIEDAD_SHA = "silice_sha256_origen"
NOMBRES_TOTALES = ("SUBSINIVA", "IVATOT", "TOTALGEN")
NS_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
NS_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
NS_PKG_REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"

_RE_REF = re.compile(r"^\$?([A-Z]{1,3})\$?(\d+)")
_RE_DIV12 = re.compile(r"/\s*12(?!\d)")
_RE_TITULO_APU = re.compile(r"PRESUPUESTO'?!\$?[A-Z]{1,3}\$?(\d+)", re.IGNORECASE)
_RE_VINCULO_FORMULA = re.compile(r"\[\d+\]")
_RE_REF_PURA = re.compile(r"^=?\+?(?:'?([^'!]+)'?!)?\$?([A-Z]{1,3})\$?(\d+)$")


# --- Identificación de los libros -----------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class Firma:
    """Lo que distingue a un libro por su contenido (causas raíz, tabla de alcance)."""

    etiqueta: str
    partidas: int
    subtotal: Decimal
    apus: int | None = None
    conflicto_apu: bool | None = None


FIRMAS: tuple[Firma, ...] = (
    Firma("A", 5, Decimal("9955.65"), apus=5),
    Firma("B-1", 16, Decimal("12917.54")),
    Firma("B-2", 16, Decimal("16491.80")),
    Firma("C-1", 9, Decimal("24311.92"), conflicto_apu=True),
    Firma("C-2", 9, Decimal("24311.92"), conflicto_apu=False),
)
ETIQUETAS: tuple[str, ...] = tuple(f.etiqueta for f in FIRMAS)

#: Versión anterior del mismo libro, para las alertas "contra otra versión" (spec §4.3).
OTRA_VERSION: dict[str, str] = {"B-2": "B-1", "C-2": "C-1"}

#: Subtotal sin IVA documentado de lo entregado (causas raíz, tabla de alcance). En C-2, Silice
#: corrige H1 (F-D7) y conserva la utilidad del 2 % (H2): 24 311,92 + 205,92 = 24 517,85
#: (memoria: 1,82565 x 100,53 = 183,5326; x 1,10 x 1,02 = 205,9236; 24 311,9242 + 205,9236).
#: C-1 da lo mismo que C-2: sus hojas que no son APU son idénticas y Silice une cada partida con
#: su APU por el título, así que el APU-09 de C-1 (el de H1) es el de la fila 24, como el APU-07
#: de C-2.
SUBTOTAL_SILICE: dict[str, tuple[Decimal, Decimal]] = {
    "A": (Decimal("9955.65"), TOL_CELDA + Decimal("0.001")),
    "B-1": (Decimal("12917.54"), TOL_CELDA + Decimal("0.001")),
    "B-2": (Decimal("16491.80"), TOL_CELDA + Decimal("0.001")),
    "C-1": (Decimal("24517.85"), TOL_IMPACTO),
    "C-2": (Decimal("24517.85"), TOL_IMPACTO),
}


@dataclass(slots=True)
class Caso:
    """Un libro de la carpeta ya identificado, con lo que `expectativas.json` dice de él."""

    etiqueta: str
    ruta: Path
    datos: bytes
    como: str  # "expectativas.json" o "contenido"
    pdf: Path | None = None
    pdf_extra: list[Path] = field(default_factory=list)
    #: Clave de `TOTALES_DOCUMENTADOS` del PDF emparejado y de cada extra (None: sin identificar).
    pdf_id: str | None = None
    pdf_extra_ids: list[str | None] = field(default_factory=list)
    #: Multiconjunto congelado de (regla, severidad, celdas); None si no hay expectativa.
    hallazgos_esperados: Counter[tuple[str, str, str]] | None = None
    congelado_ahora: bool = False
    lista_blanca: frozenset[str] = frozenset()


def _norm(texto: object) -> str:
    """Mayúsculas, sin acentos, sin NBSP y con espacios simples: para comparar rótulos y hojas."""
    plano = unicodedata.normalize("NFKD", str(texto).replace("\xa0", " "))
    sin_acentos = "".join(c for c in plano if not unicodedata.combining(c))
    return " ".join(sin_acentos.upper().split())


def _dec(valor: object) -> Decimal | None:
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, Decimal):
        return valor
    if isinstance(valor, int | float):
        return Decimal(str(valor))
    try:
        return Decimal(str(valor).strip())
    except (InvalidOperation, ValueError):
        return None


def _fmt(valor: Decimal | None, decimales: int = 2, signo: bool = False) -> str:
    if valor is None:
        return "-"
    return f"{valor:+.{decimales}f}" if signo else f"{valor:.{decimales}f}"


def _cerca(a: Decimal | None, b: Decimal | None, tol: Decimal) -> bool:
    return a is not None and b is not None and abs(a - b) <= tol


def _firma_de(libro: Any) -> tuple[int, int, Decimal | None, bool]:
    subtotal = _dec(libro.totales.get("SUBSINIVA"))
    filas = Counter(a.fila_presupuesto for a in libro.apus if a.fila_presupuesto is not None)
    conflicto = any(n > 1 for n in filas.values())
    centimos = subtotal.quantize(Decimal("0.01")) if subtotal is not None else None
    return len(libro.apus), len(libro.partidas), centimos, conflicto


def _identificar(libro: Any) -> str | None:
    apus, partidas, subtotal, conflicto = _firma_de(libro)
    for f in FIRMAS:
        if f.partidas != partidas or subtotal != f.subtotal:
            continue
        if f.apus is not None and f.apus != apus:
            continue
        if f.conflicto_apu is not None and f.conflicto_apu != conflicto:
            continue
        return f.etiqueta
    return None


_CLAVES_EXPECTATIVAS = frozenset(
    {
        "libros",
        "libros_sha256",
        "pdf_sha256",
        "hallazgos",
        "valores_conservados",
        "valores_conservados_vistos",
    }
)


@dataclass(slots=True)
class Expectativas:
    """Lo que dice el `expectativas.json` local (vacío si no existe)."""

    ruta: Path
    existe: bool = False
    libros: dict[str, tuple[str, list[str]]] = field(default_factory=dict)  # nombre -> etiqueta
    libros_sha: dict[str, str] = field(default_factory=dict)
    pdf_sha: dict[str, str] = field(default_factory=dict)
    hallazgos: dict[str, Counter[tuple[str, str, str]]] = field(default_factory=dict)
    lista_blanca: dict[str, frozenset[str]] = field(default_factory=dict)
    crudo: dict[str, Any] = field(default_factory=dict)


def _etiqueta_valida(etiqueta: str, validas: Iterable[str], donde: str) -> str:
    if etiqueta not in validas:
        raise ValueError(f"expectativas.json ({donde}): etiqueta desconocida {etiqueta!r}")
    return etiqueta


def _leer_expectativas(ruta: Path) -> Expectativas:
    exp = Expectativas(ruta)
    if not ruta.is_file():
        return exp
    crudo = json.loads(ruta.read_text(encoding="utf-8"))
    if not isinstance(crudo, dict):
        raise ValueError("expectativas.json: se esperaba un objeto JSON")
    exp.existe, exp.crudo = True, crudo
    # Formato antiguo: el objeto entero es el mapa de libros por nombre.
    libros = crudo.get("libros", {}) if _CLAVES_EXPECTATIVAS & set(crudo) else crudo
    for archivo, valor in libros.items():
        if isinstance(valor, str):
            etiqueta, pdfs = valor, []
        else:
            etiqueta = str(valor.get("etiqueta", ""))
            pdf = valor.get("pdf", [])
            pdfs = [pdf] if isinstance(pdf, str) else list(pdf)
        exp.libros[archivo.lower()] = (_etiqueta_valida(etiqueta, ETIQUETAS, "libros"), pdfs)
    for sha, etiqueta in crudo.get("libros_sha256", {}).items():
        exp.libros_sha[sha.lower()] = _etiqueta_valida(etiqueta, ETIQUETAS, "libros_sha256")
    for sha, clave_pdf in crudo.get("pdf_sha256", {}).items():
        exp.pdf_sha[sha.lower()] = _etiqueta_valida(clave_pdf, TOTALES_DOCUMENTADOS, "pdf_sha256")
    for etiqueta, lista in crudo.get("hallazgos", {}).items():
        _etiqueta_valida(etiqueta, ETIQUETAS, "hallazgos")
        exp.hallazgos[etiqueta] = Counter(
            (str(r), str(s), str(c)) for r, s, c in (tuple(x) for x in lista)
        )
    for etiqueta, lista in crudo.get("valores_conservados", {}).items():
        _etiqueta_valida(etiqueta, ETIQUETAS, "valores_conservados")
        exp.lista_blanca[etiqueta] = frozenset(_norm_origen(str(o)) for o in lista)
    return exp


def _versionable(ruta: Path) -> bool:
    """True si `ruta` está dentro de un repositorio de Git y Git NO la ignora (se podría añadir
    al repositorio por error). Fuera de un repositorio, o ignorada, False."""
    import subprocess

    carpeta = ruta.resolve().parent
    if not carpeta.is_dir():
        return False
    try:
        r = subprocess.run(
            ["git", "-C", str(carpeta), "check-ignore", "-q", ruta.name],
            capture_output=True,
            check=False,
            timeout=30,
        )
    except (OSError, subprocess.SubprocessError):
        return False
    return r.returncode == 1  # 0: ignorada; 1: no ignorada; 128: no es un repositorio


# --- Identificación de los PDF por literal documentado (sin el lector del adaptador) -----------


def _patron_monto(valor: Decimal) -> re.Pattern[str]:
    """El monto tal como puede imprimirse: miles con punto, coma, espacio o nada; decimales con
    coma o punto (`24.311,92`, `24,311.92`, `24 311,92`, `24311.92`)."""
    entero, cent = f"{valor:.2f}".split(".")
    grupos = []
    while entero:
        grupos.insert(0, entero[-3:])
        entero = entero[:-3]
    miles = r"[.,\s  ]?".join(grupos)
    return re.compile(rf"(?<![\d.,]){miles}[.,]{cent}(?!\d)")


def _texto_pdf(datos: bytes) -> str:
    import pdfplumber

    with pdfplumber.open(io.BytesIO(datos)) as doc:
        return "\n".join((p.extract_text() or "") for p in doc.pages)


def _identificar_pdf(datos: bytes, exp: Expectativas) -> tuple[str | None, str]:
    """(clave de `TOTALES_DOCUMENTADOS`, cómo). Por SHA-256 si `expectativas.json` lo nombra; si
    no, el literal cuyo TOTAL con IVA aparece en el texto del PDF (debe ser uno solo)."""
    sha = hashlib.sha256(datos).hexdigest()
    if sha in exp.pdf_sha:
        return exp.pdf_sha[sha], "sha256"
    texto = _texto_pdf(datos)
    hallados = [k for k, t in TOTALES_DOCUMENTADOS.items() if _patron_monto(t.total).search(texto)]
    if len(hallados) == 1:
        return hallados[0], "literal"
    return None, ("ningun literal" if not hallados else f"ambiguo ({', '.join(hallados)})")


# --- Hallazgos: dónde y cuánto --------------------------------------------------------------------


def _celda_de(origen: str) -> tuple[str, str] | None:
    """`HOJA!CELDA` -> (HOJA normalizada, CELDA sin `$`). None si el origen no es una celda."""
    if origen.startswith("PDF:") or "!" not in origen:
        return None
    hoja, ref = origen.rsplit("!", 1)
    return _norm(hoja.strip().strip("'")), ref.replace("$", "").strip().upper()


def _fila_de(ref: str) -> int | None:
    m = _RE_REF.match(ref)
    return int(m.group(2)) if m else None


def _celdas(h: Any) -> set[tuple[str, str]]:
    return {c for o in h.origen_ids if (c := _celda_de(str(o))) is not None}


def _norm_origen(origen: str) -> str:
    """`HOJA!CELDA` normalizada (hoja sin acentos en mayúsculas, celda sin `$`); otro origen, como
    `~` + 10 cifras de su SHA-256 (así no se escriben rutas ni nombres de archivo)."""
    c = _celda_de(origen)
    if c is not None:
        return f"{c[0]}!{c[1]}"
    return "~" + hashlib.sha256(origen.encode("utf-8")).hexdigest()[:10]


def _clave_hallazgo(h: Any) -> tuple[str, str, str]:
    """(regla, severidad, celdas de origen en su orden, unidas por `|`)."""
    celdas = "|".join(_norm_origen(str(o)) for o in h.origen_ids) or "~sin-origen"
    return str(h.regla), str(h.severidad.name), celdas


def _de(hallazgos: Iterable[Any], regla: str, severidad: str | None = None) -> list[Any]:
    return [
        h
        for h in hallazgos
        if h.regla == regla and (severidad is None or h.severidad.name == severidad)
    ]


def _en_hoja(hallazgos: Iterable[Any], hoja: str) -> list[Any]:
    objetivo = _norm(hoja)
    return [h for h in hallazgos if any(c[0] == objetivo for c in _celdas(h))]


def _filas_en(hallazgos: Iterable[Any], hoja: str) -> set[int]:
    objetivo = _norm(hoja)
    return {
        f
        for h in hallazgos
        for c in _celdas(h)
        if c[0] == objetivo and (f := _fila_de(c[1])) is not None
    }


def _menciona(h: Any, texto: str) -> bool:
    patron = re.compile(rf"(?<![A-Za-z0-9_]){re.escape(texto)}(?![A-Za-z0-9_])", re.IGNORECASE)
    return bool(patron.search(h.descripcion)) or any(patron.search(str(o)) for o in h.origen_ids)


def _lista(valores: Iterable[object], maximo: int = 6) -> str:
    vistos = sorted({str(v) for v in valores}, key=lambda s: (len(s), s))
    resto = len(vistos) - maximo
    return ", ".join(vistos[:maximo]) + (f" (+{resto})" if resto > 0 else "")


# --- Oráculos sobre el libro RECIBIDO (openpyxl, aparte del lector) ------------------------------


def _abrir_wb(datos: bytes, data_only: bool) -> Any:
    return openpyxl.load_workbook(io.BytesIO(datos), data_only=data_only)


def _hoja(wb: Any, nombre: str) -> Any | None:
    objetivo = _norm(nombre)
    return next((ws for ws in wb.worksheets if _norm(ws.title) == objetivo), None)


def _formulas(wb: Any) -> Iterable[tuple[str, str, str]]:
    for ws in wb.worksheets:
        for fila in ws.iter_rows():
            for c in fila:
                v = c.value
                texto = getattr(v, "text", v)  # ArrayFormula
                if isinstance(texto, str) and texto.startswith("="):
                    yield ws.title, c.coordinate, texto


def _refs_rotas(wb_f: Any) -> list[tuple[str, str]]:
    return [(_norm(h), ref) for h, ref, f in _formulas(wb_f) if "#REF!" in f.upper()]


def _nombres_definidos(wb: Any) -> list[tuple[str, str]]:
    nombres = [(n, d.attr_text or "") for n, d in wb.defined_names.items()]
    for ws in wb.worksheets:
        nombres.extend((n, d.attr_text or "") for n, d in ws.defined_names.items())
    return nombres


def _nombres_externos(wb_f: Any) -> list[str]:
    return [n for n, texto in _nombres_definidos(wb_f) if "[" in texto]


def _numeros_duplicados(wb_v: Any) -> dict[int, list[int]]:
    """Número de partida (PRESUPUESTO col. B) -> filas, solo los repetidos. Una fila es partida si
    B es un entero y E (unidad) tiene texto: así no cuentan los capítulos."""
    ws = _hoja(wb_v, "PRESUPUESTO")
    if ws is None:
        return {}
    por_numero: dict[int, list[int]] = {}
    for fila in range(1, ws.max_row + 1):
        numero = _dec(ws[f"B{fila}"].value)
        unidad = ws[f"E{fila}"].value
        if numero is None or numero != numero.to_integral_value() or not str(unidad or "").strip():
            continue
        por_numero.setdefault(int(numero), []).append(fila)
    return {n: filas for n, filas in por_numero.items() if len(filas) > 1}


def _apus_en_conflicto(wb_f: Any) -> dict[int, list[str]]:
    """Fila de PRESUPUESTO -> hojas APU cuyo título (A2) la nombra, solo si son dos o más."""
    por_fila: dict[int, list[str]] = {}
    for ws in wb_f.worksheets:
        if not _norm(ws.title).startswith("APU"):
            continue
        m = _RE_TITULO_APU.search(str(ws["A2"].value or ""))
        if m:
            por_fila.setdefault(int(m.group(1)), []).append(ws.title)
    return {f: hojas for f, hojas in por_fila.items() if len(hojas) > 1}


def _celdas_apu(wb_f: Any, hoja: str) -> tuple[list[str], list[str]]:
    """En la hoja APU: celdas con el literal 0,02 (utilidad escrita a mano, H2) y fórmulas con
    `/12` (rendimiento literal, H12)."""
    ws = _hoja(wb_f, hoja)
    if ws is None:
        return [], []
    literales, divisiones = [], []
    for fila in ws.iter_rows():
        for c in fila:
            v = c.value
            if (
                isinstance(v, int | float)
                and not isinstance(v, bool)
                and _cerca(_dec(v), Decimal("0.02"), Decimal("0.0000001"))
            ):
                literales.append(c.coordinate)
            if isinstance(v, str) and v.startswith("=") and _RE_DIV12.search(v):
                divisiones.append(c.coordinate)
    return literales, divisiones


def _valor_nombre(wb_f: Any, wb_v: Any, nombre: str) -> Decimal | None:
    definido = wb_f.defined_names.get(nombre)
    if definido is None:
        return None
    for hoja, coord in definido.destinations:
        ws = _hoja(wb_v, hoja)
        if ws is not None:
            return _dec(ws[coord.replace("$", "")].value)
    return None


# --- Lo que se evalúa de un libro -----------------------------------------------------------------


@dataclass(slots=True)
class Evaluacion:
    caso: Caso
    resultado: Any = None
    hallazgos: list[Any] = field(default_factory=list)
    error: str = ""
    orig_f: Any = None
    orig_v: Any = None
    devolucion: bytes | None = None
    ruta_devolucion: Path | None = None
    error_devolucion: str = ""
    dev_f: Any = None
    dev_v: Any = None
    error_reabrir: str = ""
    #: (nombre neutro, clave documentada del PDF o None, hallazgos de `cotejar_con_libro`).
    pdf_extra: list[tuple[str, str | None, list[Any]]] = field(default_factory=list)

    @property
    def etiqueta(self) -> str:
        return self.caso.etiqueta


Comprobacion = Callable[[Evaluacion], tuple[bool | Estado, str]]


def _precios_silice(r: Any) -> dict[int, Decimal] | None:
    """P.U. de Silice por fila de PRESUPUESTO, donde sea que `ResultadoFusion` lo guarde."""
    for fuente in (r, getattr(r, "datos", None), getattr(r, "calculos", None)):
        pu = getattr(fuente, "precios_unitarios", None)
        if isinstance(pu, Mapping) and pu:
            return {int(k): v for k, v in pu.items() if _dec(v) is not None}
    calculos = getattr(r, "calculos", None)
    if isinstance(calculos, Mapping) and calculos:
        por_hoja = {str(k): _dec(getattr(v, "precio_unitario", None)) for k, v in calculos.items()}
        salida = {
            p.fila: pu
            for p in r.libro.partidas
            if p.hoja_apu and (pu := por_hoja.get(p.hoja_apu)) is not None
        }
        return salida or None
    return None


def _totales_silice(r: Any) -> dict[str, Decimal]:
    totales = getattr(r, "totales_silice", None) or {}
    return {str(k): d for k, v in totales.items() if (d := _dec(v)) is not None}


# Comprobaciones de reglas ---------------------------------------------------------------------


def _l1_en(hoja: str, minimo_filas: int) -> Comprobacion:
    def comprobar(ev: Evaluacion) -> tuple[bool, str]:
        en = _en_hoja(_de(ev.hallazgos, "L1", "ERROR"), hoja)
        filas = _filas_en(en, hoja)
        ok = len(en) > 0 and len(filas) >= minimo_filas
        return ok, f"{len(en)} L1 ERROR en {hoja}; filas {_lista(filas) or 'ninguna'}"

    return comprobar


def _l1_rend_b2(ev: Evaluacion) -> tuple[bool, str]:
    esperadas = {("PRESUPUESTO", "F26"), ("PRESUPUESTO", "F27")}
    celdas = {c for h in _de(ev.hallazgos, "L1", "ERROR") for c in _celdas(h)}
    faltan = esperadas - celdas
    vistas = {f"{c[1]}" for c in celdas if c[0] == "PRESUPUESTO"}
    return not faltan, (
        f"L1 en PRESUPUESTO: {_lista(vistas) or 'ninguna'}; faltan {_lista(c[1] for c in faltan)}"
        if faltan
        else "L1 en PRESUPUESTO!F26 y F27 (REND de las partidas 15 y 16)"
    )


def _l5_refs(hoja: str | None, documentadas: int | None) -> Comprobacion:
    """Cada `#REF!` del libro recibido (oráculo openpyxl) tiene un L5 con esa celda."""

    def comprobar(ev: Evaluacion) -> tuple[bool, str]:
        oraculo = [c for c in _refs_rotas(ev.orig_f) if hoja is None or c[0] == _norm(hoja)]
        cubiertas = {c for h in _de(ev.hallazgos, "L5") for c in _celdas(h)}
        faltan = [f"{c[0]}!{c[1]}" for c in oraculo if c not in cubiertas]
        donde = hoja or "el libro"
        texto = f"{len(oraculo) - len(faltan)}/{len(oraculo)} #REF! de {donde} con L5"
        ok = not faltan and len(oraculo) > 0
        if documentadas is not None and len(oraculo) != documentadas:
            ok = False
            texto += f"; el libro tiene {len(oraculo)} y las causas raiz dicen {documentadas}"
        if faltan:
            texto += f"; sin L5: {_lista(faltan)}"
        return ok, texto

    return comprobar


def _l5_nombres(documentados: int | None) -> Comprobacion:
    def comprobar(ev: Evaluacion) -> tuple[bool, str]:
        oraculo = _nombres_externos(ev.orig_f)
        l5 = _de(ev.hallazgos, "L5")
        faltan = [n for n in oraculo if not any(_menciona(h, n) for h in l5)]
        ok = not faltan
        texto = f"{len(oraculo) - len(faltan)}/{len(oraculo)} nombres con vinculo externo con L5"
        if documentados is not None and len(oraculo) != documentados:
            ok = False
            texto += f"; el libro tiene {len(oraculo)} y las causas raiz dicen {documentados}"
        if faltan:
            texto += f"; sin L5: {_lista(faltan)}"
        return ok, texto

    return comprobar


def _l5_todo(ev: Evaluacion) -> tuple[bool, str]:
    ok_refs, t_refs = _l5_refs(None, None)(ev)
    ok_nom, t_nom = _l5_nombres(None)(ev)
    if not _refs_rotas(ev.orig_f):  # un libro sin #REF! no exige L5 por celdas
        ok_refs, t_refs = True, "sin #REF! en el libro"
    return ok_refs and ok_nom, f"{t_refs}; {t_nom}"


#: H12 (causas raíz): medidas escritas dentro de las fórmulas de MEDICIONES. Las tres formas
#: documentadas, cada una en su sección: `=9.73*2.75`, `=2.75*3` y un factor `*0.28`.
_FORMAS_H12: tuple[tuple[str, Callable[[str], bool]], ...] = (
    ("=9.73*2.75", lambda f: f == "=9.73*2.75"),
    ("=2.75*3", lambda f: f == "=2.75*3"),
    ("*0.28", lambda f: re.search(r"\*\s*0?\.28(?!\d)", f) is not None),
)


def _l7_mediciones(ev: Evaluacion) -> tuple[bool, str]:
    """Oráculo: las fórmulas de MEDICIONES leídas con openpyxl. Cada forma documentada de H12
    debe existir en el libro y TODAS sus celdas deben estar en el origen de algún L7."""
    ws = _hoja(ev.orig_f, "MEDICIONES")
    if ws is None:
        return False, "el libro no tiene MEDICIONES"
    l7 = _en_hoja(_de(ev.hallazgos, "L7"), "MEDICIONES")
    partes, ok = [], True
    cubren: set[int] = set()
    for nombre, es in _FORMAS_H12:
        celdas = [
            c.coordinate
            for fila in ws.iter_rows()
            for c in fila
            if isinstance(c.value, str) and es(c.value.replace(" ", ""))
        ]
        por_celda = {
            c: [k for k, h in enumerate(l7) if ("MEDICIONES", c) in _celdas(h)] for c in celdas
        }
        sin = [c for c, hs in por_celda.items() if not hs]
        cubren.update(k for hs in por_celda.values() for k in hs)
        ok = ok and bool(celdas) and not sin
        partes.append(
            f"{nombre} en {_lista(celdas, 3) or 'ninguna celda'}"
            + (f" (sin L7: {_lista(sin, 3)})" if sin else "")
        )
    sev = Counter(h.severidad.name for h in l7)
    return ok, (
        f"oraculo: {'; '.join(partes)}; {len(l7)} L7 en MEDICIONES ({dict(sev) or 'ninguno'}), "
        f"{len(cubren)} cubren las formas documentadas"
    )


def _l9_procura_b2(ev: Evaluacion) -> tuple[bool, str]:
    esperados = (
        Decimal("-6969.49"),
        Decimal("-4861.47"),
        Decimal("-246.69"),
        Decimal("191309.86"),
    )
    impactos = [_dec(h.impacto) for h in _en_hoja(_de(ev.hallazgos, "L9"), "PROCURA")]
    presentes = [i for i in impactos if i is not None]
    faltan = [e for e in esperados if not any(_cerca(i, e, TOL_CENTIMO) for i in presentes)]
    invertidos = [e for e in faltan if any(_cerca(i, -e, TOL_CENTIMO) for i in presentes)]
    texto = f"impactos L9 en PROCURA: {_lista(_fmt(i) for i in presentes) or 'ninguno'}"
    if faltan:
        texto += f"; faltan {_lista(_fmt(e) for e in faltan)}"
    if invertidos:
        texto += f" (con signo opuesto: {_lista(_fmt(e) for e in invertidos)})"
    return not faltan, texto


def _signo(impactos: Iterable[Decimal | None], esperado: Decimal, tol: Decimal) -> str:
    """Nota para la evidencia cuando el impacto solo aparece con el signo opuesto."""
    presentes = [i for i in impactos if i is not None]
    if any(_cerca(i, esperado, tol) for i in presentes):
        return ""
    return " (con signo opuesto)" if any(_cerca(i, -esperado, tol) for i in presentes) else ""


def _f1_h1(ev: Evaluacion) -> tuple[bool, str]:
    """F1 ERROR en la fila de H1 con impacto +205,92 ± 0,01: P.U. Silice - P guardado > 0 (el
    libro cobra de menos; convención de `fusion.proceso._f1`)."""
    f1 = _de(ev.hallazgos, "F1")
    en_fila = [h for h in f1 if H1_FILA in _filas_en([h], "PRESUPUESTO")]
    impactos = [_dec(h.impacto) for h in en_fila]
    ok = any(
        h.severidad.name == "ERROR" and _cerca(i, IMPACTO_H1, TOL_IMPACTO)
        for h, i in zip(en_fila, impactos, strict=True)
    )
    otros = sorted(_filas_en(f1, "PRESUPUESTO"))
    return ok, (
        f"F1 en PRESUPUESTO fila {H1_FILA}: impacto "
        f"{_lista(_fmt(i, 4) for i in impactos) or 'ninguno'} (esperado +{_fmt(IMPACTO_H1)} "
        f"+-{TOL_IMPACTO}){_signo(impactos, IMPACTO_H1, TOL_IMPACTO)}; F1 en filas "
        f"{_lista(otros) or 'ninguna'}"
    )


def _l4_h1(ev: Evaluacion) -> tuple[bool, str]:
    """L4 ERROR en el APU de H1, en la celda del CD que el oráculo (openpyxl) encuentra con la
    fórmula documentada `=G12+G21`, con impacto +205,92 ± 0,01 (la partida queda corta)."""
    hoja = HOJA_H1[ev.etiqueta]
    ws = _hoja(ev.orig_f, hoja)
    oraculo = [
        c.coordinate
        for fila in (ws.iter_rows() if ws is not None else ())
        for c in fila
        if isinstance(c.value, str) and c.value.replace(" ", "").upper() == FORMULA_CD_H1
    ]
    en = _en_hoja(_de(ev.hallazgos, "L4", "ERROR"), hoja)
    celdas = {c[1] for h in en for c in _celdas(h) if c[0] == _norm(hoja)}
    impactos = [_dec(h.impacto) for h in en]
    con_impacto = any(_cerca(i, IMPACTO_H1, TOL_IMPACTO) for i in impactos)
    en_cd = len(oraculo) == 1 and oraculo[0] in celdas
    return con_impacto and en_cd, (
        f"oraculo CD {FORMULA_CD_H1} en {hoja}!{_lista(oraculo) or '-'}; {len(en)} L4 ERROR en "
        f"{hoja} ({_lista(celdas) or 'sin celda'}); impacto "
        f"{_lista(_fmt(i, 4) for i in impactos) or '-'} (esperado +{_fmt(IMPACTO_H1)} "
        f"+-{TOL_IMPACTO}){_signo(impactos, IMPACTO_H1, TOL_IMPACTO)}"
    )


def _l6_h2(ev: Evaluacion) -> tuple[bool, str]:
    """L6 en el APU de H1/H2: la utilidad 0,02 escrita a mano (oráculo openpyxl), con impacto
    (global - local) x base x cantidad = +761,62 ± 0,01 (H2), y el `/12` literal (H12)."""
    hoja = HOJA_H1[ev.etiqueta]
    literales, divisiones = _celdas_apu(ev.orig_f, hoja)
    l6 = _en_hoja(_de(ev.hallazgos, "L6", "ADVERTENCIA"), hoja)
    celdas = {c[1] for h in l6 for c in _celdas(h) if c[0] == _norm(hoja)}
    utilidad = bool(literales) and bool(celdas & set(literales))
    div12 = bool(divisiones) and bool(celdas & set(divisiones))
    en_util = [h for h in l6 if {c[1] for c in _celdas(h)} & set(literales)]
    impactos = [_dec(h.impacto) for h in en_util]
    h2 = any(_cerca(i, IMPACTO_H2, TOL_IMPACTO) for i in impactos)
    return utilidad and div12 and h2, (
        f"L6 en {hoja}: {_lista(celdas) or 'ninguna'}; oraculo utilidad 0,02 en "
        f"{_lista(literales) or '-'} ({'OK' if utilidad else 'falta'}, impacto "
        f"{_lista(_fmt(i) for i in impactos) or '-'}, H2 +{_fmt(IMPACTO_H2)} +-{TOL_IMPACTO}"
        f"{_signo(impactos, IMPACTO_H2, TOL_IMPACTO)}), '/12' en {_lista(divisiones) or '-'} "
        f"({'OK' if div12 else 'falta'})"
    )


#: H18 (causas raíz, V5): en B-2 la utilidad del APU-09 está escrita a mano en 30 %
#: (`G59 = G58*0.3`) y no lee `=UTIL` (10 %). L6 la marca "a confirmar" con impacto informativo
#: (global - local) x base x cantidad = -746,98 USD sin IVA (los 746,98 de más frente al 10 %).
_RE_UTIL_30 = re.compile(r"^=\s*\+?\$?G\$?58\s*\*\s*0?\.30*\s*$", re.IGNORECASE)
IMPACTO_H18 = Decimal("-746.98")


def _l6_util30_b2(ev: Evaluacion) -> tuple[bool, str]:
    ws = _hoja(ev.orig_f, "APU-09")
    formula = str(ws["G59"].value or "") if ws is not None else ""
    oraculo = bool(_RE_UTIL_30.match(formula.replace(" ", "")))
    l6 = [
        h for h in _de(ev.hallazgos, "L6", "ADVERTENCIA") if (_norm("APU-09"), "G59") in _celdas(h)
    ]
    impactos = [_dec(h.impacto) for h in l6]
    con_impacto = any(_cerca(i, IMPACTO_H18, TOL_IMPACTO) for i in impactos)
    invertido = any(_cerca(i, -IMPACTO_H18, TOL_IMPACTO) for i in impactos)
    texto = (
        f"oraculo APU-09!G59 {'= G58*0.3' if oraculo else 'sin G58*0.3'}; "
        f"{len(l6)} L6 ADVERTENCIA en G59; impacto {_lista(_fmt(i) for i in impactos) or '-'} "
        f"(H18: {_fmt(IMPACTO_H18)} +-{TOL_IMPACTO}, informativo)"
    )
    if invertido and not con_impacto:
        texto += " (con signo opuesto)"
    return oraculo and bool(l6) and con_impacto, texto


def _l8_c2(ev: Evaluacion) -> tuple[bool, str]:
    duplicados = _numeros_duplicados(ev.orig_v)
    filas_l8 = _filas_en(_de(ev.hallazgos, "L8", "ADVERTENCIA"), "PRESUPUESTO")
    faltan = [n for n, filas in duplicados.items() if not filas_l8 & set(filas)]
    ok = set(duplicados) == {6, 7} and not faltan
    return ok, (
        f"oraculo: numeros repetidos {dict(sorted(duplicados.items())) or 'ninguno'} "
        f"(documentados 6 y 7); L8 en filas {_lista(filas_l8) or 'ninguna'}"
        + (f"; sin L8: {_lista(faltan)}" if faltan else "")
    )


def _l3_filas(filas: tuple[int, ...]) -> Comprobacion:
    def comprobar(ev: Evaluacion) -> tuple[bool, str]:
        vistas = _filas_en(_de(ev.hallazgos, "L3"), "PRESUPUESTO")
        faltan = [f for f in filas if f not in vistas]
        return not faltan, (
            f"L3 en PRESUPUESTO filas {_lista(vistas) or 'ninguna'}"
            + (f"; faltan {_lista(faltan)}" if faltan else "")
        )

    return comprobar


def _c2_l3_o_parche(ev: Evaluacion) -> tuple[bool, str]:
    """Spec §7 (errata RL-7): L2 ADVERTENCIA en las filas 19 y 22 de C-2 (el original pedía L3).
    RL-7: ningún libro tiene valores guardados desactualizados (lo confirma un oráculo aparte
    sobre las referencias directas de esas filas); el parche manual por fila es L2 ADVERTENCIA."""
    ok_l3, t_l3 = _l3_filas((19, 22))(ev)
    filas_l2 = _filas_en(_de(ev.hallazgos, "L2", "ADVERTENCIA"), "PRESUPUESTO")
    ok_l2 = {19, 22} <= filas_l2
    return ok_l3 or ok_l2, f"{t_l3}; L2 ADVERTENCIA en filas {_lista(filas_l2) or 'ninguna'}"


def _l9_resumen_c2(ev: Evaluacion) -> tuple[bool, str]:
    impactos = [_dec(h.impacto) for h in _en_hoja(_de(ev.hallazgos, "L9", "ERROR"), "RESUMEN")]
    ok = any(_cerca(i, IMPACTO_H3, TOL_IMPACTO) for i in impactos)
    return ok, (
        f"impactos L9 ERROR en RESUMEN: {_lista(_fmt(i) for i in impactos) or 'ninguno'} "
        f"(esperado H3 +{_fmt(IMPACTO_H3)} +-{TOL_IMPACTO}, composicion - subtotal)"
        f"{_signo(impactos, IMPACTO_H3, TOL_IMPACTO)}"
    )


def _c1_l3_o_l2(ev: Evaluacion) -> tuple[bool, str]:
    ok_l3, t_l3 = _l3_filas((23, 24))(ev)
    conflicto = _apus_en_conflicto(ev.orig_f)
    l2 = _de(ev.hallazgos, "L2")
    hojas_conflicto = {_norm(h) for hojas in conflicto.values() for h in hojas}
    ok_l2 = bool(conflicto) and any(
        (_filas_en([h], "PRESUPUESTO") & set(conflicto))
        or ({c[0] for c in _celdas(h)} & hojas_conflicto)
        for h in l2
    )
    # Spec §7 (con la errata RL-7): el conflicto de APU (L2) Y los enlaces parcheados de las
    # filas 23 y 24 (L2 ADVERTENCIA); L3 en esas filas vale en lugar del parche.
    filas_adv = _filas_en(_de(ev.hallazgos, "L2", "ADVERTENCIA"), "PRESUPUESTO")
    ok_parche = ok_l3 or {23, 24} <= filas_adv
    return ok_l2 and ok_parche, (
        f"{t_l3}; L2 ADVERTENCIA en filas {_lista(filas_adv) or 'ninguna'}; "
        f"oraculo conflicto de APU: "
        f"{ {f: len(h) for f, h in conflicto.items()} or 'ninguno' }; {len(l2)} L2"
        f" ({'cubre el conflicto' if ok_l2 else 'no cubre el conflicto'})"
    )


def _celda_nombre(wb_f: Any, nombre: str) -> str | None:
    """`HOJA!CELDA` normalizada a la que apunta un nombre definido (openpyxl)."""
    definido = wb_f.defined_names.get(nombre)
    if definido is None:
        return None
    for hoja, coord in definido.destinations:
        return f"{_norm(hoja)}!{coord.replace('$', '').upper()}"
    return None


def _l11_esperado(
    clave_pdf: str | None, ev: Evaluacion, hallazgos: Sequence[Any]
) -> tuple[bool, str]:
    """L11 contra un PDF cuyo literal documentado es `clave_pdf`: ERROR en exactamente los campos
    (subtotal, IVA, total) en que el literal difiere del total guardado del libro (openpyxl, por
    su nombre definido), cada uno en la celda de ese nombre y con impacto = PDF - libro (con
    signo) ± 0,01; ninguna ADVERTENCIA."""
    if clave_pdf is None:
        return False, "PDF sin identificar por literal ni por SHA-256: no hay oraculo"
    impreso = TOTALES_DOCUMENTADOS[clave_pdf]
    esperados: dict[str, Decimal] = {}
    for nombre, valor_pdf in zip(NOMBRES_TOTALES, impreso.campos(), strict=True):
        guardado = _valor_nombre(ev.orig_f, ev.orig_v, nombre)
        celda = _celda_nombre(ev.orig_f, nombre)
        if guardado is None or celda is None:
            return False, f"el libro no guarda {nombre} (oraculo openpyxl)"
        if abs(valor_pdf - guardado) > TOL_CELDA:
            esperados[celda] = valor_pdf - guardado
    errores = _de(hallazgos, "L11", "ERROR")
    avisos = _de(hallazgos, "L11", "ADVERTENCIA")
    por_celda: dict[str, list[Decimal | None]] = {}
    for h in errores:
        for c in _celdas(h):
            por_celda.setdefault(f"{c[0]}!{c[1]}", []).append(_dec(h.impacto))
    malos = [
        f"{c} {_fmt(d, signo=True)}"
        for c, d in esperados.items()
        if not any(_cerca(i, d, TOL_IMPACTO) for i in por_celda.get(c, []))
    ]
    ok = len(errores) == len(esperados) and not malos and not avisos
    texto = (
        f"PDF '{clave_pdf}' (sin IVA {_fmt(impreso.subtotal)}): esperados {len(esperados)} L11 "
        f"ERROR, hay {len(errores)} ERROR y {len(avisos)} ADVERTENCIA"
    )
    if malos:
        texto += f"; sin el impacto PDF-libro esperado: {_lista(malos, 3)}"
    return ok, texto


def _l11(ev: Evaluacion) -> tuple[bool, str]:
    """L11 debe dispararse si y solo si los totales impresos del PDF emparejado difieren de los
    guardados en el libro. Oráculo independiente: el literal documentado del PDF (identificado por
    SHA-256 o por el literal que contiene su texto, nunca por `leer_totales_impresos`) contra los
    totales del libro leídos con openpyxl. El PDF "del mismo nombre" puede ser de otra versión: en
    los libros B, los nombres de archivo de los PDF están cruzados respecto de los `.xlsx`, y lo
    correcto entonces es que L11 lo diga. Además, lo que Silice leyó del PDF (`pdf_totales`) debe
    ser el literal."""
    if ev.caso.pdf is None:
        return False, "no se encontro el PDF impreso del libro en la carpeta"
    ok, texto = _l11_esperado(ev.caso.pdf_id, ev, ev.hallazgos)
    leido = getattr(ev.resultado, "pdf_totales", None)
    if ev.caso.pdf_id is not None:
        doc = TOTALES_DOCUMENTADOS[ev.caso.pdf_id].campos()
        campos = (getattr(leido, n, None) for n in ("subtotal", "iva", "total"))
        lectura = [_dec(x) for x in campos]
        bien = all(_cerca(a, b, TOL_CELDA) for a, b in zip(lectura, doc, strict=True))
        ok = ok and bien
        if not bien:
            texto += f"; Silice leyo {'/'.join(_fmt(x) for x in lectura)} del PDF"
    for nombre, clave_pdf, hs in ev.pdf_extra:
        ok_e, t_e = _l11_esperado(clave_pdf, ev, hs)
        ok = ok and ok_e
        texto += f"; extra {nombre.rsplit('_', 1)[-1]}: {t_e}"
    return ok, texto


# Comprobaciones de valores de Silice -----------------------------------------------------------


def _pu_oraculo(ev: Evaluacion) -> dict[int, Decimal | None]:
    """P.U. esperado por fila de PRESUPUESTO, sin Silice: el P guardado del libro RECIBIDO
    (openpyxl) y, en la fila de H1 de C-1 y C-2, ese P más la memoria de H1 (`H1_DELTA_PU`).
    Las filas son las partidas del libro (B entero y E con texto, como en `_numeros_duplicados`)."""
    ws = _hoja(ev.orig_v, "PRESUPUESTO")
    if ws is None:
        return {}
    salida: dict[int, Decimal | None] = {}
    for fila in range(1, ws.max_row + 1):
        numero = _dec(ws[f"B{fila}"].value)
        if numero is None or numero != numero.to_integral_value():
            continue
        if not str(ws[f"E{fila}"].value or "").strip():
            continue
        p = _dec(ws[f"P{fila}"].value)
        if ev.etiqueta in HOJA_H1 and fila == H1_FILA and p is not None:
            p += H1_DELTA_PU
        salida[fila] = p
    return salida


def _pu_igual_guardado(partidas_documentadas: int) -> Comprobacion:
    """P.U. de Silice = P guardado (openpyxl) en cada partida; en C-1 y C-2, en las 8 que H1 no
    toca, y en la de H1, P + 1,82565 x 1,10 x 1,02 (memoria de H1)."""

    def comprobar(ev: Evaluacion) -> tuple[bool, str]:
        pus = _precios_silice(ev.resultado)
        if pus is None:
            return False, "ResultadoFusion no expone el P.U. de Silice por fila"
        oraculo = _pu_oraculo(ev)
        malas = [
            f"fila {fila}: {_fmt(pus.get(fila), 4)} vs {_fmt(esperado, 4)}"
            for fila, esperado in oraculo.items()
            if not _cerca(pus.get(fila), esperado, TOL_CELDA)
        ]
        sobran = sorted(set(pus) - set(oraculo))
        ok = not malas and not sobran and len(oraculo) == partidas_documentadas
        h1 = ev.etiqueta in HOJA_H1
        texto = (
            f"{len(oraculo) - len(malas)}/{len(oraculo)} partidas con |PU Silice - oraculo| <= "
            f"0,005" + (f" (fila {H1_FILA}: P + {_fmt(H1_DELTA_PU, 7)} de H1)" if h1 else "")
        )
        if len(oraculo) != partidas_documentadas:
            texto += f" (documentadas {partidas_documentadas})"
        if malas:
            texto += f"; {_lista(malas, 4)}"
        if sobran:
            texto += f"; Silice da P.U. en filas que no son partidas: {_lista(sobran)}"
        return ok, texto

    return comprobar


def _totales_coherentes(ev: Evaluacion) -> tuple[bool, str]:
    tot = _totales_silice(ev.resultado)
    sub, iva, gen = (tot.get(n) for n in NOMBRES_TOTALES)
    if sub is None or iva is None or gen is None:
        return False, f"totales_silice incompletos: {sorted(tot)}"
    problemas = []
    if not _cerca(gen, sub + iva, TOL_CENTIMO):
        problemas.append(f"TOTALGEN {_fmt(gen)} != SUB + IVA {_fmt(sub + iva)}")
    sub_o = _valor_nombre(ev.orig_f, ev.orig_v, "SUBSINIVA")
    iva_o = _valor_nombre(ev.orig_f, ev.orig_v, "IVATOT")
    if sub_o and iva_o is not None:
        razon = iva_o / sub_o
        if not _cerca(iva, sub * razon, TOL_CENTIMO):
            problemas.append(f"IVATOT {_fmt(iva)} != SUB x {_fmt(razon, 4)}")
    pus = _precios_silice(ev.resultado) or {}
    suma = sum(
        (p.cantidad * pus[p.fila] for p in ev.resultado.libro.partidas if p.fila in pus),
        Decimal(0),
    )
    if pus and not _cerca(sub, suma, TOL_CENTIMO):
        problemas.append(f"SUB {_fmt(sub)} != sum(cant x PU) {_fmt(suma)}")
    esperado = SUBTOTAL_SILICE.get(ev.etiqueta)
    if esperado is not None and not _cerca(sub, esperado[0], esperado[1]):
        problemas.append(f"SUB {_fmt(sub)} != documentado {_fmt(esperado[0])}")
    texto = f"SUB {_fmt(sub)} IVA {_fmt(iva)} TOTAL {_fmt(gen)}"
    if esperado is not None:
        texto += f" (documentado {_fmt(esperado[0])})"
    return not problemas, texto + (f"; {'; '.join(problemas)}" if problemas else "")


# Comprobaciones de la devolución ---------------------------------------------------------------


def _requiere_devolucion(comprobar: Comprobacion) -> Comprobacion:
    def envuelta(ev: Evaluacion) -> tuple[bool, str]:
        if ev.devolucion is None:
            return False, f"sin devolucion: {ev.error_devolucion or 'no se genero'}"
        if ev.dev_f is None:
            return False, f"no reabre: {ev.error_reabrir}"
        return comprobar(ev)

    return envuelta


def _d_reabre(ev: Evaluacion) -> tuple[bool, str]:
    if ev.devolucion is None:
        return False, f"sin devolucion: {ev.error_devolucion or 'no se genero'}"
    if ev.dev_f is None:
        return False, f"openpyxl no la reabre: {ev.error_reabrir}"
    return True, f"reabre ({len(ev.dev_f.sheetnames)} hojas, {len(ev.devolucion)} bytes)"


def _d_sin_ref(ev: Evaluacion) -> tuple[bool, str]:
    rotas = [f"{h}!{c}" for h, c, f in _formulas(ev.dev_f) if "#REF!" in f.upper()]
    return not rotas, f"{len(rotas)} formulas con #REF!" + (f": {_lista(rotas)}" if rotas else "")


def _d_sin_indirect(ev: Evaluacion) -> tuple[bool, str]:
    con = [f"{h}!{c}" for h, c, f in _formulas(ev.dev_f) if "INDIRECT(" in f.upper()]
    nombres = [n for n, t in _nombres_definidos(ev.dev_f) if "INDIRECT(" in t.upper()]
    malas = con + nombres
    return not malas, f"{len(con)} formulas y {len(nombres)} nombres con INDIRECT" + (
        f": {_lista(malas)}" if malas else ""
    )


def _d_sin_vinculos(ev: Evaluacion) -> tuple[bool, str]:
    nombres = _nombres_definidos(ev.dev_f)
    externos = [n for n, t in nombres if "[" in t]
    rotos = [n for n, t in nombres if "#REF!" in t.upper()]
    with zipfile.ZipFile(io.BytesIO(ev.devolucion or b"")) as z:
        partes = [n for n in z.namelist() if n.startswith("xl/externalLinks/")]
    en_formulas = [f"{h}!{c}" for h, c, f in _formulas(ev.dev_f) if _RE_VINCULO_FORMULA.search(f)]
    ok = not (externos or rotos or partes or en_formulas)
    return ok, (
        f"nombres con '[': {len(externos)}; nombres rotos: {len(rotos)}; "
        f"xl/externalLinks: {len(partes)}; formulas con [n]: {len(en_formulas)}"
        + (f" ({_lista(externos + rotos + en_formulas)})" if not ok else "")
    )


def _hojas_xml(z: zipfile.ZipFile) -> list[tuple[str, str]]:
    """(nombre de hoja, parte XML) leyendo workbook.xml y sus relaciones."""
    libro = leer_xml(z.read("xl/workbook.xml"))
    rels = leer_xml(z.read("xl/_rels/workbook.xml.rels"))
    destino = {r.get("Id"): r.get("Target", "") for r in rels.iter(f"{NS_PKG_REL}Relationship")}
    hojas = []
    for s in libro.iter(f"{NS_MAIN}sheet"):
        target = destino.get(s.get(f"{NS_REL}id"), "")
        parte = target.lstrip("/") if target.startswith("/") else f"xl/{target}"
        hojas.append((s.get("name", ""), parte))
    return hojas


def _cierre_par(texto: str, abre: int) -> int | None:
    """Índice del `)` que cierra el `(` de `texto[abre]`, saltando textos `"…"` y hojas `'…'`."""
    nivel, k, comilla = 0, abre, ""
    while k < len(texto):
        ch = texto[k]
        if comilla:
            if ch == comilla:
                if k + 1 < len(texto) and texto[k + 1] == comilla:
                    k += 1  # comilla doblada
                else:
                    comilla = ""
        elif ch in "\"'":
            comilla = ch
        elif ch == "(":
            nivel += 1
        elif ch == ")":
            nivel -= 1
            if nivel == 0:
                return k
        k += 1
    return None


def _partir_args(texto: str) -> list[str]:
    """Argumentos de nivel superior de una llamada (separados por `,` o `;`)."""
    args, nivel, inicio, k, comilla = [], 0, 0, 0, ""
    while k < len(texto):
        ch = texto[k]
        if comilla:
            if ch == comilla:
                if k + 1 < len(texto) and texto[k + 1] == comilla:
                    k += 1
                else:
                    comilla = ""
        elif ch in "\"'":
            comilla = ch
        elif ch in "({":
            nivel += 1
        elif ch in ")}":
            nivel -= 1
        elif ch in ",;" and nivel == 0:
            args.append(texto[inicio:k].strip())
            inicio = k + 1
        k += 1
    args.append(texto[inicio:].strip())
    return args


def _ramas_if(a: list[str]) -> list[str]:
    return a[1:3] if len(a) >= 3 else [*a[1:2], "FALSE"]


def _ramas_switch(a: list[str]) -> list[str]:
    resto = a[1:]
    return [*resto[1::2], *([resto[-1]] if len(resto) % 2 == 1 else [])]


#: Funciones cuyo resultado es uno de sus argumentos (las "ramas"): IF(c, si, no), IFERROR(x, e),
#: IFNA(x, e), CHOOSE(i, a, b, …), IFS(c1, r1, c2, r2, …), SWITCH(x, v1, r1, …, [otro]).
_RAMAS: dict[str, Callable[[list[str]], list[str]]] = {
    "IF": _ramas_if,
    "IFERROR": lambda a: a[:2],
    "IFNA": lambda a: a[:2],
    "CHOOSE": lambda a: a[1:],
    "IFS": lambda a: a[1::2],
    "SWITCH": _ramas_switch,
}
_RE_LLAMADA = re.compile(r"^(?:_xlfn\.|_xlws\.)?([A-Za-z][A-Za-z0-9.]*)\(")


def _resultados_posibles(expr: str, profundidad: int = 0) -> list[str]:
    """Las expresiones que una fórmula puede DEVOLVER, por su forma: se baja por las ramas de
    IF/IFERROR/IFNA/CHOOSE/IFS/SWITCH y por los paréntesis que envuelven todo; lo que no es una
    de esas formas es una hoja del árbol. Un `""` que solo aparece en una condición (`$B12=""`)
    no es un resultado posible."""
    e = expr.strip()
    if e.startswith("="):
        e = e[1:].strip()
    while e.startswith("+"):
        e = e[1:].strip()
    if profundidad > 40 or not e:
        return [e]
    if e.startswith("(") and _cierre_par(e, 0) == len(e) - 1:
        return _resultados_posibles(e[1:-1], profundidad + 1)
    m = _RE_LLAMADA.match(e)
    if m is not None and _cierre_par(e, m.end() - 1) == len(e) - 1:
        ramas = _RAMAS.get(m.group(1).upper())
        if ramas is not None:
            args = _partir_args(e[m.end() : -1])
            return [r for a in ramas(args) for r in _resultados_posibles(a, profundidad + 1)]
    return [e]


def _texto_vacio_original(datos: bytes) -> set[tuple[str, str]]:
    """(hoja normalizada, celda) del libro RECIBIDO cuyo valor guardado es texto vacío: cadena
    compartida vacía, `t="str"`/`inlineStr` sin texto (XML crudo, no openpyxl, que confunde
    texto vacío con celda vacía; una celda vacía en Excel vale 0, no `""`)."""
    salida: set[tuple[str, str]] = set()
    with zipfile.ZipFile(io.BytesIO(datos)) as z:
        compartidas: list[str] = []
        if "xl/sharedStrings.xml" in z.namelist():
            raiz = leer_xml(z.read("xl/sharedStrings.xml"))
            compartidas = [
                "".join(t.text or "" for t in si.iter(f"{NS_MAIN}t"))
                for si in raiz.iter(f"{NS_MAIN}si")
            ]
        for nombre, parte in _hojas_xml(z):
            if parte not in z.namelist():
                continue
            for c in leer_xml(z.read(parte)).iter(f"{NS_MAIN}c"):
                tipo, v = c.get("t", "n"), c.find(f"{NS_MAIN}v")
                texto: str | None = None
                if tipo == "s" and v is not None and (v.text or "").strip().isdigit():
                    k = int(v.text or "0")
                    texto = compartidas[k] if k < len(compartidas) else None
                elif tipo == "str":
                    texto = (v.text or "") if v is not None else ""
                elif tipo == "inlineStr":
                    texto = "".join(t.text or "" for t in c.iter(f"{NS_MAIN}t"))
                if texto == "":
                    salida.add((_norm(nombre), str(c.get("r", "")).upper()))
    return salida


def _d_valores_guardados(ev: Evaluacion) -> tuple[bool, str]:
    """Cada `<c>` con `<f>` tiene `<v>`. Un `<v>` vacío es texto vacío: XlsxWriter escribe así un
    resultado `""` (sin `t="str"`). Se acepta solo si la fórmula DEVUELVE texto vacío por su
    forma: entre sus resultados posibles (`_resultados_posibles`, las ramas de IF/IFERROR/…) hay
    un `""` literal, o una referencia pura a una celda que es texto vacío en el libro RECIBIDO o a
    otra celda de la devolución ya aceptada (punto fijo). Un `""` que solo está en una condición
    no cuenta. Un `<v>` vacío en cualquier otra fórmula es un valor que falta."""
    total = 0
    vacias: dict[tuple[str, str], str] = {}  # (hoja, celda) -> fórmula, con <v> vacío o sin <v>
    sin_v: set[tuple[str, str]] = set()
    with zipfile.ZipFile(io.BytesIO(ev.devolucion or b"")) as z:
        for nombre, parte in _hojas_xml(z):
            raiz = leer_xml(z.read(parte))
            for c in raiz.iter(f"{NS_MAIN}c"):
                f = c.find(f"{NS_MAIN}f")
                if f is None:
                    continue
                total += 1
                v = c.find(f"{NS_MAIN}v")
                if v is not None and (v.text or ""):
                    continue
                celda = (_norm(nombre), str(c.get("r", "")).upper())
                vacias[celda] = f.text or ""
                if v is None:
                    sin_v.add(celda)
    vacio_original = _texto_vacio_original(ev.caso.datos)
    resultados = {k: _resultados_posibles(f) for k, f in vacias.items() if k not in sin_v}
    aceptadas = {k for k, rs in resultados.items() if '""' in rs}
    por_original = 0
    cambio = True
    while cambio:
        cambio = False
        for k, rs in resultados.items():
            if k in aceptadas:
                continue
            for r in rs:
                m = _RE_REF_PURA.match(r)
                if m is None:
                    continue
                destino = (_norm(m.group(1)) if m.group(1) else k[0], f"{m.group(2)}{m.group(3)}")
                if destino in aceptadas or destino in vacio_original:
                    por_original += destino in vacio_original
                    aceptadas.add(k)
                    cambio = True
                    break
    sin_valor = [f"{h}!{r}" for h, r in vacias if (h, r) not in aceptadas]
    return not sin_valor and total > 0, (
        f"{total - len(sin_valor)}/{total} formulas con valor guardado "
        f"({len(aceptadas)} texto vacio por su forma, {por_original} por referencia a texto "
        "vacio del recibido)" + (f"; sin valor: {_lista(sin_valor)}" if sin_valor else "")
    )


def _d_valores_conservados(ev: Evaluacion) -> tuple[bool, str]:
    """TRAZABILIDAD de la devolución: cada fila "Valor guardado conservado" es una fórmula que
    el calculador de Silice no evaluó y cuyo valor viene del libro recibido. Se exige 0, o que la
    celda (columna Origen) esté en la lista blanca local `valores_conservados` de
    `expectativas.json`."""
    origenes = _conservados_vistos(ev)
    if origenes is None:
        return False, "la devolucion no tiene TRAZABILIDAD"
    fuera = [o for o in origenes if o not in ev.caso.lista_blanca]
    return not fuera, (
        f"{len(origenes)} valores guardados conservados; {len(origenes) - len(fuera)} en la "
        f"lista blanca local ({len(ev.caso.lista_blanca)} celdas)"
        + (f"; fuera de la lista: {_lista(fuera)}" if fuera else "")
    )


def _d_p_igual_silice(ev: Evaluacion) -> tuple[bool, str]:
    pus = _precios_silice(ev.resultado)
    if pus is None:
        return False, "ResultadoFusion no expone el P.U. de Silice por fila"
    ws = _hoja(ev.dev_v, "PRESUPUESTO")
    if ws is None:
        return False, "la devolucion no tiene PRESUPUESTO"
    malas = []
    for fila, pu in sorted(pus.items()):
        guardado = _dec(ws[f"P{fila}"].value)
        if not _cerca(guardado, pu, TOL_CELDA):
            malas.append(f"P{fila}: {_fmt(guardado, 4)} vs {_fmt(pu, 4)}")
    # Oráculo aparte de Silice: el P del libro recibido (y, en la fila de H1, P + memoria H1).
    oraculo = _pu_oraculo(ev)
    contra_oraculo = [
        f"P{fila}: {_fmt(_dec(ws[f'P{fila}'].value), 4)} vs oraculo {_fmt(esperado, 4)}"
        for fila, esperado in sorted(oraculo.items())
        if not _cerca(_dec(ws[f"P{fila}"].value), esperado, TOL_CELDA)
    ]
    ok = not malas and not contra_oraculo and bool(oraculo)
    return ok, (
        f"{len(pus) - len(malas)}/{len(pus)} P de PRESUPUESTO = PU Silice; "
        f"{len(oraculo) - len(contra_oraculo)}/{len(oraculo)} = oraculo (P recibido"
        + (f", fila {H1_FILA} + H1" if ev.etiqueta in HOJA_H1 else "")
        + ")"
        + (f"; {_lista(malas + contra_oraculo, 4)}" if malas or contra_oraculo else "")
    )


def _totales_oraculo(etiqueta: str) -> tuple[Totales, Decimal] | None:
    """Totales que debe guardar la devolución, sin Silice: en A y B, los documentados del libro
    (Silice no cambia ningún P.U.); en C, el subtotal documentado de Silice (24 517,85) con el IVA
    documentado del 16 %. Tolerancia: medio céntimo más el redondeo del literal en A y B; un
    céntimo en C (el literal 24 517,85 ya está redondeado)."""
    if etiqueta in ("A", "B-1", "B-2"):
        return TOTALES_DOCUMENTADOS[etiqueta], TOL_CELDA + Decimal("0.001")
    esperado = SUBTOTAL_SILICE.get(etiqueta)
    if esperado is None:
        return None
    sub = esperado[0]
    iva = sub * IVA_DOCUMENTADO
    return Totales(sub, iva, sub + iva), TOL_IMPACTO


def _d_nombres_totales(ev: Evaluacion) -> tuple[bool, str]:
    tot = _totales_silice(ev.resultado)
    oraculo = _totales_oraculo(ev.etiqueta)
    partes, ok = [], oraculo is not None
    campos = oraculo[0].campos() if oraculo else (None, None, None)
    for n, esperado in zip(NOMBRES_TOTALES, campos, strict=True):
        valor = _valor_nombre(ev.dev_f, ev.dev_v, n)
        bien = _cerca(valor, tot.get(n), TOL_CELDA)
        bien_o = oraculo is not None and _cerca(valor, esperado, oraculo[1])
        ok = ok and bien and bien_o
        partes.append(
            f"{n} {_fmt(valor)}"
            + ("" if bien else f" vs Silice {_fmt(tot.get(n))}")
            + ("" if bien_o else f" vs oraculo {_fmt(esperado)}")
        )
    return ok, "; ".join(partes) + (
        f" (oraculo documentado +-{oraculo[1]})" if oraculo else " (sin oraculo)"
    )


def _d_hojas_nuevas(ev: Evaluacion) -> tuple[bool, str]:
    faltan = [h for h in HOJAS_NUEVAS if h not in ev.dev_f.sheetnames]
    return not faltan, "AUDITORIA, TRAZABILIDAD y ALERTAS USD-Bs presentes" if not faltan else (
        f"faltan {_lista(faltan)}; hojas: {_lista(ev.dev_f.sheetnames, 30)}"
    )


def _d_rotulo_pu(ev: Evaluacion) -> tuple[bool, str]:
    objetivo = _norm(ROTULO_PU)
    por_hoja: dict[str, str | None] = {}
    for ws in ev.dev_f.worksheets:
        if not _norm(ws.title).startswith("APU"):
            continue
        por_hoja[ws.title] = next(
            (
                c.coordinate
                for fila in ws.iter_rows()
                for c in fila
                if isinstance(c.value, str) and objetivo in _norm(c.value)
            ),
            None,
        )
    sin = [h for h, c in por_hoja.items() if c is None]
    distintas = Counter(c for c in por_hoja.values() if c is not None)
    destinos = {
        coord
        for n, d in ev.dev_f.defined_names.items()
        if re.fullmatch(r"APU_\w+_PU", n)
        for _, coord in d.destinations
    }
    ok = bool(por_hoja) and not sin and len(distintas) == 1
    return ok, (
        f"{len(por_hoja)} APU; rotulo PU en {dict(distintas) or 'ninguna'}"
        + (f"; sin rotulo: {_lista(sin)}" if sin else "")
        + f"; nombres APU_*_PU apuntan a {_lista(destinos) or 'ninguno'}"
    )


def _d_huella(ev: Evaluacion) -> tuple[bool, str]:
    esperada = hashlib.sha256(ev.caso.datos).hexdigest()
    with zipfile.ZipFile(io.BytesIO(ev.devolucion or b"")) as z:
        if "docProps/custom.xml" not in z.namelist():
            return False, "la devolucion no tiene docProps/custom.xml"
        raiz = leer_xml(z.read("docProps/custom.xml"))
    valor = next(
        ("".join(p.itertext()).strip() for p in raiz if p.get("name") == PROPIEDAD_SHA),
        None,
    )
    if valor is None:
        return False, f"sin propiedad {PROPIEDAD_SHA}"
    ok = valor.lower() == esperada
    return ok, f"{PROPIEDAD_SHA} {'=' if ok else '!='} SHA-256 del recibido ({esperada[:12]}...)"


def _conjunto_actual(ev: Evaluacion) -> Counter[tuple[str, str, str]]:
    return Counter(_clave_hallazgo(h) for h in ev.hallazgos)


def _conjunto_exacto(ev: Evaluacion) -> tuple[Estado, str]:
    """Control de falsos positivos: el multiconjunto (regla, severidad, celdas) de los hallazgos
    del libro debe ser IDÉNTICO al congelado y revisado en `expectativas.json`. Sin expectativa,
    PENDIENTE (no FALLA); recién congelado, PENDIENTE (congelar no demuestra nada)."""
    actual = _conjunto_actual(ev)
    if ev.caso.congelado_ahora:
        return Estado.PENDIENTE, (
            f"congelado en esta corrida ({sum(actual.values())} hallazgos): revisarlo a mano y "
            "volver a correr sin --congelar"
        )
    esperado = ev.caso.hallazgos_esperados
    if esperado is None:
        return Estado.PENDIENTE, (
            f"sin conjunto congelado para el libro en expectativas.json ({sum(actual.values())} "
            "hallazgos); generarlo con --congelar y revisarlo"
        )
    sobran = actual - esperado
    faltan = esperado - actual
    if not sobran and not faltan:
        conteo = Counter(f"{r}/{s[:3]}" for r, s, _ in actual.elements())
        return Estado.OK, f"identico al congelado: {sum(actual.values())} hallazgos " + (
            f"({', '.join(f'{k} {v}' for k, v in sorted(conteo.items()))})"
        )

    def fmt(c: Counter[tuple[str, str, str]]) -> str:
        return _lista((f"{r} {s} {celdas}" for r, s, celdas in c.elements()), 6)

    return Estado.FALLA, (
        f"{sum(actual.values())} hallazgos, congelados {sum(esperado.values())}"
        + (f"; sobran (posible falso positivo): {fmt(sobran)}" if sobran else "")
        + (f"; faltan: {fmt(faltan)}" if faltan else "")
    )


# --- Tabla de metas -------------------------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class MetaLibro:
    codigo: str
    titulo: str
    comprobar: Comprobacion


METAS_REGLAS: dict[str, tuple[MetaLibro, ...]] = {
    "A": (
        MetaLibro("L1-CRONO", "L1 en CRONOGRAMA (>= 5 partidas)", _l1_en("CRONOGRAMA", 5)),
        MetaLibro("L1-AUX", "L1 en AUX", _l1_en("AUX", 1)),
        MetaLibro("L5-RESUMEN", "L5 en los 4 #REF! de RESUMEN", _l5_refs("RESUMEN", 4)),
        MetaLibro("L5-EXT", "L5 en los 11 nombres externos", _l5_nombres(11)),
    ),
    "B-1": (),
    "B-2": (
        MetaLibro("L1-REND", "L1 en REND de partidas 15 y 16 (F26, F27)", _l1_rend_b2),
        MetaLibro("L1-CRONO", "L1 en CRONOGRAMA", _l1_en("CRONOGRAMA", 1)),
        MetaLibro("L1-AUX", "L1 en AUX", _l1_en("AUX", 1)),
        MetaLibro("L9-PROCURA", "L9 en PROCURA con los 4 Dif. de H7", _l9_procura_b2),
        MetaLibro("L5-RESUMEN", "L5 en los 4 #REF! de RESUMEN", _l5_refs("RESUMEN", 4)),
        MetaLibro("L5-EXT", "L5 en los 11 nombres externos", _l5_nombres(11)),
        MetaLibro("L7", "L7 en MEDICIONES (=9.73*2.75, =2.75*3, *0.28)", _l7_mediciones),
        MetaLibro("L6", "L6 utilidad 30 % en APU-09!G59 (H18, -746,98)", _l6_util30_b2),
    ),
    "C-1": (
        MetaLibro("L2", "L2 del conflicto y parches en filas 23 y 24", _c1_l3_o_l2),
        MetaLibro("F1", "F1 en fila 24, impacto +205,92 (H1)", _f1_h1),
        MetaLibro("L4", "L4 en APU-09 (CD =G12+G21), +205,92", _l4_h1),
        MetaLibro("L6", "L6 utilidad 2 % (+761,62) y /12 en APU-09", _l6_h2),
    ),
    "C-2": (
        MetaLibro("F1", "F1 en fila 24, impacto +205,92 (H1)", _f1_h1),
        MetaLibro("L4", "L4 en APU-07 (CD =G12+G21), +205,92", _l4_h1),
        MetaLibro("L6", "L6 utilidad 2 % (+761,62) y /12 en APU-07", _l6_h2),
        MetaLibro("L8", "L8 en los numeros 6 y 7", _l8_c2),
        MetaLibro("L3-RL7", "L3 o (RL-7) L2 ADV en filas 19 y 22", _c2_l3_o_parche),
        MetaLibro("L9-RESUMEN", "L9 en RESUMEN, +183,53 (H3)", _l9_resumen_c2),
        MetaLibro("L5-CRONO", "L5 en los 6 #REF! de CRONOGRAMA", _l5_refs("CRONOGRAMA", 6)),
    ),
}

#: Partidas documentadas de cada libro (P.U. Silice = P guardado; en C, salvo la fila de H1, que
#: lleva la memoria de H1).
PARTIDAS_DOCUMENTADAS: dict[str, int] = {"A": 5, "B-1": 16, "B-2": 16, "C-1": 9, "C-2": 9}


def _totales_del_libro(ev: Evaluacion) -> tuple[bool, str]:
    """Oráculo de identificación: los totales guardados del libro recibido (openpyxl, por sus
    nombres) son los documentados de ese libro."""
    doc = TOTALES_DOCUMENTADOS[TOTALES_DEL_LIBRO[ev.etiqueta]]
    guardados = [_valor_nombre(ev.orig_f, ev.orig_v, n) for n in NOMBRES_TOTALES]
    ok = all(_cerca(g, d, TOL_CELDA) for g, d in zip(guardados, doc.campos(), strict=True))
    return ok, (
        f"guardados {'/'.join(_fmt(g) for g in guardados)}; documentados "
        f"{'/'.join(_fmt(d) for d in doc.campos())}"
    )


METAS_COMUNES: tuple[MetaLibro, ...] = (
    MetaLibro("TOT-LIBRO", "Totales guardados = documentados", _totales_del_libro),
    MetaLibro("CONJUNTO-EXACTO", "Hallazgos = conjunto congelado", _conjunto_exacto),
    MetaLibro("L5-TODO", "L5 en todo #REF! y nombre externo", _l5_todo),
    MetaLibro("L11", "L11 = literal del PDF vs libro, con signo", _l11),
    MetaLibro("TOT", "Totales de Silice coherentes y documentados", _totales_coherentes),
)

METAS_DEVOLUCION: tuple[MetaLibro, ...] = (
    MetaLibro("D-REABRE", "Devolucion reabre con openpyxl", _d_reabre),
    MetaLibro("D-REF", "Devolucion sin #REF!", _requiere_devolucion(_d_sin_ref)),
    MetaLibro("D-INDIRECT", "Devolucion sin INDIRECT", _requiere_devolucion(_d_sin_indirect)),
    MetaLibro(
        "D-VINCULOS",
        "Devolucion sin vinculos ni nombres rotos",
        _requiere_devolucion(_d_sin_vinculos),
    ),
    MetaLibro(
        "D-VALORES", "Cada formula con valor guardado", _requiere_devolucion(_d_valores_guardados)
    ),
    MetaLibro(
        "VALORES-CONSERVADOS",
        "Valores guardados conservados: 0 o lista blanca",
        _requiere_devolucion(_d_valores_conservados),
    ),
    MetaLibro(
        "D-P", "P de PRESUPUESTO = PU Silice = oraculo", _requiere_devolucion(_d_p_igual_silice)
    ),
    MetaLibro(
        "D-TOTALES",
        "SUBSINIVA/IVATOT/TOTALGEN = Silice = documentado",
        _requiere_devolucion(_d_nombres_totales),
    ),
    MetaLibro(
        "D-HOJAS", "Hojas AUDITORIA, TRAZABILIDAD, ALERTAS", _requiere_devolucion(_d_hojas_nuevas)
    ),
    MetaLibro(
        "D-PU-CELDA", "Rotulo PU en la misma celda en todo APU", _requiere_devolucion(_d_rotulo_pu)
    ),
    MetaLibro("D-SHA", "Propiedad silice_sha256_origen", _requiere_devolucion(_d_huella)),
)


def _metas_de(etiqueta: str) -> tuple[MetaLibro, ...]:
    propias = METAS_REGLAS.get(etiqueta, ())
    pu = MetaLibro(
        "PU", "PU Silice = P guardado (+H1)", _pu_igual_guardado(PARTIDAS_DOCUMENTADAS[etiqueta])
    )
    return (*propias, *METAS_COMUNES, pu, *METAS_DEVOLUCION)


# --- PDF simples (L10) ----------------------------------------------------------------------------

_SUBTIPOS_L10: tuple[tuple[str, str, str], ...] = (
    ("unidades", "suma de unidades distintas", ""),
    ("incoherente", "unidad × cantidad × precio incoherente", ""),
    ("pendiente", "precio pendiente", ""),
    ("MO fuera del total", "nota que contradice el total", "mano de obra"),
)


def _con_subtipo(hallazgos: Iterable[Any], prefijo: str, contiene: str = "") -> list[Any]:
    p, c = _norm(prefijo), _norm(contiene)
    return [
        h
        for h in hallazgos
        if h.regla == "L10" and _norm(h.descripcion).startswith(p) and c in _norm(h.descripcion)
    ]


def _revisar_pdf_simple(datos: bytes, nombre: str) -> tuple[Any, list[Any]]:
    """Por la entrada del producto (`fusion.proceso.auditar_pdf_simple`, la que usan la UI y la
    API); si esa versión de `fusion` aún no la tiene, por el adaptador directamente."""
    try:
        from fusion.proceso import auditar_pdf_simple
    except ImportError:
        from adapters.hoja_empresa.pdf import leer_pdf_simple, reglas_pdf_simple

        pdf = leer_pdf_simple(datos, nombre)
        return pdf, list(reglas_pdf_simple(pdf))
    resultado = auditar_pdf_simple(datos, nombre)
    return resultado.pdf, list(resultado.hallazgos)


def _metas_pdf_simples(rutas: Sequence[Path]) -> list[Fila]:
    filas: list[Fila] = []
    for k, ruta in enumerate(rutas, start=1):
        try:
            pdf, hallazgos = _revisar_pdf_simple(ruta.read_bytes(), f"pdf_simple_{k}.pdf")
        except Exception as e:  # noqa: BLE001 - la meta informa la falla, no la oculta
            filas.append(Fila(f"PDF-{k}.L10", "L10 sobre el PDF simple", Estado.FALLA, _falla(e)))
            continue
        numero = pdf.numero if pdf.numero in ("001", "002") else f"{k:03d}"
        faltan = [n for n, pre, con in _SUBTIPOS_L10 if not _con_subtipo(hallazgos, pre, con)]
        filas.append(
            Fila(
                f"PDF-{numero}.L10",
                "L10: unidades, incoherente, pendiente, MO",
                Estado.FALLA if faltan else Estado.OK,
                f"{len(hallazgos)} L10; "
                + (f"faltan {_lista(faltan)}" if faltan else "los 4 subtipos presentes"),
            )
        )
        if numero != "002":
            continue
        arit = _con_subtipo(hallazgos, "aritmética de fila")
        items = {i for i in (14, 18) if any(f"ITEM {i} " in _norm(h.descripcion) for h in arit)}
        suma = sum((_dec(h.impacto) or Decimal(0) for h in arit), Decimal(0))
        ok = items == {14, 18} and _cerca(suma, Decimal("4.65"), TOL_CELDA)
        filas.append(
            Fila(
                "PDF-002.L10-ARIT",
                "L10 aritmetica en items 14 y 18 (+4,65)",
                Estado.OK if ok else Estado.FALLA,
                f"items {_lista(items) or 'ninguno'}; impacto total {_fmt(suma)} (H17: 4,65)",
            )
        )
        comp = _con_subtipo(hallazgos, "cómputos ≠ presupuesto")
        pares = {(_dec(h.valor_observado), _dec(h.valor_esperado)) for h in comp}
        esperados = {
            (Decimal(90), Decimal(80)),
            (Decimal(4), Decimal(3)),
            (Decimal(5), Decimal(10)),
        }
        faltan_p = [f"{e[1]}->{e[0]}" for e in esperados if e not in pares]
        filas.append(
            Fila(
                "PDF-002.L10-COMP",
                "L10 computos != presupuesto (H16)",
                Estado.FALLA if faltan_p else Estado.OK,
                f"{len(comp)} hallazgos; "
                + (
                    f"faltan {_lista(faltan_p)}" if faltan_p else "80->90 m, 3->4 y 10->5 presentes"
                ),
            )
        )
    return filas


# --- Orquestación ---------------------------------------------------------------------------------


def _falla(e: BaseException) -> str:
    marco = traceback.extract_tb(e.__traceback__)[-1] if e.__traceback__ else None
    donde = f" en {Path(marco.filename).name}:{marco.lineno}" if marco else ""
    return f"{type(e).__name__}{donde}: {str(e)[:160]}"


@dataclass(slots=True)
class PdfCarpeta:
    """Un PDF de la carpeta con su identificación por literal documentado (sin el adaptador)."""

    ruta: Path
    sha: str
    clave: str | None
    como: str  # "sha256", "literal", o por qué no se identificó


def _pdfs(carpeta: Path, exp: Expectativas) -> list[PdfCarpeta]:
    salida = []
    for ruta in sorted(carpeta.glob("*.pdf")):
        datos = ruta.read_bytes()
        try:
            clave, como = _identificar_pdf(datos, exp)
        except Exception as e:  # noqa: BLE001 - un PDF ilegible queda sin identificar
            clave, como = None, f"no se pudo leer ({type(e).__name__})"
        salida.append(PdfCarpeta(ruta, hashlib.sha256(datos).hexdigest(), clave, como))
    return sorted(salida, key=lambda p: p.sha)


def _casos(
    carpeta: Path, exp: Expectativas, pdfs: Sequence[PdfCarpeta]
) -> tuple[dict[str, Caso], list[str]]:
    """Identifica los libros. Devuelve {etiqueta: Caso} y avisos (sin nombres de archivo)."""
    from adapters.hoja_empresa.lector import leer_libro

    id_pdf = {p.ruta: p.clave for p in pdfs}
    casos: dict[str, Caso] = {}
    avisos: list[str] = []
    for ruta in sorted(carpeta.glob("*.xlsx")):
        if ruta.name.startswith("~$"):
            continue
        datos = ruta.read_bytes()
        sha = hashlib.sha256(datos).hexdigest()
        huella = sha[:8]
        pdf_nombres: list[str] = []
        if sha in exp.libros_sha:
            etiqueta, como = exp.libros_sha[sha], "expectativas.json (sha256)"
            pdf_nombres = exp.libros.get(ruta.name.lower(), ("", []))[1]
        elif ruta.name.lower() in exp.libros:
            etiqueta, pdf_nombres = exp.libros[ruta.name.lower()]
            como = "expectativas.json"
        else:
            try:
                encontrada = _identificar(leer_libro(datos, ruta.name))
            except Exception as e:  # noqa: BLE001
                avisos.append(f"libro sha {huella}: el lector fallo ({_falla(e)})")
                continue
            como = "contenido"
            if encontrada is None:
                avisos.append(f"libro sha {huella}: no coincide con ninguna firma (ignorado)")
                continue
            etiqueta = encontrada
        if etiqueta in casos:
            avisos.append(f"libro sha {huella}: otro libro {etiqueta} (se usa el primero)")
            continue
        caso = Caso(etiqueta, ruta, datos, como)
        if pdf_nombres:
            rutas_pdf = [carpeta / n for n in pdf_nombres if (carpeta / n).is_file()]
            caso.pdf, caso.pdf_extra = (rutas_pdf[0], rutas_pdf[1:]) if rutas_pdf else (None, [])
        else:
            caso.pdf = next(
                (p.ruta for p in pdfs if p.ruta.stem.lower() == ruta.stem.lower()), None
            )
            if etiqueta == "C-2":
                # Los otros PDF identificados como "C" (por SHA-256 o literal, no por el lector).
                caso.pdf_extra = [p.ruta for p in pdfs if p.clave == "C" and p.ruta != caso.pdf]
        caso.pdf_id = id_pdf.get(caso.pdf) if caso.pdf else None
        caso.pdf_extra_ids = [id_pdf.get(r) for r in caso.pdf_extra]
        caso.hallazgos_esperados = exp.hallazgos.get(etiqueta)
        caso.lista_blanca = exp.lista_blanca.get(etiqueta, frozenset())
        casos[etiqueta] = caso
    return casos, avisos


def _metas_pdf_entregados(pdfs: Sequence[PdfCarpeta]) -> list[Fila]:
    """Cada PDF de la carpeta: la lectura del producto (`leer_totales_impresos`, el último
    bloque) debe dar el literal documentado con que se identificó el PDF."""
    from adapters.hoja_empresa.pdf import leer_totales_impresos

    filas: list[Fila] = []
    vistos: Counter[str] = Counter()
    for p in pdfs:
        clave = p.clave or "?"
        vistos[clave] += 1
        codigo = f"PDF-{clave}#{vistos[clave]}.LECTURA"
        titulo = "Totales impresos leidos = literal documentado"
        if p.clave is None:
            filas.append(
                Fila(codigo, titulo, Estado.PENDIENTE, f"sha {p.sha[:10]}: sin oraculo ({p.como})")
            )
            continue
        doc = TOTALES_DOCUMENTADOS[p.clave]
        try:
            t = leer_totales_impresos(p.ruta.read_bytes(), "pdf_entregado.pdf")
            leido = [_dec(t.subtotal), _dec(t.iva), _dec(t.total)]
        except Exception as e:  # noqa: BLE001 - la meta informa la falla, no la oculta
            filas.append(Fila(codigo, titulo, Estado.FALLA, _falla(e)))
            continue
        ok = all(_cerca(a, b, TOL_CELDA) for a, b in zip(leido, doc.campos(), strict=True))
        filas.append(
            Fila(
                codigo,
                titulo,
                Estado.OK if ok else Estado.FALLA,
                f"sha {p.sha[:10]} (por {p.como}): leido {'/'.join(_fmt(x) for x in leido)}; "
                f"documentado {'/'.join(_fmt(x) for x in doc.campos())}",
            )
        )
    return filas


def _evaluar_libro(
    caso: Caso,
    procesar: Callable[..., Any],
    devolver: Callable[[Any], bytes],
    args: Any,
    otro: bytes | None = None,
) -> Evaluacion:
    """Procesa un libro con nombres neutros (`libro_A.xlsx`): así ningún nombre de archivo real
    entra en los hallazgos ni en la devolución que se imprime. `otro` es la otra versión del mismo
    libro (B-1 para B-2), para las alertas "contra otra versión" (spec §4.3)."""
    ev = Evaluacion(caso)
    ev.orig_f, ev.orig_v = _abrir_wb(caso.datos, False), _abrir_wb(caso.datos, True)
    pdf = caso.pdf.read_bytes() if caso.pdf else None
    try:
        ev.resultado = procesar(
            caso.datos,
            f"libro_{caso.etiqueta}.xlsx",
            pdf=pdf,
            pdf_nombre=f"libro_{caso.etiqueta}.pdf" if pdf else "",
            contexto=None,
            fecha_libro=args.fecha_libro,
            fecha_devolucion=args.fecha_devolucion,
            otro=otro,
        )
        ev.hallazgos = list(ev.resultado.informe.hallazgos)
    except Exception as e:  # noqa: BLE001
        ev.error = _falla(e)
        return ev
    if caso.pdf_extra:
        from adapters.hoja_empresa.pdf import cotejar_con_libro, leer_totales_impresos

        for k, (ruta, clave) in enumerate(
            zip(caso.pdf_extra, caso.pdf_extra_ids, strict=True), start=1
        ):
            nombre = f"libro_{caso.etiqueta}_extra{k}.pdf"
            t = leer_totales_impresos(ruta.read_bytes(), nombre)
            ev.pdf_extra.append((nombre, clave, cotejar_con_libro(t, ev.resultado.libro)))
    try:
        ev.devolucion = devolver(ev.resultado)
    except Exception as e:  # noqa: BLE001
        ev.error_devolucion = _falla(e)
        return ev
    args.salida.mkdir(parents=True, exist_ok=True)
    ev.ruta_devolucion = args.salida / f"devolucion_libro_{caso.etiqueta}.xlsx"
    ev.ruta_devolucion.write_bytes(ev.devolucion)
    try:
        ev.dev_f = openpyxl.load_workbook(ev.ruta_devolucion, data_only=False)
        ev.dev_v = openpyxl.load_workbook(ev.ruta_devolucion, data_only=True)
    except Exception as e:  # noqa: BLE001
        ev.dev_f = ev.dev_v = None
        ev.error_reabrir = _falla(e)
    return ev


def _conservados_vistos(ev: Evaluacion) -> list[str] | None:
    """Origen (`HOJA!CELDA` normalizada) de cada fila "Valor guardado conservado" de la hoja
    TRAZABILIDAD de la devolución (columna A el dato, C el origen). None si no hay TRAZABILIDAD."""
    hojas = ev.dev_v.worksheets if ev.dev_v is not None else []
    ws = next((w for w in hojas if _norm(w.title).startswith("TRAZABILIDAD")), None)
    if ws is None:
        return None
    objetivo = _norm("Valor guardado conservado")
    return sorted(
        _norm_origen(str(ws.cell(f, 3).value or ""))
        for f in range(1, ws.max_row + 1)
        if _norm(ws.cell(f, 1).value or "") == objetivo
    )


def _congelar(
    exp: Expectativas,
    evaluaciones: Sequence[Evaluacion],
    pdfs: Sequence[PdfCarpeta],
) -> None:
    """Escribe en `exp.ruta` (fusionando con lo que ya tenga) el conjunto actual de hallazgos, la
    identificación por SHA-256 de libros y PDF y los valores conservados vistos. No toca la lista
    blanca `valores_conservados` (la escribe una persona)."""
    crudo = dict(exp.crudo) if exp.existe else {}
    if crudo and not (_CLAVES_EXPECTATIVAS & set(crudo)):
        crudo = {"libros": crudo}  # formato antiguo: se conserva bajo "libros"
    libros_sha = dict(crudo.get("libros_sha256", {}))
    hallazgos = dict(crudo.get("hallazgos", {}))
    vistos = dict(crudo.get("valores_conservados_vistos", {}))
    for ev in evaluaciones:
        if ev.error:
            continue
        libros_sha[hashlib.sha256(ev.caso.datos).hexdigest()] = ev.etiqueta
        hallazgos[ev.etiqueta] = sorted([list(k) for k in _conjunto_actual(ev).elements()])
        vistos[ev.etiqueta] = _conservados_vistos(ev) or []
        ev.caso.congelado_ahora = True
    pdf_sha = dict(crudo.get("pdf_sha256", {}))
    pdf_sha.update({p.sha: p.clave for p in pdfs if p.clave is not None})
    crudo.update(
        {
            "libros_sha256": dict(sorted(libros_sha.items())),
            "pdf_sha256": dict(sorted(pdf_sha.items())),
            "hallazgos": dict(sorted(hallazgos.items())),
            "valores_conservados": crudo.get("valores_conservados", {}),
            "valores_conservados_vistos": dict(sorted(vistos.items())),
        }
    )
    exp.ruta.parent.mkdir(parents=True, exist_ok=True)
    exp.ruta.write_text(json.dumps(crudo, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")


def _filas_de(ev: Evaluacion) -> list[Fila]:
    prefijo = f"{ev.etiqueta}."
    filas = [
        Fila(
            prefijo + "PROCESO",
            f"Libro {ev.etiqueta}: procesar_libro (por {ev.caso.como})",
            Estado.FALLA if ev.error else Estado.OK,
            ev.error or f"{len(ev.hallazgos)} hallazgos",
        )
    ]
    for meta in _metas_de(ev.etiqueta):
        if ev.error:
            filas.append(Fila(prefijo + meta.codigo, meta.titulo, Estado.FALLA, "no se proceso"))
            continue
        try:
            resultado, evidencia = meta.comprobar(ev)
        except Exception as e:  # noqa: BLE001
            resultado, evidencia = False, f"la comprobacion fallo: {_falla(e)}"
        estado = (
            resultado
            if isinstance(resultado, Estado)
            else (Estado.OK if resultado else Estado.FALLA)
        )
        filas.append(Fila(prefijo + meta.codigo, meta.titulo, estado, evidencia))
    return filas


def _prohibidos(carpeta: Path, evaluaciones: Iterable[Evaluacion]) -> list[str]:
    """Textos que no pueden salir por pantalla: nombres de archivo y datos de PORTADA."""
    textos = {p.stem for p in carpeta.iterdir() if p.is_file()}
    for ev in evaluaciones:
        portada = getattr(getattr(ev.resultado, "libro", None), "portada", None)
        for campo in ("cliente", "obra", "contratista", "ubicacion"):
            textos.add(str(getattr(portada, campo, "") or ""))
    return sorted((t for t in textos if len(t.strip()) >= 4), key=len, reverse=True)


def _limpiar(texto: str, prohibidos: Sequence[str]) -> str:
    for p in prohibidos:
        texto = re.sub(re.escape(p), "[...]", texto, flags=re.IGNORECASE)
    return texto


def _resumen(evaluaciones: Sequence[Evaluacion], avisos: Sequence[str], salida: Path) -> str:
    lineas = ["", "RESUMEN"]
    for ev in evaluaciones:
        conteo = Counter(f"{h.regla}/{h.severidad.name[:3]}" for h in ev.hallazgos)
        detalle = ", ".join(f"{k} {v}" for k, v in sorted(conteo.items())) or "sin hallazgos"
        lineas.append(f"  Libro {ev.etiqueta} ({ev.caso.como}): {detalle}")
    for aviso in avisos:
        lineas.append(f"  Aviso: {aviso}")
    if any(ev.ruta_devolucion for ev in evaluaciones):
        lineas.append(f"  Devoluciones en: {salida}")
    return "\n".join(lineas)


def _fecha(texto: str) -> date:
    return date.fromisoformat(texto)


def _parsear_args(argv: list[str] | None) -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__.split("\n", 1)[0])
    p.add_argument("carpeta", type=Path, help="carpeta local con los libros .xlsx y sus PDF")
    p.add_argument("--salida", type=Path, default=SALIDA_POR_DEFECTO)
    p.add_argument("--json", action="store_true", help="metas en JSON en vez de tabla")
    p.add_argument("--fecha-libro", type=_fecha, default=FECHA_LIBRO)
    p.add_argument("--fecha-devolucion", type=_fecha, default=date.today())
    p.add_argument("--pdf-simple", type=Path, nargs="*", default=list(PDF_SIMPLES))
    p.add_argument(
        "--expectativas",
        type=Path,
        default=None,
        help="expectativas.json LOCAL a leer (y a escribir con --congelar); por defecto "
        "CARPETA/expectativas.json",
    )
    p.add_argument(
        "--congelar",
        action="store_true",
        help="escribe el conjunto actual de hallazgos en las expectativas (revisarlo a mano)",
    )
    return p.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    args = _parsear_args(argv)
    if not args.carpeta.is_dir():
        print(f"No existe la carpeta de libros: {args.carpeta}", file=sys.stderr)
        return 2
    try:
        from fusion.proceso import devolver, procesar_libro
    except ImportError as e:
        print(
            f"Falta fusion.proceso: corre esto cuando V2 este integrada ({type(e).__name__}: {e}).",
            file=sys.stderr,
        )
        return 2
    ruta_exp = args.expectativas or args.carpeta / "expectativas.json"
    if args.congelar and _versionable(ruta_exp):
        print(
            "--congelar no escribe en una ruta que Git podria versionar (esta dentro de un "
            "repositorio y no esta ignorada); usa --expectativas con una ruta fuera del "
            "repositorio o ignorala en .gitignore.",
            file=sys.stderr,
        )
        return 2
    try:
        exp = _leer_expectativas(ruta_exp)
    except (ValueError, json.JSONDecodeError) as e:
        print(f"expectativas.json invalido: {e}", file=sys.stderr)
        return 2

    pdfs = _pdfs(args.carpeta, exp)
    casos, avisos = _casos(args.carpeta, exp, pdfs)
    if exp.existe and _versionable(exp.ruta):
        avisos.append(
            "expectativas.json esta en una ruta que Git no ignora: no lo anadas al repositorio"
        )
    for p in pdfs:
        if p.clave is None:
            avisos.append(f"PDF sha {p.sha[:8]}: sin literal documentado ({p.como})")
    evaluaciones = [
        _evaluar_libro(
            casos[e],
            procesar_libro,
            devolver,
            args,
            otro=casos[OTRA_VERSION[e]].datos if OTRA_VERSION.get(e) in casos else None,
        )
        for e in ETIQUETAS
        if e in casos
    ]
    if args.congelar:
        _congelar(exp, evaluaciones, pdfs)
        avisos.append(
            f"conjunto de hallazgos congelado en {exp.ruta} (LOCAL; revisarlo a mano antes de "
            "usarlo como oraculo)"
        )
    filas: list[Fila] = []
    for etiqueta in ETIQUETAS:
        if etiqueta not in casos:
            filas.append(
                Fila(
                    f"{etiqueta}.PRESENTE",
                    f"Libro {etiqueta} en la carpeta",
                    Estado.FALLA,
                    "no se identifico en la carpeta",
                )
            )
    for ev in evaluaciones:
        filas.extend(_filas_de(ev))
    filas.extend(_metas_pdf_entregados(pdfs))
    filas.extend(_metas_pdf_simples(args.pdf_simple))

    prohibidos = _prohibidos(args.carpeta, evaluaciones)
    filas = [Fila(f.codigo, f.titulo, f.estado, _limpiar(f.evidencia, prohibidos)) for f in filas]
    if args.json:
        print(_filas_a_json(filas))
    else:
        print(render_tabla(filas))
        print(_limpiar(_resumen(evaluaciones, avisos, args.salida), prohibidos))
    return 1 if any(f.estado is Estado.FALLA for f in filas) else 0


if __name__ == "__main__":
    raise SystemExit(main())
